"""Public synthetic table fixtures, rendered by the same compiler as the app."""
import os,sys,subprocess
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root))
from backend.document_preview import prepare
from backend import officecli_bridge

folder=root/'dist/qa/word-table/synthetic';folder.mkdir(parents=True,exist_ok=True)
for case in ['paragraphs','single-paragraph','nested-table','repeating-header','vertical-merge','long-body-paragraph']:
    doc=Document();doc.add_heading('Word table pagination fixture',0)
    if case=='long-body-paragraph':
        paragraph=doc.add_paragraph()
        for i in range(400):
            run=paragraph.add_run(f'Long body sentence {i} retains its complete text and emphasis. ')
            run.bold=i%3==0
    table=doc.add_table(rows=2,cols=2);table.style='Table Grid'
    table.cell(0,0).text='Department';table.cell(0,1).text='Future Display Expansion'
    if case=='repeating-header':table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    if case=='vertical-merge':
        table.cell(0,0).merge(table.cell(1,0)).text='Merged label'
    if case in {'paragraphs','single-paragraph','nested-table'}:
        cell=table.cell(1,0).merge(table.cell(1,1))
        if case=='single-paragraph':cell.text=' '.join(f'Sentence {i} preserves complete styled text.' for i in range(300))
        else:
            for i in range(55):
                p=cell.add_paragraph(f'Paragraph {i}: complete source content must survive every table page break. '*3)
                if i%3==0:p.runs[0].bold=True
            if case=='nested-table':
                nested=cell.add_table(rows=12,cols=2)
                for i,row in enumerate(nested.rows):row.cells[0].text=f'Action {i}';row.cells[1].text='Preserve the nested cell and column alignment.'
    else:
        for i in range(55):
            cells=table.add_row().cells;cells[0].text=f'Row {i}';cells[1].text='A complete table row with a long description. '*4
    source=folder/(case+'.docx');raw=folder/(case+'-raw.html');output=folder/(case+'.html')
    doc.save(source)
    p=subprocess.run([str(officecli_bridge.ensure_component('qa_preview')),'view',str(source),'html','--out',str(raw)],
        capture_output=True,timeout=90,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0,
        env={**os.environ,'OFFICECLI_NO_AUTO_RESIDENT':'1','OFFICECLI_SKIP_UPDATE':'1'})
    if p.returncode:raise RuntimeError('Fixture render failed: '+case)
    output.write_text(prepare(raw.read_text('utf-8-sig'),'.docx'),'utf-8')
    print(output,flush=True)
