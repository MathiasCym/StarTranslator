"""从《深圳展厅讲解稿.xlsx》抽取中英术语对，生成术语库。

用法：
    python tools/build_glossary.py

源数据是段落级对齐（第7列中文 / 第8列英文），不是术语列表，
因此需要让 LLM 从成对段落中抽取「术语级」的对照关系。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from backend.glossary.loader import Glossary
from backend.llm.client import LLMClient, LLMUnavailable
from backend.models import GlossaryEntry

SRC_XLSX = r"C:\Users\Administrator\Desktop\深圳展厅讲解稿.xlsx"
SHEET = "讲解稿V3.0"
ZH_COL, EN_COL = 7, 8          # 商务部版本 / 商务版英文翻译
HEADER_ROWS = 4

EXTRACT_PROMPT = """你是 LED 显示行业的技术翻译专家。下面给你一段中文原文和它对应的英文译文。

请从中抽取**术语级**的中英对照，即那些在 LED 显示行业里有固定专业译法的名词或名词短语。

要求：
1. 只抽取专业术语，不要抽整句、不要抽无专业含义的普通词（如 "welcome"、"we"）
2. 优先抽取：产品/部件（接收卡、箱体、驱动IC）、技术概念（点间距、逐点校正、灰度等级）、
   行业方案（xR、分布式、可视化）等
3. 中文和英文必须能在原文中找到对应关系，不要自己编造
4. **每条只给一个标准英文译法，不要列举同义写法**——同义表达（如 COB screen
   与 COB LED display）本来就不是错误，收录为"别名"只会在检验时制造噪音

只输出 JSON 数组，格式：
[{{"zh":"中文术语","en":"标准英文译法","note":"简短说明"}}]

中文原文：
{zh}

英文译文：
{en}
"""


def load_paragraph_pairs() -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(SRC_XLSX)
    ws = wb[SHEET]
    pairs = []
    for r in range(HEADER_ROWS + 1, ws.max_row + 1):
        zh = ws.cell(row=r, column=ZH_COL).value
        en = ws.cell(row=r, column=EN_COL).value
        if isinstance(zh, str) and isinstance(en, str) and zh.strip() and en.strip():
            pairs.append((zh.strip(), en.strip()))
    return pairs


def main() -> None:
    pairs = load_paragraph_pairs()
    print(f"读到 {len(pairs)} 个中英段落对")

    client = LLMClient()
    glossary = Glossary([])
    failed = 0

    for i, (zh, en) in enumerate(pairs, 1):
        # 段落可能很长，截断避免超 token（术语通常出现在前段）
        prompt = EXTRACT_PROMPT.format(zh=zh[:2500], en=en[:6000])
        try:
            result = client.chat_json(prompt, max_tokens=4000)
        except LLMUnavailable as e:
            print(f"  [{i}/{len(pairs)}] 失败：{e}")
            failed += 1
            continue

        if not isinstance(result, list):
            print(f"  [{i}/{len(pairs)}] 返回格式非数组，跳过")
            failed += 1
            continue

        added = 0
        for item in result:
            if not isinstance(item, dict):
                continue
            zh_t = str(item.get("zh", "")).strip()
            en_t = str(item.get("en", "")).strip()
            if not zh_t or not en_t:
                continue
            glossary.add(GlossaryEntry(
                zh=zh_t, en=en_t,
                variants=[],        # 自动抽取的别名不入库，它们是误报之源
                priority=2,
                note=str(item.get("note", "")).strip(),
                source=f"展厅讲解稿:{SHEET}",
            ))
            added += 1
        print(f"  [{i}/{len(pairs)}] 抽到 {added} 条")

    print(f"\n原始 {len(glossary.entries)} 条（{failed} 段失败）")
    glossary = dedupe(glossary)
    print(f"去重后 {len(glossary.entries)} 条")

    out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "glossary", "led_glossary.json")
    glossary.save(out)
    print(f"已保存到 {out}")


# 同一中文概念被多个段落重复抽出时，保留更完整的那条英文译法。
#
# 注意：**不再把其余写法并入 variants。** 原实现这么做，结果是
# 「COB screen」被收进「COB LED display」的 variants，检验时命中即报错——
# 而两者本来同义。别名不是错误清单，判决权已交给模型（见 matcher/judge.py）。
def dedupe(glossary: Glossary) -> Glossary:
    best: dict[str, GlossaryEntry] = {}
    for e in glossary.entries:
        key = e.zh
        if key not in best:
            best[key] = GlossaryEntry(
                zh=e.zh, en=e.en, variants=[],
                priority=e.priority, note=e.note, source=e.source,
            )
            continue
        cur = best[key]
        # 保留更长的英文（通常更完整、更专业）
        if len(e.en) > len(cur.en):
            cur.en = e.en
        if not cur.note and e.note:
            cur.note = e.note
    return Glossary(list(best.values()))


if __name__ == "__main__":
    main()
