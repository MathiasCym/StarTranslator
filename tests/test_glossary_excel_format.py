import json,os,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from copy import copy
import openpyxl
from openpyxl.styles import Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend.glossary import exchange,writer,store,excel_format


class ExcelFormatTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});self.env.start()
        self.active=self.root/'active.json';store.atomic_json(self.active,[])
        store.atomic_json(self.root/'profile/lexicon-settings.json',{'path':str(self.active)})
    def tearDown(self):self.env.stop();self.temp.cleanup()
    def test_standard_export_matches_reference_and_wraps_without_changing_values(self):
        entries=store.validate([{'zh':'很长的中文术语名称'*4,'en':'Long English terminology '*8,'note':'长段说明'*80,
                                'variants':['one alias','another alias'],'priority':1,'source':'test'}])
        store.atomic_json(self.active,entries)
        target=self.root/'standard.xlsx';exchange.export_file(str(target),'')
        book=openpyxl.load_workbook(target);sheet=book.active
        self.assertEqual(sheet.title,'术语表')
        for col,field in enumerate(writer.HEADER_FIELDS[:6],1):
            spec=excel_format.preset()['columns'][field]
            self.assertAlmostEqual(sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width,spec['width'])
            self.assertEqual(sheet.cell(2,col).font.name,'微软雅黑')
            self.assertTrue(sheet.cell(2,col).alignment.wrap_text)
            self.assertEqual(sheet.cell(2,col).border.left.style,'thin')
        self.assertTrue(sheet['A2'].font.bold);self.assertTrue(sheet['B2'].font.bold)
        self.assertFalse(sheet['C2'].font.bold)
        self.assertEqual(sheet['A2'].fill.fgColor.theme,4);self.assertEqual(sheet['B2'].fill.fgColor.theme,8)
        self.assertGreater(sheet.row_dimensions[2].height,120)
        self.assertEqual(sheet['C2'].value,entries[0]['note'])
        self.assertEqual(sheet['D2'].value,'one alias\nanother alias')
        self.assertTrue(sheet.column_dimensions['G'].hidden);self.assertTrue(sheet.column_dimensions['H'].hidden)
        book.close();self.assertEqual(exchange.load_import(target),entries)
    def test_existing_styles_follow_the_term_when_rows_are_deleted_and_added(self):
        target=self.root/'team.xlsx';columns=writer.HEADER_FIELDS[:6]
        writer.save_table(str(target),[['甲','alpha','first','','',''],['乙','beta','second','','','']],None,columns)
        book=openpyxl.load_workbook(target);sheet=book.active;sheet.title='团队术语';sheet.freeze_panes='B2'
        sheet.column_dimensions['C'].width=18;sheet.row_dimensions[3].height=44
        sheet['A3'].font=Font(name='Arial',size=13,bold=True,color='FF223344')
        sheet['A3'].fill=PatternFill('solid',fgColor='FFFFEEDD')
        original=copy(sheet['A3']._style)
        table=Table(displayName='TeamTerms',ref='A1:F3');table.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True);sheet.add_table(table)
        book.save(target);book.close()
        current=writer.load_table(str(target))
        writer.save_table(str(target),[['乙','beta','long note '*80,'','',''],['丙','=literal','','','','']],current['fingerprint'],columns,current['header'])
        book=openpyxl.load_workbook(target);sheet=book.active
        self.assertEqual(sheet.title,'团队术语');self.assertEqual(sheet.freeze_panes,'B2')
        self.assertEqual(sheet.column_dimensions['C'].width,18)
        self.assertEqual(sheet['A2'].font,Font(name='Arial',size=13,bold=True,color='FF223344'))
        self.assertEqual(sheet['A2'].fill.fgColor.rgb,'FFFFEEDD')
        self.assertTrue(sheet['C2'].alignment.wrap_text);self.assertGreater(sheet.row_dimensions[2].height,120)
        self.assertEqual(sheet['B3'].value,'=literal');self.assertEqual(sheet['B3'].data_type,'s')
        self.assertEqual(sheet.tables['TeamTerms'].ref,'A1:F3');self.assertEqual(sheet.tables['TeamTerms'].tableStyleInfo.name,'TableStyleLight9')
        book.close()
    def test_import_into_shared_workbook_keeps_styles_and_appends_hidden_metadata(self):
        target=self.root/'team.xlsx';writer.save_table(str(target),[['甲','alpha','','','','']],None,writer.HEADER_FIELDS[:6])
        book=openpyxl.load_workbook(target);book.active.column_dimensions['B'].width=38;book.save(target);book.close()
        current=store.select(str(target))
        incoming=self.root/'incoming.json';store.atomic_json(incoming,[{'zh':'乙','en':'beta','note':'new'}])
        result=exchange.import_files([str(incoming)],current['fingerprint'])
        self.assertEqual(result['added'],1)
        book=openpyxl.load_workbook(target);sheet=book.active
        self.assertEqual(sheet.column_dimensions['B'].width,38)
        self.assertTrue(sheet['A3'].font.bold);self.assertTrue(sheet['C3'].alignment.wrap_text)
        self.assertTrue(sheet.column_dimensions['G'].hidden);self.assertTrue(sheet.column_dimensions['H'].hidden)
        self.assertEqual(sheet.max_row,3);book.close()
    def test_new_shared_workbook_uses_the_same_standard_as_export(self):
        store.atomic_json(self.active,store.validate([{'zh':'甲','en':'alpha'}]))
        target=self.root/'new.xlsx';store.select(str(target),create=True)
        book=openpyxl.load_workbook(target);sheet=book.active
        self.assertEqual(sheet['A2'].font.name,'微软雅黑');self.assertTrue(sheet['A2'].font.bold)
        self.assertAlmostEqual(sheet.column_dimensions['C'].width,60.7272727272727)
        self.assertTrue(sheet.column_dimensions['G'].hidden);book.close()
    def test_complex_workbooks_are_rejected_without_overwriting(self):
        target=self.root/'complex.xlsx';book=openpyxl.Workbook();sheet=book.active
        sheet.append(['中文','英文']);sheet.append(['甲','alpha']);sheet.merge_cells('A3:B3');book.save(target);book.close()
        before=target.read_bytes()
        with self.assertRaises(writer.TableTooComplex):writer.save_table(str(target),[['乙','beta']],writer.file_fingerprint(target),['zh','en'])
        self.assertEqual(target.read_bytes(),before)
    def test_wrapping_respects_words_cjk_and_existing_newlines(self):
        font=Font(name='微软雅黑',size=11)
        self.assertEqual(excel_format.wrapped_lines('short',44,font),1)
        self.assertEqual(excel_format.wrapped_lines('one\ntwo',44,font),2)
        self.assertGreater(excel_format.wrapped_lines('中文长句'*20,17,font),2)
        self.assertGreater(excel_format.wrapped_lines('W'*120,24,font),2)
    def test_fragmented_notes_flow_into_paragraphs_without_changing_words_or_aliases(self):
        note='中文定义来自课件；英文为补充建议，待确认。\nGamma\n均匀分布\nGamma校准：把灰度更多地用在人眼敏感的低亮度区域。\n黑色\n白色\nG1\nG2\nG3\nG4'
        displayed=excel_format.presentation_note(note)
        self.assertIn('黑色 白色 G1 G2 G3 G4',displayed)
        self.assertEqual(displayed.split(),note.split())
        self.assertLess(len(displayed.splitlines()),len(note.splitlines()))
        self.assertEqual(excel_format.presentation_note('第一段\n第二段'),'第一段\n第二段')
        self.assertEqual(exchange._combine_text(note,displayed),note)
        self.assertNotEqual(exchange._combine_text('photo electric','photoelectric'),'photo electric')
        original=store.validate([{'zh':'测试','en':'test','note':note,'variants':['alias one','alias two']}])
        store.atomic_json(self.active,original);target=self.root/'fragments.xlsx';exchange.export_file(str(target),'')
        imported=exchange.load_import(target);self.assertEqual(imported[0]['variants'],original[0]['variants'])
        self.assertEqual(imported[0]['note'],displayed)
        result=exchange.import_files([str(target)],store.read()['fingerprint'])
        self.assertEqual(result['updated'],0);self.assertEqual(store.read()['entries'],original)


if __name__=='__main__':unittest.main()
