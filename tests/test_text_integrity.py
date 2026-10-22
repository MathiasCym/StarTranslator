"""Saved English, opaque links and mixed Word objects must agree end to end."""
import json,os,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
sys.path.insert(0,str(Path(__file__).resolve().parent))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from backend.extract import extract
from backend.glossary.store import Glossary,fingerprint
from backend.translation import translate_file
from backend.translation_residual import repair_residual
from backend.translation_audit import saved_residuals
from backend.text_integrity import contains_chinese,validate_identifiers
from backend.writing_policy import validate_translation
from test_word_coverage import Model,field,text_run,signature

class FixedModel:
    def __init__(self,values):self.values=iter(values);self.prompts=[]
    def chat_json(self,prompt,**kwargs):
        self.prompts.append(prompt)
        return {'translation':next(self.values)}

class TextIntegrityTests(unittest.TestCase):
    def test_equivalent_written_counts_dimensions_and_acronym_plurals(self):
        for source,target in [('配置4套10.5K系统。','Four 10.5K systems are installed.'),
                ('三维环境和GPU支持。','3D environments and GPUs are supported.'),
                ('亮度500nits。','Brightness: five hundred nits.'),
                ('部署24块屏幕。','Twenty-four screens are installed.'),
                ('预算为4 million。','The budget is four million.'),
                ('两套2K系统。','2 sets of 2K systems.'),('四百台设备。','400 devices.')]:
            with self.subTest(source=source):validate_translation(source,target)
        for source,target in [('配置4套系统。','Three systems are installed.'),
                ('配置4套系统。','Four systems and four backups.'),
                ('配置4套系统。','Twenty-four systems are installed.'),
                ('三维环境。','4D environments.'),('GPU支持。','CPUs are supported.'),
                ('VX1000支持。','VX1000s are supported.'),('亮度500nits。','Brightness: five hundred Hz.'),
                ('两套系统。','3 systems.'),('两套系统。','2 systems and 2 backups.')]:
            with self.subTest(source=source,target=target),self.assertRaises(ValueError):validate_translation(source,target)

    def test_generated_toc_labels_translate_without_touching_fields_or_links(self):
        doc=Document();doc.add_heading('目录',0)
        start=doc.add_paragraph();field(start,' TOC \\o "1-3" \\h \\z \\u ',label='')
        # Move the TOC end control below all rows, as in Word-generated TOCs.
        ending=start._p[-1];start._p.remove(ending)
        p=doc.add_paragraph();link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'_TocExample')
        run=text_run(link,'正文前文');run.append(OxmlElement('w:tab'));p._p.append(link);field(p,' PAGEREF _TocExample \\h ',label='4')
        end=doc.add_paragraph();end._p.append(ending)
        heading=doc.add_heading('正文前文',1)
        bookmark=OxmlElement('w:bookmarkStart');bookmark.set(qn('w:id'),'77');bookmark.set(qn('w:name'),'_TocExample');heading._p.append(bookmark)
        bookmark=OxmlElement('w:bookmarkEnd');bookmark.set(qn('w:id'),'77');heading._p.append(bookmark)
        doc.save(self.source);before=signature(self.source)
        class TocModel(Model):
            def chat_json(self,prompt,**kwargs):
                return {'translation':'Contents'} if prompt.split('【待译原文】\n')[-1].split('\n【')[0]=='目录' else super().chat_json(prompt,**kwargs)
        result=self.translate(doc,TocModel());self.assertTrue(result['translation_complete'])
        self.assertEqual(result['untranslated'],0);self.assertEqual(signature(result['path']),before)
        out=Document(result['path']);self.assertIn('Body before',out.paragraphs[2].text);self.assertTrue(out.paragraphs[2].text.endswith('\t4'))
        self.assertEqual(out.paragraphs[2]._p.find(qn('w:hyperlink')).get(qn('w:anchor')),'_TocExample')
        self.assertEqual(len(out.paragraphs[2]._p.findall('.//'+qn('w:tab'))),1)
    def test_dimension_shorthand_is_also_valid_during_residual_repair(self):
        class Model:
            def chat_json(self,*args,**kwargs):return {'translations':[{'source':'三维','translation':'3D'}]}
        value=repair_residual(Model(),'三维环境',{'translation':'三维 environments'})
        self.assertEqual(value['translation'],'3D environments')
        validate_translation('三维环境',value['translation'])
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.source=self.root/'original.docx'
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')})
        env.start();self.addCleanup(env.stop)
    def translate(self,doc,model=None):
        doc.save(self.source);before=fingerprint(self.source)
        result=translate_file(self.source,Glossary([]),None,client=model or Model(),workers=1)
        self.assertEqual(fingerprint(self.source),before)
        return result
    def test_rare_han_and_compatibility_characters_cannot_bypass_translation(self):
        for word in ('𠀀','﨑','〇'):
            with self.subTest(word=word):
                self.assertTrue(contains_chinese(word));doc=Document();doc.add_paragraph(word)
                result=self.translate(doc,FixedModel(['Name']))
                self.assertTrue(result['translation_complete']);self.assertEqual(result['untranslated'],0)
                self.assertEqual(Document(result['path']).paragraphs[0].text,'Name')
    def test_simple_hyperlink_field_preserves_command_and_bookmark_target(self):
        doc=Document();p=doc.add_paragraph('正文前文');simple=OxmlElement('w:fldSimple')
        command=' HYPERLINK "https://example.test/中文原始地址?q=1" \\o "中文提示" '
        simple.set(qn('w:instr'),command);text_run(simple,'链接标签');p._p.append(simple);p.add_run('正文后文')
        result=self.translate(doc);out=Document(result['path'])
        self.assertTrue(result['translation_complete'])
        self.assertEqual(out.paragraphs[0]._p.find(qn('w:fldSimple')).get(qn('w:instr')),command)
        self.assertFalse(any(contains_chinese(s.text) for s in extract(result['path']) if s.source!='image'))
    def test_controls_and_visible_text_in_same_run_do_not_block_translation(self):
        doc=Document();p=doc.add_paragraph();field(p);combined=OxmlElement('w:r')
        for run in list(p._p):
            for child in list(run):combined.append(child)
            p._p.remove(run)
        p._p.append(combined);doc.save(self.source);before=signature(self.source)
        result=self.translate(doc)
        self.assertTrue(result['translation_complete']);self.assertEqual(signature(result['path']),before)
        self.assertIn('Reference label',Document(result['path']).paragraphs[0].text)
    def test_mixed_picture_run_keeps_original_object_and_translates_text_children(self):
        doc=Document();p=doc.add_paragraph();run=OxmlElement('w:r');text_run(p._p,'正文前文')
        a=OxmlElement('w:t');a.text='文本框文字';run.append(a);pict=OxmlElement('w:pict')
        etree.SubElement(pict,'{urn:schemas-microsoft-com:vml}shape',id='untouched',style='width:100pt;height:50pt')
        run.append(pict);p._p.append(run);doc.save(self.source);before=signature(self.source)
        result=self.translate(doc)
        self.assertTrue(result['translation_complete']);self.assertEqual(signature(result['path']),before)
        out=Document(result['path']);self.assertEqual(len(list(out.element.iter(qn('w:pict')))),1)
        self.assertIn('Text box wording',out.paragraphs[0].text)
    def test_valid_paragraphs_inside_alternate_content_are_translated(self):
        doc=Document();wrapper=etree.Element('{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent')
        choice=etree.SubElement(wrapper,'{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice',Requires='w')
        paragraph=OxmlElement('w:p');text_run(paragraph,'保留文字');choice.append(paragraph)
        doc.element.body.insert(0,wrapper);result=self.translate(doc)
        self.assertTrue(result['translation_complete']);self.assertEqual(saved_residuals(result['path']),[])
    def test_word_drawingml_text_is_translated_without_geometry_changes(self):
        doc=Document();p=doc.add_paragraph('正文前文');drawing=OxmlElement('w:drawing')
        a='http://schemas.openxmlformats.org/drawingml/2006/main'
        body=etree.SubElement(drawing,'{'+a+'}txBody');para=etree.SubElement(body,'{'+a+'}p')
        run=etree.SubElement(para,'{'+a+'}r');node=etree.SubElement(run,'{'+a+'}t');node.text='文本框文字'
        host=OxmlElement('w:r');host.append(drawing);p._p.append(host);doc.save(self.source);before=signature(self.source)
        result=self.translate(doc)
        self.assertTrue(result['translation_complete']);self.assertEqual(signature(result['path']),before)
        self.assertEqual(next(Document(result['path']).element.iter('{'+a+'}t')).text,'Text box wording')
    def test_prose_is_translated_while_visible_chinese_url_is_kept_exactly(self):
        address='https://example.test/中文地址?q=12';doc=Document();doc.add_paragraph('正文前文 '+address)
        result=self.translate(doc,FixedModel(['Body before '+address]))
        self.assertTrue(result['translation_complete']);self.assertEqual(saved_residuals(result['path']),[])
        self.assertEqual(Document(result['path']).paragraphs[0].text,'Body before '+address)
    def test_changed_url_is_rejected_and_repaired_before_export(self):
        address='https://example.test/中文地址?q=12';doc=Document();doc.add_paragraph('正文前文 '+address)
        model=FixedModel(['Body before https://example.test/changed?q=12','Body before '+address])
        result=self.translate(doc,model)
        self.assertTrue(result['translation_complete']);self.assertEqual(len(model.prompts),2)
        self.assertIn(address,Document(result['path']).paragraphs[0].text)
    def test_residual_repair_does_not_replace_same_phrase_inside_url_or_split_runs(self):
        address='https://example.test/画布?q=12';source='Existing '+address+' 画布'
        class Repair:
            def chat_json(self,prompt,**kwargs):
                phrases=json.loads(prompt.split('【只需补译的中文】\n')[1].split('\n')[0])
                return {'translations':[{'source':p,'translation':'canvas'} for p in phrases]}
        result=repair_residual(Repair(),source,{'translation':source,'runs':[
            {'id':0,'text':'Existing https://example.test/画'}, {'id':1,'text':'布?q=12 画布'}]},allow_source=True)
        self.assertEqual(result['translation'],'Existing '+address+' canvas')
        self.assertEqual(''.join(r['text'] for r in result['runs']),result['translation'])
        self.assertFalse(contains_chinese(result['translation']))
    def test_relationships_and_untouched_package_parts_remain_byte_identical(self):
        doc=Document();p=doc.add_paragraph('正文前文');field(p);doc.save(self.source)
        with ZipFile(self.source) as z:before={n:z.read(n) for n in z.namelist() if n!='word/document.xml'}
        result=self.translate(doc)
        with ZipFile(result['path']) as z:self.assertEqual({n:z.read(n) for n in z.namelist() if n!='word/document.xml'},before)
    def test_disk_residual_cannot_be_hidden_by_successful_in_memory_translation(self):
        from backend.replacer.pptx_package import save_office
        doc=Document();doc.add_paragraph('正文前文')
        def faulty_writer(source,document,destination,changed):
            save_office(source,document,destination,changed)
            output=Document(destination);output.add_paragraph('意外残留中文');output.save(destination)
        with patch('backend.translation.save_office',side_effect=faulty_writer):result=self.translate(doc)
        self.assertFalse(result['translation_complete']);self.assertEqual(result['status'],'needs_review')
        self.assertEqual(result['untranslated'],1);self.assertTrue(Path(result['path']).is_file())
    def test_xml_inventory_finds_math_text_not_exposed_by_paragraph_reader(self):
        doc=Document();doc.add_paragraph('正文前文');math=OxmlElement('m:oMath');run=OxmlElement('m:r')
        node=OxmlElement('m:t');node.text='未覆盖文字';run.append(node);math.append(run);doc.paragraphs[0]._p.append(math)
        result=self.translate(doc)
        self.assertFalse(result['translation_complete'])
        self.assertTrue(any(item['source']=='package_text' for item in result['remaining_chinese']))
    def test_word_writer_rejects_lost_hyperlink_fields(self):
        from backend.extract.office import open_office
        from backend.replacer.pptx_package import save_office
        doc=Document();p=doc.add_paragraph('正文前文');field(p);doc.save(self.source)
        modified,paragraphs,_=open_office(self.source)
        node=next(modified.element.iter(qn('w:instrText')));node.getparent().remove(node)
        with self.assertRaisesRegex(ValueError,'Word 链接'):
            save_office(self.source,modified,self.root/'bad.docx',[loc for loc,*_ in paragraphs])
    def test_review_saved_file_reports_chinese_inserted_by_a_manual_edit(self):
        from backend import workflow
        from backend.models import Finding
        doc=Document();doc.add_paragraph('Original text');doc.save(self.source)
        seg=extract(self.source)[0]
        hit=Finding(segment_id=seg.id,file=seg.file,file_path=str(self.source),slide=seg.slide,para=seg.para,source=seg.source,found_text='Original',
                    suggested='Revised',kind='grammar',method='llm',editable=True)
        report={'segments':[seg],'hashes':{str(self.source):fingerprint(self.source)},'findings':[hit]}
        result=workflow._apply_report({'path':str(self.source),'edits':[{'index':0,'text':'修订'}]},report)
        self.assertEqual(result['untranslated'],1);self.assertTrue(result['content_warnings'])
        self.assertEqual(Document(result['output']).paragraphs[0].text,'修订 text')
    def test_review_never_changes_url_while_other_suggestions_can_be_applied(self):
        from backend import workflow
        from backend.models import Finding
        address='https://example.test/中文?q=12'
        doc=Document();doc.add_paragraph('Original text '+address);doc.save(self.source)
        seg=extract(self.source)[0]
        hits=[Finding(segment_id=seg.id,file=seg.file,file_path=str(self.source),slide=seg.slide,para=seg.para,source=seg.source,found_text=old,
            suggested=new,kind='grammar',method='llm',editable=True)
            for old,new in [('Original','Revised'),('example','changed')]]
        report={'segments':[seg],'hashes':{str(self.source):fingerprint(self.source)},'findings':hits}
        result=workflow._apply_report({'path':str(self.source),'selected':{seg.id:[['Original','Revised'],['example','changed']]}},report)
        self.assertEqual(result['applied_count'],1);self.assertEqual(len(result['failed']),1)
        self.assertEqual(Document(result['output']).paragraphs[0].text,'Revised text '+address)
        self.assertEqual(result['untranslated'],0)
    def test_identifier_validation_keeps_paths_emails_and_case(self):
        for value in ('https://example.test/中文?X=1','C:\\资料\\方案.docx','user@example.test'):
            validate_translation('说明 '+value,'Description '+value)
            with self.assertRaisesRegex(ValueError,'网址、邮箱或文件路径'):
                validate_identifiers(value,value.replace('example','changed').replace('资料','data').replace('方案','plan'))

if __name__=='__main__':unittest.main()
