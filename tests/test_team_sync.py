import base64,copy,json,os,tempfile,unittest,sys
from pathlib import Path
from unittest.mock import patch
from urllib.error import URLError
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend import team_sync as sync
from backend.glossary import store
from backend.llm import secret

def term(zh,en,**kw):return store.validate([{'zh':zh,'en':en,**kw}])[0]
class Remote:
    token='synthetic-maintainer'
    def __init__(self,entries):self.entries=copy.deepcopy(entries);self.sha='a'*40;self.calls=[]
    def read(self):return {'entries':copy.deepcopy(self.entries),'sha':self.sha,'missing':not self.sha}
    def commit(self,entries,sha,message):
        if sha!=self.sha:raise sync.RemoteChanged('concurrent edit')
        self.calls.append((entries,message));self.entries=copy.deepcopy(entries);self.sha='b'*40;return self.sha

class TeamSyncTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.path=self.root/'lexicon.json'
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root)});self.env.start()
        self.active=patch('backend.glossary.store.active_path',return_value=self.path);self.active.start()
        self.base=[term('色深','color depth'),term('点间距','pixel pitch',variants=['pixel spacing'])]
        store.atomic_json(self.path,self.base);store.atomic_json(sync.folder()/'base.json',{'path':str(self.path),'entries':self.base,'sha':'a'*40})
        self.remote=Remote(self.base);self.client=patch('backend.team_sync.GitHub',return_value=self.remote);self.client.start()
    def tearDown(self):self.client.stop();self.active.stop();self.env.stop();self.temp.cleanup()
    def preview(self,action='pull'):return sync.preview({'action':action,'fingerprint':store.read()['fingerprint']})
    def test_preview_does_not_write_and_independent_additions_are_merged(self):
        local=self.base+[term('本机新词','local term')];store.atomic_json(self.path,local);before=self.path.read_bytes()
        self.remote.entries.append(term('团队新词','team term'));p=self.preview()
        self.assertEqual(before,self.path.read_bytes());sync.apply({'id':p['id']})
        self.assertEqual(len(store.read()['entries']),4);self.assertEqual(self.remote.calls,[])
    def test_scalar_conflict_requires_choice_and_keeps_original_until_confirmed(self):
        local=copy.deepcopy(self.base);local[0]['en']='local depth';store.atomic_json(self.path,local);self.remote.entries[0]['en']='team depth'
        p=self.preview();self.assertEqual(len(p['conflicts']),1);before=self.path.read_bytes()
        with self.assertRaises(ValueError):sync.apply({'id':p['id']})
        self.assertEqual(before,self.path.read_bytes());sync.apply({'id':p['id'],'choices':{'0':'remote'}})
        self.assertEqual(next(e for e in store.read()['entries'] if e['zh']=='色深')['en'],'team depth')
    def test_alias_changes_merge_deletions_and_additions(self):
        local=copy.deepcopy(self.base);local[1]['variants']=['local alias'];store.atomic_json(self.path,local)
        self.remote.entries[1]['variants'].append('remote alias');p=self.preview();self.assertFalse(p['conflicts']);sync.apply({'id':p['id']})
        e=next(e for e in store.read()['entries'] if e['zh']=='点间距');self.assertEqual(set(e['variants']),{'local alias','remote alias'})
    def test_deleted_term_and_parallel_edit_are_not_silently_lost(self):
        store.atomic_json(self.path,self.base[1:]);self.remote.entries[0]['note']='new definition';p=self.preview()
        self.assertEqual(p['conflicts'][0]['field'],'entry');sync.apply({'id':p['id'],'choices':{'0':'remote'}})
        self.assertEqual(len(store.read()['entries']),2)
    def test_remote_changed_since_preview_refuses_write(self):
        p=self.preview('push');self.remote.sha='c'*40;before=self.path.read_bytes()
        with self.assertRaises(sync.RemoteChanged):sync.apply({'id':p['id'],'message':'update'})
        self.assertEqual(before,self.path.read_bytes());self.assertFalse(self.remote.calls)
    def test_local_changed_since_preview_refuses_remote_commit(self):
        p=self.preview('push');store.atomic_json(self.path,self.base+[term('新词','new')])
        with self.assertRaises(store.Conflict):sync.apply({'id':p['id'],'message':'update'})
        self.assertFalse(self.remote.calls)
    def test_put_conflict_does_not_retry_or_overwrite(self):
        store.atomic_json(self.path,self.base+[term('新词','new')]);p=self.preview('push');before=self.path.read_bytes()
        with patch.object(self.remote,'commit',side_effect=sync.RemoteChanged('parallel commit')):
            with self.assertRaises(sync.RemoteChanged):sync.apply({'id':p['id'],'message':'add term'})
        self.assertEqual(before,self.path.read_bytes())
    def test_push_is_one_normal_commit_and_persists_base(self):
        store.atomic_json(self.path,self.base+[term('新词','new')]);p=self.preview('push');self.assertEqual(p['changes'][0]['kind'],'added')
        result=sync.apply({'id':p['id'],'message':'补充新词'});self.assertTrue(result['committed']);self.assertEqual(len(self.remote.calls),1)
        self.assertEqual(sync.settings()['revision'],'b'*40);self.assertFalse(self.preview()['conflicts'])
    def test_initial_publication_keeps_local_library_and_reports_additions(self):
        self.remote.entries=[];self.remote.sha='';(sync.folder()/'base.json').unlink();p=self.preview('push')
        self.assertTrue(p['remote_missing']);self.assertEqual(len(p['changes']),2);sync.apply({'id':p['id'],'message':'初始化词库'})
        self.assertEqual(len(self.remote.entries),2)
    def test_credential_is_encrypted_never_returned_and_plain_fallback_refused(self):
        key='github_pat_synthetic_only_123456789';value=sync.configure({'token':key});self.assertNotIn(key,json.dumps(value));self.assertTrue(secret.is_encrypted(sync.read_json('settings.json',{})['token']))
        with patch('backend.team_sync.secret.encrypt',return_value='plain:'+key):
            with self.assertRaises(ValueError):sync.configure({'token':key})
        self.assertEqual(secret.decrypt(sync.read_json('settings.json',{})['token']),key)
        sync.configure({'clear_token':True});self.assertFalse(sync.settings()['has_token'])
    def test_network_failure_is_readable_and_keeps_local_bytes(self):
        self.client.stop();before=self.path.read_bytes()
        with patch('backend.team_sync.request.urlopen',side_effect=URLError('offline')):
            with self.assertRaisesRegex(ValueError,'本机词库保持不变'):self.preview()
        self.assertEqual(before,self.path.read_bytes());self.client.start()
    def test_excel_sync_preserves_custom_columns_and_column_width(self):
        from openpyxl import Workbook,load_workbook
        p=self.root/'team.xlsx';w=Workbook();s=w.active;s.append(['中文术语','标准英文','部门']);s.append(['色深','color depth','研发']);s.column_dimensions['C'].width=42;w.save(p)
        current=store.read(p);save=sync.save_local(str(p),[term('色深','colour depth')],current['fingerprint'])
        self.assertEqual(save['entries'][0]['en'],'colour depth');w=load_workbook(p)
        self.assertEqual(w.active['C2'].value,'研发');self.assertEqual(w.active.column_dimensions['C'].width,42)

if __name__=='__main__':unittest.main()
