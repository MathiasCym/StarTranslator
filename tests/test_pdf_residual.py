"""PDF physical wraps, untranslated text recovery and confirmation-only review."""
import json,os,re,sys,tempfile,threading,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from test_file_formats import pdf_fixture
from backend.translation import translate_file,Cancelled
from backend.translation_residual import repair_residual,CHINESE,PHRASES
from backend.glossary.loader import Glossary
from backend.glossary import store
from backend.llm.client import LLMUnavailable
from backend.matcher import check_segments
from backend.extract import extract
from backend import workflow
from zipfile import ZipFile
from lxml import etree

class Model:
    def __init__(self,text='The main screen is 87 feet across with a 12Kx10K canvas.',missing=None,failures=0):
        self.text=text;self.missing=missing or {};self.failures=failures;self.prompts=[];self.repairs=0
    def chat_json(self,prompt,**kwargs):
        self.prompts.append((prompt,kwargs))
        if '【只需补译的中文】' in prompt:
            self.repairs+=1
            if self.repairs<=self.failures:return {'translations':[]}
            phrases=json.loads(prompt.split('【只需补译的中文】\n')[1].split('\n')[0])
            return {'translations':[{'source':p,'translation':self.missing.get(p,'term')} for p in phrases]}
        if '【待译原文】' in prompt:return {'translation':self.text}
        return []

class PdfResidualTests(unittest.TestCase):
    def test_reflow_identifier_policy_does_not_leak_into_word_or_powerpoint(self):
        from backend.translation_policy import validate_for_file
        source,target='LEDDome和300TBDellPowerScale','LED Dome and Dell PowerScale with 300 TB'
        validate_for_file('source.PDF',source,target)
        for filename in ['source.docx','source.pptx','source.xlsx','source.txt',None]:
            with self.subTest(filename=filename),self.assertRaises(ValueError):
                validate_for_file(filename,source,target)
        for filename in ['source.pdf','source.docx','source.pptx']:
            with self.subTest(filename=filename),self.assertRaises(ValueError):
                validate_for_file(filename,'型号VX1000','Model VX2000')

    def test_pdf_reflow_does_not_confuse_joined_words_with_changed_identifiers(self):
        from backend.writing_policy import validate_translation
        pairs=[('场馆地址1252DistrictDrive。','Venue address: 1252 District Drive.'),
            ('LEDDome采用NVIDIARTXPROGPU。','The LED Dome uses NVIDIA RTX PRO GPUs.'),
            ('使用200Mbit和300TBDellPowerScale。','Use 200 Mbit and 300 TB Dell PowerScale.'),
            ('使用300TBDellPowerScale。','Use Dell PowerScale with 300 TB of storage.'),
            ('COSM.CosmiscomingtoCleveland。','COSM. "Cosm is coming to Cleveland."'),
            ('COSM.CosmiscomingtoCleveland。','COSM. Cosm is coming to Cleveland.')]
        for source,target in pairs:
            with self.subTest(source=source):validate_translation(source,target,allow_reflow=True)
        for source,target in [('VX1000。','VX 2000.'),('NVIDIARTXPROGPU。','NVIDIA RTX PRO CPUs.'),
                ('200Mbit。','200 Mbyte.'),('VX1000。','VX1000s.'),
                ('300TBDellPowerScale。','Dell PowerScale with 300 GB.'),('300TBDellPowerScale。','300 TB of storage.')]:
            with self.subTest(source=source),self.assertRaises(ValueError):validate_translation(source,target,allow_reflow=True)
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.g=Glossary([])
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'data')});self.env.start()
    def tearDown(self):self.env.stop();self.tmp.cleanup()
    def pdf(self,text):
        p=self.root/'source.pdf';pdf_fixture(p,text);return p
    def test_pdf_annotation_targets_survive_reflow_and_final_review_save(self):
        from pypdf import PdfReader,PdfWriter
        from pypdf.annotations import Link
        source=self.pdf('主屏直径为87英尺，整体画布为12Kx10K。')
        reader=PdfReader(source);writer=PdfWriter();writer.append(reader)
        uri='https://example.test/原始目标?q=unchanged'
        writer.add_annotation(0,Link(rect=(0,0,100,20),url=uri))
        writer.add_annotation(0,Link(rect=(0,30,100,50),target_page_index=0))
        writer.write(source);before=store.fingerprint(source)
        result=translate_file(source,self.g,None,client=Model())
        self.assertTrue(result['translation_complete']);self.assertEqual(store.fingerprint(source),before)
        output=Document(result['path'])
        self.assertIn(uri,[r.target_ref for r in output.part.rels.values() if r.is_external])
        w='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        bookmarks={n.get(w+'name') for n in output.element.iter(w+'bookmarkStart')}
        self.assertEqual(output.paragraphs[0]._p[0].tag,w+'pPr')
        targets={n.get(w+'anchor') for n in output.element.iter(w+'hyperlink') if n.get(w+'anchor')}
        self.assertEqual(targets,{'_StarPDFPage_1'});self.assertLessEqual(targets,bookmarks)
        job={'id':'pdf-link-save','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            workflow.run(job,dict(kind='check',paths=[result['path']],use_llm=False))
        saved=workflow.apply_report(dict(report_id=job['id'],path=result['path'],save_unchanged=True))
        self.assertIn(uri,[r.target_ref for r in Document(saved['output']).part.rels.values() if r.is_external])
        self.assertEqual(store.fingerprint(source),before)
    def test_physical_pdf_line_wraps_are_translated_together(self):
        source=self.pdf('主屏直径为87英尺，\n整体画布为12Kx10K。');before=store.fingerprint(source);model=Model()
        result=translate_file(source,self.g,None,client=model)
        self.assertEqual(len(model.prompts),1);self.assertIn('87英尺，\n整体画布',model.prompts[0][0])
        self.assertEqual(result['untranslated'],0);self.assertEqual(Document(result['path']).paragraphs[0].text,model.text)
        self.assertEqual(before,store.fingerprint(source))
    def test_pdf_last_pass_repairs_partial_english_without_discarding_the_file(self):
        source=self.pdf('主屏直径为87英尺，整体画布为12Kx10K。')
        model=Model('The main screen is 87 feet across with a 12Kx10K 画布.',{'画布':'canvas'},2)
        result=translate_file(source,self.g,None,client=model)
        self.assertEqual(model.repairs,3);self.assertEqual(result['untranslated'],0)
        self.assertEqual(result['status'],'complete');self.assertEqual(Document(result['path']).paragraphs[0].text,'The main screen is 87 feet across with a 12Kx10K canvas.')
    def test_many_residual_phrases_are_batched_and_keep_existing_english(self):
        text='Existing English '+ ' / '.join(chr(0x4e10+i) for i in range(40))+' 87 feet'
        model=Model();result=repair_residual(model,text,{'translation':text},self.g,allow_source=True)
        self.assertEqual(model.repairs,4);self.assertFalse(CHINESE.search(result['translation']))
        self.assertTrue(result['translation'].startswith('Existing English '));self.assertTrue(result['translation'].endswith(' 87 feet'))
    def test_late_supplement_failure_retains_the_successful_batches(self):
        text='Existing English '+ ' / '.join(chr(0x4e10+i) for i in range(15))
        class LaterFailure(Model):
            def chat_json(self,prompt,**kwargs):
                if self.repairs:raise LLMUnavailable('temporary timeout')
                return super().chat_json(prompt,**kwargs)
        with self.assertRaises(LLMUnavailable) as caught:repair_residual(LaterFailure(),text,{'translation':text},allow_source=True)
        partial=caught.exception.partial_result
        self.assertEqual(len(partial['_residual_replacements']),12);self.assertEqual(len(PHRASES.findall(partial['translation'])),3)
    def test_review_lists_every_residual_and_only_applies_after_confirmation(self):
        source=self.root/'mixed.docx';doc=Document();doc.add_paragraph('A 12Kx10K 全景显示.');doc.save(source);before=store.fingerprint(source)
        model=Model(missing={'全景显示':'panoramic display'});job={'id':'residual-check','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=self.g):workflow.run(job,{'kind':'check','paths':[str(source)],'use_llm':True,'_client':model})
        hits=job['result']['findings'];self.assertEqual(len(hits),1);self.assertEqual(hits[0]['suggested'],'panoramic display')
        self.assertEqual(store.fingerprint(source),before)
        saved=workflow.apply_report({'report_id':job['id'],'path':str(source),'selected':{hits[0]['segment_id']:[['全景显示','panoramic display']]}})
        self.assertEqual(Document(saved['output']).paragraphs[0].text,'A 12Kx10K panoramic display.')
        self.assertEqual(store.fingerprint(source),before)
    def test_failed_supplement_remains_visible_and_is_not_marked_complete(self):
        source=self.root/'mixed.docx';doc=Document();doc.add_paragraph('A 全景显示.');doc.save(source)
        class Offline(Model):
            def chat_json(self,prompt,**kwargs):
                if '【只需补译的中文】' in prompt:raise LLMUnavailable('model unavailable')
                return []
        hits,stats=check_segments(extract(source),self.g,True,client=Offline(),expect_english=True)
        self.assertEqual(stats['model_status'],'partial');self.assertTrue(stats['warnings']);self.assertEqual(stats['completed_files'],[])
        self.assertEqual(hits[0].found_text,'全景显示');self.assertTrue(hits[0].editable);self.assertEqual(hits[0].suggested,'')
    def test_rules_only_detection_does_not_call_a_model(self):
        source=self.root/'mixed.docx';doc=Document();doc.add_paragraph('A 全景显示.');doc.save(source)
        model=Model();hits,stats=check_segments(extract(source),self.g,False,client=model,expect_english=True)
        self.assertEqual(model.prompts,[]);self.assertEqual(hits[0].found_text,'全景显示');self.assertTrue(hits[0].editable);self.assertEqual(hits[0].suggested,'')
    def test_compression_mode_does_not_apply_to_pdf(self):
        source=self.pdf('主屏直径为87英尺，整体画布为12Kx10K。');model=Model()
        result=translate_file(source,self.g,None,client=model,translation_mode='compressed')
        self.assertEqual(result['translation_mode'],'faithful');self.assertIn('【忠实翻译】',model.prompts[0][1]['system'])
    def test_supplement_cannot_duplicate_neighbor_numbers(self):
        text='预计 300+';model=Model(missing={'预计':'Estimated 300+'})
        with self.assertRaisesRegex(ValueError,'数字'):repair_residual(model,text,{'translation':text},self.g,allow_source=True)
    def test_cancelled_supplement_makes_no_model_call(self):
        model=Model()
        def stop():raise Cancelled('cancelled')
        with self.assertRaises(Cancelled):repair_residual(model,'预计',{'translation':'Estimated 预计'},check_cancel=stop)
        self.assertEqual(model.prompts,[])

if __name__=='__main__':unittest.main()
