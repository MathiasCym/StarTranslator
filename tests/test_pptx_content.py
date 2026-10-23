"""Adjacent advertising labels must not leak into another translated line."""
import os,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from backend.pptx_content import validate_labels,cross_line_repetitions
from backend.translation import translate_file
from backend.glossary.loader import Glossary
from backend.llm.client import LLMClient


class PptxContentTests(unittest.TestCase):
    def test_only_new_occurrences_are_rejected_with_word_boundaries(self):
        for label,source,bad in [('CINITY LED','助力《阿凡达3》','Avatar 3 on CINITY LED'),
                                 ('HeyLED','进驻县城','HeyLED Enters County-Level Cities'),
                                 ('VLED','技术落地','VLED Technology Deployment')]:
            original=label+'\v'+source
            with self.assertRaisesRegex(ValueError,'独立标题'):validate_labels(original,source,bad)
        validate_labels('CINITY LED\vCINITY LED技术','CINITY LED技术','CINITY LED technology')
        validate_labels('CINITY LED\vCINITY LED技术','CINITY LED\vCINITY LED技术','CINITY LED\vCINITY LED technology')
        validate_labels('LED\v显示屏','显示屏','LED display')
        validate_labels('VLED\v说明','说明','AVLEDX')
        validate_labels('Meeting Minutes\v会议纪要','会议纪要','Meeting Minutes')
        validate_labels('TITLE\v标题','标题','TITLE')
        self.assertEqual(cross_line_repetitions('2024年12月23日\v中国北京','23 December 2024, Beijing China\vBeijing China'),['beijing china'])
        self.assertFalse(cross_line_repetitions('Beijing China\vBeijing China','Beijing China\vBeijing China'))
        self.assertFalse(cross_line_repetitions('技术说明\v技术说明','Technical description\vTechnical description'))
    def test_suspected_cross_line_prose_is_repaired_semantically_and_cached(self):
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(Path(temp)/'profile')}):
            source=Path(temp)/'date.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
            shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2));shape.text='2024年12月23日\v中国北京';deck.save(source)
            class Model(LLMClient):
                def __init__(self):
                    self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='date-fixture';self.api_format='openai';self.calls=[]
                def chat_json(self,prompt,**kwargs):
                    self.calls.append(prompt)
                    return {'translation':('23 December 2024\nBeijing, China' if '【疑似跨行重复】' in prompt else
                        '23 December 2024, Beijing China' if '【待译原文】\n2024年12月23日' in prompt else 'Beijing China')}
            model=Model();result=translate_file(source,Glossary([]),None,client=model,workers=1,defer_visual=True)
            text=Presentation(result['path']).slides[0].shapes[0].text
            self.assertTrue(result['translation_complete']);self.assertEqual(text.count('Beijing'),1)
            self.assertTrue(any('【疑似跨行重复】' in p for p in model.calls))
            model.calls.clear();resumed=translate_file(source,Glossary([]),None,client=model,workers=1,defer_visual=True)
            self.assertTrue(resumed['translation_complete']);self.assertFalse(model.calls)
    def test_translation_retries_locally_preserves_styles_and_resumes_without_duplicate(self):
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(Path(temp)/'profile')}):
            source=Path(temp)/'source.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
            shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2))
            shape.text='CINITY LED\v助力《阿凡达3》'
            first=shape.text_frame.paragraphs[0].runs[0];first.font.bold=True;first.font.size=Pt(24);first.font.color.rgb=RGBColor(30,100,200)
            shape.text_frame.paragraphs[0].runs[-1].font.size=Pt(20);deck.save(source)
            class Model(LLMClient):
                def __init__(self,answers):
                    self.answers=iter(answers);self.prompts=[]
                    self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='ppt-label-fixture';self.api_format='openai'
                def chat_json(self,prompt,**kwargs):self.prompts.append(prompt);return {'translation':next(self.answers)}
            model=Model(['Avatar 3 on CINITY LED','Supports Avatar 3'])
            result=translate_file(source,Glossary([]),None,client=model,workers=1,defer_visual=True)
            output=Presentation(result['path']);paragraph=output.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertTrue(result['translation_complete']);self.assertEqual(len(model.prompts),2)
            self.assertIn('独立标题',model.prompts[-1]);self.assertEqual(paragraph.text.count('CINITY LED'),1)
            self.assertIn('Supports Avatar 3',paragraph.text);self.assertTrue(paragraph.runs[0].font.bold)
            self.assertEqual(paragraph.runs[0].font.color.rgb,RGBColor(30,100,200))
            never=Model([]);resumed=translate_file(source,Glossary([]),None,client=never,workers=1,defer_visual=True)
            self.assertTrue(resumed['translation_complete']);self.assertFalse(never.prompts)
            self.assertEqual(Presentation(resumed['path']).slides[0].shapes[0].text,paragraph.text)
    def test_word_manual_break_translation_is_unchanged(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(Path(temp)/'profile')}):
            source=Path(temp)/'word.docx';doc=Document();doc.add_paragraph('VLED\n技术落地');doc.save(source)
            class Model:
                def chat_json(self,prompt,**kwargs):return {'translation':'VLED Technology Deployment'}
            result=translate_file(source,Glossary([]),None,client=Model(),workers=1)
            self.assertTrue(result['translation_complete']);self.assertIn('VLED Technology Deployment',Document(result['path']).paragraphs[0].text)
    def test_proofreading_cannot_reintroduce_the_neighbor_label(self):
        from backend.document_context import DocumentContext
        from backend.models import Segment
        from backend.matcher import check_segments
        seg=Segment('one','source.pptx','CINITY LED\vSupports Avatar 3','slide1/shape0/p0',slide=1)
        context=DocumentContext([seg])
        self.assertTrue(context.duplicates_neighbor_label(seg,'Supports Avatar 3','Avatar 3 on CINITY LED'))
        self.assertFalse(context.duplicates_neighbor_label(seg,'Supports Avatar 3','Supports Avatar 3 worldwide'))
        class Model:
            def chat_json(self,prompt,**kwargs):return [{'found':'Supports Avatar 3','suggested':'Avatar 3 on CINITY LED','kind':'grammar','reason':'complete heading'}]
        findings,stats=check_segments([seg],Glossary([]),client=Model())
        self.assertFalse(findings);self.assertEqual(stats['model_status'],'complete')
        self.assertEqual(stats['discarded_context_hits'],1)
    def test_reordered_slide_part_names_write_to_the_actual_page(self):
        from zipfile import ZipFile
        from lxml import etree
        from backend.extract.office import open_office
        from backend.replacer.pptx_package import save_office
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(Path(temp)/'profile')}):
            folder=Path(temp);source=folder/'source.pptx';deck=Presentation()
            for text in ['品牌说明','操作说明']:
                slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text=text
            deck.save(source)
            with ZipFile(source) as z:parts={n:z.read(n) for n in z.namelist()}
            root=etree.fromstring(parts['ppt/presentation.xml']);ids=root.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst');ids.insert(0,ids[-1])
            parts['ppt/presentation.xml']=etree.tostring(root)
            with ZipFile(source,'w') as z:
                for name,data in parts.items():z.writestr(name,data)
            document,paragraphs,_=open_office(source)
            self.assertEqual(str(document.slides[0].part.partname),'/ppt/slides/slide2.xml')
            self.assertEqual(paragraphs[0][1].text,'操作说明');paragraphs[0][1].text='Operating instructions'
            output=folder/'manual.pptx';save_office(source,document,output,['slide1/shape0/p0'])
            with ZipFile(output) as z:
                changed=[name for name,data in parts.items() if z.read(name)!=data]
            self.assertEqual(changed,['ppt/slides/slide2.xml'])
            self.assertEqual(Presentation(output).slides[0].shapes[0].text,'Operating instructions')
            class Model:
                def chat_json(self,prompt,**kwargs):return {'translation':'Operating instructions' if '【待译原文】\n操作说明' in prompt else 'Brand description'}
            result=translate_file(source,Glossary([]),None,client=Model(),workers=1,defer_visual=True)
            self.assertTrue(result['translation_complete']);self.assertFalse(result['remaining_chinese'])
            self.assertEqual([s.shapes[0].text for s in Presentation(result['path']).slides],['Operating instructions','Brand description'])


if __name__=='__main__':unittest.main()
