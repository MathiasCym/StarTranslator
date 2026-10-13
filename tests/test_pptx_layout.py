"""Translation coverage, real font fitting and native package preservation."""
import os, sys, tempfile, unittest, threading
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from io import BytesIO
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
if os.environ.get('QA_PYTHON_LIB'):
    sys.path.append(os.environ['QA_PYTHON_LIB'])  # fixture authoring only; not a shipped dependency
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from backend.translation import translate_file, CHINESE
from backend.extract import extract
from backend.extract.office import open_office
from backend.glossary.loader import Glossary
from backend.models import GlossaryEntry
from backend.replacer.pptx_layout import fit_document, capture_layout, Metrics, estimate, _paragraph_data
from backend import workflow, officecli_bridge
from backend.llm.client import LLMUnavailable


class Model:
    def __init__(self,answers): self.answers=iter(answers); self.prompts=[]
    def chat_json(self,prompt,**kwargs):
        self.prompts.append(prompt)
        result=next(self.answers)
        if isinstance(result,Exception): raise result
        return result if isinstance(result,dict) else {'translation':result}


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.folder=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.folder/'profile')}); self.env.start()
        self.glossary=Glossary([])
    def tearDown(self): self.env.stop(); self.temp.cleanup()
    def deck(self,text,width=5,height=1,size=18):
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(width),Inches(height))
        shape.text=text
        for p in shape.text_frame.paragraphs:
            for r in p.runs: r.font.name='Arial'; r.font.size=Pt(size)
        source=self.folder/'source.pptx'; deck.save(source)
        return source,deck,shape
    def translate(self,source,model):
        return translate_file(str(source),self.glossary,None,client=model)

    def test_soft_breaks_translate_dates_and_labels_without_losing_line_structure(self):
        source,deck,shape=self.deck('HDR\v亮度500nits。\v刷新率60Hz。')
        original=source.read_bytes()
        self.assertTrue(extract(str(source))[0].editable)
        model=Model(['500 nits.','60 Hz.']); result=self.translate(source,model)
        paragraph=Presentation(result['path']).slides[0].shapes[0].text_frame.paragraphs[0]
        self.assertEqual(paragraph.text,'HDR\v500 nits.\v60 Hz.')
        self.assertEqual(len(paragraph._p.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}br')),2)
        self.assertEqual(result['remaining_chinese'],[])
        self.assertEqual(source.read_bytes(),original); self.assertEqual(len(model.prompts),2)

    def test_leading_and_trailing_breaks_remain_native_break_nodes(self):
        source,_,_=self.deck('\v亮度500nits。\v')
        result=self.translate(source,Model(['500 nits.']))
        self.assertEqual(Presentation(result['path']).slides[0].shapes[0].text,'\v500 nits.\v')
        with ZipFile(result['path']) as archive:
            self.assertNotIn(b'_x000B_',archive.read('ppt/slides/slide1.xml'))

    def test_word_manual_line_breaks_are_retained(self):
        source=self.folder/'source.docx'; doc=Document(); doc.add_paragraph('亮度500nits。\n刷新率60Hz。'); doc.save(source)
        result=self.translate(source,Model(['500 nits.','60 Hz.']))
        self.assertEqual(Document(result['path']).paragraphs[0].text,'500 nits.\n60 Hz.')

    def test_residual_chinese_is_patched_without_rewriting_correct_english(self):
        source,_,_=self.deck('支持逐点校正，亮度500nits。')
        model=Model(['Supports 逐点校正 at 500 nits.',{'translations':[{'source':'逐点校正','translation':'pixel-level calibration'}]}])
        result=self.translate(source,model)
        self.assertEqual(Presentation(result['path']).slides[0].shapes[0].text,'Supports pixel-level calibration at 500 nits.')
        self.assertEqual(result['remaining_chinese'],[]); self.assertEqual(len(model.prompts),2)
        self.assertIn('只需补译的中文',model.prompts[1]); self.assertNotIn('重新生成要求',model.prompts[1])

    def test_confirmed_glossary_repairs_residual_without_extra_model_call(self):
        source,_,_=self.deck('支持逐点校正，亮度500nits。')
        self.glossary=Glossary([GlossaryEntry('逐点校正','pixel-level calibration',priority=1)])
        model=Model(['Supports 逐点校正 at 500 nits.']); result=self.translate(source,model)
        self.assertEqual(result['needs_review'],0); self.assertEqual(len(model.prompts),1)

    def test_supplement_network_failure_preserves_generated_english_and_continues(self):
        source,deck,shape=self.deck('支持逐点校正，亮度500nits。\n刷新率60Hz。')
        model=Model(['Supports 逐点校正 at 500 nits.',LLMUnavailable('连接超时'),'60 Hz.'])
        result=self.translate(source,model)
        output=Presentation(result['path']).slides[0].shapes[0]
        self.assertEqual(output.text,'Supports 逐点校正 at 500 nits.\n60 Hz.')
        self.assertEqual(result['untranslated'],1); self.assertEqual(result['status'],'needs_review')
        self.assertIn('局部补译',result['review_items'][0]['reason'])

    def test_fields_are_protected_and_remaining_chinese_cannot_claim_completion(self):
        source,deck,shape=self.deck('亮度500nits。')
        other=deck.slides[0].shapes.add_textbox(Inches(1),Inches(3),Inches(5),Inches(1))
        p=other.text_frame.paragraphs[0]; field=OxmlElement('a:fld'); field.set('id','{00000000-0000-0000-0000-000000000001}'); field.set('type','datetime')
        node=OxmlElement('a:t'); node.text='中文日期'; field.append(node); p._p.append(field); deck.save(source)
        result=self.translate(source,Model(['500 nits.']))
        self.assertEqual(result['status'],'needs_review'); self.assertEqual(result['untranslated'],1)
        self.assertEqual(result['remaining_chinese'][0]['text'],'中文日期')
        self.assertIn('English_Draft',result['path'])
        self.assertEqual(Presentation(result['path']).slides[0].shapes[1].text,'中文日期')

    def test_style_failure_uses_confirmed_term_anchor_and_preserves_color_and_bold(self):
        source,deck,shape=self.deck('')
        p=shape.text_frame.paragraphs[0]
        for text,bold,color in [('采用',False,'FFFFFF'),('逐点校正',True,'48CCFF'),('。',False,'FFFFFF')]:
            r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(18); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
        deck.save(source)
        self.glossary=Glossary([GlossaryEntry('逐点校正','pixel-level calibration',priority=1)])
        result=self.translate(source,Model(['Use pixel-level calibration.',LLMUnavailable('样式映射未完成')]))
        runs=Presentation(result['path']).slides[0].shapes[0].text_frame.paragraphs[0].runs
        emphasis=next(r for r in runs if r.text=='pixel-level calibration')
        self.assertTrue(emphasis.font.bold); self.assertEqual(str(emphasis.font.color.rgb),'48CCFF')
        self.assertFalse(any('局部颜色' in w['reason'] for w in result['layout_warnings']))

    def test_paragraphs_fit_original_slots_without_moving_or_changing_emphasis(self):
        source,deck,shape=self.deck('中文标题\n中文说明',width=3,height=1,size=28)
        shape.text_frame.paragraphs[1].runs[0].font.size=Pt(20)
        shape.text_frame.paragraphs[0].runs[0].font.bold=True
        shape.text_frame.paragraphs[1].runs[0].font.color.rgb=RGBColor.from_string('48CCFF')
        deck.save(source); original=(shape.left,shape.top,shape.width,shape.height)
        result=self.translate(source,Model(['Technical display solution','Accurate calibration for cinema screens']))
        output=Presentation(result['path']); fitted=output.slides[0].shapes[0]
        self.assertEqual((fitted.left,fitted.top,fitted.width,fitted.height),original)
        self.assertTrue(result['layout_adjustments']); self.assertEqual(fitted.text_frame.auto_size,MSO_AUTO_SIZE.NONE)
        a,b=[p.runs[0] for p in fitted.text_frame.paragraphs]
        self.assertGreater(a.font.size.pt,b.font.size.pt)
        self.assertLess(a.font.size.pt,28); self.assertLess(b.font.size.pt,20)
        self.assertTrue(a.font.bold); self.assertEqual(str(b.font.color.rgb),'48CCFF')
        self.assert_frame_fits(output,fitted)

    def assert_frame_fits(self,deck,shape):
        frame=shape.text_frame
        width=(shape.width-frame.margin_left-frame.margin_right)/12700
        height=(shape.height-frame.margin_top-frame.margin_bottom)/12700
        with Metrics() as metrics:
            h,w,clipped=estimate(_paragraph_data(frame,shape,deck.slides[0],deck),width,metrics,wrap=frame.word_wrap is not False)
        self.assertLessEqual(h,height); self.assertFalse(clipped)

    def test_inherited_placeholder_font_is_resolved_before_fitting(self):
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[0]); shape=slide.shapes.title
        shape.width,shape.height=Inches(4),Inches(.7); shape.text='Long technical display solution title'
        self.assertIsNone(shape.text_frame.paragraphs[0].runs[0].font.size)
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'])
        run=shape.text_frame.paragraphs[0].runs[0]
        self.assertIsNotNone(run.font.size); self.assertGreater(run.font.size.pt,18)
        self.assertTrue(adjustments); self.assertEqual(warnings,[]); self.assert_frame_fits(deck,shape)

    def test_merged_table_cell_is_measured_as_one_frame(self):
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
        table=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(4),Inches(.55)).table
        cell=table.cell(0,0).merge(table.cell(0,1)) or table.cell(0,0)
        cell.text='Pixel-level brightness and chroma calibration'
        for r in cell.text_frame.paragraphs[0].runs: r.font.size=Pt(22); r.font.bold=True
        adjustments,warnings=fit_document(deck,['slide1/shape0/table/0/0/0'])
        self.assertTrue(adjustments); self.assertEqual(warnings,[])
        self.assertLess(cell.text_frame.paragraphs[0].runs[0].font.size.pt,22)
        self.assertTrue(cell.text_frame.paragraphs[0].runs[0].font.bold)

    def test_crowded_frame_stops_at_readable_size_and_reports_location(self):
        _,deck,shape=self.deck('A long explanation '*40,width=1,height=.4,size=20)
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'])
        self.assertTrue(warnings); self.assertEqual(warnings[0]['slide'],1)
        self.assertGreaterEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt,9)
        self.assertEqual(shape.width,Inches(1))

    def test_existing_autofit_scale_is_baked_even_when_run_font_is_inherited(self):
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[0]); shape=slide.shapes.title
        shape.text='Short title'; shape.width=Inches(6); shape.height=Inches(1)
        shape.text_frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        shape.text_frame._txBody.bodyPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit').set('fontScale','50000')
        with Metrics() as metrics:
            original=_paragraph_data(shape.text_frame,shape,slide,deck)[0][1][0][1][1]
        fit_document(deck,['slide1/shape0/p0'])
        self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt,original)
        self.assertEqual(shape.text_frame.auto_size,MSO_AUTO_SIZE.NONE)

    def test_blank_paragraphs_count_towards_available_height(self):
        _,deck,shape=self.deck('Technical display\n\nAccurate calibration',width=3,height=1,size=22)
        shape.text_frame.paragraphs[1]._p.get_or_add_endParaRPr().set('sz','2200')
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0','slide1/shape0/p2'])
        self.assertTrue(adjustments); self.assertFalse(warnings); self.assert_frame_fits(deck,shape)

    def test_standard_quantity_unit_spacing_is_valid_but_missing_units_are_rejected(self):
        from backend.writing_policy import validate_translation
        validate_translation('刷新率60Hz，帧率120fps。','A refresh rate of 60 Hz and a frame rate of 120 fps.')
        with self.assertRaises(ValueError): validate_translation('60Hz，60Hz。','60 Hz and 60 pixels.')

    def test_zero_height_uses_source_ink_budget_without_changing_geometry(self):
        _,deck,shape=self.deck('显示方案',width=3,height=1,size=18)
        shape.height=1; shape.text_frame.auto_size=MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        original=(shape.left,shape.top,shape.width,shape.height)
        snapshot=capture_layout(deck,['slide1/shape0/p0'])
        shape.text_frame.paragraphs[0].runs[0].text='Technical display solutions for cinema screens'
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'],snapshot)
        self.assertFalse(warnings); self.assertTrue(adjustments[0]['original_bounds'])
        self.assertEqual((shape.left,shape.top,shape.width,shape.height),original)
        self.assertEqual(shape.text_frame.auto_size,MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT)

    def test_zero_width_and_height_keep_source_geometry_and_visual_text_bounds(self):
        _,deck,shape=self.deck('显示控制器',width=3,height=1,size=12)
        shape.width=shape.height=1; shape.text_frame.auto_size=MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        original=(shape.left,shape.top,shape.width,shape.height)
        snapshot=capture_layout(deck,['slide1/shape0/p0'])
        shape.text_frame.paragraphs[0].runs[0].text='Display controller'
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'],snapshot)
        self.assertFalse(warnings); self.assertTrue(adjustments[0]['original_bounds'])
        self.assertEqual((shape.left,shape.top,shape.width,shape.height),original)
        with Metrics() as metrics:
            before=estimate(snapshot['slide1/shape0']['data'],10000,metrics,wrap=False)
            after=estimate(_paragraph_data(shape.text_frame,shape,deck.slides[0],deck),10000,metrics,wrap=False)
        self.assertLessEqual(after[1],before[1]+.5); self.assertLessEqual(after[0],before[0]+.5)

    def test_long_card_title_does_not_push_date_and_body_below_the_original_slots(self):
        _,deck,shape=self.deck('显示方案\v技术落地\n2024.11.20\n中国首家LED影厅正式启用。',width=1.7,height=2.5,size=11)
        shape.text_frame.word_wrap=True
        for run in shape.text_frame.paragraphs[0].runs:run.font.size=Pt(20);run.font.bold=True
        original=(shape.left,shape.top,shape.width,shape.height)
        snapshot=capture_layout(deck,['slide1/shape0/p0','slide1/shape0/p2'])
        shape.text_frame.paragraphs[0].runs[0].text='VLED'
        shape.text_frame.paragraphs[0].runs[1].text='Commercial deployment'
        shape.text_frame.paragraphs[2].runs[0].text="China's first commercial LED cinema screen is now in operation."
        with Metrics() as metrics:
            source_title=estimate([snapshot['slide1/shape0']['data'][0]],(shape.width-shape.text_frame.margin_left-shape.text_frame.margin_right)/12700*.975,metrics)[0]
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0','slide1/shape0/p2'],snapshot)
        self.assertFalse(warnings);self.assertTrue(adjustments)
        with Metrics() as metrics:
            new_title=estimate([_paragraph_data(shape.text_frame,shape,deck.slides[0],deck)[0]],(shape.width-shape.text_frame.margin_left-shape.text_frame.margin_right)/12700*.975,metrics)[0]
        self.assertLess(abs(source_title-new_title),2)
        self.assertEqual((shape.left,shape.top,shape.width,shape.height),original)
        self.assert_frame_fits(deck,shape)

    def test_excess_line_spacing_is_compacted_before_extreme_font_shrinking(self):
        _,deck,shape=self.deck('Cinema display calibration and professional color management',width=3,height=.65,size=18)
        shape.text_frame.word_wrap=True;shape.text_frame.paragraphs[0].line_spacing=2.0
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'])
        self.assertFalse(warnings);self.assertTrue(adjustments[0]['spacing_compacted'])
        self.assertEqual(shape.text_frame.paragraphs[0].line_spacing,1.0)
        self.assert_frame_fits(deck,shape)

    def test_font_substitution_retains_color_bold_and_hyperlink_and_all_geometry(self):
        _,deck,shape=self.deck('Professional technical display calibration for cinema screens',width=1.5,height=.8,size=20)
        shape.text_frame.word_wrap=True
        run=shape.text_frame.paragraphs[0].runs[0];run.font.name='Arial Black';run.font.bold=True
        run.font.color.rgb=RGBColor.from_string('48CCFF');run.hyperlink.address='https://example.test/reference'
        original=(shape.left,shape.top,shape.width,shape.height)
        adjustments,warnings=fit_document(deck,['slide1/shape0/p0'])
        self.assertTrue(any(a['font'] for a in adjustments));self.assertTrue(run.font.bold)
        self.assertEqual(str(run.font.color.rgb),'48CCFF');self.assertEqual(run.hyperlink.address,'https://example.test/reference')
        self.assertEqual((shape.left,shape.top,shape.width,shape.height),original)
        if not warnings:self.assert_frame_fits(deck,shape)

    def test_chart_labels_and_embedded_workbook_change_together_without_touching_numbers(self):
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6]); data=CategoryChartData()
        data.categories=['2022年','2023年']; data.add_series('数量',[25,44])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(1),Inches(6),Inches(3),data)
        source=self.folder/'chart.pptx'; deck.save(source); original=source.read_bytes()
        model=Model(['Quantity','2022','2023']); result=self.translate(source,model)
        self.assertEqual(result['remaining_chinese'],[])
        with ZipFile(source) as before,ZipFile(result['path']) as after:
            chart=etree.fromstring(after.read('ppt/charts/chart1.xml'))
            self.assertEqual(chart.xpath('//*[local-name()="numCache"]//*[local-name()="v"]/text()'),['25','44'])
            self.assertEqual(chart.xpath('//*[local-name()="strCache"]//*[local-name()="v"]/text()'),['Quantity','2022','2023'])
            book=next(n for n in after.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx'))
            with ZipFile(BytesIO(after.read(book))) as workbook:
                strings=etree.fromstring(workbook.read('xl/sharedStrings.xml'))
                text=''.join(strings.xpath('//*[local-name()="t"]/text()'))
                self.assertFalse(CHINESE.search(text)); self.assertIn('Quantity',text)
            changed=[n for n in before.namelist() if before.read(n)!=after.read(n)]
            self.assertEqual(set(changed),{'ppt/charts/chart1.xml',book})
        self.assertEqual(source.read_bytes(),original)
        self.assertEqual(officecli_bridge.compare(officecli_bridge.inspect(source),officecli_bridge.inspect(result['path'])),[])

    def test_shared_master_footer_is_translated_and_its_original_style_survives(self):
        from copy import deepcopy
        source,deck,shape=self.deck('LED影院解决方案',height=.5,size=16)
        shape.text_frame.paragraphs[0].runs[0].font.bold=True
        deck.slide_master._element.spTree.append(deepcopy(shape._element))
        deck.slides[0].shapes._spTree.remove(shape._element); deck.save(source)
        result=self.translate(source,Model(['LED cinema solution']))
        self.assertEqual(result['remaining_chinese'],[])
        master=Presentation(result['path']).slide_master
        footer=next(s for s in master.shapes if s.has_text_frame and s.text=='LED cinema solution')
        self.assertTrue(footer.text_frame.paragraphs[0].runs[0].font.bold)

    def test_proofreading_export_fits_longer_replacement_and_preserves_soft_break(self):
        source,deck,shape=self.deck('HDR\vdot spacing',width=2,height=.65,size=22)
        deck.save(source)
        glossary=Glossary([GlossaryEntry('点间距','pixel pitch',wrong_variants=['dot spacing'],priority=1)])
        job={'id':'layout-review','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=glossary):
            workflow.run(job,{'kind':'check','paths':[str(source)],'use_llm':False})
        self.assertEqual(job['status'],'complete'); finding=job['result']['findings'][0]
        result=workflow.apply_report({'report_id':job['result']['report_id'],'path':str(source),'selected':{finding['segment_id']:[[finding['found_text'],finding['suggested']]]}})
        output=Presentation(result['output']); shape=output.slides[0].shapes[0]
        self.assertEqual(shape.text,'HDR\vpixel pitch'); self.assertTrue(result['layout_adjustments'])
        self.assert_frame_fits(output,shape)


if __name__=='__main__': unittest.main()
