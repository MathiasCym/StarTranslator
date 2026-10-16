"""No-network integration tests for bilingual translation, rules and shared editing."""
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from backend.extract import extract
from backend.glossary import store, writer
from backend.glossary.loader import Glossary, _load_any, _rows_to_entries, DEFAULT_PATH
from backend.models import GlossaryEntry, Segment
from backend.matcher import check_segments
from backend.matcher.rules import check_rules
from backend.matcher.judge import judge_text
from backend.references import ReferenceIndex
from backend.translation import translate_file, Cancelled
from backend.llm.client import LLMUnavailable
from backend import workflow
from backend.server import app


class FakeModel:
    def __init__(self, response=None):
        self.response = response or {'translation':'The pixel pitch is 2.5 mm.'}
        self.prompts = []
    def chat_json(self,prompt,**kwargs):
        self.prompts.append(prompt)
        return self.response


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        # These tests exercise text/package behavior. Image review has its own suite.
        for target in ('backend.translation.review_pptx','backend.workflow.review_pptx'):
            stub=patch(target,return_value=([],[]));stub.start();self.addCleanup(stub.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.environment = patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'data')})
        self.environment.start()
        self.g = Glossary([GlossaryEntry('点间距','pixel pitch',variants=['dot pitch'],
            machine_variants=['pixel distance'],wrong_variants=['dot spacing'],priority=1)])

    def tearDown(self):
        self.environment.stop()
        self.temp.cleanup()

    def document(self,name,text):
        path = self.root/name
        path.parent.mkdir(parents=True,exist_ok=True)
        doc = Document()
        for part in text.split('\n'):
            doc.add_paragraph(part)
        doc.save(path)
        return str(path)

    def check_job(self,paths,**kw):
        job={'id':'test-report','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            workflow.run(job,{'kind':'check','paths':paths,'use_llm':False,**kw})
        self.assertEqual(job['status'],'complete',job)
        return job['result']

    def test_aliases_candidates_and_word_boundaries_do_not_trigger_rules(self):
        segment=Segment('one','x','dot pitch; pixel distance; dot spacingx; dot spacing.','a')
        hits=check_rules([segment],self.g)
        self.assertEqual([h.found_text for h in hits],['dot spacing'])
        # 已确认错译在 priority 1 与 2 都参与确定性替换（脚本先兜底）。
        self.g.entries[0].priority=2
        self.assertEqual([h.found_text for h in check_rules([segment],self.g)],['dot spacing'])
        # 仅供参考（priority=3）不参与。
        self.g.entries[0].priority=3
        self.assertEqual(check_rules([segment],self.g),[])

    def test_ambiguous_and_accepted_variants_do_not_trigger_rules(self):
        self.g.entries.append(GlossaryEntry('别的概念','other term',wrong_variants=['dot spacing'],priority=1))
        self.assertEqual(check_rules([Segment('a','x','dot spacing','a')],self.g),[])
        self.g.entries.pop()
        self.g.entries.append(GlossaryEntry('合法词','dot spacing'))
        self.assertEqual(check_rules([Segment('a','x','dot spacing','a')],self.g),[])

    def test_legacy_false_positives_remain_clean_without_model(self):
        glossary=Glossary.load()
        fixtures=json.loads((Path(__file__).parent/'fixtures/false_positives.json').read_text(encoding='utf-8'))
        for case in fixtures['cases']:
            self.assertEqual(check_rules([Segment('a','x',case['context'],'a')],glossary),[])

    def test_rules_only_never_calls_model(self):
        with patch('backend.matcher.LLMClient') as model:
            findings,stats=check_segments([Segment('a','x','dot spacing','a')],self.g,use_llm=False)
        model.assert_not_called()
        self.assertEqual(len(findings),1)
        self.assertEqual(stats['model_status'],'not_requested')

    def test_offline_keeps_rules_and_reports_incomplete(self):
        model=FakeModel()
        with patch.object(model,'chat_json',side_effect=LLMUnavailable('offline')), \
             patch('backend.matcher.time.sleep'):
            findings,stats=check_segments([Segment('a','x','dot spacing','a')],self.g,client=model)
        self.assertEqual(len(findings),1)
        self.assertEqual(stats['model_status'],'unavailable')
        self.assertTrue(stats['warnings'])

    def test_one_failing_segment_does_not_abort_the_rest(self):
        # 单段持续失败应被跳过并记录，其余段落照常复核（不再 break 整份文档）。
        class FlakyModel(FakeModel):
            def chat_json(self,prompt,**kwargs):
                if 'BROKEN' in prompt.split('【待审校段落】\n', 1)[1].split('\n\n输出 JSON', 1)[0]:
                    raise LLMUnavailable('bad segment')
                return [{'found':'dot spacing','suggested':'pixel pitch','concept_zh':'点间距',
                         'kind':'term','confidence':'high','reason':'术语译错'}]
        segments=[Segment('a','x','BROKEN text','a',file_path='f'),
                  Segment('b','x','dot spacing here','b',file_path='f')]
        with patch('backend.matcher.time.sleep'):
            findings,stats=check_segments(segments,self.g,client=FlakyModel())
        self.assertEqual(stats['model_status'],'partial')
        self.assertEqual(stats['errors'],1)
        self.assertEqual(stats['unreviewed'],1)
        self.assertEqual(stats['judged'],1)
        self.assertTrue(any(f.segment_id=='b' and f.found_text=='dot spacing' for f in findings))
        self.assertTrue(any('未能完成模型复核' in w for w in stats['warnings']))

    def test_same_filenames_get_unique_ids_and_all_occurrences_export(self):
        first=self.document('first/same.docx','dot spacing and dot spacing\ndot spacing')
        second=self.document('second/same.docx','dot spacing')
        original=Path(first).read_bytes()
        result=self.check_job([first,second])
        ids=[s['id'] for s in result['segments']]
        self.assertEqual(len(ids),len(set(ids)))
        self.assertEqual(len(result['findings']),3)
        selected={f['segment_id']:[[f['found_text'],f['suggested']]] for f in result['findings'] if f['file_path']==first}
        output=workflow.apply_report({'report_id':result['report_id'],'path':first,'selected':selected})
        self.assertEqual(output['applied_count'],3)
        self.assertNotIn('dot spacing',' '.join(p.text for p in Document(output['output']).paragraphs))
        self.assertEqual(Path(first).read_bytes(),original)
        self.assertIn('dot spacing',Document(second).paragraphs[0].text)

    def test_source_change_blocks_export(self):
        path=self.document('a.docx','dot spacing')
        result=self.check_job([path])
        self.document('a.docx','changed')
        with self.assertRaisesRegex(ValueError,'不一致'):
            workflow.apply_report({'report_id':result['report_id'],'path':path,'selected':{}})

    def test_overwrite_is_direct_without_backup_or_recheck(self):
        path = self.document('overwrite.docx', 'dot spacing')
        original = Path(path).read_bytes()
        result = self.check_job([path]); f = result['findings'][0]
        data = {'report_id':result['report_id'],'path':path,'selected':{f['segment_id']:[[f['found_text'],f['suggested']]]},'overwrite':True}
        output = workflow.apply_report(data)
        self.assertEqual(output['output'], path)
        self.assertTrue(output['overwritten'])
        self.assertEqual(Document(path).paragraphs[0].text, 'pixel pitch')
        self.assertNotEqual(Path(path).read_bytes(), original)
        self.assertNotIn('backup', output)
        self.assertTrue(output['finding_updates'][0]['finding']['applied'])
        self.assertEqual(list(self.root.glob('overwrite*')), [Path(path)])
        self.assertEqual(list(self.root.glob('.star-review-*')), [])
        with self.assertRaisesRegex(ValueError, '已应用'): workflow.apply_report(data)

    def test_overwrite_remaining_choice_uses_same_report_without_model(self):
        self.g.entries.append(GlossaryEntry('刷新率','refresh rate',wrong_variants=['frame change frequency'],priority=1))
        path = self.document('remaining.docx', 'dot spacing and frame change frequency')
        result = self.check_job([path])
        first = next(f for f in result['findings'] if f['found_text']=='dot spacing')
        second = next(f for f in result['findings'] if f['found_text']=='frame change frequency')
        def apply(f):
            return workflow.apply_report({'report_id':result['report_id'],'path':path,'selected':{f['segment_id']:[[f['found_text'],f['suggested']]]},'overwrite':True})
        with patch('backend.workflow.check_segments') as model:
            output = apply(first)
            remaining = next(u['finding'] for u in output['finding_updates'] if u['finding']['found_text']==second['found_text'])
            self.assertTrue(remaining['editable'])
            self.assertEqual(remaining['context'], 'pixel pitch and frame change frequency')
            apply(second)
            model.assert_not_called()
        self.assertEqual(Document(path).paragraphs[0].text, 'pixel pitch and refresh rate')
        self.assertEqual(list(self.root.glob('remaining*')), [Path(path)])

    def test_failed_atomic_overwrite_leaves_original_untouched(self):
        path = self.document('locked.docx', 'dot spacing'); original = Path(path).read_bytes()
        result = self.check_job([path]); f = result['findings'][0]
        with patch('backend.workflow.os.replace', side_effect=PermissionError('Office file locked')):
            with self.assertRaises(PermissionError):
                workflow.apply_report({'report_id':result['report_id'],'path':path,'selected':{f['segment_id']:[[f['found_text'],f['suggested']]]},'overwrite':True,'confirm_overwrite':True})
        self.assertEqual(Path(path).read_bytes(), original)
        self.assertEqual(list(self.root.glob('.star-review-*')), [])
        self.assertEqual(list(self.root.glob('locked_before_review*.docx')), [])
        self.assertFalse(workflow.REPORTS[result['report_id']]['findings'][0].applied)

    def test_ppt_overwrite_and_stale_source_guard(self):
        path = self.root/'cinema.pptx'; deck = Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6]); slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text='dot spacing'; deck.save(path)
        original=path.read_bytes(); result=self.check_job([str(path)]); f=result['findings'][0]
        data={'report_id':result['report_id'],'path':str(path),'selected':{f['segment_id']:[[f['found_text'],f['suggested']]]},'overwrite':True,'confirm_overwrite':True}
        output=workflow.apply_report(data)
        self.assertEqual(Presentation(path).slides[0].shapes[0].text, 'pixel pitch')
        self.assertNotEqual(path.read_bytes(), original)
        self.assertNotIn('backup', output)
        self.assertEqual(list(self.root.glob('cinema*')), [path])
        deck=Presentation(path); deck.slides[0].shapes[0].text='external change'; deck.save(path)
        with self.assertRaisesRegex(ValueError, '不一致'): workflow.apply_report(data)

    def test_cross_run_rule_preserves_unaffected_format(self):
        path=self.root/'styled.docx'; doc=Document(); p=doc.add_paragraph()
        p.add_run('Prefix ').bold=True; p.add_run('dot '); p.add_run('spacing'); p.add_run(' suffix').italic=True
        doc.save(path)
        result=self.check_job([str(path)]); f=result['findings'][0]
        output=workflow.apply_report({'report_id':result['report_id'],'path':str(path),'selected':{f['segment_id']:[[f['found_text'],f['suggested']]]}})
        para=Document(output['output']).paragraphs[0]
        self.assertEqual(para.text,'Prefix pixel pitch suffix')
        self.assertTrue(para.runs[0].bold); self.assertTrue(para.runs[-1].italic)

    def test_reference_pair_is_used_in_model_prompt(self):
        reference=self.document('cn.docx','点间距为2.5毫米。')
        target=self.document('en.docx','The dot spacing is 2.5 mm.')
        refs=ReferenceIndex([reference],self.g,{target:reference})
        model=FakeModel([]); model.response=[]
        segment=extract(target)[0]
        judge_text(segment.text,self.g,model,reference=refs.context(segment))
        self.assertIn('点间距为2.5毫米',model.prompts[0])
        self.assertIn('cn.docx',model.prompts[0])

    def test_translation_preserves_source_and_supports_tables(self):
        path=self.document('source.docx','点间距为2.5毫米。')
        doc=Document(path); doc.add_table(rows=1,cols=1).cell(0,0).text='点间距为2.5毫米。'; doc.save(path)
        original=Path(path).read_bytes(); model=FakeModel()
        output=translate_file(path,self.g,ReferenceIndex([],self.g),client=model)
        self.assertEqual(output['translated'],2)
        self.assertEqual(Document(output['path']).tables[0].cell(0,0).text,'The pixel pitch is 2.5 mm.')
        self.assertEqual(Path(path).read_bytes(),original)

    def test_ppt_group_and_table_translation(self):
        path=self.root/'source.pptx'; prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
        group=slide.shapes.add_group_shape(); group.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='点间距为2.5毫米。'
        slide.shapes.add_table(1,1,0,Inches(2),Inches(3),Inches(1)).table.cell(0,0).text='点间距为2.5毫米。'
        prs.save(path)
        result=translate_file(str(path),self.g,ReferenceIndex([],self.g),client=FakeModel())
        self.assertEqual(result['translated'],2)
        self.assertTrue(all('pixel pitch' in s.text for s in extract(result['path'])))

    def test_translation_failure_keeps_a_marked_draft_but_explicit_cancellation_writes_nothing(self):
        path=self.document('source.docx','点间距为2.5毫米。')
        result=translate_file(path,self.g,None,client=FakeModel({'translation':'仍然是中文'}))
        self.assertFalse(result['translation_complete']);self.assertTrue(result['review_items'])
        self.assertEqual(Document(result['path']).paragraphs[0].text,'点间距为2.5毫米。')
        outputs=list(self.root.glob('*English*'))
        cancel=threading.Event();cancel.set()
        with self.assertRaises(Cancelled):
            translate_file(path,self.g,None,client=FakeModel(),cancel=cancel)
        self.assertEqual(list(self.root.glob('*English*')),outputs)

    def test_shared_save_conflict_backup_and_deletion(self):
        current=store.read()
        shared=self.root/'shared.json'; store.select(str(shared),create=True)
        a=store.read(); b=store.read()
        saved=store.save(a['entries'][1:],a['fingerprint'])
        self.assertEqual(len(saved['entries']),len(a['entries'])-1)
        self.assertTrue(Path(str(shared)+'.bak').exists())
        with self.assertRaises(store.Conflict): store.save(b['entries'],b['fingerprint'])
        self.assertEqual(store.read()['entries'],saved['entries'])
        self.assertEqual(store.select()['entries'],current['entries'])

    def test_shared_lock_prevents_concurrent_writer(self):
        current=store.read(); lock=Path(current['path']+'.lock');lock.touch()
        with self.assertRaises(store.Conflict): store.save(current['entries'],current['fingerprint'])
        self.assertEqual(store.read()['fingerprint'],current['fingerprint'])
        lock.unlink()

    def test_priority_two_wrong_variants_trigger_and_priority_three_does_not(self):
        # ⑤ 脚本层自动改正：复用 wrong_variants，priority 1 与 2 生效。
        segment=Segment('one','x','The dot spacing here is wrong.','a')
        self.assertEqual([h.suggested for h in check_rules([segment],self.g)],['pixel pitch'])
        self.g.entries[0].priority=2
        self.assertEqual([h.suggested for h in check_rules([segment],self.g)],['pixel pitch'])
        self.g.entries[0].priority=3
        self.assertEqual(check_rules([segment],self.g),[])

    def test_excel_table_round_trip_keeps_rich_fields(self):
        path=str(self.root/'glossary.xlsx')
        entry=GlossaryEntry('点间距','pixel pitch',variants=['dot pitch'],
            machine_variants=['pixel distance'],wrong_variants=['dot spacing'],
            priority=1,note='定义',source='专业术语.pptx / 第1页')
        rows=store._entries_to_rows([entry.to_dict()],writer.HEADER_FIELDS)
        writer.save_table(path,rows,None,columns=writer.HEADER_FIELDS)
        loaded=writer.load_table(path)
        self.assertEqual(loaded['columns'],writer.HEADER_FIELDS)
        back=[e.to_dict() for e in _rows_to_entries([writer.header_for(loaded['columns']),*loaded['rows']])]
        self.assertEqual(len(back),1)
        got=back[0]
        self.assertEqual(got['variants'],['dot pitch'])
        self.assertEqual(got['machine_variants'],['pixel distance'])
        self.assertEqual(got['wrong_variants'],['dot spacing'])
        self.assertEqual(got['priority'],1)
        self.assertEqual(got['source'],'专业术语.pptx / 第1页')
        self.assertEqual(got['note'],'定义')

    def test_legacy_three_column_table_stays_three_columns(self):
        path=str(self.root/'legacy.csv')
        Path(path).write_text('中文,英文,备注\n点间距,pixel pitch,定义\n',encoding='utf-8-sig')
        loaded=writer.load_table(path)
        self.assertEqual(loaded['columns'],['zh','en','note'])
        writer.save_table(path,[['接收卡','receiving card',''],['发送卡','sending card','']],loaded['fingerprint'])
        again=writer.load_table(path)
        self.assertEqual(again['columns'],['zh','en','note'])
        self.assertEqual(len(again['rows']),2)

    def test_shared_excel_lexicon_read_save_and_conflict(self):
        shared=self.root/'shared.xlsx'
        store.select(str(shared),create=True)
        self.assertTrue(shared.exists())
        a=store.read(); b=store.read()
        self.assertEqual(len(a['entries']),len(b['entries']))
        saved=store.save(a['entries'][1:],a['fingerprint'])
        self.assertEqual(len(saved['entries']),len(a['entries'])-1)
        with self.assertRaises(store.Conflict): store.save(b['entries'],b['fingerprint'])
        self.assertEqual(len(store.read()['entries']),len(a['entries'])-1)

    def test_detail_fields_can_extend_legacy_table_without_losing_headers(self):
        shared = self.root / 'legacy-detail.csv'
        shared.write_text('中文,英文,备注,部门,负责人\n点间距,pixel pitch,定义,显示部,张三\n', encoding='utf-8-sig')
        current = writer.load_table(str(shared))
        columns = current['columns'] + ['variants', 'wrong_variants', 'machine_variants', 'priority']
        rows = [current['rows'][0] + ['dot pitch', 'dot spacing', 'pixel distance', '2']]
        store.edit_table(str(shared), rows, columns, current['fingerprint'])
        saved = writer.load_table(str(shared))
        self.assertEqual(saved['header'][:5], ['中文', '英文', '备注', '部门', '负责人'])
        self.assertEqual(saved['rows'][0][3:5], ['显示部', '张三'])
        entry = store.read(shared)['entries'][0]
        self.assertEqual(entry['variants'], ['dot pitch'])
        self.assertEqual(entry['wrong_variants'], ['dot spacing'])
        self.assertEqual(entry['machine_variants'], ['pixel distance'])
        self.assertEqual(entry['priority'], 2)

    def test_grid_and_detail_changes_share_one_conflict_checked_save(self):
        shared = self.root / 'unified.xlsx'
        store.select(str(shared), create=True)
        current = store.read()
        rows, columns = store.table_rows(current), store.table_columns(current)
        rows[0][columns.index('en')] = 'Updated standard wording'
        rows[0][columns.index('note')] = '从详情编辑的备注'
        rows[0][columns.index('wrong_variants')] = 'confirmed wrong wording'
        store.edit_table(str(shared), rows, columns, current['fingerprint'])
        fresh = store.read()['entries'][0]
        self.assertEqual(fresh['en'], 'Updated standard wording')
        self.assertEqual(fresh['note'], '从详情编辑的备注')
        self.assertEqual(fresh['wrong_variants'], ['confirmed wrong wording'])
        with self.assertRaises(store.FileChanged):
            store.edit_table(str(shared), rows, columns, current['fingerprint'])

    def test_grid_edit_table_preserves_fields_and_rejects_stale(self):
        shared=self.root/'grid.xlsx'
        store.select(str(shared),create=True)
        current=store.read()
        rows=store.table_rows(current); columns=store.table_columns(current)
        self.assertEqual(columns,writer.HEADER_FIELDS)
        zc=columns.index('zh'); wc=columns.index('wrong_variants')
        rows[0][zc]='改过的中文'; rows[0][wc]='dot spacing'
        result=store.edit_table(str(shared),rows,columns,current['fingerprint'])
        self.assertEqual(result['count'],len(rows))
        fresh=store.read()
        got=next(e for e in fresh['entries'] if e['zh']=='改过的中文')
        self.assertEqual(got['wrong_variants'],['dot spacing'])
        with self.assertRaises(store.FileChanged):
            store.edit_table(str(shared),rows,columns,current['fingerprint'])

    def test_builtin_lexicon_is_not_editable_through_table_api(self):
        client=app.test_client()
        result=client.get('/api/glossary/table?path='+DEFAULT_PATH)
        self.assertEqual(result.status_code,403)
        result=client.post('/api/glossary/table',json={'path':DEFAULT_PATH,'rows':[]})
        self.assertEqual(result.status_code,403)

    def test_table_api_reports_conflict(self):
        client=app.test_client()
        shared=self.root/'api.xlsx'
        store.select(str(shared),create=True)
        current=store.read()
        rows=store.table_rows(current)
        result=client.post('/api/glossary/table',json={'path':str(shared),'rows':rows,
            'columns':store.table_columns(current),'fingerprint':'stale-fingerprint'})
        self.assertEqual(result.status_code,409)
        self.assertTrue(result.json.get('conflict'))

    def test_department_content_and_csv_column_detection(self):
        entries=store.seed_entries()
        self.assertTrue(any('专业术语.pptx' in e['source'] for e in entries))
        corrected=next(e for e in entries if e['zh']=='被动式驱动')
        self.assertEqual(corrected['en'],'passive matrix driving')
        self.assertIn('Passive Mactive',corrected['wrong_variants'])
        csv=self.root/'import.csv';csv.write_text('序号,名词（中文）,英文,解释\n1,点间距,pixel pitch,定义\n',encoding='utf-8-sig')
        entry=_load_any(str(csv)).entries[0]
        self.assertEqual((entry.zh,entry.en),('点间距','pixel pitch'))

    def test_report_without_suggestions_can_explicitly_save_an_unchanged_copy(self):
        source=self.document('already-correct.docx','The pixel pitch is 2.5 mm.')
        before=Path(source).read_bytes();report=self.check_job([source])
        self.assertEqual(report['findings'],[])
        data={'report_id':report['report_id'],'path':source,'selected':{},'output_dir':str(self.root/'out')}
        with self.assertRaises(ValueError): workflow.apply_report(data)
        result=workflow.apply_report({**data,'save_unchanged':True})
        self.assertEqual(result['applied_count'],0)
        self.assertNotEqual(result['output'],source)
        self.assertEqual(Path(result['output']).read_bytes(),before)
        self.assertEqual(Path(source).read_bytes(),before)

    def test_api_save_requires_fingerprint_and_model_settings_persist(self):
        client=app.test_client()
        result=client.get('/api/lexicon');self.assertEqual(result.status_code,200)
        self.assertEqual(client.post('/api/lexicon',json={'entries':result.json['entries']}).status_code,409)
        result=client.post('/api/model/settings',json={'provider':'compatible','base_url':'http://127.0.0.1:1234/v1','model':'test'})
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(client.get('/api/model/settings').json['model'],'test')


if __name__=='__main__': unittest.main()
