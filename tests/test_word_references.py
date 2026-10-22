"""Broken input references remain visible as warnings, never invented targets."""
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
sys.path[:0]=[str(Path(__file__).resolve().parents[1]),str(Path(__file__).resolve().parent)]
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from backend.word_references import inspect_file,printed_toc
from backend.replacer.pptx_package import save_office
from backend.extract.office import open_office
from backend.glossary.loader import Glossary
from backend import workflow
from test_word_coverage import field,text_run
from test_file_formats import pdf_fixture


class ReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});env.start();self.addCleanup(env.stop)

    def source(self,valid=False,simple=False,cached='7'):
        doc=Document();p=doc.add_paragraph('1. Example heading\t')
        if simple:
            node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),' PAGEREF section_one \\h ');text_run(node,cached);p._p.append(node)
        else:field(p,' PAGEREF section_one \\h ',label=cached)
        target=doc.add_heading('Example heading',1)
        if valid:
            begin=OxmlElement('w:bookmarkStart');begin.set(qn('w:name'),'section_one');begin.set(qn('w:id'),'42');target._p.append(begin)
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'42');target._p.append(end)
        # External addresses must not participate in internal bookmark repair.
        field(doc.add_paragraph(),' HYPERLINK "https://example.test/原地址" ',label='Original link')
        path=self.root/'source.docx';doc.save(path);return path

    def save(self,path):
        doc,paragraphs,_=open_office(path);out=self.root/(path.stem+'-saved.docx')
        save_office(path,doc,out,[]);return out,getattr(doc,'_star_reference_warnings',[])

    def test_missing_reference_locks_only_that_field_keeps_cache_and_preserves_source(self):
        source=self.source();before=source.read_bytes();out,warnings=self.save(source)
        self.assertEqual(len(warnings),1);self.assertIn('缓存页码未重新核实',warnings[0]['reason'])
        doc=Document(out);fields=list(doc.element.iter(qn('w:fldChar')))
        self.assertEqual(fields[0].get(qn('w:fldLock')),'true')
        self.assertIsNone(fields[-3].get(qn('w:fldLock')))
        self.assertTrue(doc.paragraphs[0].text.endswith('7'))
        self.assertEqual(source.read_bytes(),before)
        with ZipFile(source) as a,ZipFile(out) as b:
            for name in a.namelist():
                if name.endswith('.rels') or name.startswith('word/media/'):self.assertEqual(a.read(name),b.read(name))
        again,second=self.save(out);self.assertEqual(len(second),1);self.assertEqual(Document(again).paragraphs[0].text,doc.paragraphs[0].text)

    def test_valid_reference_and_all_original_bookmarks_are_unchanged(self):
        source=self.source(valid=True);changes,warnings=inspect_file(source)
        self.assertEqual(changes,{});self.assertEqual(warnings,[])
        out,warnings=self.save(source)
        with ZipFile(source) as a,ZipFile(out) as b:
            from backend.replacer.word_package import signature
            self.assertEqual(signature(a.read('word/document.xml')),signature(b.read('word/document.xml')))

    def test_simple_reference_and_baked_error_have_explicit_unknown_page(self):
        source=self.source(simple=True,cached='Error! Bookmark not defined.')
        out,warnings=self.save(source);doc=Document(out)
        self.assertEqual(len(warnings),1)
        self.assertTrue(''.join(n.text or '' for n in doc.paragraphs[0]._p.iter(qn('w:t'))).endswith('—'))
        node=next(doc.element.iter(qn('w:fldSimple')))
        self.assertEqual(node.get(qn('w:instr')),' PAGEREF section_one \\h ')
        self.assertEqual(node.get(qn('w:fldLock')),'true')

    def test_printed_error_rule_does_not_remove_prose_or_legitimate_page_numbers(self):
        for text in ['Error! Bookmark not defined.','Use Error! Bookmark not defined. as the diagnostic.','1. Explain\tthis error: Error! Bookmark not defined.','1. Good heading........7']:
            self.assertEqual(printed_toc(text),text)
        for error in ['错误！未定义书签。','Error! Bookmark not defined.','错误！未找到引用源。']:
            self.assertEqual(printed_toc('2.4 Example heading........'+error),'2.4 Example heading........—')

    def test_static_toc_error_across_runs_retains_heading_style_and_instruction(self):
        doc=Document();p=doc.add_paragraph();p.add_run('2.4 Heading').bold=True;p.add_run('\tError! Book');p.add_run('mark not defined.')
        source=self.root/'source.docx';doc.save(source);out,warnings=self.save(source)
        result=Document(out).paragraphs[0]
        self.assertEqual(result.text,'2.4 Heading\t—');self.assertTrue(result.runs[0].bold);self.assertEqual(len(warnings),1)

    def test_pdf_toc_error_is_reported_before_translation_without_changing_pdf(self):
        from backend.formats import pdf_document
        source=self.root/'source.pdf';pdf_fixture(source,'2.4 Example heading........错误！未定义书签。');before=source.read_bytes()
        doc=pdf_document(source)
        self.assertEqual(doc.paragraphs[0][1].text,'2.4 Example heading........—')
        self.assertEqual(len(doc._star_reference_warnings),1);self.assertEqual(source.read_bytes(),before)
        self.assertEqual(doc._star_reference_warnings[0]['slide'],1)

    def test_rule_only_save_without_suggestions_still_contains_broken_fields(self):
        source=self.source();before=source.read_bytes();job={'id':'refs','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,{'kind':'check','paths':[str(source)],'use_llm':False})
        result=workflow.apply_report({'report_id':'refs','path':str(source),'selected':{},'save_unchanged':True})
        self.assertEqual(result['applied_count'],0);self.assertTrue(result['content_warnings'])
        self.assertEqual(Document(result['output']).element.find('.//'+qn('w:fldChar')).get(qn('w:fldLock')),'true')
        self.assertEqual(source.read_bytes(),before)

    def test_pdf_rule_only_save_reports_source_toc_error_and_keeps_unknown_page(self):
        source=self.root/'source.pdf';pdf_fixture(source,'2.4 Example heading........Error! Bookmark not defined.');before=source.read_bytes()
        job={'id':'pdf-refs','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,{'kind':'check','paths':[str(source)],'use_llm':False})
        result=workflow.apply_report({'report_id':'pdf-refs','path':str(source),'selected':{},'save_unchanged':True})
        text='\n'.join(p.text for p in Document(result['output']).paragraphs)
        self.assertIn('Example heading',text);self.assertIn('—',text);self.assertNotIn('Bookmark',text)
        self.assertTrue(any('原稿目录' in reason for reason in result['content_warnings']))
        self.assertEqual(source.read_bytes(),before)

    def test_translation_save_reports_reference_issue_without_inventing_untranslated_body(self):
        from backend.translation import translate_file
        from test_translation_recovery import Model
        source=self.source();doc=Document(source);doc.add_paragraph('你好');doc.save(source);before=source.read_bytes()
        model=Model(['Hello.'])
        result=translate_file(source,Glossary([]),None,client=model)
        self.assertTrue(result['translation_complete']);self.assertEqual(result['untranslated'],0)
        self.assertEqual(len(model.prompts),1)
        self.assertEqual(len(result['reference_warnings']),1)
        self.assertTrue(result['layout_warnings']);self.assertEqual(source.read_bytes(),before)


if __name__=='__main__':unittest.main()
