"""Translated Word titles must fit their actual DrawingML and VML frames."""
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend.extract.office import open_office
from backend.replacer.pptx_package import save_office
from backend.replacer.word_package import signature
from backend.replacer.word_textboxes import fit_textboxes


class Metrics:
    def measure(self,text,style,scale=1):return len(text)*style[1]*scale*.52,style[1]*scale*1.2


def fixture(path):
    doc=Document();doc.styles['Normal'].paragraph_format.alignment=3
    xml='''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">
    <w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:spPr><a:xfrm><a:ext cx="5379720" cy="1879600"/></a:xfrm></wps:spPr><wps:txbx><w:txbxContent>{text}</w:txbxContent></wps:txbx><wps:bodyPr lIns="91440" rIns="91440" tIns="45720" bIns="45720"/></wps:wsp></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape style="width:423.6pt;height:148pt"><v:textbox><w:txbxContent>{text}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>'''
    text=''.join('<w:p><w:pPr><w:spacing w:before="240" w:afterLines="100"/></w:pPr><w:r><w:rPr><w:sz w:val="72"/><w:b/><w:color w:val="2F75B5"/></w:rPr><w:t>'+s+'</w:t></w:r></w:p>' for s in ['Brand','Customer Project Research',''])
    doc.element.body.insert(0,parse_xml(xml.format(text=text)))
    doc.add_paragraph('Unchanged body paragraph').runs[0].bold=True
    doc.save(path)


class WordTextBoxTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.source=self.root/'source.docx';fixture(self.source)
    def open(self):
        document,paragraphs,_=open_office(str(self.source))
        changed=[loc for loc,p,*_ in paragraphs if p.text=='Customer Project Research']
        return document,paragraphs,changed
    def test_fit_changes_typography_only_in_both_alternate_text_frames(self):
        document,paragraphs,changed=self.open()
        before=self.source.read_bytes()
        adjustments,warnings=fit_textboxes(document,changed,Metrics())
        self.assertEqual(len(adjustments),2);self.assertFalse(warnings)
        for loc,para,*_ in paragraphs:
            if loc in changed:
                self.assertEqual(para._p.find(qn('w:pPr')+'/'+qn('w:jc')).get(qn('w:val')),'left')
                self.assertTrue(para.runs[0].bold)
                self.assertEqual(para.runs[0].font.color.rgb.__str__(),'2F75B5')
                self.assertGreaterEqual(para.runs[0].font.size.pt,28)
                self.assertEqual(para.text,'Customer Project Research')
        output=self.root/'output.docx';save_office(self.source,document,output,changed)
        with ZipFile(self.source) as a,ZipFile(output) as b:
            self.assertEqual(signature(a.read('word/document.xml')),signature(b.read('word/document.xml')))
            for name in a.namelist():
                if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
            for raw in [a.read('word/document.xml'),b.read('word/document.xml')]:
                texts=etree.fromstring(raw).xpath('//w:t/text()',namespaces={'w':qn('w:t').split('}')[0][1:]})
                if raw==a.read('word/document.xml'):original=texts
                else:self.assertEqual(texts,original)
        self.assertEqual(self.source.read_bytes(),before)
    def test_no_selected_text_keeps_all_original_typography(self):
        document,_,_=self.open();before=etree.tostring(document.element)
        self.assertEqual(fit_textboxes(document,[],Metrics()),([],[]))
        self.assertEqual(etree.tostring(document.element),before)
    def test_regular_body_text_never_enters_textbox_fitting(self):
        document,paragraphs,_=self.open();before=etree.tostring(document.element)
        changed=[loc for loc,p,*_ in paragraphs if p.text=='Unchanged body paragraph']
        self.assertEqual(fit_textboxes(document,changed,Metrics()),([],[]))
        self.assertEqual(etree.tostring(document.element),before)
    def test_unfittable_text_is_preserved_and_warned_instead_of_shrunk_to_tiny_text(self):
        document,paragraphs,changed=self.open()
        for loc,para,*_ in paragraphs:
            if loc in changed:para.runs[0].text='LongUnbreakableIdentifier'*20
        adjustments,warnings=fit_textboxes(document,changed,Metrics())
        self.assertEqual(len(warnings),2)
        for loc,para,*_ in paragraphs:
            if loc in changed:
                self.assertEqual(para.text,'LongUnbreakableIdentifier'*20)
                self.assertGreaterEqual(para.runs[0].font.size.pt,28)
    def test_invalid_measurement_does_not_prevent_saving_the_valid_document(self):
        document,paragraphs,changed=self.open()
        output=self.root/'valid-contents.docx'
        with patch('backend.replacer.word_textboxes.fit_textboxes',side_effect=ValueError('unmeasurable font metadata')):
            save_office(self.source,document,output,changed)
        self.assertTrue(output.is_file())
        self.assertTrue(document._star_word_layout_result[1])
        self.assertIn('Customer Project Research',''.join(Document(output).element.itertext()))


if __name__=='__main__':unittest.main()
