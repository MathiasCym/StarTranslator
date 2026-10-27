"""English-to-Chinese format writers and English-only policy isolation."""
import asyncio
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt
from openpyxl import Workbook, load_workbook
from backend import workflow, task_queue, comparison
from backend.translation import translate_file
from backend.reverse_translation import audit_saved
from backend.extract import extract
from backend.extract.office import open_office
from backend.glossary.loader import Glossary
from backend.glossary import store
from backend.llm.client import LLMClient

SOURCE = 'The pixel pitch is 2.5 mm.'
TARGET = '点间距为 2.5 mm。'


class Model(LLMClient):
    def __init__(self, suggestions=False):
        self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='reverse-fixture';self.api_format='openai'
        self.calls=[];self.suggestions=suggestions
    def chat_json(self, prompt, system='', **options):
        self.calls.append((prompt,system))
        data=json.loads(prompt)
        if '复核' in system and '"ok"' in system:return {'ok':True,'issues':[]}
        if '审校员' in system:
            return {'reviewed':True,'suggestions':[{'original':data['中文译文'],'replacement':'像素间距为 2.5 mm。','reason':'统一中文术语'}] if self.suggestions and data['中文译文']==TARGET else []}
        source=data['待译原文']
        target=TARGET if source==SOURCE else '屏幕支持 60 Hz。' if source=='The display supports 60 Hz.' else source
        result={'translation':target}
        if data.get('原文样式片段'):
            groups=data['原文样式片段']
            if len(groups)==2 and groups[0]['text']=='The pixel pitch':
                result['runs']=[{'id':groups[0]['id'],'text':'点间距'},{'id':groups[1]['id'],'text':'为 2.5 mm。'}]
            else:result['runs']=[{'id':g['id'],'text':target if i==0 else ''} for i,g in enumerate(groups)]
        return result


class ReverseTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.addCleanup(self.tmp.cleanup)
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});env.start();self.addCleanup(env.stop)
        self.glossary=Glossary([])
        stub=patch('backend.workflow.store.load_glossary',return_value=self.glossary);stub.start();self.addCleanup(stub.stop)
        self.previous=dict(workflow.JOBS);self.reports=dict(workflow.REPORTS)
        workflow.JOBS.clear();workflow.REPORTS.clear();self.addCleanup(self.restore)
    def restore(self):
        workflow.JOBS.clear();workflow.JOBS.update(self.previous);workflow.REPORTS.clear();workflow.REPORTS.update(self.reports)
    def doc(self):
        file=self.root/'source.docx';doc=Document();p=doc.add_paragraph()
        run=p.add_run('The pixel pitch');run.bold=True;run.font.color.rgb=RGBColor(255,0,0)
        p.add_run(' is 2.5 mm.')
        table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='The display supports 60 Hz.'
        h=doc.add_paragraph();link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'keep')
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='https://example.com/中文';r.append(t);link.append(r);h._p.append(link)
        doc.save(file);return file
    def test_word_styles_tables_links_and_direction_specific_cache(self):
        source=self.doc();before=store.fingerprint(source);model=Model()
        result=translate_file(source,self.glossary,None,client=model,direction='en-zh')
        self.assertTrue(result['translation_complete']);self.assertEqual(result['output_audit']['status'],'complete')
        doc=Document(result['path']);self.assertEqual(doc.paragraphs[0].text,TARGET)
        self.assertTrue(doc.paragraphs[0].runs[0].bold);self.assertEqual(str(doc.paragraphs[0].runs[0].font.color.rgb),'FF0000')
        self.assertEqual(doc.tables[0].cell(0,0).text,'屏幕支持 60 Hz。')
        self.assertEqual(store.fingerprint(source),before)
        record=comparison.read(result['comparison_id']);self.assertEqual(record['direction'],'en-zh')
        second=Model();again=translate_file(source,self.glossary,None,client=second,direction='en-zh')
        self.assertEqual(again['resumed_segments'],2);self.assertEqual(second.calls,[])
        with patch('backend.translation.mapped_translation',side_effect=AssertionError('English policy must not run')):
            translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
    def test_ppt_preserves_color_bold_and_passes_chinese_target_to_visual_review(self):
        source=self.root/'source.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2));p=shape.text_frame.paragraphs[0]
        a=p.add_run();a.text='The pixel pitch';a.font.bold=True;a.font.size=Pt(24)
        b=p.add_run();b.text=' is 2.5 mm.';b.font.size=Pt(24);deck.save(source)
        with (patch('backend.pptx_visual_review.review_pptx',return_value=([],[])) as visual,
                patch('backend.officecli_bridge.inspect',return_value={}),patch('backend.officecli_bridge.compare',return_value=[])):
            result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
        self.assertEqual(visual.call_args.kwargs['target_language'],'zh')
        output=Presentation(result['path']);self.assertEqual(output.slides[0].shapes[0].text,TARGET)
        self.assertTrue(output.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold)
    def test_all_editable_text_containers(self):
        fixtures={'.txt':SOURCE+'\r\n', '.md':'# '+SOURCE+'\n\n[Link](https://example.com/中文)',
            '.html':'<p>'+SOURCE+'</p><a href="https://example.com/中文">NOVA</a>',
            '.json':json.dumps({'keep-key':SOURCE,'url':'https://example.com/中文'}),
            '.srt':'1\n00:00:01,000 --> 00:00:03,000\n'+SOURCE+'\n',
            '.ass':'[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: 0,0:00:01.00,0:00:03.00,Default,,0,0,0,,'+SOURCE+'\n'}
        for extension,content in fixtures.items():
            with self.subTest(extension=extension):
                source=self.root/('source'+extension);source.write_text(content,encoding='utf-8')
                result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
                self.assertIn(TARGET,'\n'.join(s.text for s in extract(result['path'])))
                self.assertTrue(result['translation_complete'])
                self.assertEqual(result['output_audit']['status'],'complete')
                if extension in ('.md','.html'):self.assertIn('https://example.com/中文',Path(result['path']).read_text(encoding='utf-8'))
                if extension=='.json':self.assertEqual(json.loads(Path(result['path']).read_text(encoding='utf-8'))['url'],'https://example.com/中文')
        source=self.root/'source.xlsx';book=Workbook();book.active['A1']=SOURCE;book.active['B1']='=1+2';book.save(source)
        result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
        out=load_workbook(result['path']);self.assertEqual(out.active['A1'].value,TARGET);self.assertEqual(out.active['B1'].value,'=1+2');out.close()
        source=self.root/'source.epub'
        with ZipFile(source,'w') as package:
            package.writestr('mimetype','application/epub+zip')
            package.writestr('META-INF/container.xml','<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/content.opf"/></rootfiles></container>')
            package.writestr('OPS/content.opf','<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="page" href="page.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="page"/></spine></package>')
            package.writestr('OPS/page.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><body><p>'+SOURCE+'</p></body></html>')
        result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
        self.assertIn(TARGET,'\n'.join(s.text for s in extract(result['path'])))
    def test_pdf_text_layer_uses_chinese_output_without_english_residual_audit(self):
        from test_file_formats import pdf_fixture
        source=self.root/'source.pdf';pdf_fixture(source,SOURCE)
        with patch('backend.pdf_quality.audit',side_effect=AssertionError('English-only PDF audit must not run')):
            result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
        self.assertEqual(Path(result['path']).suffix,'.docx')
        self.assertIn(TARGET,'\n'.join(s.text for s in extract(result['path'])))
        self.assertTrue(any('图片内英文' in w['reason'] for w in result['warnings']))
    def test_combined_review_save_unchanged_and_apply_chinese_suggestion(self):
        source=self.doc();model=Model(suggestions=True)
        with patch('backend.workflow.LLMClient',return_value=model),patch('backend.translation.LLMClient',return_value=model):
            job=workflow.start({'kind':'translate_review','direction':'en-zh','paths':[str(source)],'combined_output':'final'})
            final=asyncio.run(task_queue.wait(job['id'],.005))
        self.assertEqual(final['status'],'complete',final.get('error'))
        report=final['result']['report'];self.assertTrue(report['findings'])
        path=report['files'][0]['path'];data={'report_id':job['id'],'path':path,'save_unchanged':True}
        with patch('backend.translation_audit.saved_residuals',side_effect=AssertionError('Chinese is intended')):
            saved=workflow.apply_report(data)
        self.assertIn('Chinese_checked',saved['output']);self.assertEqual(saved['untranslated'],0)
        applied=workflow.apply_report({**data,'edits':[{'index':0,'text':'像素间距为 2.5 mm。'}]})
        self.assertEqual(Document(applied['output']).paragraphs[0].text,'像素间距为 2.5 mm。')
        self.assertEqual(comparison.read(applied['comparison_id'])['output_path'],applied['output'])
        options=__import__('backend.history',fromlist=['resume_options']).resume_options(job['id'])
        self.assertEqual(options['direction'],'en-zh')
    def test_chinese_proofreading_does_not_route_back_to_english(self):
        source=self.root/'chinese.docx';doc=Document();doc.add_paragraph(TARGET);doc.save(source)
        with (patch('backend.workflow.LLMClient',return_value=Model()),
              patch('backend.workflow.predominantly_chinese',side_effect=AssertionError('Do not translate intended Chinese back to English'))):
            job=workflow.start({'kind':'check','direction':'en-zh','paths':[str(source)],'use_llm':True})
            result=asyncio.run(task_queue.wait(job['id'],.005))
        self.assertEqual(result['status'],'complete');self.assertEqual(result['kind'],'check')
        self.assertEqual(result['result']['stats']['target_language'],'zh')
        self.assertEqual(Document(source).paragraphs[0].text,TARGET)

    def test_save_preserves_pending_reverse_translation_until_it_is_resolved(self):
        source=self.root/'partial.docx';doc=Document();doc.add_paragraph(SOURCE)
        doc.add_paragraph('The warranty applies only to indoor displays.');doc.save(source)
        model=Model()
        with patch('backend.workflow.LLMClient',return_value=model):
            job=workflow.start({'kind':'translate_review','direction':'en-zh','paths':[str(source)],'combined_output':'final'})
            result=asyncio.run(task_queue.wait(job['id'],.005))
        self.assertEqual(result['status'],'partial');self.assertTrue(result['can_resume'])
        output=result['result']['outputs'][0]
        self.assertEqual(output['untranslated'],1)
        self.assertEqual(output['output_audit']['status'],'partial')
        data={'report_id':job['id'],'path':output['path'],'save_unchanged':True}
        first=workflow.apply_report(data)
        self.assertEqual(first['untranslated'],1)
        self.assertGreaterEqual(first['needs_review'],1)
        self.assertTrue(any('翻译待核对' in warning for warning in first['content_warnings']))
        self.assertEqual(__import__('backend.history',fromlist=['record']).record(job['id'])['status'],'partial')
        report=result['result']['report'];index=next(i for i,f in enumerate(report['findings']) if f['reason'].startswith('英译中未完成：'))
        fixed=workflow.apply_report({**data,'edits':[{'index':index,'text':'保修仅适用于室内显示屏。'}]})
        self.assertFalse(fixed['failed']);self.assertEqual(fixed['untranslated'],0)
        self.assertEqual(fixed['output_audit']['status'],'complete')
        recorded=__import__('backend.history',fromlist=['record']).record(job['id'])
        self.assertEqual(recorded['status'],'complete');self.assertFalse(recorded['can_resume'])
        again=workflow.apply_report(data)
        self.assertEqual(again['untranslated'],0)
        self.assertIn('保修仅适用于室内显示屏。',Document(again['output']).paragraphs[1].text)

    def test_chinese_review_accepts_only_uniquely_locatable_substring(self):
        from backend.reverse_translation import check_segments
        class Local(Model):
            def chat_json(self,prompt,system='',**options):
                self.calls.append((prompt,system))
                return {'reviewed':True,'suggestions':[{'original':'点间距','replacement':'像素间距','reason':'统一术语'}]}
        source=self.root/'target.docx';doc=Document();doc.add_paragraph(TARGET);doc.save(source)
        findings,stats=check_segments(extract(source),self.glossary,client=Local())
        self.assertEqual(stats['model_status'],'complete');self.assertEqual(len(findings),1)
        self.assertEqual(findings[0].found_text,'点间距');self.assertEqual(findings[0].suggested,'像素间距')
        doc=Document();doc.add_paragraph('点间距与点间距必须一致。');doc.save(source)
        findings,stats=check_segments(extract(source),self.glossary,client=Local())
        self.assertEqual(findings,[]);self.assertEqual(stats['model_status'],'partial')

    def test_review_retries_bad_location_without_hiding_or_repeating_transport_failure(self):
        from backend.reverse_translation import check_segments
        class Repair(Model):
            def chat_json(self,prompt,system='',**options):
                self.calls.append((prompt,system))
                return {'reviewed':True,'suggestions':[{'original':'不存在' if len(self.calls)==1 else TARGET,'replacement':'像素间距为 2.5 mm。','reason':'统一术语'}]}
        source=self.root/'target.docx';doc=Document();doc.add_paragraph(TARGET);doc.save(source);model=Repair()
        findings,stats=check_segments(extract(source),self.glossary,client=model)
        self.assertEqual(len(model.calls),2);self.assertEqual(stats['model_status'],'complete');self.assertEqual(len(findings),1)

    def test_combined_review_checkpoints_survive_new_draft_paths(self):
        source=self.doc()
        for iteration in range(2):
            model=Model()
            with patch('backend.workflow.LLMClient',return_value=model):
                job=workflow.start({'kind':'translate_review','direction':'en-zh','paths':[str(source)],'combined_output':'final'})
                result=asyncio.run(task_queue.wait(job['id'],.005))
            self.assertEqual(result['status'],'complete',result.get('error'))
            if iteration:
                self.assertEqual(model.calls,[])
                self.assertEqual(result['result']['stats']['resumed_segments'],2)
                self.assertEqual(result['result']['outputs'][0]['resumed_segments'],2)

    def test_pdf_recovers_only_whitespace_before_reverse_numeric_validation(self):
        from test_file_formats import pdf_fixture
        from backend.reverse_translation import pdf_source_text
        class PdfModel(Model):
            def chat_json(self,prompt,system='',**options):
                data=json.loads(prompt)
                if 'PDF英文文字层' in data:
                    self.calls.append((prompt,system));return {'text':SOURCE}
                return super().chat_json(prompt,system,**options)
        source=self.root/'joined.pdf';pdf_fixture(source,'The pixel pitch is2.5 mm.');model=PdfModel()
        result=translate_file(source,self.glossary,None,client=model,direction='en-zh')
        self.assertTrue(result['translation_complete'],result['review_items'])
        self.assertIn(TARGET,'\n'.join(s.text for s in extract(result['path'])))
        again=PdfModel();result=translate_file(source,self.glossary,None,client=again,direction='en-zh')
        self.assertEqual(result['resumed_segments'],1);self.assertEqual(again.calls,[])
        class Unsafe(PdfModel):
            def chat_json(self,*a,**kw):return {'text':'The pixel pitch is 3.5 mm.'}
        with self.assertRaisesRegex(ValueError,'空格未能安全'):
            pdf_source_text('The pixel pitch is2.5 mm.',Unsafe(),threading.Event())

    def test_inline_markup_is_translated_as_one_sentence_in_each_container(self):
        english='Do not disconnect the power supply while the LED display is updating.'
        class Inline(Model):
            def chat_json(self,prompt,system='',**options):
                data=json.loads(prompt)
                if data.get('待译原文')==english:
                    self.calls.append((prompt,system));groups=data['原文样式片段']
                    return {'translation':'请勿在LED显示屏更新期间断开电源。','runs':[{'id':groups[0]['id'],'text':'请勿'},{'id':groups[1]['id'],'text':'在LED显示屏更新期间断开电源。'}]}
                return super().chat_json(prompt,system,**options)
        tail=english[len('Do not'):]
        fixtures={'.html':'<p><strong>Do not</strong>'+tail+'</p>',
            '.md':'**Do not**'+tail+'\n',
            '.srt':'1\n00:00:01,000 --> 00:00:03,000\n<i>Do not</i>'+tail+'\n',
            '.ass':'[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: 0,0:00:01.00,0:00:03.00,Default,,0,0,0,,{\\b1}Do not{\\b0}'+tail+'\n'}
        for ext,text in fixtures.items():
            with self.subTest(ext=ext):
                source=self.root/('inline'+ext);source.write_text(text,'utf-8');model=Inline()
                result=translate_file(source,self.glossary,None,client=model,direction='en-zh')
                self.assertTrue(result['translation_complete'],result['review_items'])
                actual=''.join(s.text for s in extract(result['path']) if s.editable)
                self.assertEqual(actual,'请勿在LED显示屏更新期间断开电源。')
                self.assertEqual(result['translated'],result['total'])
                second=Inline();again=translate_file(source,self.glossary,None,client=second,direction='en-zh')
                self.assertEqual(second.calls,[]);self.assertEqual(again['resumed_segments'],again['total'])

    def test_semantic_emphasis_repair_keeps_fixed_translation_and_negative_scope(self):
        english='Do not disconnect the power supply while the LED display is updating.'
        chinese='在LED显示屏更新时，请勿断开电源。'
        class Emphasis(Model):
            def chat_json(self,prompt,system='',**options):
                self.calls.append((prompt,system));data=json.loads(prompt)
                if '译文样式片段' in data:
                    return {'ok':True,'issues':[]}  # An optimistic model cannot bypass the deterministic negation anchor.
                if '已确定中文' in data:return {'anchors':[{'id':0,'text':'请勿'}]}
                if '原文' in data:return {'ok':True,'issues':[]}
                return {'translation':chinese,'runs':[{'id':0,'text':'在LED显示屏更新时，请勿'},{'id':1,'text':'断开电源。'}]}
        source=self.root/'warning.docx';doc=Document();p=doc.add_paragraph();r=p.add_run('Do not');r.bold=True;r.font.color.rgb=RGBColor(255,0,0);p.add_run(english[6:]);doc.save(source)
        model=Emphasis();result=translate_file(source,self.glossary,None,client=model,direction='en-zh')
        self.assertTrue(result['translation_complete'],result['review_items'])
        paragraph=Document(result['path']).paragraphs[0]
        self.assertEqual(paragraph.text,chinese)
        self.assertEqual(''.join(r.text for r in paragraph.runs if r.bold),'请勿')
        self.assertEqual(sum('待译原文' in prompt for prompt,_ in model.calls),1)

    def test_unknown_person_names_are_not_guessed_in_translation_or_review(self):
        class Names(Model):
            translations=0;reviews=0
            def chat_json(self,prompt,system='',**options):
                data=json.loads(prompt.split('\n上次检查：')[0]);self.calls.append((prompt,system))
                if '候选及英文原文语境' in data:return {'people':['Chen Yiming']}
                if '中文译文' in data:
                    self.reviews+=1
                    return {'reviewed':True,'suggestions':[{'original':'Chen Yiming','replacement':'陈一鸣','reason':'翻译人名'}] if self.reviews==1 else []}
                if '待译原文' in data:
                    self.translations+=1
                    text='陈一鸣批准了报告。' if self.translations==1 else 'Chen Yiming批准了报告。'
                    return {'translation':text,'runs':[{'id':0,'text':text}]}
                return {'ok':True,'issues':[]}
        source=self.root/'person.docx';doc=Document();doc.add_paragraph('Chen Yiming approves the report.');doc.save(source)
        model=Names()
        with patch('backend.workflow.LLMClient',return_value=model):
            job=workflow.start({'kind':'translate_review','direction':'en-zh','paths':[str(source)],'combined_output':'final'})
            result=asyncio.run(task_queue.wait(job['id'],.005))
        self.assertEqual(result['status'],'complete',result.get('error'))
        self.assertEqual(model.translations,2);self.assertEqual(model.reviews,2)
        self.assertEqual(result['result']['report']['findings'],[])
        output=result['result']['outputs'][0]
        self.assertIn('Chen Yiming',Document(output['path']).paragraphs[0].text)

    def test_pure_ppt_layout_changes_do_not_inflate_translation_count(self):
        source=self.root/'partial.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(3));box.text_frame.paragraphs[0].text=SOURCE
        box.text_frame.add_paragraph().text='The warranty applies only to indoor displays.';deck.save(source)
        def visual(src,out,document,changed,*args,**options):
            changed.append('slide1/shape0/p1');return [],[]
        with patch('backend.pptx_visual_review.review_pptx',side_effect=visual),patch('backend.officecli_bridge.inspect',return_value={}),patch('backend.officecli_bridge.compare',return_value=[]):
            result=translate_file(source,self.glossary,None,client=Model(),direction='en-zh')
        self.assertEqual(result['translated'],1);self.assertEqual(result['untranslated'],1);self.assertEqual(result['total'],2)

    def test_pdf_logical_paragraphs_keep_headings_without_sentence_tail_fragments(self):
        from backend.reverse_translation import pdf_source_text
        source='Verification and H2 Strategy\nH2 strategy: improve the 10-bit grayscale\naccuracy.\nThe refresh rate is 60 Hz.'
        class Logical(Model):
            def chat_json(self,*args,**options):return {'paragraphs':source.splitlines()}
        target=pdf_source_text(source,Logical(),threading.Event())
        self.assertEqual(target,'Verification and H2 Strategy\nH2 strategy: improve the 10-bit grayscale accuracy.\nThe refresh rate is 60 Hz.')

    def test_person_substring_and_wrapped_name_are_not_mistaken_for_a_title(self):
        from types import SimpleNamespace
        from backend.reverse_entities import identify,validate_names
        class Person(Model):
            def chat_json(self,*args,**options):return {'people':['Chen Yiming']}
        terms=identify([SimpleNamespace(text='Chen Yiming\nTechnical Report',editable=True)],self.glossary,Person(),threading.Event())
        self.assertEqual(terms,{'Chen Yiming':'Chen Yiming'})
        validate_names('Chen\nYiming approves this.','Chen Yiming批准此事。',terms)
        with self.assertRaises(ValueError):validate_names('Chen\nYiming approves this.','陈一鸣批准此事。',terms)


if __name__=='__main__':unittest.main()
