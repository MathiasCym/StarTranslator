import unittest,sys,os,tempfile,threading
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend import bit_depth
from backend.models import Segment
from backend.glossary.loader import Glossary
from backend.glossary.bit_depth import merge
from backend.writing_policy import validate_translation
from backend.translation import translate_file
from docx import Document

class BitDepthTests(unittest.TestCase):
    def table(self):
        return [Segment('0','x','Color Depth','body/0/table/0/0/0',source='table'),Segment('1','x','8bit / 10bit / 12-bit','body/0/table/1/0/0',source='table'),Segment('2','x','Packet Size','body/0/table/0/1/0',source='table'),Segment('3','x','8bit','body/0/table/1/1/0',source='table')]
    def test_table_depth_column_only_and_manual_findings(self):
        segs=self.table();hits=bit_depth.check(segs);self.assertEqual([(h.found_text,h.suggested) for h in hits],[('8bit','8-bit'),('10bit','10-bit')])
        self.assertTrue(all(not h.applied for h in hits))
    def test_counts_rates_codes_and_correct_forms_are_preserved(self):
        for text in ['8 bits','8 bit/s','code_8bit','MODEL8BIT','8-bit input','100bit input']:
            self.assertEqual(bit_depth.replacements(text,Segment('0','x',text,'p'),[]),[])
        self.assertEqual(bit_depth.replacements('8bit input',Segment('0','x','8bit input','p'),[]),[(0,4,'8-bit')])
        validate_translation('8BIT输入','8-bit input')
        mixed=Segment('0','x','8bit input, with a packet of 8bit','p')
        self.assertEqual([(f.found_text,f.suggested) for f in bit_depth.check([mixed])],[('8bit input','8-bit input')])
    def test_glossary_adds_convention_without_overwriting_custom_terms(self):
        values=[{'zh':'位深','en':'bit','variants':[]},{'zh':'12bit','en':'custom label','note':'team note'}]
        merged,_=merge(values);self.assertEqual(next(e for e in merged if e['zh']=='位深')['en'],'bit depth')
        self.assertEqual(next(e for e in merged if e['zh']=='12bit')['en'],'custom label')
        self.assertEqual(next(e for e in merged if e['zh']=='8bit')['en'],'8-bit')
        self.assertEqual(merge(merged)[1],[])
    def test_translation_normalizes_preexisting_english_depth_cells_and_preserves_runs(self):
        with tempfile.TemporaryDirectory() as temp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':temp}):
            source=Path(temp)/'source.docx';d=Document();d.add_paragraph('位深参数');t=d.add_table(2,2)
            t.cell(0,0).text='Color Depth';t.cell(0,1).text='Packet Size'
            p=t.cell(1,0).paragraphs[0];p.add_run('8').bold=True;p.add_run('bit / 10bit');t.cell(1,1).text='8bit';d.save(source);before=source.read_bytes()
            class Model:
                def chat_json(self,*args,**kw): return {'translation':'Bit depth parameters'}
            result=translate_file(str(source),Glossary([]),None,client=Model())
            output=Document(result['path']);self.assertEqual(output.tables[0].cell(1,0).text,'8-bit / 10-bit')
            self.assertTrue(output.tables[0].cell(1,0).paragraphs[0].runs[0].bold)
            self.assertEqual(output.tables[0].cell(1,1).text,'8bit');self.assertEqual(before,source.read_bytes())

if __name__=='__main__':unittest.main()
