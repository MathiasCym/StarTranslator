"""Translation quality problems are repaired locally or disclosed in a draft."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from backend.translation import translate_file
from backend.writing_policy import validate_translation, repair_numeric_translation
from backend.glossary.loader import Glossary, _load_any
from backend.extract import extract
from backend.references import ReferenceIndex
from backend import workflow
from backend.llm.client import LLMResponseInvalid, LLMUnavailable


class Model:
    def __init__(self, answers):
        self.answers, self.prompts = iter(answers), []
    def chat_json(self, prompt, **kwargs):
        self.prompts.append(prompt)
        answer = next(self.answers)
        if isinstance(answer, Exception):
            raise answer
        return {'translation': answer}


class RecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.environment = patch.dict(os.environ, {'NOVATERM_DATA_DIR': str(self.root/'profile')})
        self.environment.start()
    def tearDown(self):
        self.environment.stop(); self.temp.cleanup()
    def document(self, *paragraphs):
        path = self.root/'source.docx'; doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(path)
        return path
    def translate(self, source, model):
        return translate_file(str(source), Glossary([]), None, client=model)

    def test_activity_reports_actual_requests_and_never_claims_cache_hits_are_model_work(self):
        from backend.llm.client import LLMClient
        class ResumeModel(Model,LLMClient):
            def __init__(self,answers):
                Model.__init__(self,answers)
                self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='activity-fixture';self.api_format='openai'
        source=self.document('你好','谢谢');events=[]
        model=ResumeModel(['Hello.','Thank you.'])
        first=translate_file(source,Glossary([]),None,client=model,workers=3,
            activity=lambda *v:events.append(v))
        self.assertTrue(first['translation_complete'])
        active=[e for e in events if e[0]=='translate']
        self.assertTrue(any('等待模型响应' in e[1] for e in active))
        self.assertTrue(any('校验模型返回内容' in e[1] for e in active))
        self.assertEqual([e[2] for e in active],sorted(e[2] for e in active))
        self.assertEqual(active[-1][2:],(2,2))
        self.assertFalse(any('模型正在翻译第' in e[1] for e in events))
        events.clear();never=ResumeModel([])
        second=translate_file(source,Glossary([]),None,client=never,workers=1,activity=lambda *v:events.append(v))
        self.assertTrue(second['translation_complete']);self.assertGreater(second['resumed_segments'],0)
        self.assertEqual(never.prompts,[])
        self.assertFalse(any('等待模型响应' in e[1] for e in events))

    def test_cinema_quantities_decimal_formats_and_chinese_units_are_equivalent(self):
        for source, text in [('21.4万块，市场潜力巨大', '214,000 screens offer significant market potential.'),
                             ('8.6万块，规模持续增长', '86,000 screens, with continued growth.'),
                             ('投资2亿元', 'An investment of 200 million yuan.'),
                             ('亮度500.0nits，点间距０２．５０毫米', '500 nits with a pixel pitch of 2.5 mm.')]:
            validate_translation(source, text)
        with self.assertRaises(ValueError):
            validate_translation('21.4万块', '21,400 screens')

    def test_one_brightness_value_is_corrected_without_another_model_call(self):
        source = self.document('亮度为500nits。')
        model = Model(['The brightness is 1000 nits.'])
        result = self.translate(source, model)
        self.assertEqual(Document(result['path']).paragraphs[0].text, 'The brightness is 500 nits.')
        self.assertEqual(len(model.prompts), 1)
        self.assertEqual(result['needs_review'], 0)

    def test_ambiguous_units_and_model_digits_are_not_blindly_replaced(self):
        for source, text in [('亮度500nits', '5000 pixels'), ('VX1000输出500nits', 'VX2000 outputs 500 nits'),
                             ('分辨率3840×2160', '1920 x 1080')]:
            self.assertEqual(repair_numeric_translation(source, text), text)

    def test_local_model_repair_preserves_good_content_then_continues(self):
        source = self.document('支持HDR，分辨率为3840×2160。', '亮度500nits。')
        model = Model(['Supports HDR at 1920 x 1080.', 'Supports HDR at 3840 x 2160.', '500 nits.'])
        result = self.translate(source, model)
        self.assertEqual(result['needs_review'], 0)
        self.assertEqual(result['translated'], 2)
        self.assertIn('仅修正有差异的位置', model.prompts[1])

    def test_unresolved_first_paragraph_keeps_draft_and_translates_later_paragraph(self):
        source = self.document('分辨率为3840×2160。', '亮度500nits。'); original = source.read_bytes()
        model = Model(['1920 x 1080.', '1920 x 1080.', '500 nits.', '1920 x 1080.', '1920 x 1080.'])
        result = self.translate(source, model)
        self.assertEqual(result['status'], 'needs_review')
        self.assertEqual(result['needs_review'], 1)
        self.assertEqual([p.text for p in Document(result['path']).paragraphs], ['分辨率为3840×2160。', '500 nits.'])
        self.assertFalse(result['translation_complete'])
        self.assertEqual(result['review_items'][0]['source_text'], '分辨率为3840×2160。')
        self.assertEqual(source.read_bytes(), original)

    def test_invalid_json_preserves_one_original_paragraph_without_dropping_other_translations(self):
        source = self.document('亮度500nits。', '刷新率60Hz。', '点间距2.5毫米。')
        model = Model(['500 nits.', LLMResponseInvalid('JSON 无效'), LLMResponseInvalid('JSON 无效'), '2.5 mm.', LLMResponseInvalid('JSON 无效'), LLMResponseInvalid('JSON 无效')])
        result = self.translate(source, model)
        self.assertEqual(result['translated'], 2)
        self.assertEqual(result['untranslated'], 1)
        self.assertTrue(result['review_items'][0]['retained_source'])
        self.assertEqual(Document(result['path']).paragraphs[1].text, '刷新率60Hz。')

    def test_service_failure_saves_completed_content_and_discloses_remaining_original(self):
        source = self.document('亮度500nits。', '刷新率60Hz。', '点间距2.5毫米。', '对比度1000:1。')
        model = Model(['500 nits.', LLMUnavailable('连接超时'), LLMUnavailable('连接超时')])
        result = self.translate(source, model)
        self.assertEqual(result['translated'], 1)
        self.assertEqual(result['untranslated'], 3)
        self.assertEqual(result['needs_review'], 3)
        self.assertEqual(len(model.prompts), 3)

    def test_owner_files_are_rejected_before_reading_or_model_calls(self):
        locked = self.root/'~$原稿.pptx'  # Deliberately nonexistent: classify before filesystem access.
        model = Model([])
        for call in [lambda: extract(str(locked)), lambda: self.translate(locked, model),
                     lambda: workflow.start({'kind':'translate','paths':[str(locked)]}),
                     lambda: ReferenceIndex([str(locked)], Glossary([])),
                     lambda: _load_any(str(self.root/'~$术语.xlsx'))]:
            with self.assertRaisesRegex(ValueError, '临时锁文件'):
                call()
        self.assertEqual(model.prompts, [])

    def test_original_location_is_first_reference_even_after_first_four_textboxes(self):
        source = self.document(*['中文原稿' + str(i) for i in range(7)])
        target = self.root/'English.docx'; doc = Document(source); doc.paragraphs[6].text = 'English paragraph 6'; doc.save(target)
        refs = ReferenceIndex([str(source)], Glossary([]), {str(target):str(source)})
        context = refs.context(extract(str(target))[-1], 6)
        self.assertIn('对应原稿位置', context.split('\n')[0])
        self.assertIn('中文原稿6', context.split('\n')[1])


if __name__ == '__main__':
    unittest.main()
