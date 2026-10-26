import json,os,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend.glossary import store,writer,exchange
from backend.server import app
from openpyxl import load_workbook,Workbook


def entry(zh='逐点校正',en='pixel-level calibration',**kwargs):
    return store.validate([{'zh':zh,'en':en,**kwargs}])[0]


class ExchangeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.profile=self.root/'profile'
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.profile)});self.env.start()
        self.active=self.root/'active.json';store.atomic_json(self.active,[])
        store.atomic_json(self.profile/'lexicon-settings.json',{'path':str(self.active)})
        self.client=app.test_client()
    def tearDown(self):self.env.stop();self.temp.cleanup()
    def current(self,entries):store.atomic_json(self.active,entries);return store.read()
    def source(self,entries,name='team.json'):
        path=self.root/name;store.atomic_json(path,entries);return str(path)
    def test_json_and_excel_roundtrip_every_field_and_literal_text(self):
        original=[entry(variants=['brightness, static','全称\n解释'],machine_variants=['point calibration'],wrong_variants=['point proofreading'],
                        priority=1,note='两行定义\n保留说明',source='人工校对'),entry('测试文字','=literal',note='=not a formula',priority=3)]
        current=self.current(original);before=self.active.read_bytes()
        for ext in ('json','xlsx'):
            target=self.root/f'export.{ext}';result=exchange.export_file(str(target),'')
            self.assertEqual(result['count'],2)
            self.assertEqual(exchange.load_import(target),original)
            self.assertEqual(self.active.read_bytes(),before)
            if ext=='xlsx':
                book=load_workbook(target,data_only=False);sheet=book.active
                self.assertEqual(sheet['B3'].data_type,'s');self.assertEqual(sheet['B3'].value,'=literal')
                self.assertTrue(sheet.column_dimensions['G'].hidden);self.assertTrue(sheet.column_dimensions['H'].hidden)
                self.assertEqual(sheet.freeze_panes,'C2');book.close()
    def test_export_includes_unsaved_rows_without_changing_the_active_file(self):
        current=self.current([entry()]);before=self.active.read_bytes()
        rows=store.table_rows(current);columns=store.table_columns(current);rows[0][columns.index('note')]='尚未保存的修改'
        result=self.client.post('/api/lexicon/export',json={'path':str(self.root/'draft.json'),'destination_fingerprint':'','rows':rows,'columns':columns})
        self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(exchange.load_import(self.root/'draft.json')[0]['note'],'尚未保存的修改')
        self.assertEqual(self.active.read_bytes(),before)
    def test_editor_export_does_not_split_a_single_alias_containing_commas(self):
        current=self.current([entry(variants=['calibration, pixel-level'],wrong_variants=['wrong, in this context'])])
        rows=store.table_rows(current);columns=store.table_columns(current)
        for ext in ('json','xlsx'):
            target=self.root/f'editor.{ext}'
            exchange.export_file(str(target),'',rows,columns)
            self.assertEqual(exchange.load_import(target),current['entries'])
    def test_same_term_merges_all_lists_and_keeps_existing_confirmation(self):
        current=self.current([entry(variants=['calibration'],priority=1,note='原定义')])
        path=self.source([entry(variants=['calibration','pixel calibration'],machine_variants=['dot calibration'],wrong_variants=['proofreading'],note='补充定义')])
        result=exchange.import_files([path],current['fingerprint'])
        merged=result['entries'][0]
        self.assertEqual((result['added'],result['updated']),(0,1));self.assertEqual(merged['priority'],1)
        self.assertEqual(merged['variants'],['calibration','pixel calibration'])
        self.assertEqual(merged['machine_variants'],['dot calibration']);self.assertEqual(merged['wrong_variants'],['proofreading'])
        self.assertEqual(merged['note'],'原定义\n补充定义')
        before=self.active.read_bytes();again=exchange.import_files([path],result['fingerprint'])
        self.assertEqual(again['updated'],0);self.assertEqual(self.active.read_bytes(),before)
    def test_different_translation_is_reported_without_duplicate_or_alias_promotion(self):
        current=self.current([entry(priority=1)])
        path=self.source([entry(en='point correction'),entry('新术语','new term',priority=1)])
        result=exchange.import_files([path],current['fingerprint'])
        self.assertEqual(result['added'],1);self.assertEqual(len(result['conflicts']),1)
        self.assertEqual(len(result['entries']),2);self.assertEqual(result['entries'][0]['en'],'pixel-level calibration')
        self.assertEqual(result['entries'][0]['variants'],[]);self.assertEqual(result['entries'][1]['priority'],1)
    def test_alias_and_wrong_translation_disagreements_are_not_silently_merged(self):
        current=self.current([entry(variants=['accepted'],wrong_variants=['rejected'])])
        path=self.source([entry(variants=['rejected','new alias'],wrong_variants=['accepted','new wrong'])])
        result=exchange.import_files([path],current['fingerprint']);e=result['entries'][0]
        self.assertEqual(e['variants'],['accepted','new alias']);self.assertEqual(e['wrong_variants'],['rejected','new wrong'])
        self.assertEqual(len(result['conflicts']),2)
    def test_invalid_later_file_leaves_active_and_earlier_sources_untouched(self):
        current=self.current([entry()]);first=self.source([entry('新术语','new term')]);bad=self.root/'broken.json';bad.write_text('{',encoding='utf8')
        before=self.active.read_bytes()
        result=self.client.post('/api/lexicon/import',json={'paths':[first,str(bad)],'fingerprint':current['fingerprint']})
        self.assertEqual(result.status_code,400);self.assertEqual(self.active.read_bytes(),before)
        self.assertEqual(exchange.load_import(first)[0]['zh'],'新术语')
    def test_partial_table_row_is_an_error_instead_of_being_silently_skipped(self):
        current=store.read();path=self.root/'partial.csv';path.write_text('中文,英文\n测试,test\n未完成,\n',encoding='utf-8-sig')
        before=self.active.read_bytes()
        result=self.client.post('/api/lexicon/import',json={'paths':[str(path)],'fingerprint':current['fingerprint']})
        self.assertEqual(result.status_code,400);self.assertIn('同时填写',result.json['error']);self.assertEqual(self.active.read_bytes(),before)
    def test_stale_active_file_version_does_not_overwrite_another_editor(self):
        self.current([entry()]);path=self.source([entry('新的','new')]);before=self.active.read_bytes()
        result=self.client.post('/api/lexicon/import',json={'paths':[path],'fingerprint':'stale'})
        self.assertEqual(result.status_code,409);self.assertEqual(self.active.read_bytes(),before)
    def test_export_refuses_active_file_and_existing_unconfirmed_target(self):
        self.current([entry()]);before=self.active.read_bytes()
        with self.assertRaises(ValueError):exchange.export_file(str(self.active),store.fingerprint(self.active))
        target=self.root/'already.json';target.write_text('untouched')
        with self.assertRaises(store.Conflict):exchange.export_file(str(target),'')
        self.assertEqual(target.read_text(),'untouched');self.assertEqual(self.active.read_bytes(),before)
        exchange.export_file(str(target),store.fingerprint(target));self.assertEqual(exchange.load_import(target)[0]['zh'],'逐点校正')
    def test_import_preserves_unsaved_edits_and_cancelled_file_list_is_rejected(self):
        current=self.current([entry()]);columns=store.table_columns(current);rows=store.table_rows(current)
        rows[0][columns.index('note')]='工作中的备注';path=self.source([entry('新增','added')])
        result=exchange.import_files([path],current['fingerprint'],rows,columns)
        self.assertEqual(result['entries'][0]['note'],'工作中的备注');self.assertEqual(result['added'],1)
        with self.assertRaises(ValueError):exchange.import_files([],result['fingerprint'])
    def test_legacy_shared_table_is_expanded_without_losing_custom_columns(self):
        shared=self.root/'shared.csv';shared.write_text('中文,英文,负责人\n逐点校正,pixel-level calibration,Alice\n',encoding='utf-8-sig')
        store.atomic_json(self.profile/'lexicon-settings.json',{'path':str(shared)})
        current=store.read();path=self.source([entry(variants=['pixel calibration'],wrong_variants=['point proofreading'],note='定义')])
        result=exchange.import_files([path],current['fingerprint'])
        self.assertEqual(result['entries'][0]['variants'],['pixel calibration'])
        table=writer.load_table(str(shared));self.assertEqual(table['rows'][0][2],'Alice')
        self.assertTrue(set(writer.HEADER_FIELDS)<=set(table['columns']));self.assertEqual(table['header'][2],'负责人')
    def test_malformed_excel_and_empty_import_have_clear_errors(self):
        current=store.read();file=self.root/'invalid.xlsx';file.write_bytes(b'not Excel')
        for path in [str(file),self.source([],name='empty.json')]:
            response=self.client.post('/api/lexicon/import',json={'paths':[path],'fingerprint':current['fingerprint']})
            self.assertEqual(response.status_code,400,response.json)
    def test_multiline_notes_stay_idempotent(self):
        original=[entry(note='第一行\n第二行')];incoming=[entry(note='新增一行\n再一行')]
        once,_=exchange.merge_entries(original,incoming);twice,report=exchange.merge_entries(once,incoming)
        self.assertEqual(once,twice);self.assertEqual(report['updated'],0)
    def test_locale_headers_import_only_chinese_and_english_in_any_order(self):
        for headers,row in [
            (['zh-CN','en-US','es-ES','fr-FR','ja-JP'],['示例术语','example term','ejemplo','exemple','例']),
            (['fr-FR','EN_gb','zh-Hans','ko'],['exemple','example term','示例术语','예'])]:
            book=Workbook();sheet=book.active;sheet.append(headers);sheet.append(row)
            source=self.root/'languages.xlsx';book.save(source);book.close();before=source.read_bytes()
            imported=exchange.load_import(source)
            self.assertEqual(len(imported),1)
            self.assertEqual((imported[0]['zh'],imported[0]['en']),('示例术语','example term'))
            self.assertEqual(imported[0]['note'],'')
            self.assertEqual(source.read_bytes(),before)
    def test_locale_workbook_merge_retains_confirmation_and_is_repeatable(self):
        original=[entry('示例姓名','Example Name',category='person',priority=1,note='已确认姓名',variants=['E. Name'])]
        book=Workbook();sheet=book.active;sheet.append(['zh-CN','en-US','es-ES'])
        sheet.append(['示例姓名','Example Name','otro']);sheet.append(['新增术语','new term','otro'])
        source=self.root/'languages.xlsx';book.save(source);book.close()
        incoming=exchange.load_import(source)
        merged,report=exchange.merge_entries(original,incoming)
        self.assertEqual(report['added'],1)
        self.assertEqual(merged[0],{**original[0],'source':'用户导入'})
        repeated,report=exchange.merge_entries(merged,incoming)
        self.assertEqual(merged,repeated);self.assertEqual(report['added'],0);self.assertEqual(report['updated'],0)
    def test_multiple_english_columns_need_explicit_selection(self):
        source=self.root/'multiple.csv';source.write_text('zh-CN,en-US,en-GB\n示例,example,sample\n',encoding='utf-8-sig')
        with self.assertRaisesRegex(ValueError,'重复字段'):exchange.load_import(source)
    def test_empty_export_is_a_valid_shareable_file(self):
        for ext in ('json','xlsx'):
            target=self.root/f'empty.{ext}';self.assertEqual(exchange.export_file(str(target),'')['count'],0)
            self.assertEqual(exchange.load_import(target),[])


if __name__=='__main__':unittest.main()
