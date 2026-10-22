"""Offline end-to-end format, structural preservation and reviewed export checks."""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_STORED
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, ArrayObject, TextStringObject, DecodedStreamObject
from backend.extract import extract
from backend.extract.office import open_office
from backend.glossary.loader import Glossary
from backend.glossary import store
from backend.models import GlossaryEntry
from backend.translation import translate_file, Cancelled
from backend.matcher import check_segments
from backend import formats, workflow, comparison, document_preview, lan_service
from werkzeug.datastructures import FileStorage

SOURCE = '点间距为2.5mm。'
ENGLISH = 'The pixel pitch is 2.5 mm.'


class Model:
    def __init__(self, text=ENGLISH): self.text, self.calls = text, 0
    def chat_json(self, prompt, **kwargs): self.calls += 1; return {'translation':self.text}


def pdf_fixture(path, text=SOURCE):
    writer = PdfWriter(); page = writer.add_blank_page(600,800)
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type0'),
        NameObject('/BaseFont'):NameObject('/STSong-Light'),NameObject('/Encoding'):NameObject('/UniGB-UCS2-H')})
    cid = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/CIDFontType0'),
        NameObject('/BaseFont'):NameObject('/STSong-Light'),NameObject('/CIDSystemInfo'):DictionaryObject({
            NameObject('/Registry'):TextStringObject('Adobe'),NameObject('/Ordering'):TextStringObject('GB1'),NameObject('/Supplement'):NumberObject(4)})})
    font[NameObject('/DescendantFonts')] = ArrayObject([writer._add_object(cid)])
    cmap = DecodedStreamObject(); pairs = '\n'.join(f'<{ord(c):04x}> <{ord(c):04x}>' for c in set(text))
    cmap.set_data(('/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def /CMapName /Test def /CMapType 2 def 1 begincodespacerange <0000> <FFFF> endcodespacerange '+str(len(set(text)))+' beginbfchar '+pairs+' endbfchar endcmap CMapName currentdict /CMap defineresource pop end end').encode())
    font[NameObject('/ToUnicode')] = writer._add_object(cmap)
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(('BT /F1 12 Tf 30 750 Td <'+text.encode('utf-16-be').hex()+'> Tj ET').encode())
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'data')}); self.env.start()
        self.glossary = Glossary([GlossaryEntry('点间距','pixel pitch',wrong_variants=['dot spacing'],priority=1)])
    def tearDown(self): self.env.stop(); self.temp.cleanup()

    def file(self, extension, text=None):
        path = self.root/('sample'+extension)
        if extension == '.xlsx':
            book = Workbook(); sheet = book.active; sheet.title = '参数'
            sheet['A1'] = SOURCE; sheet['A1'].font = Font(bold=True,color='FF0000'); sheet['A1'].number_format = '@'
            sheet['B1'] = 2.5; sheet['C1'] = '=B1*2'; sheet['A2'] = 'dot spacing'
            sheet.merge_cells('A3:B3'); sheet['A3'] = 'unchanged'; sheet.column_dimensions['A'].width = 30
            book.save(path)
            with ZipFile(path,'a') as z: z.writestr('custom/untouched.bin',b'\x00\x01keep binary\xff')
        elif extension == '.epub':
            with ZipFile(path,'w') as z:
                z.writestr('mimetype','application/epub+zip',compress_type=ZIP_STORED)
                z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
                z.writestr('OPS/book.opf','<package><manifest><item id="ch1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="ch1"/></spine></package>')
                z.writestr('OPS/chapter.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><head><link href="style.css"/></head><body><p>'+SOURCE+'</p><p>dot spacing</p><img src="picture.png"/></body></html>')
                z.writestr('OPS/picture.png',b'untouched image'); z.writestr('OPS/style.css','body{color:blue}')
        elif extension == '.pdf': pdf_fixture(path)
        else:
            contents = {
                '.txt':SOURCE+'\r\n\r\ndot spacing\r\n',
                '.md':'---\ntitle: 原始元数据\n---\n# '+SOURCE+'\n\n[dot spacing](https://example.com/原始链接)\n```python\nx="不要翻译"\n```\n',
                '.html':'<!doctype html><html><head><style>body{color:blue}</style><script>const value="不执行";</script></head><body><p>'+SOURCE+'</p><a href="/原始地址">dot spacing</a><!-- 不执行 --></body></html>',
                '.json':json.dumps({'原始键':SOURCE,'text':'dot spacing','number':2.5,'flag':True,'url':'https://example.com/原始地址'},ensure_ascii=False,indent=2),
                '.srt':'1\r\n00:00:01,000 --> 00:00:03,000\r\n<i>'+SOURCE+'</i>\r\n\r\n2\r\n00:00:04,000 --> 00:00:06,000\r\ndot spacing\r\n',
                '.ass':'[Script Info]\nTitle: 原始标题\n[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: 0,0:00:01.00,0:00:03.00,Default,,0,0,0,,{\\b1}'+SOURCE+'{\\b0}\\Ndot spacing\n'
            }
            path.write_bytes((contents[extension] if text is None else text).encode('utf-8'))
        return path

    def translate(self,path,**kwargs):
        return translate_file(str(path),self.glossary,None,client=Model(),**kwargs)

    def test_every_format_translates_and_supports_manual_reviewed_replacements(self):
        for extension in ['.txt','.md','.html','.json','.srt','.ass','.xlsx','.epub']:
            with self.subTest(extension=extension):
                source = self.file(extension); before = store.fingerprint(source)
                result = self.translate(source)
                self.assertEqual(before,store.fingerprint(source)); self.assertEqual(Path(result['path']).suffix,extension)
                self.assertEqual(result['untranslated'],0)
                self.assertTrue(result['comparison_id']); self.assertIn(ENGLISH,[s.text for s in extract(result['path'])])
                job = {'id':'format-'+extension,'cancel':threading.Event(),'status':'running'}
                with patch('backend.workflow.store.load_glossary',return_value=self.glossary):
                    workflow.run(job,{'kind':'check','paths':[result['path']],'use_llm':False})
                self.assertEqual(job['status'],'complete',job)
                finding = next(f for f in job['result']['findings'] if f['found_text']=='dot spacing')
                saved = workflow.apply_report({'report_id':job['id'],'path':result['path'],'selected':{finding['segment_id']:[['dot spacing','pixel pitch']]}})
                self.assertEqual(saved['applied_count'],1)
                self.assertNotIn('dot spacing',[s.text for s in extract(saved['output'])])
                self.assertIn('pixel pitch',[s.text for s in extract(saved['output'])])
                self.assertTrue(saved['finding_updates'][0]['finding']['applied'])

    def test_text_and_markup_structure_is_preserved(self):
        for extension in ['.txt','.md','.html','.json','.srt','.ass']:
            with self.subTest(extension=extension):
                source = self.file(extension); raw = source.read_bytes()
                output = Path(self.translate(source)['path']).read_bytes()
                self.assertEqual(output,raw.replace(SOURCE.encode(),ENGLISH.encode()))

    def test_excel_preserves_formulas_styles_links_and_untouched_package_parts(self):
        source = self.file('.xlsx'); output = self.translate(source)['path']
        book = load_workbook(output); sheet = book.active
        self.assertEqual(sheet['A1'].value,ENGLISH); self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet['B1'].value,2.5); self.assertEqual(sheet['C1'].value,'=B1*2')
        self.assertEqual(str(sheet.merged_cells),'A3:B3'); self.assertEqual(sheet.column_dimensions['A'].width,30)
        with ZipFile(source) as a,ZipFile(output) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name != 'xl/worksheets/sheet1.xml': self.assertEqual(a.read(name),b.read(name),name)

    def test_epub_retains_images_styles_manifest_and_mimetype_position(self):
        source = self.file('.epub'); output = self.translate(source)['path']
        with ZipFile(source) as a,ZipFile(output) as b:
            self.assertEqual(a.namelist(),b.namelist()); self.assertEqual(b.infolist()[0].compress_type,ZIP_STORED)
            for name in a.namelist():
                expected = a.read(name).replace(SOURCE.encode(),ENGLISH.encode()) if name.endswith('.xhtml') else a.read(name)
                self.assertEqual(b.read(name),expected)

    def test_pdf_text_layer_exports_real_word_and_rejects_overwrite(self):
        source = self.file('.pdf'); before = store.fingerprint(source)
        self.assertEqual(extract(str(source))[0].text,SOURCE)
        result = self.translate(source); output = Path(result['path'])
        self.assertEqual(output.suffix,'.docx'); self.assertEqual(Document(output).paragraphs[0].text,ENGLISH)
        self.assertIn('重新排版',result['layout_note']); self.assertTrue(result['warnings'])
        self.assertEqual(before,store.fingerprint(source)); self.assertTrue(result['comparison_id'])
        self.assertEqual(comparison.read(result['comparison_id'])['rows'][0]['translation'],ENGLISH)
        with self.assertRaisesRegex(ValueError,'不能覆盖'):
            workflow._apply_report({'path':str(source),'overwrite':True},{})

    def test_scan_and_encrypted_pdf_have_actionable_messages(self):
        source = self.root/'scan.pdf'; writer = PdfWriter(); writer.add_blank_page(600,800); writer.write(source)
        with self.assertRaisesRegex(ValueError,'OCR'): extract(str(source))
        writer.encrypt('test-password'); writer.write(source)
        with self.assertRaisesRegex(ValueError,'加密'): extract(str(source))

    def test_direct_pdf_proofreading_save_produces_word_even_without_changes(self):
        source = self.root/'english.pdf'; pdf_fixture(source,'The pixel pitch is 2.5 mm.')
        job = {'id':'pdf-check','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=self.glossary): workflow.run(job,{'kind':'check','paths':[str(source)],'use_llm':False})
        output = workflow.apply_report({'report_id':job['id'],'path':str(source),'selected':{},'save_unchanged':True})['output']
        self.assertEqual(Path(output).suffix,'.docx'); self.assertEqual(Document(output).paragraphs[0].text,'The pixel pitch is 2.5 mm.')

    def test_final_only_translation_and_review_keeps_structure_and_one_public_output(self):
        def rule_review(segments,glossary,*args,**kwargs):
            found,stats = check_segments(segments,glossary,False,cancel=kwargs.get('cancel'))
            stats['model_status'] = 'complete'; return found,stats
        for extension in ['.txt','.md','.html','.json','.srt','.ass','.xlsx','.epub','.pdf']:
            with self.subTest(extension=extension):
                source = self.file(extension); folder = self.root/('public-'+extension); folder.mkdir()
                job = {'id':'final-'+extension,'cancel':threading.Event(),'status':'running'}
                with patch('backend.workflow.store.load_glossary',return_value=self.glossary),patch('backend.workflow.check_segments',side_effect=rule_review):
                    workflow.run(job,{'kind':'translate_review','combined_output':'final','paths':[str(source)],'output_dir':str(folder),'_client':Model()})
                self.assertEqual(job['status'],'partial' if extension=='.pdf' else 'complete',job); self.assertEqual(list(folder.iterdir()),[])
                report = job['result']['report']; path = report['files'][0]['path']
                saved = workflow.apply_report({'report_id':job['id'],'path':path,'selected':{},'save_unchanged':True})
                self.assertEqual(len(list(folder.iterdir())),1); output = Path(saved['output'])
                self.assertEqual(output.suffix,'.docx' if extension == '.pdf' else extension)
                self.assertIn(ENGLISH,[s.text for s in extract(str(output))])
                self.assertNotIn('_Draft',output.name)

    def test_cancel_does_not_create_partial_output(self):
        for extension in ['.txt','.xlsx','.pdf']:
            path = self.file(extension); cancel = threading.Event(); cancel.set(); folder = self.root/'cancel'
            with self.assertRaises(Cancelled): self.translate(path,output_dir=folder,cancel=cancel)
            self.assertFalse(folder.exists())

    def test_utf16_bom_and_gb18030_encoding_survive(self):
        for encoding in ['utf-8-sig','utf-16','gb18030']:
            source = self.file('.txt'); source.write_bytes((SOURCE+'\r\n').encode(encoding))
            output = Path(self.translate(source)['path']).read_bytes()
            self.assertEqual(output,(ENGLISH+'\r\n').encode(encoding))

    def test_html_entities_and_script_do_not_enter_translation(self):
        path = self.file('.html','<p>屏幕 &amp; 显示</p><script>window.secret="中文"</script><pre>保留代码</pre>')
        segments = extract(str(path)); self.assertEqual([s.text for s in segments],['屏幕 & 显示'])
        document,paras,_ = open_office(str(path)); paras[0][1].text = 'Screen & display <value>'
        destination = self.root/'safe.html'; document.save(destination)
        self.assertEqual(destination.read_text(encoding='utf-8'),'<p>Screen &amp; display &lt;value&gt;</p><script>window.secret="中文"</script><pre>保留代码</pre>')
        path.write_text('<p>屏幕 &copy 显示</p>',encoding='utf-8')
        document,paras,_ = open_office(str(path)); self.assertEqual(paras[0][1].text,'屏幕 © 显示')
        paras[0][1].text = 'Screen'; document.save(destination)
        self.assertEqual(destination.read_text(encoding='utf-8'),'<p>Screen</p>')

    def test_json_keys_and_paths_are_not_modified_and_duplicate_keys_rejected(self):
        path = self.file('.json','{"中文键":"中文值","path":"C:\\\\目录\\\\文件.txt","url":"https://example.com/中文"}')
        self.assertEqual([s.text for s in extract(str(path))],['中文值'])
        path.write_text('{"a":1,"a":2}')
        with self.assertRaisesRegex(ValueError,'重复键名'): extract(str(path))

    def test_invalid_subtitles_are_not_silently_corrupted(self):
        path = self.file('.srt','1\nnot a timestamp\n中文\n')
        with self.assertRaisesRegex(ValueError,'时间轴'): extract(str(path))
        path = self.file('.ass','[Events]\nFormat: Text, Start\nDialogue: 中文,0:00:01.00')
        with self.assertRaisesRegex(ValueError,'最后'): extract(str(path))

    def test_markdown_reference_targets_and_embedded_code_are_immutable(self):
        source = self.file('.md','[中文标题][中文标识]\n\n[中文标识]: https://example.com/中文\n<script>const value="保留代码";</script>')
        self.assertEqual([s.text for s in extract(str(source))],['中文标题'])

    def test_subtitle_model_line_breaks_cannot_create_new_cues_or_events(self):
        for kind in ['.srt','.ass']:
            source = self.file(kind); document,paragraphs,_ = open_office(str(source))
            paragraphs[0][1].text = 'New\r\ntext'
            output = self.root/('output'+kind); document.save(output)
            self.assertIn('New text',[s.text for s in extract(str(output))])
            self.assertEqual(len(extract(str(output))),len(extract(str(source))))

    def test_lan_accepts_the_same_formats_without_requiring_office_zip_for_text(self):
        for extension in ['.txt','.md','.html','.json','.srt','.ass','.xlsx','.epub','.pdf']:
            path = self.file(extension)
            saved = lan_service.save_upload(FileStorage(stream=io.BytesIO(path.read_bytes()),filename=path.name),self.root/'uploads')
            self.assertEqual(Path(saved).read_bytes(),path.read_bytes())
        with self.assertRaisesRegex(ValueError,'不支持'): lan_service.save_upload(FileStorage(stream=io.BytesIO(b'bad'),filename='bad.exe'),self.root/'uploads')

    def test_text_comparison_and_preview_are_escaped_and_keep_original_snapshot(self):
        source = self.file('.html','<p>'+SOURCE+'</p><script>evil()</script>')
        result = self.translate(source); source.unlink()
        preview = document_preview.preview(result['comparison_id'],'source')
        self.assertIn('正文阅读预览',preview); self.assertIn(SOURCE,preview); self.assertNotIn('evil()',preview)
        self.assertIn('不还原',preview)

    def test_frontend_supported_formats_match_backend(self):
        import re
        ui = (Path(__file__).resolve().parents[1]/'renderer/document-input.js').read_text(encoding='utf-8')
        names = re.search(r'const extensions = \[(.*?)\]',ui)[1]
        self.assertEqual({'.'+s for s in re.findall(r"'([^']+)'",names)},formats.SUPPORTED)


if __name__ == '__main__': unittest.main()
