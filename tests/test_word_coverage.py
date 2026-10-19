"""Visible Word text and immutable fields/objects must both survive export."""
import base64,json,os,re,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile,ZIP_DEFLATED
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import RGBColor
from lxml import etree
from backend.extract import extract
from backend.glossary.store import Glossary,fingerprint
from backend.translation import translate_file

CN=re.compile(r'[\u3400-\u9fff]')
WORDS={'正文前文':'Body before ','链接标签':'Reference label','正文后文':' body after',
       '文本框文字':'Text box wording','页眉文字':'Header wording','嵌套文字':'Nested wording',
       '脚注文字':'Footnote wording','保留文字':'Retained wording'}

def text_run(parent,text):
    run=OxmlElement('w:r');node=OxmlElement('w:t');node.text=text
    run.append(node);parent.append(run);return run

def field(paragraph,command='HYPERLINK "https://example.test/原始地址"',label='链接标签'):
    for kind in ('begin','instruction','separate','result','end'):
        run=OxmlElement('w:r')
        if kind=='instruction':node=OxmlElement('w:instrText');node.text=command
        elif kind=='result':node=OxmlElement('w:t');node.text=label
        else:node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind)
        run.append(node);paragraph._p.append(run)

def signature(path):
    with ZipFile(path) as z:
        roots=[etree.fromstring(z.read(n)) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml')]
        fields=[(n.tag,tuple(sorted(n.attrib.items())),n.text) for root in roots for n in root.iter()
                if n.tag in {qn('w:fldChar'),qn('w:instrText'),qn('w:bookmarkStart'),qn('w:bookmarkEnd')}]
        protected={n:z.read(n) for n in z.namelist() if n.startswith('word/media/') or n.endswith('.rels')}
    return fields,protected

class Model:
    def __init__(self):self.prompts=[]
    def chat_json(self,prompt,**kwargs):
        self.prompts.append(prompt)
        def english(text):
            for old,new in WORDS.items():text=text.replace(old,new)
            if CN.search(text):raise AssertionError('Unexpected fixture: '+text)
            return text
        result={'translation':english(prompt.split('【待译原文】\n')[-1].split('\n【')[0])}
        marker='【需保留的原文样式片段】\n'
        if marker in prompt:
            groups,_=json.JSONDecoder().raw_decode(prompt.split(marker)[-1])
            result['runs']=[{'id':g['id'],'text':english(g['text'])} for g in groups]
        return result

class WordCoverageTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.path=self.root/'source.docx'
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});env.start();self.addCleanup(env.stop)
    def translate(self,doc):
        doc.save(self.path);before=fingerprint(self.path);model=Model()
        result=translate_file(str(self.path),Glossary([]),None,client=model,workers=1)
        self.assertEqual(fingerprint(self.path),before)
        return result,model
    def test_hyperlink_field_translates_prose_label_and_retains_full_paragraph_context(self):
        doc=Document();p=doc.add_paragraph('正文前文');field(p);p.add_run('正文后文');doc.save(self.path)
        before=signature(self.path);result,model=self.translate(doc)
        self.assertTrue(result['translation_complete']);self.assertEqual(result['untranslated'],0)
        self.assertEqual(Document(result['path']).paragraphs[0].text,'Body before Reference label body after')
        self.assertEqual(signature(result['path']),before)
        self.assertTrue(all('正文前文' in p and '链接标签' in p and '正文后文' in p for p in model.prompts))
    def test_validated_2016_checkpoints_keep_identity_and_remain_available(self):
        from backend.checkpoints import Checkpoint
        identity={'source':str(self.path),'fingerprint':'fixture','model':{'model':'fixture'}}
        old='0b812f78dc872c18d916ec6b046062e9c7dfa7bd8372a133ba4343244fc86cc0'
        with patch('backend.checkpoints.policy',return_value=old):
            Checkpoint('translation',identity).put(['body/0'],{'translation':'Verified English'})
        self.assertEqual(Checkpoint('translation',identity).get(['body/0'])['translation'],'Verified English')
        self.assertIsNone(Checkpoint('translation',{**identity,'fingerprint':'changed'}).get(['body/0']))
    def test_native_hyperlink_preserves_target_bold_and_color(self):
        doc=Document();p=doc.add_paragraph('正文前文');link=OxmlElement('w:hyperlink')
        link.set(qn('r:id'),p.part.relate_to('https://example.test/原始地址',RT.HYPERLINK,is_external=True))
        run=text_run(link,'链接标签');prop=OxmlElement('w:rPr');prop.append(OxmlElement('w:b'))
        color=OxmlElement('w:color');color.set(qn('w:val'),'2080FF');prop.append(color);run.insert(0,prop)
        p._p.append(link);p.add_run('正文后文');result,_=self.translate(doc)
        output=Document(result['path']);node=output.paragraphs[0]._p.find(qn('w:hyperlink'))
        self.assertEqual(''.join(n.text or '' for n in node.iter(qn('w:t'))).strip(),'Reference label')
        self.assertIsNotNone(node.find('.//'+qn('w:b')))
        self.assertEqual(node.find('.//'+qn('w:color')).get(qn('w:val')),'2080FF')
        self.assertEqual(output.part.rels[node.get(qn('r:id'))].target_ref,'https://example.test/原始地址')
    def test_dynamic_date_protects_only_result_and_does_not_claim_completion(self):
        doc=Document();p=doc.add_paragraph('正文前文');field(p,'DATE \\@ "yyyy年M月"','中文日期');p.add_run('正文后文')
        doc.save(self.path);before=signature(self.path);result,_=self.translate(doc)
        self.assertEqual(Document(result['path']).paragraphs[0].text,'Body before中文日期body after')
        self.assertFalse(result['translation_complete']);self.assertEqual(result['untranslated'],1)
        self.assertEqual(signature(result['path']),before)
    def test_numeric_page_field_does_not_block_surrounding_translation(self):
        doc=Document();p=doc.add_paragraph('正文前文');field(p,'PAGE','3');p.add_run(' Existing English.')
        result,_=self.translate(doc);self.assertTrue(result['translation_complete'])
        self.assertEqual(Document(result['path']).paragraphs[0].text,'Body before 3 Existing English.')
    def test_textbox_and_host_prose_translate_without_changing_geometry(self):
        doc=Document();p=doc.add_paragraph('正文前文');run=OxmlElement('w:r');pict=OxmlElement('w:pict')
        shape=etree.SubElement(pict,'{urn:schemas-microsoft-com:vml}shape',style='width:100pt;height:50pt')
        box=OxmlElement('w:txbxContent');inner=OxmlElement('w:p');text_run(inner,'文本框文字')
        box.append(inner);shape.append(box);run.append(pict);p._p.append(run);p.add_run('正文后文')
        result,_=self.translate(doc);out=Document(result['path'])
        self.assertTrue(result['translation_complete']);self.assertEqual(result['translated'],3)
        self.assertEqual(out.paragraphs[0].text,'Body before body after')
        self.assertEqual(''.join(n.text or '' for box in out.element.iter(qn('w:txbxContent')) for n in box.iter(qn('w:t'))),'Text box wording')
        self.assertEqual(next(out.element.iter('{urn:schemas-microsoft-com:vml}shape')).get('style'),'width:100pt;height:50pt')
    def test_picture_is_not_deleted_or_duplicated_when_prose_changes(self):
        doc=Document();p=doc.add_paragraph('正文前文');png=self.root/'pixel.png'
        png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aWQ0AAAAASUVORK5CYII='))
        p.add_run().add_picture(str(png));p.add_run('正文后文');doc.save(self.path);before=signature(self.path)
        result,_=self.translate(doc)
        self.assertEqual(len(Document(result['path']).inline_shapes),1);self.assertEqual(signature(result['path']),before)
    def test_content_control_nested_tables_and_header_are_covered(self):
        doc=Document();sdt=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');p=OxmlElement('w:p')
        text_run(p,'嵌套文字');content.append(p);sdt.append(content);doc.element.body.insert(0,sdt)
        table=doc.add_table(1,1);table.cell(0,0).text='保留文字';table.cell(0,0).add_table(1,1).cell(0,0).text='嵌套文字'
        doc.sections[0].header.paragraphs[0].text='页眉文字';result,_=self.translate(doc)
        values=[s.text for s in extract(result['path']) if s.source!='image']
        self.assertFalse(any(CN.search(v) for v in values));self.assertEqual(values.count('Nested wording'),2)
    def test_merged_cells_are_processed_once(self):
        doc=Document();doc.add_table(1,2).cell(0,0).merge(doc.tables[0].cell(0,1)).text='保留文字'
        result,_=self.translate(doc);self.assertEqual(result['total'],1)
    def test_bookmarks_and_ordinary_inline_emphasis_keep_their_positions(self):
        doc=Document();p=doc.add_paragraph('正文前文');begin=OxmlElement('w:bookmarkStart')
        begin.set(qn('w:id'),'7');begin.set(qn('w:name'),'source-reference');p._p.append(begin)
        run=p.add_run('链接标签');run.bold=True;run.font.color.rgb=RGBColor(0x20,0x80,0xFF)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'7');p._p.append(end);p.add_run('正文后文')
        doc.save(self.path);before=signature(self.path);result,_=self.translate(doc)
        out=Document(result['path']).paragraphs[0];styled=next(r for r in out.runs if 'Reference label' in r.text)
        self.assertTrue(styled.bold);self.assertEqual(str(styled.font.color.rgb),'2080FF')
        self.assertEqual(signature(result['path']),before)
        self.assertEqual([n.tag.rsplit('}',1)[-1] for n in out._p],['r','bookmarkStart','r','bookmarkEnd','r'])
    def test_original_english_token_split_by_a_link_is_never_split_with_spaces(self):
        doc=Document();p=doc.add_paragraph('Micro');field(p,label='LED');p.add_run(' 正文后文')
        result,_=self.translate(doc)
        self.assertTrue(Document(result['path']).paragraphs[0].text.startswith('MicroLED'))
    def test_cross_paragraph_hyperlink_field_does_not_hide_following_prose(self):
        doc=Document();doc.add_paragraph('正文前文');p=doc.add_paragraph()
        field(p,'HYPERLINK "https://example.test"','保留文字');end=p._p[-1];p._p.remove(end)
        tail=doc.add_paragraph();tail._p.append(end);tail.add_run('正文后文')
        result,_=self.translate(doc)
        self.assertTrue(result['translation_complete'])
        self.assertEqual(Document(result['path']).paragraphs[1].text,'Retained wording')
        self.assertEqual(Document(result['path']).paragraphs[2].text,'body after')
    def test_saved_locators_work_for_later_review_replacement(self):
        doc=Document();p=doc.add_paragraph('正文前文');field(p);p.add_run('正文后文');doc.save(self.path)
        before=extract(str(self.path));result,_=self.translate(doc);after=extract(result['path'])
        self.assertEqual([s.locator for s in before],[s.locator for s in after])
        from backend.replacer.applier import apply_replacements
        target=next(s for s in after if s.text.strip()=='Reference label')
        output,done,failed=apply_replacements(result['path'],after,{target.id:[('Reference label','Reference wording')]})
        self.assertFalse(failed);self.assertEqual(len(done),1);self.assertIn('Reference wording',Document(output).paragraphs[0].text)
        self.assertEqual(signature(output),signature(result['path']))
    def test_untraversed_visible_chinese_is_reported_not_claimed_complete(self):
        doc=Document();doc.add_paragraph('正文前文');unknown=OxmlElement('w:customUnknown');text_run(unknown,'保留文字')
        doc.element.body.insert(1,unknown);result,_=self.translate(doc)
        self.assertFalse(result['translation_complete']);self.assertEqual(result['untranslated'],1)
        self.assertTrue(any(i['source']=='word_other' for i in result['remaining_chinese']))
    def test_footnote_translation_keeps_numbering(self):
        doc=Document();doc.add_paragraph('正文前文');doc.save(self.path)
        with ZipFile(self.path) as z:files={n:z.read(n) for n in z.namelist()}
        root=etree.fromstring(files['[Content_Types].xml'])
        etree.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
        files['[Content_Types].xml']=etree.tostring(root);root=etree.fromstring(files['word/_rels/document.xml.rels'])
        etree.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='noteFixture',Type=RT.FOOTNOTES,Target='footnotes.xml')
        files['word/_rels/document.xml.rels']=etree.tostring(root)
        files['word/footnotes.xml']=('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t>脚注文字</w:t></w:r></w:p></w:footnote></w:footnotes>').encode()
        with ZipFile(self.path,'w',ZIP_DEFLATED) as z:
            for name,data in files.items():z.writestr(name,data)
        result=translate_file(str(self.path),Glossary([]),None,client=Model(),workers=1)
        self.assertTrue(result['translation_complete'])
        with ZipFile(result['path']) as z:root=etree.fromstring(z.read('word/footnotes.xml'))
        self.assertEqual(next(root.iter(qn('w:t'))).text,'Footnote wording')
        self.assertEqual(len(list(root.iter(qn('w:footnoteRef')))),1)
        self.assertEqual(next(root.iter(qn('w:footnote'))).get(qn('w:id')),'1')

if __name__=='__main__':unittest.main()
