"""Offline regression for persistent history, editable choices and complete tasks."""
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from backend import workflow, history
from backend.document_context import DocumentContext
from backend.extract import extract
from backend.glossary import store
from backend.glossary.loader import Glossary
from backend.models import GlossaryEntry, Segment
from backend.translation import translate_file
from backend.server import app


class TrialTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {'NOVATERM_DATA_DIR': str(self.root / 'data')}); self.env.start()
        self.glossary = Glossary([GlossaryEntry('点间距', 'pixel pitch', priority=1, wrong_variants=['dot spacing']),
                                 GlossaryEntry('刷新率', 'refresh rate', priority=1, wrong_variants=['bad refresh'])])
    def tearDown(self):
        self.env.stop(); self.temp.cleanup()
    def doc(self, text):
        path = self.root / 'source.docx'; doc = Document(); doc.add_paragraph(text); doc.save(path)
        return str(path)
    def report(self, path):
        job = {'id': 'trial', 'status': 'running', 'cancel': threading.Event()}
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary):
            workflow.run(job, {'kind': 'check', 'paths': [path], 'use_llm': False})
        self.assertEqual(job['status'], 'complete', job)
        return job['result']
    def test_edited_replacements_save_cumulatively_without_changing_original(self):
        path = self.doc('dot spacing and bad refresh'); original = Path(path).read_bytes()
        report = self.report(path)
        first = next(i for i, f in enumerate(report['findings']) if f['found_text'] == 'dot spacing')
        second = 1 - first
        a = workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': first, 'text': 'pixel pitch (P2.5)'}]})
        self.assertTrue(a['finding_updates'][first]['finding']['applied'])
        self.assertEqual(Document(a['output']).paragraphs[0].text, 'pixel pitch (P2.5) and bad refresh')
        b = workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': second, 'text': 'refresh rate'}]})
        self.assertEqual(Document(b['output']).paragraphs[0].text, 'pixel pitch (P2.5) and refresh rate')
        self.assertEqual(Path(path).read_bytes(), original)
        self.assertTrue(all(item['finding']['applied'] for item in b['finding_updates']))
        rows = history.query(kind='check')['records']
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(rows[0]['exports']), 2)
        self.assertEqual(len(rows[0]['outputs']), 2)
        self.assertEqual(rows[0]['counts']['replaced'], 2)
        self.assertEqual(rows[0]['exports'][-1]['path'], b['output'])
    def test_overwrite_after_save_as_preserves_previous_edits(self):
        path = self.doc('dot spacing and bad refresh'); report = self.report(path)
        first = next(i for i, f in enumerate(report['findings']) if f['found_text'] == 'dot spacing')
        workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': first, 'text': 'pixel pitch'}]})
        result = workflow.apply_report({'report_id': 'trial', 'path': path, 'overwrite': True, 'edits': [{'index': 1-first, 'text': 'refresh rate'}]})
        self.assertTrue(result['overwritten']); self.assertEqual(Document(path).paragraphs[0].text, 'pixel pitch and refresh rate')
        self.assertFalse(list(self.root.glob('*.bak')))
    def test_changed_output_and_invalid_edits_are_rejected(self):
        path = self.doc('dot spacing and bad refresh'); report = self.report(path)
        for edit in [{'index': 99, 'text': 'x'}, {'index': 0, 'text': ''}, {'index': '0', 'text': 'x'}]:
            with self.assertRaises(ValueError): workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [edit]})
        first = next(i for i, f in enumerate(report['findings']) if f['found_text'] == 'dot spacing')
        saved = workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': first, 'text': 'pixel pitch'}]})
        doc = Document(saved['output']); doc.add_paragraph('external edit'); doc.save(saved['output'])
        with self.assertRaisesRegex(ValueError, '上次输出'):
            workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': 1-first, 'text': 'refresh rate'}]})
        # A user edit made while the next export is being prepared must also
        # survive; don't replace the source or keep an outdated new result.
        for overwrite in (False, True):
            self.report(path)
            previous = workflow.apply_report({'report_id': 'trial', 'path': path, 'edits': [{'index': first, 'text': 'pixel pitch'}]})
            original_save = workflow.save_office
            def save_then_external_edit(*args, **kwargs):
                original_save(*args, **kwargs)
                doc = Document(previous['output']); doc.add_paragraph('concurrent edit'); doc.save(previous['output'])
            original_bytes = Path(path).read_bytes()
            with patch('backend.workflow.save_office', side_effect=save_then_external_edit), self.assertRaisesRegex(ValueError, '上次输出'):
                workflow.apply_report({'report_id': 'trial', 'path': path, 'overwrite': overwrite, 'edits': [{'index': 1-first, 'text': 'refresh rate'}]})
            self.assertEqual(Path(path).read_bytes(), original_bytes)
            self.assertEqual(Document(previous['output']).paragraphs[-1].text, 'concurrent edit')
    def test_history_filters_dates_and_survives_new_connections(self):
        now = datetime.now(); old = now - timedelta(days=370)
        for ident, stamp in [('today', now), ('old', old)]:
            history.begin(ident, {'kind': 'translate', 'paths': ['/qa/original.pptx']})
            history.finish({'id': ident, 'status': 'complete', 'result': {'outputs': [{'path': '/qa/out.pptx'}]}}, {'kind': 'translate', 'paths': []})
            with history.connection() as db: db.execute('UPDATE tasks SET started=? WHERE id=?', (stamp.timestamp(), ident))
        self.assertEqual(history.query('today')['total'], 1)
        self.assertEqual(history.query('month')['total'], 1); self.assertEqual(history.query('year')['total'], 1)
        result = history.query('custom', old.date().isoformat(), old.date().isoformat())
        self.assertEqual(result['records'][0]['id'], 'old')
        self.assertEqual(result['records'][0]['outputs'][0]['path'], '/qa/out.pptx')
        self.assertEqual(history.query(text='original.pptx')['total'], 2)
        self.assertEqual(history.query(kind='check')['total'], 0)
        with self.assertRaises(ValueError): history.query('custom', now.date().isoformat(), old.date().isoformat())
        self.assertEqual(app.test_client().get('/api/history?preset=all').json['total'], 2)
    def test_interrupted_history_does_not_mark_active_job_as_stopped(self):
        for ident in ['active', 'stopped']: history.begin(ident, {'kind': 'check', 'paths': []})
        rows = {row['id']: row for row in history.query(active_ids=['active'])['records']}
        self.assertEqual(rows['active']['status'], 'running'); self.assertEqual(rows['stopped']['status'], 'interrupted')
    def test_context_resolves_screens_without_changing_cabinets_or_modules(self):
        rows = [Segment('title', 'x', '全球及中国影院市场', 'a', file_path='f', slide=1),
                Segment('count', 'x', '21.4万块', 'b', file_path='f', slide=1),
                Segment('cabinets', 'x', '8块箱体', 'c', file_path='f', slide=2)]
        context = DocumentContext(rows)
        self.assertEqual(context.repair_quantity(rows[1], '21.4万块', '214,000 cabinets'), '214,000 screens')
        self.assertEqual(context.repair_quantity(rows[2], '8块箱体', '8 cabinets'), '8 cabinets')
        self.assertIn('全球及中国影院市场', context.prompt(rows[1]))
        self.assertNotEqual(context.cache_key(rows[1], '块'), context.cache_key(rows[2], '块'))
    def test_ppt_split_screen_counts_use_whole_page_and_consistent_nouns(self):
        path = self.root / 'counts.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        for index, text in enumerate(['全球及中国影院市场', '全球银幕总量21.4万块', '中国银幕数量8.6万块', '2025年落地超过200块']):
            slide.shapes.add_textbox(Inches(1), Inches(index+0.2), Inches(8), Inches(.8)).text = text
        deck.save(path); source_bytes = path.read_bytes()
        responses = {'全球及中国影院市场': 'Global and Chinese cinema markets', '全球银幕总量21.4万块': 'Global screen count: 214,000 units',
                     '中国银幕数量8.6万块': 'China: 86,000 cabinets', '2025年落地超过200块': 'Over 200 units deployed in 2025'}
        class Model:
            prompts = []
            def chat_json(self, prompt, **kwargs):
                self.prompts.append(prompt); source = prompt.split('【待译原文】\n', 1)[1].split('\n【', 1)[0]
                return {'translation': responses[source]}
        model = Model()
        output = translate_file(str(path), self.glossary, None, client=model)
        text = '\n'.join(s.text for s in extract(output['path']))
        self.assertIn('214,000 screens', text); self.assertIn('86,000 screens', text); self.assertIn('200 screens', text)
        self.assertEqual(path.read_bytes(), source_bytes)
        self.assertTrue(all('全球及中国影院市场' in prompt for prompt in model.prompts))
    def test_combined_job_leaves_all_suggestions_pending_until_confirmation(self):
        path = self.doc('点间距为2.5毫米。'); source_bytes = Path(path).read_bytes(); prompts = []
        phases = []
        class Model:
            def chat_json(self, prompt, **kwargs):
                prompts.append((prompt, kwargs))
                phases.append(job['phase'])
                if '【待译原文】' in prompt: return {'translation': 'The flexibile dot spacing is 2.5 mm.'}
                assert '【待审校段落】' in prompt
                return [{'found':'flexibile','suggested':'flexible','kind':'typo','reason':'拼写错误'}]
        job = {'id': 'combined', 'status': 'running', 'cancel': threading.Event()}
        finish = history.finish
        def save_ready(snapshot, data):
            self.assertEqual(job['status'], 'running')
            self.assertIn('report', job['result'])
            self.assertIn('combined', workflow.REPORTS)
            finish(snapshot, data)
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary), patch('backend.translation.LLMClient', return_value=Model()), patch('backend.matcher.LLMClient', return_value=Model()), patch('backend.workflow.history.finish', side_effect=save_ready):
            workflow.run(job, {'kind': 'translate_review', 'paths': [path]})
        self.assertEqual(job['status'], 'complete', job)
        result = job['result']['outputs'][0]
        self.assertTrue(result['auto_reviewed']); self.assertEqual(result['review_status'], 'complete')
        self.assertEqual(Document(result['path']).paragraphs[0].text, 'The flexibile dot spacing is 2.5 mm.')
        self.assertEqual(result['corrected'], 0); self.assertEqual(Path(path).read_bytes(), source_bytes)
        self.assertEqual(phases, ['translate','check'])
        report = job['result']['report']
        self.assertEqual(report['stats']['llm_hits'], 1)
        self.assertEqual(report['stats']['judged'], 1)
        self.assertTrue(all(not f['applied'] for f in report['findings']))
        self.assertEqual(report['files'][0]['path'], result['path'])
        self.assertEqual(job['result']['completed_paths'], [path])
        self.assertEqual(workflow.REPORTS['combined']['hashes'][result['path']], store.fingerprint(result['path']))
        self.assertEqual(history.query(kind='translate_review')['records'][0]['counts']['replaced'], 0)
        self.assertEqual(history.query(kind='translate_review')['records'][0]['outputs'][0]['path'], result['path'])
        self.assertTrue(any('the Chinese market' in kwargs.get('system', '') for _, kwargs in prompts))
        saved = workflow.apply_report({'report_id':'combined','path':result['path'],
            'edits':[{'index':i,'text':f['suggested']} for i,f in enumerate(report['findings'])]})
        self.assertEqual(Document(saved['output']).paragraphs[0].text, 'The flexible pixel pitch is 2.5 mm.')
        self.assertEqual(Document(result['path']).paragraphs[0].text, 'The flexibile dot spacing is 2.5 mm.')
        self.assertEqual(history.query(kind='translate_review')['records'][0]['counts']['replaced'],2)
    def test_combined_export_failure_retains_report_for_manual_application(self):
        path = self.doc('点间距为2.5毫米。'); original = Path(path).read_bytes()
        class Model:
            def chat_json(self, prompt, **kwargs):
                return {'translation':'The dot spacing is 2.5 mm.'} if '【待译原文】' in prompt else []
        job = {'id':'recover','status':'running','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary), patch('backend.translation.LLMClient', return_value=Model()), patch('backend.matcher.LLMClient', return_value=Model()):
            workflow.run(job, {'kind':'translate_review','paths':[path]})
        report = job['result']['report']; output = job['result']['outputs'][0]['path']
        self.assertEqual(report['stats']['model_status'], 'complete')
        with patch('backend.workflow.save_office', side_effect=PermissionError('文件暂不可写')), self.assertRaises(PermissionError):
            workflow.apply_report({'report_id':'recover','path':output,'edits':[{'index':0,'text':'pixel pitch'}]})
        self.assertFalse(report['findings'][0]['applied'])
        self.assertEqual(job['result']['completed_paths'], [path])
        repaired = workflow.apply_report({'report_id':'recover','path':output,'overwrite':True,'edits':[{'index':0,'text':'pixel pitch'}]})
        self.assertNotEqual(repaired['output'], output)
        self.assertEqual(Document(repaired['output']).paragraphs[0].text, 'The pixel pitch is 2.5 mm.')
        self.assertEqual(Document(output).paragraphs[0].text, 'The dot spacing is 2.5 mm.')
        self.assertEqual(Path(path).read_bytes(), original)
        self.assertEqual(history.query()['total'], 1)
    def test_final_only_stages_then_saves_one_file_with_cumulative_confirmed_changes(self):
        path = self.doc('点间距为2.5毫米。'); original=Path(path).read_bytes(); target=self.root/'outputs';target.mkdir()
        class Model:
            def chat_json(self,prompt,**kwargs):
                return {'translation':'The flexibile dot spacing is 2.5 mm.'} if '【待译原文】' in prompt else [{'found':'flexibile','suggested':'flexible','kind':'typo','reason':'拼写错误'}]
        job={'id':'final-only','status':'running','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=self.glossary),patch('backend.translation.LLMClient',return_value=Model()),patch('backend.matcher.LLMClient',return_value=Model()):
            workflow.run(job,{'kind':'translate_review','paths':[path],'output_dir':str(target),'combined_output':'final'})
        report=job['result']['report']; staged=job['result']['outputs'][0]['path']
        self.assertEqual(list(target.iterdir()),[]);self.assertTrue(Path(staged).exists())
        self.assertEqual(history.query()['records'][0]['outputs'],[])
        self.assertTrue(history.query()['records'][0]['awaiting_confirmation'])
        first=next(i for i,f in enumerate(report['findings']) if f['method']=='rule')
        a=workflow.apply_report({'report_id':job['id'],'path':staged,'edits':[{'index':first,'text':'pixel pitch'}]})
        self.assertEqual(Document(a['output']).paragraphs[0].text,'The flexibile pixel pitch is 2.5 mm.')
        b=workflow.apply_report({'report_id':job['id'],'path':staged,'edits':[{'index':1-first,'text':'flexible'}]})
        self.assertEqual(a['output'],b['output']);self.assertTrue(b['overwritten'])
        self.assertEqual(len(list(target.iterdir())),1)
        self.assertEqual(Document(b['output']).paragraphs[0].text,'The flexible pixel pitch is 2.5 mm.')
        self.assertEqual(Path(path).read_bytes(),original)
        row=history.query()['records'][0];self.assertEqual(len(row['outputs']),1);self.assertFalse(row['awaiting_confirmation'])
        self.assertEqual(row['counts']['replaced'],2)
    def test_final_only_can_save_without_accepting_any_suggestion_and_detects_external_changes(self):
        path=self.doc('点间距为2.5毫米。');target=self.root/'outputs';target.mkdir()
        class Model:
            def chat_json(self,prompt,**kwargs):
                return {'translation':'The dot spacing is 2.5 mm.'} if '【待译原文】' in prompt else []
        job={'id':'no-changes','status':'running','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary',return_value=self.glossary),patch('backend.translation.LLMClient',return_value=Model()),patch('backend.matcher.LLMClient',return_value=Model()):
            workflow.run(job,{'kind':'translate_review','paths':[path],'output_dir':str(target),'combined_output':'final'})
        staged=job['result']['outputs'][0]['path']
        saved=workflow.apply_report({'report_id':job['id'],'path':staged,'save_unchanged':True})
        self.assertEqual(saved['applied_count'],0);self.assertFalse(saved['finding_updates'][0]['finding']['applied'])
        self.assertEqual(Document(saved['output']).paragraphs[0].text,'The dot spacing is 2.5 mm.')
        doc=Document(saved['output']);doc.add_paragraph('external edit');doc.save(saved['output'])
        with self.assertRaisesRegex(ValueError,'上次输出'):
            workflow.apply_report({'report_id':job['id'],'path':staged,'edits':[{'index':0,'text':'pixel pitch'}]})
        self.assertEqual(len(list(target.iterdir())),1)
    def test_combined_cancel_keeps_output_and_report_but_not_completed_path(self):
        path = self.doc('点间距为2.5毫米。')
        job = {'id':'cancel','status':'running','cancel':threading.Event()}
        class Model:
            def chat_json(self, prompt, **kwargs):
                if '【待译原文】' in prompt: return {'translation':'The dot spacing is 2.5 mm.'}
                job['cancel'].set(); return []
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary), patch('backend.translation.LLMClient', return_value=Model()), patch('backend.matcher.LLMClient', return_value=Model()):
            workflow.run(job, {'kind':'translate_review','paths':[path]})
        self.assertEqual(job['status'], 'cancelled')
        self.assertEqual(job['result']['report']['stats']['model_status'], 'cancelled')
        self.assertEqual(job['result']['completed_paths'], [])
        self.assertTrue(Path(job['result']['outputs'][0]['path']).is_file())
        self.assertFalse(job['result']['outputs'][0]['review_exported'])
    def test_completion_status_waits_for_report_and_history(self):
        path = self.doc('dot spacing'); job = {'id':'timing','status':'running','cancel':threading.Event()}
        finish = history.finish
        def verify_ready(snapshot, data):
            self.assertEqual(job['status'], 'running')
            self.assertIn('timing', workflow.REPORTS)
            self.assertIn('findings', job['result'])
            finish(snapshot, data)
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary), patch('backend.workflow.history.finish', side_effect=verify_ready):
            workflow.run(job, {'kind':'check','paths':[path],'use_llm':False})
        self.assertEqual(job['status'], 'complete')
    def test_only_fully_checked_files_are_consumed_in_a_partial_batch(self):
        first = self.doc('First file dot spacing')
        second = str(self.root / 'second.docx'); doc = Document(); doc.add_paragraph('Second file dot spacing'); doc.save(second)
        class Model:
            def chat_json(self, prompt, **kwargs):
                if '【待审校段落】\nSecond' in prompt: raise ValueError('fixture unavailable')
                return []
        job = {'id':'partial','status':'running','cancel':threading.Event()}
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary), patch('backend.matcher.LLMClient', return_value=Model()):
            workflow.run(job, {'kind':'check','paths':[first,second]})
        self.assertEqual(job['result']['completed_paths'], [first])
        self.assertEqual(job['result']['stats']['model_status'], 'partial')
    def test_legacy_exports_merge_once_and_preserve_parent_date_and_outputs(self):
        history.begin('review', {'kind':'check','paths':['original.pptx']})
        history.finish({'id':'review','status':'complete','result':{'stats':{'judged':12}}}, {'kind':'check','paths':[]})
        before = history.query()['records'][0]
        for index in range(2):
            row = history.begin('old-export-'+str(index), {'kind':'export','paths':['original.pptx']})
            row.update(report_id='review', status='complete', finished_at=row['started_at'],
                       outputs=[{'path':'result-'+str(index)+'.pptx','source':'original.pptx'}], counts={'replaced':index+1})
            with history.connection() as db:
                db.execute('UPDATE tasks SET status=?,payload=? WHERE id=?', ('complete',json.dumps(row),row['id']))
        for _ in range(2):
            records = history.query()['records']; self.assertEqual(len(records), 1)
            row = records[0]
            self.assertEqual(row['id'], 'review'); self.assertEqual(row['started_at'], before['started_at'])
            self.assertEqual(row['finished_at'], before['finished_at'])
            self.assertEqual(row['counts'], {'judged':12,'translated':0,'needs_review':0,'replaced':3})
            self.assertEqual(len(row['exports']), 2); self.assertEqual(len(row['outputs']), 2)
        self.assertEqual(history.query(kind='check',text='result-1.pptx')['total'], 1)
    def test_legacy_exports_with_missing_parent_keep_all_file_metadata(self):
        for index in range(2):
            row = history.begin('orphan-'+str(index), {'kind':'export','paths':['source-'+str(index)+'.pptx']})
            row.update(report_id='missing-parent',status='complete',finished_at=row['started_at'],
                       outputs=[{'path':'saved-'+str(index)+'.pptx','source':row['paths'][0]}],counts={'replaced':2})
            with history.connection() as db:
                db.execute('UPDATE tasks SET status=?,payload=? WHERE id=?', ('complete',json.dumps(row),row['id']))
        result = history.query(); self.assertEqual(result['total'], 1)
        row = result['records'][0]
        self.assertTrue(row['review_metadata_missing']); self.assertEqual(row['kind'],'check')
        self.assertEqual(row['counts']['replaced'],4); self.assertEqual(len(row['paths']),2)
        self.assertEqual(len(row['outputs']),2); self.assertEqual(len(row['exports']),2)
    def test_team_confirmed_terms_are_present_on_first_use(self):
        entries = store.read()['entries']
        for zh, en in [('缩放显示', 'Scaling display'), ('认证机构', 'Certification Institution')]:
            self.assertEqual(next(entry['en'] for entry in entries if entry['zh'] == zh), en)


if __name__ == '__main__': unittest.main()
