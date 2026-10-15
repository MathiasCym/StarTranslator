"""Local HTTP failures and batched/serial document equivalence; no external service."""
import http.client
import io
import json
import os
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from backend.llm.client import LLMClient, LLMUnavailable, LLMRequestFailed
from backend.glossary.loader import Glossary
from backend.translation import translate_file
from backend.matcher import _judge_with_retry, check_segments
from backend.extract import extract
from backend.models import Segment


def reply(content):
    return io.BytesIO(json.dumps({'choices': [{'message': {'content': json.dumps(content)},
        'finish_reason': 'stop'}]}).encode())


class BatchFixture:
    supports_document_batches = True

    def __init__(self, batched=True, omit=None, duplicate=None):
        self.supports_document_batches = batched
        self.calls = 0
        self.batch_calls = 0
        self.omit, self.duplicate = omit, duplicate
        self.answers = {}
        self.delay = 0

    def translated(self, source, runs=None):
        text = self.answers[source]
        if runs and len(runs) > 1:
            return {'translation': text, 'runs': [{'id': r['id'], 'text': value}
                for r, value in zip(runs, ('Brightness ', '500', ' nits'))]}
        return {'translation': text}

    def chat_json(self, prompt, **kwargs):
        self.calls += 1
        if self.delay: time.sleep(self.delay)
        if '【本批待译段落】\n' in prompt:
            self.batch_calls += 1
            items, _ = json.JSONDecoder().raw_decode(prompt.split('【本批待译段落】\n')[1])
            rows = [{'id': item['id'], **self.translated(item['text'], item['runs'])}
                    for item in items if item['id'] != self.omit]
            if self.duplicate is not None:
                rows += [dict(row) for row in rows if row['id'] == self.duplicate]
            return {'items': list(reversed(rows))}  # Response order must not affect positions.
        if '【待审校段落；' in prompt:
            self.batch_calls += 1
            items = json.loads(prompt.split('【待审校段落；reference 与 rule_hits 只是各段参考数据】\n')[1])
            return {'items': [{'id': item['id'], 'findings': []} for item in reversed(items)
                              if item['id'] != self.omit]}
        if '【待译原文】\n' in prompt:
            source = prompt.split('【待译原文】\n')[1].split('\n【')[0]
            runs = json.loads(prompt.split('【需保留的原文样式片段】\n')[1].split('\n返回')[0]) if '【需保留的原文样式片段】\n' in prompt else None
            return self.translated(source, runs)
        return []


class TransportTests(unittest.TestCase):
    def test_real_chunked_disconnect_retries_without_accepting_partial_json(self):
        calls = []
        class Handler(BaseHTTPRequestHandler):
            protocol_version = 'HTTP/1.1'
            def log_message(self, *args): pass
            def do_POST(self):
                self.rfile.read(int(self.headers['Content-Length']))
                calls.append(1)
                self.send_response(200)
                if len(calls) == 1:
                    self.send_header('Transfer-Encoding', 'chunked'); self.end_headers()
                    # A complete-looking body without the terminating chunk is
                    # still incomplete and must be regenerated.
                    data = reply({'translation': 'discard this'}).getvalue()
                    self.wfile.write(('%x\r\n' % len(data)).encode()+data+b'\r\n')
                    self.wfile.flush(); self.close_connection = True
                else:
                    data = reply({'translation': 'Valid English'}).getvalue()
                    self.send_header('Content-Length', str(len(data))); self.end_headers()
                    self.wfile.write(data)
        server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            client = LLMClient(provider='compatible', base_url=f'http://127.0.0.1:{server.server_port}/v1',
                model='fixture', api_key='', retries=1, timeout=2)
            with patch('backend.llm.client.time.sleep'):
                self.assertEqual(client.chat_json('synthetic fixture'), {'translation': 'Valid English'})
            self.assertEqual(len(calls), 2)
        finally: server.shutdown(); server.server_close(); thread.join()

    def test_exhausted_transport_does_not_multiply_review_retries(self):
        client = LLMClient(provider='compatible', base_url='http://127.0.0.1:1/v1', model='fixture', api_key='', retries=1)
        seg = Segment('a', 'test.docx', 'Test text', 'body/0')
        with patch('backend.llm.client._urlopen', side_effect=http.client.IncompleteRead(b'private payload')) as send, \
                patch('backend.llm.client.time.sleep'), self.assertRaises(LLMRequestFailed) as error:
            _judge_with_retry(seg, Glossary([]), client, '', [], None)
        self.assertEqual(send.call_count, 2)
        self.assertNotIn('private', str(error.exception))
        self.assertIn('传输中断', str(error.exception))

    def test_cancel_interrupts_backoff_and_prevents_another_send(self):
        client = LLMClient(provider='compatible', base_url='http://127.0.0.1:1/v1', model='fixture', api_key='')
        client.cancel = threading.Event()
        def broken(*args, **kwargs):
            client.cancel.set(); raise http.client.IncompleteRead(b'')
        with patch('backend.llm.client._urlopen', side_effect=broken) as send, self.assertRaisesRegex(LLMRequestFailed, '取消'):
            client.chat_json('fixture')
        self.assertEqual(send.call_count, 1)

    def test_persistent_disconnect_keeps_completed_translation(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {'NOVATERM_DATA_DIR': tmp}):
            source = Path(tmp)/'source.docx'; doc = Document()
            doc.add_paragraph('亮度500尼特'); doc.add_paragraph('支持图像处理'); doc.save(source)
            original = source.read_bytes()
            client = LLMClient(provider='compatible', base_url='http://127.0.0.1:1/v1', model='fixture', api_key='', retries=1)
            with patch('backend.llm.client._urlopen', side_effect=[reply({'translation': 'Brightness 500 nits'}),
                    http.client.IncompleteRead(b''), http.client.IncompleteRead(b'')]), patch('backend.llm.client.time.sleep'):
                result = translate_file(str(source), Glossary([]), None, client=client, workers=1)
            self.assertEqual([p.text for p in Document(result['path']).paragraphs], ['Brightness 500 nits', '支持图像处理'])
            self.assertEqual(result['status'], 'needs_review')
            self.assertEqual(result['needs_review'], 1)
            self.assertEqual(source.read_bytes(), original)


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        self.environment = patch.dict(os.environ, {'NOVATERM_DATA_DIR': str(self.root/'profile')})
        self.environment.start(); self.addCleanup(self.environment.stop)
        for target in ('backend.translation.review_pptx', 'backend.workflow.review_pptx'):
            stub = patch(target, return_value=([], [])); stub.start(); self.addCleanup(stub.stop)

    def document(self, count=8):
        doc = Document(); answers = {}
        for i in range(count):
            text = f'技术说明第{i+1}项'; doc.add_paragraph(text); answers[text] = f'Technical item {i+1}'
        source = self.root/'source.docx'; doc.save(source)
        return source, answers

    def test_missing_or_duplicate_batch_id_falls_back_only_for_that_paragraph(self):
        source, answers = self.document()
        for option in ({'omit': '2'}, {'duplicate': '2'}):
            model = BatchFixture(**option); model.answers = answers
            result = translate_file(str(source), Glossary([]), None, client=model)
            self.assertEqual(model.calls, 2)
            self.assertEqual([p.text for p in Document(result['path']).paragraphs], list(answers.values()))
            self.assertEqual(result['needs_review'], 0)

    def test_batch_and_serial_pptx_parts_match_including_styles_geometry_and_images(self):
        deck = Presentation(); answers = {}
        from PIL import Image
        image = self.root/'picture.png'; Image.new('RGB', (80, 60), 'blue').save(image)
        for page in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(image), Inches(7), Inches(1), width=Inches(1))
            for i in range(8):
                box = slide.shapes.add_textbox(Inches(.2), Inches(.2+i*.7), Inches(6), Inches(.6))
                paragraph = box.text_frame.paragraphs[0]
                if i == 0:
                    for text, bold in [('亮度', False), ('500', True), ('尼特', False)]:
                        run = paragraph.add_run(); run.text = text; run.font.bold = bold
                        run.font.name = 'Arial'; run.font.size = Pt(18)
                    answers['亮度500尼特'] = 'Brightness 500 nits'
                else:
                    text = f'技术说明第{page*8+i}项'; paragraph.text = text
                    answers[text] = f'Technical item {page*8+i}'
        source = self.root/'source.pptx'; deck.save(source); original = source.read_bytes()
        outputs, calls = [], []
        for enabled in (False, True):
            model = BatchFixture(enabled); model.answers = answers
            result = translate_file(str(source), Glossary([]), None, self.root/str(enabled), client=model)
            outputs.append(result['path']); calls.append(model.calls)
            self.assertEqual(result['remaining_chinese'], [])
            self.assertEqual(result['needs_review'], 0)
        with ZipFile(outputs[0]) as serial, ZipFile(outputs[1]) as batch:
            self.assertEqual(set(serial.namelist()), set(batch.namelist()))
            self.assertTrue(all(serial.read(name) == batch.read(name) for name in serial.namelist()))
        self.assertEqual(source.read_bytes(), original)
        self.assertLess(calls[1], calls[0])

    def test_batch_review_has_explicit_coverage_and_missing_id_is_reviewed_separately(self):
        source, answers = self.document()
        doc = Document(source)
        for paragraph, text in zip(doc.paragraphs, answers.values()): paragraph.text = text
        doc.save(source); segments = extract(str(source))
        model = BatchFixture(omit=segments[2].id)
        findings, stats = check_segments(segments, Glossary([]), client=model)
        self.assertEqual(model.calls, 2)
        self.assertEqual(stats['judged'], len(segments))
        self.assertEqual(stats['model_status'], 'complete')
        self.assertEqual(findings, [])

    def test_batch_review_cannot_move_a_finding_to_another_paragraph(self):
        from backend.matcher.batch import review_batch
        segments = [Segment(str(i), 'test.pptx', text, f'slide1/shape{i}/p0', slide=1)
                    for i, text in enumerate(['Correct heading', 'pythsical output'])]
        class Model:
            def chat_json(self, *args, **kwargs):
                finding = {'found': 'pythsical', 'suggested': 'physical', 'kind': 'typo'}
                return {'items': [{'id': '0', 'findings': [finding]}, {'id': '1', 'findings': [finding]}]}
        values = review_batch([(seg, '', []) for seg in segments], Glossary([]), Model(), '')
        self.assertEqual(values['0'], ([], False))
        self.assertEqual(values['1'][0][0]['suggested'], 'physical')

    def test_thinking_is_used_only_for_failed_ppt_repairs(self):
        from backend.translation import mapped_translation
        group=[{'id':0,'text':'分辨率3840×2160'}]
        segment=Segment('p','fixture.pptx',group[0]['text'],'slide1/shape1/p0',slide=1)
        class Repair:
            supports_thinking_mode=True
            def __init__(self):self.calls=[]
            def chat_json(self,prompt,**kwargs):
                self.calls.append(kwargs)
                return {'translation':'Resolution 1920 x 1080' if len(self.calls)==1 else 'Resolution 3840 x 2160'}
        client=Repair()
        result,_=mapped_translation(client,'fixture',group[0]['text'],group,None,segment=segment)
        self.assertEqual(''.join(r['text'] for r in result),'Resolution 3840 x 2160')
        self.assertNotIn('thinking',client.calls[0]);self.assertTrue(client.calls[1]['thinking'])
        segment.file='fixture.docx';client=Repair()
        mapped_translation(client,'fixture',group[0]['text'],group,None,segment=segment)
        self.assertTrue(all('thinking' not in call for call in client.calls))

    def test_batch_keeps_each_paragraphs_full_reference_and_glossary(self):
        from backend.translation_batch import translate_batch
        from backend.matcher.batch import review_batch
        from backend.document_context import DocumentContext
        segments=[Segment(str(i),'fixture.pptx','技术说明'+str(i),'slide1/shape'+str(i)+'/p0',slide=1) for i in range(2)]
        class References:
            def context(self,seg,ordinal):return ('reference-'+str(ordinal)+' ')*750
        class Recorder:
            def __init__(self):self.prompts=[]
            def chat_json(self,prompt,**kwargs):self.prompts.append(prompt);return {'items':[]}
        entries=[(i,seg,[{'id':0,'text':seg.text}]) for i,seg in enumerate(segments)]
        references=References();client=Recorder()
        with patch('backend.translation_batch._build_glossary_text',side_effect=['term-one','term-two']):
            translate_batch(client,entries,Glossary([]),references,segments,DocumentContext(segments),'fixture system')
        for i in range(2):self.assertIn(references.context(segments[i],i),client.prompts[0])
        self.assertIn('term-one',client.prompts[0]);self.assertIn('term-two',client.prompts[0])
        client=Recorder()
        with patch('backend.matcher.batch._build_glossary_text',side_effect=['term-one','term-two']):
            review_batch([(seg,references.context(seg,i),[]) for i,seg in enumerate(segments)],Glossary([]),client,'page')
        for i in range(2):self.assertIn(references.context(segments[i],i),client.prompts[0])
        client=Recorder()
        self.assertEqual(review_batch([(segments[0],'x'*40000,[]),(segments[1],'y',[])],Glossary([]),client,'page'),{})
        self.assertEqual(client.prompts,[])

    def test_content_and_style_failure_use_at_most_one_thinking_repair(self):
        from backend.translation import mapped_translation
        groups=[{'id':0,'text':'亮度'},{'id':1,'text':'500尼特'}]
        segment=Segment('p','fixture.pptx','亮度500尼特','slide1/shape1/p0',slide=1)
        class Repair:
            supports_thinking_mode=True
            def __init__(self):self.calls=[]
            def chat_json(self,prompt,**kwargs):
                self.calls.append(kwargs)
                if len(self.calls)==1:return {'translation':'Brightness 600 nits'}
                if len(self.calls)==2:return {'translation':'Brightness 500 nits'}
                return {'runs':[{'id':0,'text':'Brightness '},{'id':1,'text':'500 nits'}]}
        client=Repair()
        runs,_=mapped_translation(client,'fixture',segment.text,groups,None,segment=segment)
        self.assertEqual(''.join(run['text'] for run in runs),'Brightness 500 nits')
        self.assertEqual(sum(bool(call.get('thinking')) for call in client.calls),1)


class CheckpointTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')})
        env.start();self.addCleanup(env.stop)
        self.source=self.root/'original.docx';doc=Document()
        doc.add_paragraph('亮度500尼特');doc.add_paragraph('支持图像处理');doc.save(self.source)
        self.glossary=Glossary([])

    def client(self,model='fixture'):
        return LLMClient(provider='compatible',base_url='http://127.0.0.1:1/v1',model=model,api_key='',retries=0)

    def translate(self,client):
        return translate_file(str(self.source),self.glossary,None,client=client,workers=1,cancel=getattr(client,'cancel',None))

    def interrupted(self,cancel=False):
        client=self.client();client.cancel=threading.Event()
        def send(*args,**kwargs):
            if not getattr(send,'done',False):
                send.done=True
                if cancel:client.cancel.set()
                return reply({'translation':'Brightness 500 nits'})
            raise http.client.IncompleteRead(b'')
        with patch('backend.llm.client._urlopen',side_effect=send):
            if cancel:
                from backend.translation import Cancelled
                with self.assertRaises(Cancelled):self.translate(client)
            else:self.translate(client)

    def test_new_client_restores_successful_translation_after_disconnect_or_cancel(self):
        for cancelled in (False,True):
            with self.subTest(cancelled=cancelled):
                # Separate profiles make both cases exercise their first request.
                with patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/str(cancelled))}):
                    self.interrupted(cancelled)
                    with patch('backend.llm.client._urlopen',return_value=reply({'translation':'Supports image processing'})) as send:
                        result=self.translate(self.client())
                    self.assertEqual(send.call_count,1)
                    self.assertEqual(result['resumed_segments'],1)
                    self.assertEqual(result['needs_review'],0)
                    self.assertEqual([p.text for p in Document(result['path']).paragraphs],['Brightness 500 nits','Supports image processing'])

    def test_changed_source_glossary_or_model_invalidates_old_translation(self):
        from backend.models import GlossaryEntry
        for change in ('source','glossary','model'):
            with self.subTest(change=change),patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/change)}):
                self.glossary=Glossary([]);self.interrupted()
                if change=='source':
                    doc=Document(self.source);doc.paragraphs[0].text='亮度600尼特';doc.save(self.source)
                if change=='glossary':self.glossary=Glossary([GlossaryEntry('图像处理','image processing',priority=1)])
                brightness='600' if change=='source' else '500'
                with patch('backend.llm.client._urlopen',side_effect=[reply({'translation':'Brightness '+brightness+' nits'}),
                        reply({'translation':'Supports image processing'})]) as send:
                    result=self.translate(self.client('new-model' if change=='model' else 'fixture'))
                self.assertEqual(send.call_count,2)
                self.assertEqual(result['resumed_segments'],0)
                if change=='source':
                    doc=Document(self.source);doc.paragraphs[0].text='亮度500尼特';doc.save(self.source)

    def test_review_resumes_confirmed_hits_and_clean_paragraphs_without_applying_suggestions(self):
        doc=Document(self.source);doc.paragraphs[0].text='pythsical output';doc.paragraphs[1].text='Image processing';doc.save(self.source)
        segments=extract(str(self.source))
        finding={'found':'pythsical','suggested':'physical','kind':'typo','reason':'拼写错误'}
        with patch('backend.llm.client._urlopen',side_effect=[reply([finding]),http.client.IncompleteRead(b'')]):
            hits,stats=check_segments(segments,self.glossary,client=self.client(),workers=1)
        self.assertEqual(stats['model_status'],'partial')
        self.assertEqual(stats['judged'],1)
        with patch('backend.llm.client._urlopen',return_value=reply([])) as send:
            hits,stats=check_segments(segments,self.glossary,client=self.client(),workers=1)
        self.assertEqual(send.call_count,1)
        self.assertEqual(stats['resumed_segments'],1)
        self.assertEqual(stats['model_status'],'complete')
        self.assertEqual(hits[0].suggested,'physical');self.assertFalse(hits[0].applied)
        with patch('backend.llm.client._urlopen') as send:
            _,stats=check_segments(segments,self.glossary,client=self.client(),workers=1)
        send.assert_not_called();self.assertEqual(stats['resumed_segments'],2)
        self.assertEqual(Document(self.source).paragraphs[0].text,'pythsical output')

    def test_history_restores_settings_but_never_credentials(self):
        from backend import history
        data={'kind':'translate_review','paths':[str(self.source)],'references':[],
              'pairs':{},'combined_output':'final','workers':3,'translation_mode':'faithful',
              '_client':self.client(),'api_key':'must-not-persist'}
        history.begin('resume-fixture',data)
        restored=history.resume_options('resume-fixture')
        self.assertEqual(restored['combined_output'],'final');self.assertEqual(restored['workers'],3)
        self.assertNotIn('must-not-persist',json.dumps(history.query()))
        self.assertNotIn('_client',restored)

    def test_combined_final_only_resume_reuses_translation_and_review_in_new_job(self):
        from backend import workflow
        before=self.source.read_bytes()
        def run(ident,client):
            job={'id':ident,'kind':'translate_review','status':'running','percent':0,'error':'','cancel':threading.Event()}
            data={'kind':'translate_review','paths':[str(self.source)],'references':[],
                'output_dir':str(self.root/'output'),'combined_output':'final','workers':1,'_client':client}
            with patch('backend.workflow.store.load_glossary',return_value=self.glossary):workflow.run(job,data)
            return job
        responses=[reply({'translation':'Brightness 500 nits'}),reply({'translation':'Supports image processing'}),reply([]),http.client.IncompleteRead(b'')]
        with patch('backend.llm.client._urlopen',side_effect=responses):first=run('first-attempt',self.client())
        self.assertEqual(first['result']['report']['stats']['model_status'],'partial')
        with patch('backend.llm.client._urlopen',return_value=reply([])) as send:second=run('resume-attempt',self.client())
        self.assertEqual(send.call_count,1)
        self.assertEqual(second['result']['outputs'][0]['resumed_segments'],2)
        self.assertEqual(second['result']['report']['stats']['resumed_segments'],1)
        self.assertEqual(second['result']['report']['stats']['model_status'],'complete')
        self.assertNotEqual(first['result']['outputs'][0]['path'],second['result']['outputs'][0]['path'])
        self.assertEqual(self.source.read_bytes(),before)
        self.assertFalse(list((self.root/'output').glob('*')))  # Final-only remains private until confirmation.


if __name__ == '__main__': unittest.main()
