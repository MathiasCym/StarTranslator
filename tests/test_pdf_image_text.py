"""PDF figure translation is transactional and must not alter other content."""
import io,json,os,sys,tempfile,threading,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
sys.path[:0]=[str(Path(__file__).resolve().parents[1]),str(Path(__file__).resolve().parent)]
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from backend.pdf_images import append_images
from backend import pdf_image_text as module
from backend.glossary.loader import Glossary
from backend.llm.client import LLMClient,LLMUnavailable


def picture():
    im=Image.new('RGBA',(600,100),(0,0,0,0));draw=ImageDraw.Draw(im)
    font=module._font(40);draw.text((18,20),'source label',font=font,fill='white',stroke_width=1,stroke_fill='gray')
    stream=io.BytesIO();im.save(stream,format='PNG');return stream.getvalue()


OCR={'width':600,'height':100,'angle':0,'lines':[{'id':0,'text':'方寸之间体验非凡','box':[12,10,580,90]}]}


class Model:
    def __init__(self,scan=None,final=None,text='Extraordinary experiences in every space'):
        self.scan=scan or {'reviewed':True,'lines':[{'id':0,'text':'方寸之间体验非凡'}],'unlocated_chinese':[]}
        self.final=final or {'reviewed':True,'remaining_chinese':[],'meaning_preserved':True,'other_content_preserved':True,'readable':True}
        self.text=text;self.images=[];self.prompts=[]
    def chat_images_json(self,prompt,images,**kw):self.images.append((prompt,images));return self.scan if len(images)==1 else self.final
    def chat_json(self,prompt,**kw):self.prompts.append(prompt);return {'translation':self.text}


class PdfImageTextTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});self.env.start()
        self.source=self.root/'source.pdf';self.source.write_bytes(b'local-fixture-not-sent')
        self.g=Glossary([])
    def tearDown(self):self.env.stop();self.tmp.cleanup()
    def document(self,raw=None):
        word=Document();raw=raw or picture()
        count,issues=append_images(word,SimpleNamespace(images=[SimpleNamespace(data=raw)]),3)
        self.assertEqual((count,issues),(1,[]))
        return SimpleNamespace(document=word)
    def blob(self,doc):
        shape=doc.document.inline_shapes[0]
        return doc.document.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob
    def review(self,doc,model,ocr=None,**kw):
        with patch.object(module,'view_image',side_effect=lambda raw:(raw,1)),patch.object(module,'recognize',return_value=ocr or OCR):
            return module.review_images(doc,self.source,model,self.g,[],**kw)

    def test_verified_translation_changes_only_selected_pixels_and_keeps_geometry(self):
        doc=self.document();before=self.blob(doc);shape=doc.document.inline_shapes[0];geometry=(shape.width,shape.height)
        result=self.review(doc,Model())
        self.assertEqual(result['status'],'complete');self.assertEqual(result['translated_lines'],1)
        after=self.blob(doc);self.assertNotEqual(before,after)
        self.assertEqual((shape.width,shape.height),geometry)
        with Image.open(io.BytesIO(before)) as a,Image.open(io.BytesIO(after)) as b:
            self.assertTrue(module.outside_unchanged(a,b,[(10,8,582,92)]))
        self.assertEqual(self.source.read_bytes(),b'local-fixture-not-sent')

    def test_rejected_after_image_restores_original_not_a_half_edited_image(self):
        doc=self.document();before=self.blob(doc)
        model=Model(final={'reviewed':True,'remaining_chinese':['遗留'],'meaning_preserved':True,'other_content_preserved':True,'readable':True})
        result=self.review(doc,model)
        self.assertEqual(result['status'],'partial');self.assertEqual(result['translated_images'],0)
        self.assertEqual(self.blob(doc),before)

    def test_no_chinese_keeps_the_exact_image_bytes(self):
        doc=self.document();before=self.blob(doc);model=Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':[]})
        result=self.review(doc,model,ocr={**OCR,'lines':[{'id':0,'text':'Existing English','box':[12,10,580,90]}]})
        self.assertEqual(result['checked'],1);self.assertEqual(result['translated_images'],0)
        self.assertEqual(self.blob(doc),before);self.assertEqual(model.prompts,[])

    def test_scan_does_not_receive_page_prose_that_can_hallucinate_picture_text(self):
        model=Model();doc=self.document()
        with patch.object(module,'view_image',side_effect=lambda raw:(raw,1)),patch.object(module,'recognize',return_value=OCR):
            result=module.review_images(doc,self.source,model,self.g,
                [SimpleNamespace(slide=3,source='text',text='该页正文并不在图片里')])
        self.assertEqual(result['status'],'complete')
        self.assertNotIn('该页正文并不在图片里',model.images[0][0])
        self.assertIn('该页正文并不在图片里',model.prompts[0])

    def test_missing_vision_does_not_masquerade_as_no_chinese(self):
        doc=self.document();before=self.blob(doc)
        result=self.review(doc,object())
        self.assertEqual(result['status'],'partial');self.assertEqual(result['checked'],0)
        self.assertEqual(self.blob(doc),before)

    def test_unlocated_chinese_is_reported_without_arbitrary_box_guesses(self):
        doc=self.document();before=self.blob(doc)
        model=Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':['还有文字']})
        result=self.review(doc,model)
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def test_unlocated_visual_false_positive_requires_a_separate_confirmation(self):
        class Recheck(Model):
            def chat_images_json(self,prompt,images,**kw):
                self.images.append((prompt,images))
                return ({'reviewed':True,'lines':[],'unlocated_chinese':['大']}
                    if len(self.images)==1 else {'reviewed':True,'confirmed_chinese':[],'false_positives':['大']})
        doc=self.document();before=self.blob(doc)
        result=self.review(doc,Recheck(),ocr={**OCR,'lines':[]})
        self.assertEqual(result['status'],'complete');self.assertEqual(result['rechecked'],1)
        self.assertEqual(self.blob(doc),before)

    def test_confirmed_unlocated_chinese_is_not_dismissed_as_a_false_positive(self):
        class Recheck(Model):
            def chat_images_json(self,prompt,images,**kw):
                self.images.append((prompt,images))
                return ({'reviewed':True,'lines':[],'unlocated_chinese':['中文']}
                    if len(self.images)==1 else {'reviewed':True,'confirmed_chinese':['中文'],'false_positives':[]})
        doc=self.document();before=self.blob(doc);result=self.review(doc,Recheck(),ocr={**OCR,'lines':[]})
        self.assertEqual(result['status'],'partial');self.assertEqual(result['issues'][0]['source_text'],'中文')
        self.assertEqual(self.blob(doc),before)

    def test_bad_model_region_cannot_erase_the_image(self):
        doc=self.document();before=self.blob(doc)
        model=Model(scan={'reviewed':True,'lines':[{'id':500,'text':'方寸之间体验非凡'}],'unlocated_chinese':[]})
        result=self.review(doc,model)
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def test_model_scan_corrects_ocr_han_typos_but_cannot_change_a_model_number(self):
        ocr={'lines':[{'id':0,'text':'有寸之间伸验菲凡 VX1000','box':[0,0,400,50]}]}
        answer={'reviewed':True,'lines':[{'id':0,'text':'方寸之间体验非凡 VX1000'}],'unlocated_chinese':[]}
        self.assertEqual(module._scan(answer,ocr)[0]['text'],'方寸之间体验非凡 VX1000')
        answer['lines'][0]['text']='方寸之间体验非凡 VX2000'
        with self.assertRaises(ValueError):module._scan(answer,ocr)

    def test_translation_cannot_add_numbers_or_leave_chinese(self):
        for text in ['Extraordinary 2000 experiences','还有中文']:
            with self.subTest(text=text):
                doc=self.document();before=self.blob(doc);model=Model(text=text)
                result=self.review(doc,model)
                self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)
                self.assertEqual(len(model.prompts),2)

    def test_low_readability_is_a_failure_instead_of_forcing_tiny_text(self):
        doc=self.document();before=self.blob(doc);result=self.review(doc,Model(text='Long text '*400))
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def test_english_italic_logo_does_not_fail_the_chinese_rotation_check(self):
        doc=self.document();ocr={**OCR,'angle':-12,'lines':[{'id':0,'text':'NOVA STAR','box':[12,10,580,90]}]}
        result=self.review(doc,Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':[]}),ocr=ocr)
        self.assertEqual(result['status'],'complete');self.assertEqual(result['translated_lines'],0)

    def test_native_han_and_empty_model_scan_cannot_silently_pass(self):
        doc=self.document();before=self.blob(doc)
        result=self.review(doc,Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':[]}))
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def test_native_false_positive_requires_explicit_visual_resolution(self):
        doc=self.document();before=self.blob(doc)
        model=Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':[],'non_chinese_ids':[0]})
        result=self.review(doc,model)
        self.assertEqual(result['status'],'complete');self.assertEqual(result['rechecked'],1)
        self.assertEqual(self.blob(doc),before);self.assertEqual(len(model.images),2)

    def test_overlong_slogan_gets_one_faithful_wording_retry(self):
        class Reword(Model):
            def chat_json(self,prompt,**kw):
                self.prompts.append(prompt)
                return {'translation':'Long text '*400 if len(self.prompts)==1 else 'Extraordinary experiences in every space'}
        model=Reword();result=self.review(self.document(),model)
        self.assertEqual(result['status'],'complete');self.assertEqual(len(model.prompts),2)

    def test_changes_outside_text_mask_are_detected(self):
        a=Image.new('RGBA',(100,100),'white');b=a.copy();b.putpixel((99,99),(0,0,0,255))
        self.assertFalse(module.outside_unchanged(a,b,[(0,0,10,10)]))
        for boxes in ([],[[-1,0,10,10]],[[0,0,101,100]],[[0,0,10.5,10]],[[10,0,0,10]]):
            with self.subTest(boxes=boxes):self.assertFalse(module.outside_unchanged(a,b,boxes))

    def test_partially_selected_native_han_still_requires_remaining_line_resolution(self):
        doc=self.document();before=self.blob(doc)
        ocr={**OCR,'lines':[OCR['lines'][0],{'id':1,'text':'另一处中文','box':[12,91,580,99]}]}
        result=self.review(doc,Model(),ocr=ocr)
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def pdf_with_image(self,text=None):
        from test_file_formats import pdf_fixture
        from pypdf import PdfReader,PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,NumberObject,DecodedStreamObject
        if text is None:pdf_fixture(self.source)
        else:pdf_fixture(self.source,text)
        writer=PdfWriter(clone_from=PdfReader(self.source));page=writer.pages[0]
        with Image.open(io.BytesIO(picture())) as im:
            im=im.convert('RGB');image=DecodedStreamObject();image.set_data(im.tobytes())
            image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),
                NameObject('/Width'):NumberObject(im.width),NameObject('/Height'):NumberObject(im.height),
                NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        page['/Resources'][NameObject('/XObject')]=DictionaryObject({NameObject('/Figure'):writer._add_object(image)})
        content=DecodedStreamObject();content.set_data(page.get_contents().get_data()+b'\nq 300 0 0 50 20 100 cm /Figure Do Q')
        page[NameObject('/Contents')]=writer._add_object(content)
        writer.write(self.source)

    def test_failed_image_check_remains_visible_and_does_not_block_review_or_save(self):
        from backend import workflow
        from backend.translation_audit import saved_residuals
        from test_file_formats import ENGLISH
        self.pdf_with_image();original=self.source.read_bytes()
        class TextOnly:
            def chat_json(self,prompt,**kw):return {'translation':ENGLISH} if '【待译原文】' in prompt else []
        job={'id':'image-failure-save','kind':'translate_review','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            workflow.run(job,dict(kind='translate_review',paths=[str(self.source)],references=[],pairs={},
                output_dir=str(self.root/'output'),combined_output='final',use_llm=True,workers=1,_client=TextOnly()))
        result=job['result'];self.assertTrue(result['outputs'],result)
        output=result['outputs'][0];self.assertFalse(output['translation_complete'])
        report=result['report'];issues=[f for f in report['findings'] if f['source']=='image']
        self.assertEqual(len(issues),1);self.assertFalse(issues[0]['editable'])
        self.assertIn('原图',issues[0]['reason']);self.assertEqual(report['stats']['model_status'],'partial')
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],save_unchanged=True))
        self.assertTrue(Path(saved['output']).is_file());self.assertFalse(saved_residuals(saved['output']))
        self.assertEqual(len(Document(saved['output']).inline_shapes),1)
        self.assertEqual(self.source.read_bytes(),original)

    def test_verified_image_is_present_in_the_final_reviewed_file(self):
        from backend import workflow
        from backend.translation_audit import saved_residuals
        from test_file_formats import ENGLISH
        self.pdf_with_image();original=self.source.read_bytes()
        class FullModel(Model):
            def chat_json(self,prompt,**kw):
                if '【待译原文】' in prompt:return {'translation':ENGLISH}
                if '原文：方寸' in prompt:return {'translation':self.text}
                return []
        job={'id':'image-success-save','kind':'translate_review','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=self.g),patch.object(module,'view_image',side_effect=lambda raw:(raw,1)),patch.object(module,'recognize',return_value=OCR),patch('backend.pdf_quality.audit',return_value={'status':'complete','issues':[],'checked_pages':1}):
            workflow.run(job,dict(kind='translate_review',paths=[str(self.source)],references=[],pairs={},
                output_dir=str(self.root/'output'),combined_output='final',use_llm=True,workers=1,_client=FullModel()))
        result=job['result'];self.assertTrue(result['outputs'],result)
        output=result['outputs'][0];self.assertTrue(output['translation_complete'],output)
        self.assertEqual(output['image_review']['translated_images'],1)
        self.assertEqual(result['report']['stats']['model_status'],'complete')
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],save_unchanged=True))
        self.assertFalse(saved_residuals(saved['output']));self.assertEqual(len(Document(saved['output']).inline_shapes),1)
        self.assertEqual(self.source.read_bytes(),original)

    def test_image_only_pdf_uses_the_image_pass_without_inventing_body_text(self):
        from backend.translation import translate_file
        self.pdf_with_image('')
        with patch.object(module,'view_image',side_effect=lambda raw:(raw,1)),patch.object(module,'recognize',return_value=OCR):
            result=translate_file(self.source,self.g,None,client=Model(),output_dir=self.root/'output')
        self.assertEqual(result['translated'],0);self.assertTrue(result['translation_complete'],result)
        self.assertEqual(result['image_review']['translated_images'],1)
        word=Document(result['path']);self.assertEqual(len(word.inline_shapes),1)
        self.assertFalse(any('OCR is required' in p.text for p in word.paragraphs))

    def test_complex_background_is_preserved(self):
        im=Image.new('RGBA',(600,100));p=im.load()
        for y in range(100):
            for x in range(600):p[x,y]=((x*13)%256,(y*7)%256,(x+y)%256,255)
        raw=io.BytesIO();im.save(raw,format='PNG');doc=self.document(raw.getvalue());before=self.blob(doc)
        result=self.review(doc,Model())
        self.assertEqual(result['status'],'partial');self.assertEqual(self.blob(doc),before)

    def test_cancellation_does_not_write_any_image_change(self):
        from backend.translation import Cancelled
        doc=self.document();before=self.blob(doc)
        def stop():raise Cancelled('cancelled')
        with self.assertRaises(Cancelled):self.review(doc,Model(),check_cancel=stop)
        self.assertEqual(self.blob(doc),before)

    def test_verified_image_checkpoint_reuses_without_model_or_ocr_requests(self):
        class CachedModel(Model,LLMClient):
            def __init__(self):
                Model.__init__(self);self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='image-fixture';self.api_format='openai'
        first=self.review(self.document(),CachedModel());self.assertEqual(first['translated_images'],1)
        model=CachedModel()
        def forbidden(*a,**kw):raise AssertionError('No repeated model call')
        model.chat_json=forbidden;model.chat_images_json=forbidden
        with patch.object(module,'recognize',side_effect=forbidden):
            result=module.review_images(self.document(),self.source,model,self.g,[])
        self.assertEqual(result['resumed'],1);self.assertEqual(result['translated_images'],1)

    def test_damaged_image_checkpoint_is_rechecked_instead_of_aborting(self):
        import sqlite3
        from contextlib import closing
        class CachedModel(Model,LLMClient):
            def __init__(self):
                Model.__init__(self);self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='damaged-image';self.api_format='openai'
        self.assertEqual(self.review(self.document(),CachedModel())['translated_images'],1)
        with closing(sqlite3.connect(self.root/'profile/task-checkpoints.sqlite3')) as db,db:
            for key,value in db.execute("select item_key,value from checkpoints where stage='pdf-image'").fetchall():
                value=json.loads(value);value.pop('boxes')
                db.execute("update checkpoints set value=? where item_key=?",(json.dumps(value),key))
        model=CachedModel();result=self.review(self.document(),model)
        self.assertEqual(result['status'],'complete');self.assertEqual(result['resumed'],0)
        self.assertEqual(len(model.images),2)

    def test_unreadable_source_figure_is_not_reported_as_zero_images_success(self):
        word=Document();count,failures=append_images(word,SimpleNamespace(images=[SimpleNamespace(data=b'broken-image')]),2)
        self.assertEqual(count,0);self.assertTrue(failures)
        result=module.review_images(SimpleNamespace(document=word),self.source,Model(),self.g,[])
        self.assertEqual(result['status'],'partial');self.assertEqual(result['issues'][0]['slide'],2)

    def test_missing_native_ocr_is_not_replaced_by_a_successful_model_guess(self):
        model=Model(scan={'reviewed':True,'lines':[],'unlocated_chinese':[]})
        with patch.object(module,'recognize',side_effect=ValueError('OCR language missing')):
            result=module.review_images(self.document(),self.source,model,self.g,[])
        self.assertEqual(result['status'],'partial');self.assertEqual(model.images,[])


if __name__=='__main__':unittest.main()
