"""The normal product, not an agent's inspection script, verifies saved artifacts."""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path[:0]=[str(Path(__file__).resolve().parents[1]),str(Path(__file__).resolve().parent)]
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image,ImageDraw
from backend import pdf_quality as quality,workflow
from backend.glossary.loader import Glossary
from backend.glossary.store import fingerprint
from backend.llm.client import LLMClient,LLMUnavailable
from backend.translation import Cancelled,translate_file
from test_file_formats import pdf_fixture,ENGLISH


class Model:
    def __init__(self,answers=None):self.calls=[];self.answers=list(answers or [])
    def chat_images_json(self,prompt,images,**kw):
        self.calls.append((prompt,images))
        return self.answers.pop(0) if self.answers else dict(reviewed=True,ocr_checks=[],other_chinese=[],layout_issues=[])
    def chat_json(self,prompt,**kw):return {'translation':ENGLISH} if '【待译原文】' in prompt else []


class PdfQualityTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.addCleanup(self.tmp.cleanup)
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});env.start();self.addCleanup(env.stop)
        self.path=self.root/'output.docx';d=Document();d.add_paragraph('Hello 2026')
        raw=io.BytesIO();Image.new('RGB',(80,35),'white').save(raw,format='PNG');d.add_picture(io.BytesIO(raw.getvalue()))
        paragraph=d.add_paragraph();start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'1');start.set(qn('w:name'),'page1');paragraph._p.append(start)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'1');paragraph._p.append(end)
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),d.part.relate_to('https://example.test/原始链接',RT.HYPERLINK,is_external=True));paragraph._p.append(link)
        internal=OxmlElement('w:hyperlink');internal.set(qn('w:anchor'),'page1');paragraph._p.append(internal)
        d.save(self.path);self.expected=quality.manifest(d);self.ocr={'lines':[]}
        self.render_calls=[];self.page_texts=['Hello 2026'];self.visual_words='Hello 2026'
        render=patch.object(quality.native,'preview',side_effect=self.render);render.start();self.addCleanup(render.stop)
        ocr=patch.object(quality,'recognize',side_effect=lambda *a:self.ocr);ocr.start();self.addCleanup(ocr.stop)

    def render(self,path,folder,digest,page):
        from pypdf import PdfReader,PdfWriter
        self.render_calls.append(page);folder.mkdir(parents=True,exist_ok=True)
        writer=PdfWriter()
        for index,text in enumerate(self.page_texts):
            file=self.root/f'page-fixture-{index}.pdf';pdf_fixture(file,text)
            writer.add_page(PdfReader(file).pages[0])
        writer.write(folder/'document.pdf')
        image=Image.new('RGB',(600,800),'white');ImageDraw.Draw(image).text((20,20),self.visual_words,fill='black')
        image.save(folder/f'page-{page}.png')
        return 'native-page',len(self.page_texts)

    def audit(self,client=None,**kw):return quality.audit(self.path,self.expected,source=self.root/'original.pdf',client=client or Model(),**kw)

    def test_actual_saved_file_is_checked_without_editing_it(self):
        before=fingerprint(self.path);model=Model();events=[]
        result=self.audit(model,progress=lambda *args:events.append(args))
        self.assertEqual(result['status'],'complete');self.assertEqual(result['checked_pages'],1)
        self.assertEqual(len(model.calls),1);self.assertEqual(fingerprint(self.path),before)
        self.assertTrue(events);self.assertIn('最终输出页面',model.calls[0][0])

    def test_missing_image_link_bookmark_and_target_are_independently_detected(self):
        for key in ('images','links','bookmarks','targets'):
            with self.subTest(key=key):
                expected={**self.expected,key:['not-the-original']}
                result=quality.audit(self.path,expected,client=Model())
                self.assertEqual(result['status'],'partial');self.assertTrue(any(i['kind']=='structure' for i in result['issues']))

    def test_changed_pixels_or_geometry_do_not_pass_a_successful_model_answer(self):
        d=Document(self.path);d.inline_shapes[0].width+=12700;d.save(self.path)
        result=self.audit();self.assertTrue(any(i['kind']=='structure' for i in result['issues']))

    def test_missing_rendered_word_is_detected_even_when_model_says_no_chinese(self):
        self.page_texts=['Hello'];result=self.audit()
        self.assertTrue(any(i['kind']=='visibility' for i in result['issues']))

    def test_split_style_runs_and_wrapped_words_are_not_false_clipping(self):
        d=Document(self.path);p=d.paragraphs[0];p.clear();p.add_run('Hel');p.add_run('lo 2026');d.save(self.path)
        self.page_texts=['Hel\nlo 2026'];result=self.audit()
        self.assertFalse(any(i['kind']=='visibility' for i in result['issues']))

    def test_actual_chinese_in_saved_text_is_not_suppressed_by_vision(self):
        d=Document(self.path);d.paragraphs[0].add_run('遗漏中文');d.save(self.path)
        result=self.audit();self.assertTrue(any(i['kind']=='chinese' and i['text']=='Hello 2026遗漏中文' for i in result['issues']))

    def test_no_visual_model_does_not_claim_native_review_complete(self):
        result=self.audit(object());self.assertEqual(result['status'],'partial')
        self.assertEqual(result['checked_pages'],0);self.assertEqual(self.render_calls,[])

    def test_native_renderer_unavailable_is_an_explicit_nonblocking_issue(self):
        with patch.object(quality.native,'preview',return_value=None):result=self.audit()
        self.assertEqual(result['status'],'partial');self.assertTrue(self.path.is_file())
        self.assertEqual(result['checked_pages'],0)

    def test_ocr_unavailable_cannot_be_treated_as_an_empty_chinese_list(self):
        model=Model()
        with patch.object(quality,'recognize',side_effect=ValueError('OCR unavailable')):result=self.audit(model)
        self.assertEqual(result['status'],'partial');self.assertEqual(model.calls,[])

    def test_high_dpi_native_pages_are_normalized_before_ocr_and_model(self):
        render=self.render
        dimensions=(2160,3056)
        def large(path,folder,digest,page):
            result=render(path,folder,digest,page)
            Image.new('RGBA',dimensions,'white').save(folder/f'page-{page}.png')
            return result
        def recognize(path,check):
            with Image.open(path) as image:self.assertLessEqual(max(image.size),2500)
            return {'lines':[]}
        with patch.object(quality.native,'preview',side_effect=large),patch.object(quality,'recognize',side_effect=recognize):
            for dimensions in ((1440,2037),(2160,3056),(2880,4074)):
                with self.subTest(dimensions=dimensions):
                    result=self.audit()
                    self.assertEqual(result['status'],'complete');self.assertEqual(result['checked_pages'],1)

    def test_every_native_ocr_suspicion_requires_explicit_visual_resolution(self):
        self.ocr={'lines':[{'id':7,'text':'疑似中文','box':[0,0,60,30]}]}
        result=self.audit();self.assertEqual(result['status'],'partial')
        self.assertTrue(result['issues'][0]['retryable'])
        answer=dict(reviewed=True,ocr_checks=[dict(id=7,chinese=False,text='')],other_chinese=[],layout_issues=[])
        self.assertEqual(self.audit(Model([answer]))['status'],'complete')

    def test_model_finds_chinese_that_native_ocr_missed(self):
        result=self.audit(Model([dict(reviewed=True,ocr_checks=[],other_chinese=['图内中文'],layout_issues=[])]))
        self.assertEqual(result['status'],'partial');self.assertEqual(result['issues'][0]['slide'],1)

    def test_bad_or_duplicate_ocr_ids_do_not_count_as_checked(self):
        for checks in ([dict(id=99,chinese=False)],[dict(id=7,chinese=False)]*2,[dict(id=7,chinese='false')]):
            self.ocr={'lines':[dict(id=7,text='疑似中文',box=[0,0,60,30])]}
            result=self.audit(Model([dict(reviewed=True,ocr_checks=checks,other_chinese=[],layout_issues=[])]))
            self.assertEqual(result['checked_pages'],0);self.assertEqual(result['status'],'partial')

    def test_one_invalid_model_answer_does_not_skip_other_pages(self):
        self.page_texts=['Hello 2026','Second page'];self.visual_words='Hello 2026 Second page'
        model=Model([{}]);result=self.audit(model)
        self.assertEqual(len(model.calls),2);self.assertEqual(result['checked_pages'],1)
        self.assertEqual(result['status'],'partial')

    def test_model_outage_preserves_file_and_marks_remaining_pages(self):
        self.page_texts=['Hello 2026','Second page']
        class Offline(Model):
            def chat_images_json(self,*a,**kw):raise LLMUnavailable('model disconnected')
        result=self.audit(Offline());self.assertEqual(result['checked_pages'],0)
        self.assertEqual(len([i for i in result['issues'] if i['retryable']]),2);self.assertTrue(self.path.is_file())

    def test_cancellation_and_external_modification_are_not_cached_as_success(self):
        cancel=threading.Event();cancel.set()
        with self.assertRaises(Cancelled):self.audit(cancel=cancel)
        owner=self
        class Modify(Model):
            def chat_images_json(self,*a,**kw):
                d=Document(owner.path);d.add_paragraph('External change');d.save(owner.path)
                return super().chat_images_json(*a,**kw)
        with self.assertRaisesRegex(ValueError,'变化'):self.audit(Modify())

    def test_verified_unchanged_pages_resume_without_model_or_ocr(self):
        class Cached(Model,LLMClient):
            def __init__(self):
                Model.__init__(self);self.provider='compatible';self.base_url='http://127.0.0.1:1/v1';self.model='quality-fixture';self.api_format='openai'
        first=self.audit(Cached());self.assertEqual(first['status'],'complete')
        def forbidden(*a,**kw):raise AssertionError('Verified pages must be reused')
        model=Cached();model.chat_images_json=forbidden
        with patch.object(quality,'recognize',side_effect=forbidden):second=self.audit(model)
        self.assertEqual(second['resumed_pages'],1);self.assertEqual(second['status'],'complete')
        self.visual_words='Changed visible output'
        model=Cached();third=self.audit(model)
        self.assertEqual(third['resumed_pages'],0);self.assertEqual(len(model.calls),1)

    def test_translate_review_and_each_save_automatically_call_the_product_checker(self):
        source=self.root/'original.pdf';pdf_fixture(source);self.page_texts=[ENGLISH];self.visual_words=ENGLISH
        job=dict(id='automatic-pdf-quality',kind='translate_review',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='translate_review',combined_output='final',paths=[str(source)],
                references=[],pairs={},output_dir=str(self.root/'output'),workers=1,_client=Model()))
        result=job['result'];output=result['outputs'][0]
        self.assertEqual(output['output_audit']['checked_pages'],1);self.assertEqual(job['status'],'complete',result)
        events=[]
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],save_unchanged=True),
            client=Model(),progress=lambda *args:events.append(args))
        self.assertEqual(saved['output_audit']['checked_pages'],1)
        self.assertTrue(any('实际' in str(e) for e in events));self.assertTrue(Path(saved['output']).is_file())

    def test_saved_checks_use_latest_edited_bytes_and_retain_an_unverified_file(self):
        source=self.root/'original.pdf';pdf_fixture(source);self.page_texts=[ENGLISH];self.visual_words=ENGLISH
        job=dict(id='saved-pdf-quality',kind='translate_review',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='translate_review',combined_output='final',paths=[str(source)],
                references=[],pairs={},output_dir=str(self.root/'output'),workers=1,_client=Model()))
        output=job['result']['outputs'][0]
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],save_unchanged=True),client=object())
        self.assertTrue(Path(saved['output']).is_file());self.assertGreater(saved['needs_review'],0)
        self.assertEqual(saved['output_audit']['status'],'partial');self.assertTrue(saved['content_warnings'])

    def test_manual_replacement_with_chinese_is_detected_in_the_newly_written_file(self):
        from backend.models import Finding
        source=self.root/'original.pdf';pdf_fixture(source);self.page_texts=[ENGLISH];self.visual_words=ENGLISH
        job=dict(id='post-edit-pdf',kind='translate_review',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='translate_review',combined_output='final',paths=[str(source)],
                references=[],pairs={},output_dir=str(self.root/'output'),workers=1,_client=Model()))
        output=job['result']['outputs'][0];report=workflow.REPORTS[job['id']]
        segment=next(s for s in report['segments'] if s.editable and s.text==ENGLISH)
        replacement=ENGLISH+' 未译中文'
        report['findings'].append(Finding(segment_id=segment.id,file=segment.file,file_path=segment.file_path,
            slide=None,para=segment.para,source=segment.source,found_text=ENGLISH,suggested=replacement))
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],edits=[dict(index=0,text=replacement)]),client=Model())
        self.assertIn(replacement,[p.text for p in Document(saved['output']).paragraphs])
        self.assertEqual(saved['output_audit']['status'],'partial');self.assertGreater(saved['untranslated'],0)
        self.assertTrue(any(i['kind']=='chinese' for i in saved['output_audit']['issues']))

    def test_word_report_save_does_not_enter_pdf_output_review(self):
        job=dict(id='word-save-isolated',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])),patch.object(quality,'audit',side_effect=AssertionError('PDF only')):
            workflow.run(job,dict(kind='check',use_llm=False,paths=[str(self.path)]))
            saved=workflow.apply_report(dict(report_id=job['id'],path=str(self.path),save_unchanged=True))
        self.assertIsNone(saved['output_audit']);self.assertTrue(Path(saved['output']).is_file())

    def test_reusable_command_requires_model_authorization_before_reading_inputs(self):
        import subprocess
        tool=Path(__file__).resolve().parents[1]/'tools/check-document-quality.py'
        target=self.root/'not-created'
        result=subprocess.run([sys.executable,str(tool),'--source','missing.pdf','--evidence',str(target)],capture_output=True,text=True)
        self.assertEqual(result.returncode,2);self.assertIn('--allow-model is required',result.stderr)
        self.assertFalse(target.exists())

    def test_word_translation_does_not_enter_the_pdf_checker(self):
        source=self.root/'source.docx';word=Document();word.add_paragraph('点间距为2.5mm。');word.save(source)
        with patch.object(quality,'audit',side_effect=AssertionError('PDF checks must not alter Word')):
            result=translate_file(source,Glossary([]),None,client=Model(),output_dir=self.root/'word-output')
        self.assertIsNone(result['output_audit'])

    def test_rule_only_pdf_save_does_not_call_an_unrequested_model(self):
        source=self.root/'english.pdf';pdf_fixture(source,ENGLISH)
        job=dict(id='rules-only-pdf',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='check',use_llm=False,paths=[str(source)]))
        with patch('backend.workflow.LLMClient',side_effect=AssertionError('Rule-only mode must not start a model')):
            saved=workflow.apply_report(dict(report_id=job['id'],path=str(source),save_unchanged=True))
        self.assertEqual(saved['output_audit']['status'],'partial');self.assertTrue(Path(saved['output']).is_file())

    def test_rule_only_desktop_background_save_cannot_enable_a_model(self):
        import time
        from backend import report_exports
        source=self.root/'english.pdf';pdf_fixture(source,ENGLISH)
        job=dict(id='rules-only-async-pdf',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='check',use_llm=False,paths=[str(source)]))
        class Forbidden(Model):
            def chat_images_json(self,*a,**kw):raise AssertionError('Disabled model must stay disabled in the desktop export worker')
        with patch('backend.report_exports.LLMClient',return_value=Forbidden()):
            task=report_exports.start(dict(report_id=job['id'],path=str(source),save_unchanged=True))
            deadline=time.monotonic()+5
            while task['status'] in ('queued','running') and time.monotonic()<deadline:
                time.sleep(.02);task=report_exports.read(task['id'])
        self.assertEqual(task['status'],'complete',task)
        self.assertEqual(task['result']['output_audit']['status'],'partial')
        self.assertTrue(Path(task['result']['output']).is_file())

    def test_pdf_page_with_unreadable_vector_content_is_reported_without_losing_other_pages(self):
        from pypdf import PdfReader,PdfWriter
        from pypdf.generic import DecodedStreamObject,NameObject
        from backend.formats import pdf_document
        source=self.root/'mixed.pdf';pdf_fixture(source)
        writer=PdfWriter(clone_from=PdfReader(source));page=writer.add_blank_page(600,800)
        stream=DecodedStreamObject();stream.set_data(b'10 10 m 20 20 l S');page[NameObject('/Contents')]=writer._add_object(stream);writer.write(source)
        document=pdf_document(source)
        manifest=quality.manifest(document.document)
        self.assertEqual(len(manifest['source_issues']),1);self.assertEqual(manifest['source_issues'][0]['slide'],2)


if __name__=='__main__':unittest.main()
