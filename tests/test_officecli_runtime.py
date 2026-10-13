"""Offline regression checks for component launches and early task failures."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import officecli_bridge as bridge, workflow, history
from pptx import Presentation
from pptx.util import Inches


class OfficeRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {'NOVATERM_DATA_DIR': str(self.root / 'profile')})
        self.env.start()

    def tearDown(self):
        self.env.stop(); self.temp.cleanup()

    def records(self):
        return [json.loads(line) for line in (self.root / 'profile/logs/officecli.jsonl').read_text(encoding='utf-8').splitlines()]

    def test_real_component_checks_unicode_snapshot_without_changing_source(self):
        source = self.root / '检查样本.pptx'
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Diagnostic 300+'
        deck.save(source); original = source.read_bytes()
        result = bridge.inspect(source)
        self.assertEqual(result['errors'], [])
        self.assertEqual(source.read_bytes(), original)

    def test_explicit_packaged_resource_directory_wins(self):
        component = self.root / 'resources/officecli/officecli.exe'
        component.parent.mkdir(parents=True); component.write_bytes(b'fixture')
        with patch.dict(os.environ, {'STAR_RESOURCE_DIR': str(self.root / 'resources')}):
            self.assertEqual(bridge.binary_path(), component)

    def test_missing_component_reports_checked_location(self):
        with patch.object(bridge, 'component_candidates', return_value=[self.root / 'missing.exe']):
            with self.assertRaisesRegex(ValueError, '组件缺失.*missing.exe'):
                bridge.command('validate', self.root / 'file.pptx')
        self.assertTrue(any(record.get('failure') == 'component' for record in self.records()))
        state = next(record for record in self.records() if record.get('failure') == 'component_preflight')
        self.assertEqual(state['candidates'][0]['present'], False)

    def test_access_denied_is_not_reported_as_missing(self):
        with patch.object(Path, 'stat', side_effect=PermissionError(13, 'Denied')):
            with self.assertRaisesRegex(ValueError, '无法访问.*13'):
                bridge.binary_path()

    def test_same_size_component_corruption_is_detected_before_execution(self):
        component = self.root / 'officecli.exe'; component.write_bytes(b'bad!')
        import hashlib
        with patch.object(bridge, 'binary_path', return_value=component), patch.object(bridge, 'expected_component', return_value=(4, hashlib.sha256(b'good').hexdigest())), patch.object(bridge.subprocess, 'run') as launch:
            with self.assertRaisesRegex(ValueError, '与本版本不一致'):
                bridge.command('validate', self.root / 'source.pptx')
        launch.assert_not_called()
        self.assertEqual(self.records()[0]['failure'], 'component_preflight')

    def test_startup_and_later_disappearance_have_distinct_records(self):
        component = self.root / 'resources/officecli/officecli.exe'
        component.parent.mkdir(parents=True); component.write_bytes(b'good')
        import hashlib
        with patch.object(bridge, 'component_candidates', return_value=[component]), patch.object(bridge, 'expected_component', return_value=(4, hashlib.sha256(b'good').hexdigest())):
            self.assertTrue(bridge.audit_startup()['verified'])
            component.unlink()
            with self.assertRaisesRegex(ValueError, '组件缺失'):
                bridge.ensure_component('before_translation')
        records = self.records()
        self.assertTrue(records[0]['verified'])
        self.assertEqual(records[-1]['stage'], 'before_translation')
        self.assertFalse(records[-1]['candidates'][0]['present'])

    def test_missing_component_does_not_stop_app_startup(self):
        with patch.object(bridge, 'component_candidates', return_value=[self.root / 'missing.exe']):
            self.assertFalse(bridge.audit_startup()['verified'])

    def test_pptx_preflight_runs_before_parsing_or_model_calls(self):
        from backend import translation
        with patch.object(bridge, 'ensure_component', side_effect=ValueError('组件缺失')), patch.object(translation, 'extract') as extract, patch.object(translation, 'LLMClient') as model:
            with self.assertRaisesRegex(ValueError, '组件缺失'):
                translation.translate_file(str(self.root / 'source.pptx'), None, None)
        extract.assert_not_called(); model.assert_not_called()

    def test_launch_failure_has_system_error_and_persistent_record(self):
        with patch.object(bridge.subprocess, 'run', side_effect=PermissionError(13, 'Denied')):
            with self.assertRaisesRegex(ValueError, '无法启动.*13'):
                bridge.command('validate', self.root / 'file.pptx')
        self.assertEqual(self.records()[0]['system_error'], 13)

    def test_timeout_has_bounded_longer_deadline_and_diagnostic(self):
        with patch.object(bridge.subprocess, 'run', side_effect=subprocess.TimeoutExpired('officecli', 180)) as run:
            with self.assertRaisesRegex(ValueError, '超时（180 秒）'):
                bridge.command('validate', self.root / 'file.pptx')
        self.assertEqual(run.call_args.kwargs['timeout'], 180)
        self.assertEqual(run.call_args.kwargs['cwd'], str(self.root))
        self.assertEqual(self.records()[0]['failure'], 'timeout')

    def test_native_startup_exit_keeps_hex_code_and_stderr(self):
        result = subprocess.CompletedProcess([], -1073741515, '', 'Required runtime library could not be loaded')
        with patch.object(bridge.subprocess, 'run', return_value=result):
            with self.assertRaisesRegex(ValueError, '运行库.*0xC0000135'):
                bridge.command('validate', self.root / 'file.pptx')
        self.assertIn('runtime library', self.records()[0]['stderr'])

    def test_bom_and_diagnostic_prefix_are_accepted_but_two_payloads_are_not(self):
        envelope = json.dumps({'success': True, 'data': {'errors': []}})
        for output in ['\ufeff' + envelope, 'Diagnostic startup line\n' + envelope]:
            with patch.object(bridge.subprocess, 'run', return_value=subprocess.CompletedProcess([], 0, output, '')):
                self.assertEqual(bridge.command('validate', self.root / 'file.pptx'), {'errors': []})
        with self.assertRaises(json.JSONDecodeError):
            bridge.parse_response(envelope + '\n' + envelope)

    def test_invalid_output_records_diagnostics_without_document_text(self):
        result = subprocess.CompletedProcess([], 2, 'PRIVATE DOCUMENT WORDS', 'Startup failed')
        with patch.object(bridge.subprocess, 'run', return_value=result):
            with self.assertRaisesRegex(ValueError, '未返回有效检查结果.*退出码 2'):
                bridge.command('view', self.root / 'file.pptx')
        text = json.dumps(self.records())
        self.assertIn('Startup failed', text); self.assertNotIn('PRIVATE DOCUMENT WORDS', text)

    def test_immediate_translation_failure_is_failed_in_history_and_preserves_source(self):
        source = self.root / 'source.pptx'; source.write_bytes(b'untouched synthetic source')
        job = {'id': 'early-stop', 'cancel': threading.Event(), 'status': 'running'}
        with patch.object(workflow.store, 'load_glossary'), patch.object(workflow, 'ReferenceIndex'), patch.object(workflow, 'translate_file', side_effect=ValueError('文档检查组件无法启动（系统错误 5）。')):
            workflow.run(job, {'kind': 'translate_review', 'paths': [str(source)]})
        self.assertEqual(job['status'], 'failed'); self.assertEqual(job['result']['outputs'], [])
        self.assertIn('系统错误 5', job['result']['errors'][0]['error'])
        record = history.query()['records'][0]
        self.assertEqual(record['status'], 'failed'); self.assertTrue(record['errors'])
        self.assertEqual(source.read_bytes(), b'untouched synthetic source')

    def test_empty_exception_has_readable_message_and_history_entry(self):
        job = {'id': 'empty-error', 'cancel': threading.Event(), 'status': 'running'}
        with patch.object(workflow.store, 'load_glossary', side_effect=RuntimeError()):
            workflow.run(job, {'kind': 'translate', 'paths': ['synthetic.docx']})
        self.assertEqual(job['status'], 'failed'); self.assertIn('RuntimeError', job['error'])
        self.assertIn('RuntimeError', history.query()['records'][0]['errors'][0]['error'])

    def test_cancel_stops_only_the_private_check_child_promptly(self):
        from backend.translation import Cancelled
        cancel = threading.Event()
        timer = threading.Timer(.3, cancel.set); timer.start()
        started = time.monotonic()
        try:
            with self.assertRaises(Cancelled):
                bridge.run_check([sys.executable, '-c', 'import time; time.sleep(30)'], cancel=cancel,
                    capture_output=True, encoding='utf-8', errors='replace', timeout=180,
                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            self.assertLess(time.monotonic() - started, 3)
        finally: timer.join()


if __name__ == '__main__': unittest.main()
