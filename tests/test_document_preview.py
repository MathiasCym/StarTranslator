"""Full-page previews are read-only, isolated, refreshable and owner-scoped."""
import os,sys,unittest,tempfile,json,shutil,io
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from lxml import html
from zipfile import ZipFile
from lxml import etree
from backend import comparison,document_preview,history,workflow,lan_service
from backend.extract import extract
from backend.extract.office import open_office
from backend.glossary import store

class DocumentPreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'data')});self.env.start()
        self.picture=self.root/'picture.png';Image.new('RGB',(150,80),'#23a0ce').save(self.picture)
        # Exercise the compatible renderer independently of whichever Office
        # applications happen to be installed on this test computer.
        native=patch('backend.word_native_preview.preview',return_value=None)
        native.start();self.addCleanup(native.stop)
        native_ppt=patch('backend.pptx_native_preview.preview',return_value=None)
        native_ppt.start();self.addCleanup(native_ppt.stop)
    def tearDown(self):self.env.stop();self.temp.cleanup()
    def pair(self,kind='.docx'):
        source=self.root/('中文原稿'+kind);output=self.root/('English'+kind)
        for path,language in [(source,'中文完整页面'),(output,'Complete English page')]:
            if kind=='.pptx':
                d=Presentation();d.slide_width=Inches(12);d.slide_height=Inches(6.75)
                for i in range(2):
                    s=d.slides.add_slide(d.slide_layouts[6]);s.shapes.add_textbox(Inches(1),Inches(.5),Inches(8),Inches(1)).text=language+' '+str(i+1)
                    s.shapes.add_picture(str(self.picture),Inches(8),Inches(2),Inches(2))
                    t=s.shapes.add_table(2,2,Inches(1),Inches(2),Inches(3),Inches(1)).table;t.cell(0,0).text=language
                d.save(path)
                # Native cached chart data; no Excel writer is required at runtime.
                with ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
                slide=etree.fromstring(parts['ppt/slides/slide2.xml']);ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
                frame=etree.fromstring(b'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:nvGraphicFramePr><p:cNvPr id="990" name="Chart"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="914400" y="3657600"/><a:ext cx="5486400" cy="1828800"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart r:id="rIdPreviewChart"/></a:graphicData></a:graphic></p:graphicFrame>')
                slide.find('p:cSld/p:spTree',ns).append(frame);parts['ppt/slides/slide2.xml']=etree.tostring(slide)
                rel=etree.fromstring(parts['ppt/slides/_rels/slide2.xml.rels']);etree.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdPreviewChart',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart',Target='../charts/chart1.xml');parts['ppt/slides/_rels/slide2.xml.rels']=etree.tostring(rel)
                ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/ppt/charts/chart1.xml',ContentType='application/vnd.openxmlformats-officedocument.drawingml.chart+xml');parts['[Content_Types].xml']=etree.tostring(ct)
                parts['ppt/charts/chart1.xml']=b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:plotArea><c:layout/><c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:v>Screens</c:v></c:tx><c:cat><c:strLit><c:ptCount val="2"/><c:pt idx="0"><c:v>2025</c:v></c:pt><c:pt idx="1"><c:v>2026</c:v></c:pt></c:strLit></c:cat><c:val><c:numLit><c:formatCode>General</c:formatCode><c:ptCount val="2"/><c:pt idx="0"><c:v>261</c:v></c:pt><c:pt idx="1"><c:v>300</c:v></c:pt></c:numLit></c:val></c:ser><c:axId val="1"/><c:axId val="2"/></c:barChart><c:catAx><c:axId val="1"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="b"/><c:crossAx val="2"/></c:catAx><c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="l"/><c:crossAx val="1"/></c:valAx></c:plotArea></c:chart></c:chartSpace>'
                with ZipFile(path,'w') as z:
                    for name,content in parts.items():z.writestr(name,content)
            else:
                d=Document();d.add_heading(language,0);d.add_picture(str(self.picture),width=Inches(2));d.add_table(2,2).cell(0,0).text=language;d.add_page_break();d.add_heading(language+' 2',1);d.save(path)
        _,paras,_=open_office(str(output));ident=comparison.create(extract(str(source)),{loc:p for loc,p,*_ in paras},str(source),str(output))
        return source,output,ident
    def test_ppt_keeps_whole_slides_images_tables_and_chart(self):
        source,output,ident=self.pair('.pptx');before=store.fingerprint(source);out_before=store.fingerprint(output)
        first=document_preview.preview(ident,'source');text=document_preview.preview(ident,'source',2);dom=html.fromstring(text)
        self.assertEqual(len(dom.xpath('//*[@class="slide-container"]')),1);self.assertGreaterEqual(len(dom.xpath('//img')),1)
        self.assertIn('data-star-page="2"',text);self.assertIn('data-star-total="2"',text)
        self.assertIn('data:image/',text);self.assertIn('2025',text);self.assertIn('300',text);self.assertIn('中文完整页面',text)
        self.assertNotEqual(first,text)
        bars=dom.xpath('//svg/rect');self.assertEqual(len(bars),2)
        self.assertAlmostEqual(float(bars[1].get('height'))/float(bars[0].get('height')),300/261,places=2)
        self.assertTrue(dom.xpath('//table') or dom.xpath('//*[contains(@class,"table")]'))
        self.assertEqual(before,store.fingerprint(source));self.assertEqual(out_before,store.fingerprint(output))
        source.unlink();self.assertIn('中文完整页面',document_preview.preview(ident,'source'))
    def test_word_has_page_breaks_images_and_tables(self):
        _,_,ident=self.pair();text=document_preview.preview(ident,'source');dom=html.fromstring(text)
        self.assertEqual(len(dom.xpath('//*[@class="page-wrapper"]')),2);self.assertTrue(dom.xpath('//img'));self.assertTrue(dom.xpath('//table'))
        response=document_preview.response(ident,'source',shared=True)
        self.assertIn("connect-src 'none'",response.headers['Content-Security-Policy']);self.assertIn("frame-ancestors 'self'",response.headers['Content-Security-Policy'])
    def test_new_output_invalidates_only_its_cached_preview(self):
        source,output,ident=self.pair();original=document_preview.preview(ident,'source');old=document_preview.preview(ident,'output')
        d=Document(output);d.paragraphs[0].text='Confirmed corrected English';d.save(output);comparison.update(ident,str(output))
        new=document_preview.preview(ident,'output');self.assertIn('Confirmed corrected English',new);self.assertNotEqual(old,new)
        self.assertEqual(original,document_preview.preview(ident,'source'));self.assertFalse(comparison.read(ident)['output_changed'])
    def test_word_native_pages_do_not_pass_through_the_html_word_reflow(self):
        source,output,ident=self.pair();before=store.fingerprint(output)
        raw='<html><head></head><body data-star-renderer="native-word"><div class="page-wrapper"><div class="page"><img src="data:image/png;base64,AA=="></div></div></body></html>'
        with patch('backend.word_native_preview.preview',return_value=(raw,21)) as render,patch('backend.document_preview.subprocess.run') as officecli:
            result=document_preview.preview(ident,'output',3)
        self.assertEqual(render.call_args.args[-1],3);officecli.assert_not_called()
        self.assertIn('data-star-page="3"',result);self.assertIn('data-star-total="21"',result)
        self.assertNotIn('starSplitWordTable',result);self.assertNotIn('兼容排版预览',result)
        self.assertEqual(before,store.fingerprint(output))
    def test_compatible_word_preview_identifies_the_layout_limitation(self):
        _,_,ident=self.pair();result=document_preview.preview(ident,'output')
        self.assertIn('data-star-preview-notice="兼容排版预览',result)
    def test_ppt_native_image_is_the_page_and_bypasses_compatible_html(self):
        source,output,ident=self.pair('.pptx');before=store.fingerprint(output)
        raw='<html><head></head><body data-star-renderer="native-pptx"><div class="slide-container"><div class="slide"><img src="data:image/png;base64,AA=="></div></div></body></html>'
        with patch('backend.pptx_native_preview.preview',return_value=(raw,2)) as render,patch('backend.document_preview.subprocess.run') as officecli:
            result=document_preview.preview(ident,'output',2)
        self.assertEqual(render.call_args.args[-2:],(2,2));officecli.assert_not_called()
        self.assertIn('data-star-page="2"',result);self.assertIn('data-star-renderer="native-pptx"',result)
        self.assertNotIn('兼容排版预览',result);self.assertEqual(before,store.fingerprint(output))
    def test_ppt_compatible_preview_discloses_its_different_renderer(self):
        _,_,ident=self.pair('.pptx');result=document_preview.preview(ident,'output')
        self.assertIn('未使用 PowerPoint 实际页面',result)
    def test_legacy_record_recovers_original_from_history(self):
        source,output,ident=self.pair();value=comparison.read(ident);value={k:v for k,v in value.items() if not k.startswith('_')}
        store.atomic_json(comparison.folder()/(ident+'.json'),value);shutil.rmtree(comparison.folder()/ident)
        data={'kind':'translate','paths':[str(source)]};history.begin('legacy',data);history.finish({'id':'legacy','status':'complete','result':{'outputs':[{'path':str(output),'source':str(source),'comparison_id':ident}]}},data)
        self.assertIn('中文完整页面',document_preview.preview(ident,'source'));self.assertTrue(comparison.read(ident).get('_source_snapshot'))
    def test_report_saves_keep_latest_comparison_for_both_final_and_overwrite(self):
        from backend.models import Finding
        for mode,overwrite in [('both',False),('final',False),('both',True)]:
            with self.subTest(mode=mode,overwrite=overwrite):
                source,output,ident=self.pair();original=source.read_bytes()
                segments=extract(str(output));segment=segments[0];key='compare-save-'+mode+str(overwrite)
                findings=[Finding(segment.id,output.name,None,1,'text',old,suggested=new,file_path=str(output),method='llm') for old,new in [('Complete','Corrected'),('English','Reviewed')]]
                workflow.REPORTS[key]={'segments':segments,'hashes':{str(output):store.fingerprint(output)},'findings':findings,
                    'output_options':{str(output):{'mode':mode,'staged':mode=='final','target_dir':str(self.root),'source':str(source),'comparison_id':ident,'translation_complete':True}}}
                self.addCleanup(lambda k=key:workflow.REPORTS.pop(k,None))
                # Prime the cache before applying edits, then save twice through
                # the real report writer so neither text nor page preview is stale.
                document_preview.preview(ident,'output')
                for i in range(2):
                    saved=workflow.apply_report(dict(report_id=key,path=str(output),overwrite=overwrite,edits=[dict(index=i,text=findings[i].suggested)]))
                    self.assertEqual(saved['comparison_id'],ident)
                    value=comparison.read(saved['comparison_id']);self.assertEqual(value['output_path'],saved['output']);self.assertFalse(value['output_changed'])
                    self.assertIn('Corrected',document_preview.preview(ident,'output'))
                self.assertIn('Corrected Reviewed',document_preview.preview(ident,'output'))
                self.assertTrue(any('Corrected Reviewed' in row['translation'] for row in value['rows']))
                self.assertEqual(source.read_bytes(),original)
    def test_shared_report_maps_each_file_to_its_own_comparison_after_save(self):
        source,output,ident=self.pair();raw=str(output)
        job={'id':'compare-public','kind':'translate_review','status':'complete','phase':'complete','percent':100,'message':'完成','result':{
            'outputs':[{'path':str(self.root/'other.docx')},{'path':raw,'comparison_id':ident}],
            'report':{'files':[{'path':str(self.root/'other.docx')},{'path':raw}],'findings':[]}}}
        with patch.dict(workflow.REPORTS,{'compare-public':{'working_paths':{raw:str(output)}}}):
            value=lan_service.public(job)
        self.assertIsNone(value['result']['report']['files'][0]['comparison_index'])
        self.assertEqual(value['result']['report']['files'][1]['comparison_index'],1)
        self.assertNotIn(str(self.root),json.dumps(value))
    def test_missing_original_shows_reason_without_discarding_text(self):
        source,_,ident=self.pair();shutil.rmtree(comparison.folder()/ident);source.unlink()
        value=comparison.read(ident);value={k:v for k,v in value.items() if not k.startswith('_')};store.atomic_json(comparison.folder()/(ident+'.json'),value)
        self.assertIn('旧对照',document_preview.response(ident,'source').get_data(as_text=True));self.assertTrue(comparison.read(ident)['rows'])
    def test_frames_strip_external_resources_and_cannot_call_host(self):
        raw='<html><head><script src="https://outside.invalid/a.js"></script><link href="file:///private.css"></head><body><div class="slide-container"><div class="slide"><img src="file:///secret.png" onerror="alert(1)"><a href="https://outside.invalid">link</a><iframe src="file:///secret"></iframe><script>alert(1)</script></div></div></body></html>'
        prepared=document_preview.prepare(raw,'.pptx');self.assertNotIn('outside.invalid',prepared);self.assertNotIn('secret',prepared);self.assertNotIn('alert(1)',prepared)
        self.assertIn('star-page-preview',prepared);self.assertIn("default-src 'none'",document_preview.CSP)
    def test_shared_preview_and_cleanup_keep_ownership(self):
        source,output,ident=self.pair();app=lan_service.create_app(True);first=app.test_client();other=app.test_client()
        csrf=first.get('/shared/session').json['csrf'];other.get('/shared/session')
        with patch('backend.workflow.JOBS',{}):
            token=first.get_cookie(lan_service.COOKIE).value;session=lan_service.SESSIONS[token]
            folder=lan_service.base_dir()/token/'owned';folder.mkdir(parents=True);session['jobs']=['owned'];session['folders']={'owned':str(folder)}
            workflow.JOBS['owned']={'id':'owned','kind':'translate','status':'complete','result':{'outputs':[{'comparison_id':ident,'path':str(output)}]}}
            url='/shared/jobs/owned/comparison/0/preview/source'
            self.assertEqual(other.get(url).status_code,400);self.assertEqual(first.get(url).status_code,200)
            value=first.get('/shared/jobs/owned/comparison/0').json;self.assertFalse(any(k.startswith('_') for k in value));self.assertNotIn(str(self.root),json.dumps(value))
            self.assertEqual(first.post('/shared/jobs/owned/cleanup',headers={'X-Star-CSRF':csrf}).status_code,200)
        self.assertTrue(source.is_file());self.assertFalse((comparison.folder()/ident).exists())

if __name__=='__main__':unittest.main()
