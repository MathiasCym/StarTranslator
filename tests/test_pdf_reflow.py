"""PDF-specific extraction/recovery must not weaken editable Office checks."""
import io, os, re, sys, tempfile, unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile
sys.path[:0]=[str(Path(__file__).resolve().parents[1]),str(Path(__file__).resolve().parent)]
from docx import Document
from PIL import Image
from backend.pdf_reflow import restore_spaces, compact, page_blocks, join_link_wraps
from backend.pdf_images import append_images
from backend.formats import pdf_document
from backend.translation import translate_file
from backend.translation_audit import saved_residuals
from backend.glossary.loader import Glossary
from backend.glossary.store import fingerprint
from backend.llm.client import LLMClient
from test_file_formats import pdf_fixture


class PdfReflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});self.env.start()
    def tearDown(self):self.env.stop();self.tmp.cleanup()

    def test_positioned_word_gaps_keep_existing_spaces_and_characters(self):
        plain='Hollywood Park, USA\n[32]COSM.NHLandCosmannouncemulti-yearpartnership.'
        layout='HollywoodPark,USA\n[32] COSM. NHL and Cosm announce multi-year partnership.'
        result=restore_spaces(plain,layout)
        self.assertEqual(result,'Hollywood Park, USA\n[32] COSM. NHL and Cosm announce multi-year partnership.')
        self.assertEqual(compact(result),compact(plain))

    def test_positioned_text_cannot_replace_or_reorder_source_characters(self):
        plain='VX1000 200 Mbit\nChapter 2\nChapter 3'
        layout='VX2000 200 Mbyte\nChapter 3\nChapter 2'
        self.assertEqual(restore_spaces(plain,layout),plain)

    def test_unsupported_position_extraction_falls_back_to_complete_plain_text(self):
        from backend.pdf_reflow import extract_page
        from pypdf.errors import PdfReadError
        class Page:
            def extract_text(self,**kw):
                if kw:raise PdfReadError('unsupported font positioning')
                return '客户报告 VX1000\nOriginal English'
        self.assertEqual(extract_page(Page()),'客户报告 VX1000\nOriginal English')

    def test_model_numbers_resolution_urls_and_paths_cannot_gain_spaces(self):
        plain='VX1000 12Kx10K A8S-N https://example.test/中文\nC:\\Manual\\资料.docx'
        layout='VX 1000 12K x 10K A8S - N https://example.test/中 文\nC:\\Manual\\资 料.docx'
        self.assertEqual(restore_spaces(plain,layout),plain)

    def test_link_wrap_join_requires_an_exact_annotation_not_a_guess(self):
        uri='https://example.test/long-document'
        text='[30] '+uri[:-4]+'\n'+uri[-4:]+'\nNext paragraph'
        self.assertEqual(join_link_wraps(text,[uri]),'[30] '+uri+'\n\nNext paragraph')
        self.assertEqual(join_link_wraps(text,[]),text)
        self.assertEqual(join_link_wraps(text,[uri+'other']),text)

    def test_headers_toc_and_references_are_independent_with_body_wraps_intact(self):
        text='COSM 客户项目调研报告\n- 20 -\n参考资料\n[30] COSM. A title\nhttps://example.test/original\n[31] COSM. Another title'
        blocks=page_blocks(text,{compact('COSM 客户项目调研报告')})
        self.assertEqual([b.kind for b in blocks],['header','footer','heading','reference','reference'])
        self.assertIn('\nhttps://example.test/original',blocks[3].text)
        wrapped='主屏直径为87英尺，\n整体画布为12Kx10K。'
        self.assertEqual([b.text for b in page_blocks(wrapped)],[wrapped])
        self.assertEqual([b.kind for b in page_blocks('目录\n摘要........3\n1. 项目........4')],['heading','toc','toc'])
        self.assertEqual([b.kind for b in page_blocks('2026.08.24')],['body'])

    def test_numeric_dates_and_page_numbers_are_kept_after_block_separation(self):
        source=self.root/'source.pdf';text='客户报告\n- 20 -\n2026.08.24\n参考资料\n[30] COSM. Original title'
        pdf_fixture(source,text)
        document=pdf_document(source)
        output=self.root/'out.docx';document.save(output)
        self.assertEqual(compact(''.join(p.text for p in Document(output).paragraphs)),compact(text))
        self.assertFalse(any('- 20 -' in p.text for _,p,*_ in document.paragraphs))

    def test_existing_english_title_is_preserved_when_whole_rewrite_is_rejected(self):
        # Deliberately keep the old whole-page shape: fallback also has to work
        # when layout extraction cannot reliably split or restore a word gap.
        source=self.root/'source.pdf'
        title='COSM.NHLandCosmannouncemulti-yearpartnershipprovidingultra-highresolutioncamera'
        text='COSM 客户项目调研报告\n'+title+'\n型号VX1000，带宽200Mbit。'
        pdf_fixture(source,text);before=fingerprint(source)
        class Model(LLMClient):
            calls=0
            chat_images_json=None  # Text-recovery fixture, not a vision client.
            def __init__(self):
                self.provider='compatible';self.base_url='http://127.0.0.1:1/v1'
                self.model='pdf-recovery-fixture';self.api_format='openai'
            def chat_json(self,prompt,**kw):
                import json
                self.calls+=1
                if '【只需补译的中文】' in prompt:
                    phrases=json.loads(prompt.split('【只需补译的中文】\n')[1].split('\n')[0])
                    translations={'客户项目调研报告':'Customer Project Research Report','型号':'Model ', '带宽':'bandwidth '}
                    return {'translations':[{'source':p,'translation':translations[p]} for p in phrases]}
                return {'translation':'COSM Customer Project Research Report\nNHL and Cosm have a partnership.\nModel VX2000, bandwidth 200 Mbit.'}
        model=Model();result=translate_file(source,Glossary([]),None,client=model,workers=1)
        self.assertTrue(result['translation_complete'],result['review_items'])
        self.assertEqual(result['untranslated'],0);self.assertFalse(saved_residuals(result['path']))
        output='\n'.join(p.text for p in Document(result['path']).paragraphs)
        self.assertIn(title,output);self.assertIn('VX1000',output);self.assertNotIn('VX2000',output)
        self.assertEqual(fingerprint(source),before)
        calls=model.calls
        resumed=translate_file(source,Glossary([]),None,client=model,workers=3)
        self.assertGreater(resumed['resumed_segments'],0);self.assertEqual(model.calls,calls)

    def test_repeated_pdf_headers_resume_even_if_only_one_alias_was_checkpointed(self):
        from pypdf import PdfReader,PdfWriter
        from backend.extract import extract
        from backend.extract.office import open_office
        from backend.checkpoints import open_checkpoint,translation_key,digest
        from backend.replacer.styled_text import style_groups,slice_groups
        first=self.root/'page1.pdf';second=self.root/'page2.pdf';source=self.root/'source.pdf'
        pdf_fixture(first,'测试报告\n介绍A。');pdf_fixture(second,'测试报告\n介绍B。')
        writer=PdfWriter();writer.append(PdfReader(first));writer.append(PdfReader(second));writer.write(source)
        class Model(LLMClient):
            chat_images_json=None
            def __init__(self):
                self.provider='compatible';self.base_url='http://127.0.0.1:1/v1'
                self.model='pdf-header-fixture';self.api_format='openai';self.calls=0;self.forbid=False
            def chat_json(self,prompt,**kw):
                if self.forbid:raise AssertionError('cached header must not request the model')
                self.calls+=1;text=prompt.split('【待译原文】\n')[1].split('\n【')[0].strip()
                return {'translation':{'测试报告':'Test Report','介绍A。':'Introduction A.','介绍B。':'Introduction B.'}[text]}
        model=Model();g=Glossary([])
        result=translate_file(source,g,None,client=model)
        self.assertTrue(result['translation_complete']);self.assertEqual(model.calls,3)
        doc,ps,_=open_office(source);ps={loc:p for loc,p,*_ in ps}
        header=[s for s in extract(source) if s.text=='测试报告'][-1]
        groups=slice_groups(style_groups(ps[header.locator]),0,len(header.text),visual_keys=True)
        checkpoint=open_checkpoint('translation',{'source':os.path.normcase(os.path.abspath(source)),
            'fingerprint':fingerprint(source),'mode':'faithful'},g,None,model)
        with checkpoint.connection() as db:
            db.execute('DELETE FROM checkpoints WHERE scope=? AND item_key=?',(checkpoint.scope,digest(translation_key(header,header.text,groups))))
        model.forbid=True
        resumed=translate_file(source,g,None,client=model,workers=1)
        self.assertTrue(resumed['translation_complete']);self.assertEqual(resumed['resumed_segments'],4)
        self.assertEqual(model.calls,3)

    def test_invalid_local_recovery_cannot_add_numbers_or_hide_a_failure(self):
        source=self.root/'source.pdf';pdf_fixture(source,'型号VX1000。')
        class Model:
            def chat_json(self,prompt,**kw):
                if '【只需补译的中文】' in prompt:return {'translations':[{'source':'型号','translation':'Model 2000'}]}
                return {'translation':'Model VX2000.'}
        result=translate_file(source,Glossary([]),None,client=Model())
        self.assertFalse(result['translation_complete']);self.assertGreater(result['untranslated'],0)
        self.assertTrue(saved_residuals(result['path']))
        self.assertIn('VX1000',''.join(p.text for p in Document(result['path']).paragraphs))

    def test_invalid_recovery_json_is_a_local_issue_not_a_connection_outage(self):
        source=self.root/'source.pdf';pdf_fixture(source,'型号VX1000。')
        class Model:
            def chat_json(self,prompt,**kw):
                if '【只需补译的中文】' in prompt:return {}
                return {'translation':'Model VX2000.'}
        result=translate_file(source,Glossary([]),None,client=Model())
        self.assertFalse(result['translation_complete'])
        self.assertFalse(any(item.get('service_unavailable') for item in result['review_items']))
        self.assertNotIn('连接',''.join(item['reason'] for item in result['review_items']))

    def test_recovery_separates_english_from_adjacent_protected_chinese_url(self):
        from backend.pdf_reflow import recover_chinese
        source='客户报告https://example.test/原始链接?q=1'
        class Model:
            def chat_json(self,prompt,**kw):return {'translations':[{'source':'客户报告','translation':'Customer Report'}]}
        result,runs=recover_chinese(Model(),source,[{'id':0,'text':source}],Glossary([]),lambda:None)
        self.assertEqual(result['translation'],'Customer Report https://example.test/原始链接?q=1')
        self.assertEqual(''.join(r['text'] for r in runs),result['translation'])

    def test_contents_use_right_aligned_leaders_without_changing_page_numbers(self):
        from backend.pdf_reflow import fit_contents
        from docx.oxml.ns import qn
        document=Document();p=document.add_paragraph('2.1 A long English section title...................- 8 -')
        fit_contents(document,[p]);fit_contents(document,[p])
        self.assertEqual(p.text,'2.1 A long English section title\t- 8 -')
        tab=p._p.pPr.find(qn('w:tabs'))[0]
        self.assertEqual(tab.get(qn('w:val')),'right');self.assertEqual(tab.get(qn('w:leader')),'dot')

    def test_visible_reference_urls_are_clickable_without_duplicate_long_addresses(self):
        from pypdf import PdfReader,PdfWriter
        from pypdf.annotations import Link
        uri='https://example.test/original-reference'
        source=self.root/'source.pdf';pdf_fixture(source,'参考资料\n[1] COSM. Original English title\n'+uri)
        writer=PdfWriter();writer.append(PdfReader(source));writer.add_annotation(0,Link(rect=(0,0,100,20),url=uri));writer.write(source)
        document=pdf_document(source);out=self.root/'out.docx';document.save(out)
        word=Document(out)
        self.assertEqual('\n'.join(p.text for p in word.paragraphs).count(uri),1)
        self.assertIn(uri,[r.target_ref for r in word.part.rels.values() if r.is_external])

    def test_international_link_host_does_not_create_new_chinese_body_labels(self):
        from pypdf import PdfReader,PdfWriter
        from pypdf.annotations import Link
        from backend.text_integrity import contains_chinese
        uri='https://示例.test/原始链接'
        source=self.root/'source.pdf';pdf_fixture(source,'Original English title\n'+uri)
        writer=PdfWriter();writer.append(PdfReader(source));writer.add_annotation(0,Link(rect=(0,0,100,20),url=uri));writer.write(source)
        document=pdf_document(source);out=self.root/'out.docx';document.save(out)
        word=Document(out)
        self.assertFalse(any(contains_chinese(p.text) for p in word.paragraphs))
        self.assertIn(uri,[r.target_ref for r in word.part.rels.values() if r.is_external])

    def test_pdf_images_are_preserved_with_bounded_aspect_ratio(self):
        raw=io.BytesIO();Image.new('RGB',(1200,600),'navy').save(raw,format='JPEG')
        document=Document();page=SimpleNamespace(images=[SimpleNamespace(data=raw.getvalue())])
        count,warnings=append_images(document,page,7)
        self.assertEqual((count,warnings),(1,[]))
        shape=document.inline_shapes[0]
        self.assertAlmostEqual(shape.width/shape.height,2)
        self.assertLessEqual(shape.height/12700,270)
        self.assertIn('page 7',shape._inline.docPr.get('descr'))
        output=self.root/'images.docx';document.save(output)
        with ZipFile(output) as z:self.assertIn(raw.getvalue(),[z.read(n) for n in z.namelist() if n.startswith('word/media/')])

    def test_jpeg2000_is_converted_without_losing_alpha(self):
        raw=io.BytesIO()
        try:Image.new('RGBA',(40,20),(10,20,30,90)).save(raw,format='JPEG2000')
        except (OSError,ValueError):self.skipTest('Pillow JPEG2000 unavailable')
        document=Document();count,warnings=append_images(document,SimpleNamespace(images=[SimpleNamespace(data=raw.getvalue())]),1)
        self.assertEqual((count,warnings),(1,[]))
        output=self.root/'jp2.docx';document.save(output)
        with ZipFile(output) as z:
            files=[n for n in z.namelist() if n.startswith('word/media/')]
            with Image.open(io.BytesIO(z.read(files[0]))) as image:
                self.assertEqual(image.size,(40,20));self.assertEqual(image.convert('RGBA').getpixel((0,0)),(10,20,30,90))

    def test_image_failure_is_reported_and_does_not_discard_other_images(self):
        raw=io.BytesIO();Image.new('RGB',(20,30)).save(raw,format='PNG')
        document=Document();count,warnings=append_images(document,SimpleNamespace(images=[SimpleNamespace(data=b'not an image'),SimpleNamespace(data=raw.getvalue())]),4)
        self.assertEqual(count,1);self.assertEqual(len(warnings),1);self.assertIn('第 4 页',warnings[0])

    def test_failure_to_decode_one_pdf_xobject_keeps_later_figures(self):
        raw=io.BytesIO();Image.new('RGB',(20,30)).save(raw,format='PNG')
        class Images:
            def keys(self):return ['/Broken','/Valid']
            def __getitem__(self,key):
                if key=='/Broken':raise ValueError('invalid XObject image')
                return SimpleNamespace(data=raw.getvalue())
        document=Document();count,warnings=append_images(document,SimpleNamespace(images=Images()),2)
        self.assertEqual(count,1);self.assertEqual(len(warnings),1)


if __name__=='__main__':unittest.main()
