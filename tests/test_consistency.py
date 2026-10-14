"""Offline checks for document-wide reuse, conservative advice and manual edits."""
import json
import os
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor, TimeoutError
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from backend.consistency import DocumentTranslationMemory, check_consistency
from backend.models import Segment, GlossaryEntry
from backend.glossary.loader import Glossary
from backend.writing_policy import validate_translation, repair_numeric_translation
from backend.matcher import check_segments
from backend.references import ReferenceIndex
from backend.translation import translate_file, Cancelled
from backend.extract import extract
from backend import workflow


def segment(i, text, file='deck.pptx', **extra):
    return Segment(str(i), file, text, f'slide/{i}/p0', slide=i, para=1, file_path=file, **extra)


class ConsistencyTests(unittest.TestCase):
    def setUp(self):
        for target in ('backend.translation.review_pptx','backend.workflow.review_pptx'):
            stub=patch(target,return_value=([],[]));stub.start();self.addCleanup(stub.stop)
        self.g = Glossary([])

    def test_user_period_example_is_advisory_with_and_without_originals(self):
        values = [segment(1, 'H2 Strategy'), segment(2, 'Second half Strategy')]
        for original in [None, {'1':'下半年战略', '2':'下半年战略'}]:
            hits = check_consistency(values, self.g, aligned_sources=original)
            self.assertEqual([(f.found_text, f.suggested) for f in hits], [('Second half Strategy','H2 Strategy')])
            self.assertEqual(hits[0].method, 'consistency')
            self.assertFalse(hits[0].applied)
            self.assertIn('不代表原写法一定错误', hits[0].reason)

    def test_exact_repeated_original_without_glossary_and_confirmed_standard(self):
        values = [segment(1,'Pixel-level correction'), segment(2,'Pixel calibration')]
        originals = dict.fromkeys(['1','2'], '逐点校正技术')
        hits = check_consistency(values,self.g,aligned_sources=originals)
        self.assertEqual(hits[0].suggested, values[0].text)
        g = Glossary([GlossaryEntry('逐点校正技术','pixel-level calibration',priority=1)])
        hits = check_consistency(values,g,aligned_sources=originals)
        self.assertEqual(len(hits),2)
        self.assertTrue(all(f.suggested=='pixel-level calibration' for f in hits))

    def test_case_spacing_punctuation_file_isolation_and_single_usage(self):
        for values in [[segment(1,'H2 Strategy'),segment(2,'H2 strategy.')],
                       [segment(1,'H2 Strategy'),segment(2,'Second half Strategy',file='other.pptx')],
                       [segment(1,'H2 Strategy')]]:
            self.assertEqual(check_consistency(values,self.g),[])

    def test_confirmed_period_standard_does_not_get_an_opposite_advice(self):
        g=Glossary([GlossaryEntry('下半年战略','Second-half strategy',priority=1)])
        values=[segment(1,'H2 Strategy'),segment(2,'Second half Strategy')]
        hits=check_consistency(values,g,aligned_sources=dict.fromkeys(['1','2'],'下半年战略'))
        self.assertEqual([(f.found_text,f.suggested) for f in hits],[('H2 Strategy','Second-half strategy')])

    def test_different_periods_and_hardware_codes_are_not_equivalent(self):
        values = [segment(1,'H1 Strategy'),segment(2,'Second half Strategy'),
                  segment(3,'H2 receiver card'),segment(4,'Second half receiver card')]
        self.assertEqual(check_consistency(values,self.g),[])

    def test_numeric_mismatch_and_identifier_change_are_not_copied(self):
        for first,second,source in [('500 nits','1000 nits','亮度为500nits'),
                                    ('VX1000 display','VX2000 screen','VX1000显示设备')]:
            self.assertEqual(check_consistency([segment(1,first),segment(2,second)],self.g,
                aligned_sources=dict.fromkeys(['1','2'],source)),[])

    def test_acronym_expansion_and_ambiguous_chinese_are_kept(self):
        for g,source,one,two in [
                (Glossary([GlossaryEntry('发光二极管','LED',['light-emitting diode'],priority=1)]), '发光二极管','LED','light-emitting diode'),
                (Glossary([GlossaryEntry('发光二极管','LED',['light-emitting diode'],priority=1)]), '发光二极管','LED','light-emitting diode (LED)'),
                (Glossary([GlossaryEntry('发光二极管','LED',['light-emitting diode'],priority=1)]), '发光二极管','LED','LED (light-emitting diode)'),
                (Glossary([GlossaryEntry('控制系统','control system'),GlossaryEntry('控制系统','control platform')]), '控制系统','control system','control platform')]:
            self.assertEqual(check_consistency([segment(1,one),segment(2,two)],g,
                aligned_sources=dict.fromkeys(['1','2'],source)),[])

    def test_accepted_forms_only_advise_when_paired_with_same_chinese(self):
        g = Glossary([GlossaryEntry('逐点校正','pixel-level calibration',['pixel calibration'],priority=1)])
        values = [segment(1,'Use pixel-level calibration.'),segment(2,'Apply pixel calibration now.')]
        self.assertEqual(check_consistency(values,g),[])
        hits = check_consistency(values,g,aligned_sources={'1':'采用逐点校正。','2':'现在应用逐点校正。'})
        self.assertEqual([(f.found_text,f.suggested) for f in hits],[('pixel calibration','pixel-level calibration')])

    def test_nested_alias_cannot_duplicate_a_canonical_prefix(self):
        g = Glossary([GlossaryEntry('点间距','pixel pitch',['pitch','dot pitch'],priority=1)])
        values = [segment(1,'A pixel pitch.'),segment(2,'The dot pitch.')]
        hits = check_consistency(values,g,aligned_sources={'1':'点间距参数','2':'该点间距参数'})
        self.assertEqual([(f.found_text,f.suggested) for f in hits],[('dot pitch','pixel pitch')])

    def test_counted_object_context_is_not_merged_across_pages(self):
        values = [segment(1,'Twenty screens'),segment(2,'Twenty cabinets')]
        originals = {'1':'数量二十块','2':'数量二十块'}
        values += [segment(3,'Screen statistics'),segment(4,'Cabinet statistics')]
        values[2].slide=1;values[3].slide=2
        originals.update({'3':'屏幕统计','4':'箱体统计'})
        self.assertEqual(check_consistency(values,self.g,aligned_sources=originals),[])

    def test_advice_uses_existing_review_calls_without_auto_acceptance(self):
        class Model:
            def __init__(self): self.prompts=[]
            def chat_json(self,prompt,**kwargs): self.prompts.append(prompt);return []
        model=Model()
        values=[segment(1,'H2 Strategy'),segment(2,'Second half Strategy')]
        hits,stats=check_segments(values,self.g,client=model)
        self.assertEqual(len(model.prompts),len(values))
        self.assertEqual(stats['consistency_hits'],1)
        self.assertEqual(stats['llm_hits'],0)
        self.assertIn('全文用语统一建议',model.prompts[1])
        self.assertFalse(hits[0].applied)
        with patch('backend.matcher.LLMClient') as client:
            _,stats=check_segments(values,self.g,use_llm=False)
        client.assert_not_called();self.assertEqual(stats['consistency_hits'],1)

    def test_calendar_abbreviations_preserve_real_quantities_and_codes(self):
        for source,target in [('下半年战略','H2 Strategy'),('上半年计划','H1 Plan'),
                ('第一季度销售500台','Q1 Sales: 500 units'),('第二季度报告','Q2 Report'),
                ('第三季度报告','3Q Report'),('第四季度报告','Q4 Report'),
                ('下半年发布H2控制器','H2 launch of H2 controllers'),
                ('下半年亮度500nits','H2 brightness: 500 nits')]:
            validate_translation(source,target)
        for source,target in [('下半年战略','H1 Strategy'),('上半年战略','H2 Strategy'),
                ('第一季度报告','Q2 Report'),('下半年销售500台','H2 sales: 600 units'),
                ('VX2下半年计划','VX3 H2 Plan'),('V2.1下半年计划','V2.2 H2 Plan'),
                ('H2控制器','Second half controller')]:
            with self.assertRaises(ValueError): validate_translation(source,target)
        self.assertEqual(repair_numeric_translation('下半年亮度500nits','H2 brightness: 600 nits'),'H2 brightness: 500 nits')

    def test_single_flight_reuses_text_and_rebinds_style_ids(self):
        values=[segment(1,'下半年战略'),segment(2,'下半年战略')]
        memory=DocumentTranslationMemory(values,self.g)
        entered,release=threading.Event(),threading.Event()
        calls=[]
        def generate():
            calls.append(1);entered.set();self.assertTrue(release.wait(2))
            return [{'id':7,'text':'H2 Strategy'}],False
        with ThreadPoolExecutor(2) as pool:
            first=pool.submit(memory.translate,values[0],values[0].text,[{'id':7}],generate,lambda:None)
            self.assertTrue(entered.wait(2))
            second=pool.submit(memory.translate,values[1],values[1].text,[{'id':19}],generate,lambda:None)
            release.set();self.assertEqual(first.result()[0][0]['id'],7)
            self.assertEqual(second.result(),([{'id':19,'text':'H2 Strategy'}],False))
        self.assertEqual(len(calls),1)

    def test_mixed_styles_short_ambiguous_terms_and_unresolved_objects_bypass_reuse(self):
        for text,groups,g in [('下半年战略',[{'id':0},{'id':1}],self.g),('控制',[{'id':0}],self.g),
                ('数量二十块',[{'id':0}],self.g),('控制系统',[{'id':0}],Glossary([GlossaryEntry('控制系统','system'),GlossaryEntry('控制系统','platform')]))]:
            values=[segment(1,text),segment(2,text)];memory=DocumentTranslationMemory(values,g);calls=[]
            def generate(): calls.append(1);return [{'id':0,'text':'English'}],False
            for s in values: memory.translate(s,text,groups,generate,lambda:None)
            self.assertEqual(len(calls),2)

    def test_failure_not_cached_and_model_timeout_does_not_wait_forever(self):
        values=[segment(1,'下半年战略'),segment(2,'下半年战略')]
        memory=DocumentTranslationMemory(values,self.g)
        def fail(): raise TimeoutError('model timeout')
        with self.assertRaises(TimeoutError): memory.translate(values[0],values[0].text,[{'id':0}],fail,lambda:None)
        good=lambda:([{'id':1,'text':'H2 Strategy'}],False)
        self.assertEqual(memory.translate(values[1],values[1].text,[{'id':1}],good,lambda:None)[0][0]['text'],'H2 Strategy')

    def test_cancelled_waiter_stops_while_owner_finishes(self):
        values=[segment(1,'下半年战略'),segment(2,'下半年战略')]
        memory=DocumentTranslationMemory(values,self.g);entered=threading.Event();release=threading.Event()
        def generate(): entered.set();release.wait(2);return [{'id':0,'text':'H2 Strategy'}],False
        def cancel(): raise Cancelled('cancelled')
        with ThreadPoolExecutor(2) as pool:
            first=pool.submit(memory.translate,values[0],values[0].text,[{'id':0}],generate,lambda:None)
            self.assertTrue(entered.wait(2))
            second=pool.submit(memory.translate,values[1],values[1].text,[{'id':1}],generate,cancel)
            try:
                with self.assertRaises(Cancelled): second.result(timeout=1)
            finally: release.set()
            first.result()


class ConsistencyWorkflowTests(unittest.TestCase):
    def setUp(self):
        for target in ('backend.translation.review_pptx','backend.workflow.review_pptx'):
            stub=patch(target,return_value=([],[]));stub.start();self.addCleanup(stub.stop)
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.environment=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});self.environment.start()
    def tearDown(self): self.environment.stop();self.temp.cleanup()

    def test_word_report_manual_edit_is_preserved_with_no_extra_calls(self):
        source=self.root/'source.docx';target=self.root/'English.docx'
        original=Document();original.add_paragraph('下半年战略');original.add_paragraph('下半年战略');original.save(source)
        english=Document();english.add_paragraph('H2 Strategy');english.add_paragraph('Second half Strategy');english.save(target)
        before_source=source.read_bytes();before_english=target.read_bytes()
        job={'id':'consistency-report','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,{'kind':'check','paths':[str(target)],'references':[str(source)],'pairs':{str(target):str(source)},'use_llm':False})
        report=job['result'];self.assertEqual(report['stats']['consistency_hits'],1)
        self.assertEqual(target.read_bytes(),before_english)
        output=workflow.apply_report({'report_id':report['report_id'],'path':str(target),
            'edits':[{'index':0,'text':'H2 Strategic Plan'}]})
        self.assertEqual(Document(output['output']).paragraphs[1].text,'H2 Strategic Plan')
        self.assertEqual(source.read_bytes(),before_source);self.assertEqual(target.read_bytes(),before_english)
        self.assertTrue(output['finding_updates'][0]['finding']['applied'])

    def test_real_cross_page_translation_reduces_calls_and_preserves_images_and_bounds(self):
        source=self.root/'source.pptx';deck=Presentation()
        for size in [28,22,18]:
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            shape=slide.shapes.add_textbox(Inches(.5),Inches(.5),Inches(6),Inches(1))
            shape.text='下半年战略';shape.text_frame.paragraphs[0].runs[0].font.size=Pt(size)
            slide.shapes.add_picture(str(Path(__file__).resolve().parents[1]/'build/icon.png'),Inches(7),Inches(1),width=Inches(1))
        deck.save(source);original=source.read_bytes()
        class Model:
            def __init__(self): self.calls=[]
            def chat_json(self,prompt,**kwargs): self.calls.append(prompt);return {'translation':'H2 Strategy'}
        model=Model();result=translate_file(str(source),Glossary([]),None,client=model)
        self.assertEqual(len(model.calls),1);self.assertEqual(result['translated'],3)
        self.assertEqual(result['needs_review'],0);self.assertEqual(result['consistency_hits'],0)
        output=Presentation(result['path'])
        for a,b in zip(deck.slides,output.slides):
            self.assertEqual(a.shapes[0].left,b.shapes[0].left);self.assertEqual(a.shapes[0].width,b.shapes[0].width)
            self.assertEqual(b.shapes[0].text,'H2 Strategy')
            self.assertLessEqual(b.shapes[0].text_frame.paragraphs[0].runs[0].font.size,a.shapes[0].text_frame.paragraphs[0].runs[0].font.size)
        with ZipFile(source) as a,ZipFile(result['path']) as b:
            for name in a.namelist():
                if name.startswith('ppt/media/'): self.assertEqual(a.read(name),b.read(name))
        self.assertEqual(source.read_bytes(),original)

    def test_mixed_styles_remain_distinct_and_inconsistency_is_advisory(self):
        source=self.root/'source.docx';doc=Document()
        for _ in range(2):
            p=doc.add_paragraph();p.add_run('下半年').bold=True;p.add_run('战略')
        doc.save(source)
        class Model:
            def __init__(self): self.calls=0
            def chat_json(self,prompt,**kwargs):
                self.calls+=1
                parts=['H2 ','Strategy'] if self.calls==1 else ['Second half ','Strategy']
                return {'translation':''.join(parts),'runs':[{'id':i,'text':t} for i,t in enumerate(parts)]}
        model=Model();result=translate_file(str(source),Glossary([]),None,client=model)
        self.assertEqual(model.calls,2);self.assertEqual(result['status'],'complete')
        self.assertEqual(result['consistency_hits'],1)
        self.assertFalse(result['consistency_items'][0]['applied'])
        for p in Document(result['path']).paragraphs:
            self.assertTrue(p.runs[0].bold);self.assertFalse(bool(p.runs[1].bold))

    def test_final_only_combined_report_requires_confirmation_and_keeps_document_parts(self):
        source=self.root/'source.docx';doc=Document()
        for _ in range(2):
            p=doc.add_paragraph();p.add_run('下半年').bold=True;p.add_run('战略')
        doc.add_picture(str(Path(__file__).resolve().parents[1]/'build/icon.png'))
        doc.save(source);original=source.read_bytes();out=self.root/'out'
        class Model:
            def __init__(self): self.translations=0;self.reviews=0
            def chat_json(self,prompt,**kwargs):
                if '【待译原文】' not in prompt: self.reviews+=1;return []
                self.translations+=1
                parts=['H2 ','Strategy'] if self.translations==1 else ['Second half ','Strategy']
                return {'translation':''.join(parts),'runs':[{'id':i,'text':t} for i,t in enumerate(parts)]}
        model=Model();job={'id':'consistency-final','cancel':threading.Event(),'status':'running'}
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,{'kind':'translate_review','paths':[str(source)],'_client':model,
                'combined_output':'final','output_dir':str(out)})
        self.assertEqual(job['status'],'complete')
        self.assertEqual((model.translations,model.reviews),(2,2))
        report=job['result']['report'];self.assertEqual(report['stats']['consistency_hits'],1)
        self.assertFalse(any(f['applied'] for f in report['findings']))
        self.assertEqual(list(out.glob('*.docx')),[])
        working=report['files'][0]['path']
        self.assertEqual(Document(working).paragraphs[1].text,'Second half Strategy')
        saved=workflow.apply_report({'report_id':report['report_id'],'path':working,
            'edits':[{'index':0,'text':'H2 Strategy'}]})
        self.assertEqual(len(list(out.glob('*.docx'))),1)
        self.assertEqual(Document(saved['output']).paragraphs[1].text,'H2 Strategy')
        self.assertTrue(Document(saved['output']).paragraphs[1].runs[0].bold)
        with ZipFile(source) as a,ZipFile(saved['output']) as b:
            for name in a.namelist():
                if name.startswith('word/media/'): self.assertEqual(a.read(name),b.read(name))
        self.assertEqual(source.read_bytes(),original)


if __name__=='__main__': unittest.main()
