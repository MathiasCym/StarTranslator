"""No-network regression tests: isolation, bounded concurrency, layout and comparison."""
import asyncio
import io
import json
import os
import threading
import time
import unittest
import zipfile
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from backend import workflow, task_queue, comparison, lan_service
from backend.glossary.loader import Glossary
from backend.models import GlossaryEntry, Segment
from backend.matcher import check_segments
from backend.translation import translate_file
from backend.glossary import store


class AsyncLanTests(unittest.TestCase):
    def setUp(self):
        for target in ('backend.translation.review_pptx','backend.workflow.review_pptx'):
            stub=patch(target,return_value=([],[]));stub.start();self.addCleanup(stub.stop)
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'data')});self.env.start()
        self.g=Glossary([GlossaryEntry('点间距','pixel pitch',wrong_variants=['dot spacing'],priority=1)])
        with workflow.LOCK:
            self.assertFalse(any(j['status'] in ('running','queued') for j in workflow.JOBS.values()))
            workflow.JOBS.clear();workflow.REPORTS.clear();task_queue.MAINTENANCE=False
        lan_service.SESSIONS.clear()
    def tearDown(self):
        self.env.stop();self.temp.cleanup()
    def doc(self,name,text):
        path=self.root/name;path.parent.mkdir(parents=True,exist_ok=True);doc=Document();doc.add_paragraph(text);doc.save(path);return str(path)
    def wait(self,job_id):return asyncio.run(task_queue.wait(job_id,.01))
    def test_async_jobs_keep_separate_reports_and_exports(self):
        paths=[self.doc(f'{i}/same.docx','dot spacing') for i in range(2)]
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            async def jobs():
                jobs=await asyncio.gather(*(task_queue.submit({'kind':'check','paths':[p],'use_llm':False}) for p in paths))
                return await asyncio.gather(*(task_queue.wait(j['id'],.01) for j in jobs))
            jobs=asyncio.run(jobs())
        self.assertEqual(len(workflow.REPORTS),2)
        for job,path in zip(jobs,paths):
            self.assertEqual(job['status'],'complete');finding=job['result']['findings'][0]
            saved=workflow.apply_report({'report_id':job['id'],'path':path,'edits':[{'index':0,'text':'pixel pitch'}]})
            self.assertEqual(Document(saved['output']).paragraphs[0].text,'pixel pitch')
            self.assertEqual(Document(path).paragraphs[0].text,'dot spacing')
    def test_queue_bounds_and_cancel_pending(self):
        active=threading.Event();release=threading.Event();count=[0];lock=threading.Lock()
        def run(job,data):
            with lock:
                count[0]+=1
                if count[0]==2:active.set()
            release.wait(3);job['status']='complete'
        path=self.doc('queue.docx','sample')
        with patch('backend.workflow.run',run):
            jobs=[workflow.start({'kind':'check','paths':[path],'use_llm':False}) for _ in range(3)]
            self.assertTrue(active.wait(1));self.assertEqual(workflow.JOBS[jobs[2]['id']]['status'],'queued')
            from backend import history
            rows=history.query(active_ids=[j['id'] for j in jobs])['records']
            self.assertEqual(next(r['status'] for r in rows if r['id']==jobs[2]['id']),'queued')
            try:
                cancelled=workflow.cancel(jobs[2]['id'])
                self.assertEqual(cancelled['status'],'cancelled')
                self.assertFalse(release.is_set())
                self.assertNotIn('_future',cancelled)
                self.assertNotIn('_history_data',cancelled)
                self.assertEqual(next(r['status'] for r in history.query(active_ids=[j['id'] for j in jobs[:2]])['records'] if r['id']==jobs[2]['id']),'cancelled')
                self.assertEqual(workflow.cancel(jobs[2]['id'])['status'],'cancelled')
            finally: release.set()
            for job in jobs:self.wait(job['id'])
        self.assertEqual(count[0],2);self.assertEqual(workflow.JOBS[jobs[2]['id']]['status'],'cancelled')
        self.assertEqual(next(r['status'] for r in history.query()['records'] if r['id']==jobs[2]['id']),'cancelled')
    def test_maintenance_cannot_start_new_task(self):
        task_queue.MAINTENANCE=True
        with self.assertRaisesRegex(ValueError,'校验'):workflow.start({'kind':'check','paths':[self.doc('sample.docx','test')],'use_llm':False})
    def test_parallel_review_keeps_order_and_context(self):
        segments=[Segment(str(i),'sample','Screen count 214,000',f'slide{i}/p0',slide=i+1,file_path='sample.pptx') for i in range(6)]
        peak=[0];active=[0];lock=threading.Lock()
        def judge(*args):
            with lock: active[0]+=1;peak[0]=max(peak[0],active[0])
            time.sleep(.08)
            with lock: active[0]-=1
            return [],False
        started=time.monotonic()
        with patch('backend.matcher._judge_with_retry',judge): findings,stats=check_segments(segments,self.g,client=object())
        self.assertGreaterEqual(peak[0],2);self.assertLessEqual(peak[0],3);self.assertEqual(stats['judged'],6)
        self.assertLess(time.monotonic()-started,.42);self.assertEqual(findings,[])
    def test_parallel_ppt_translation_retains_geometry_and_context(self):
        path=self.root/'slides.pptx';deck=Presentation();frames=[]
        for i in range(3):
            slide=deck.slides.add_slide(deck.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1));box.text=f'屏幕数量 {11+i}';frames.append((box.left,box.top,box.width,box.height))
        deck.save(path);before=store.fingerprint(path)
        class Model:
            active=0;peak=0;lock=threading.Lock()
            def chat_json(self,prompt,**kw):
                import re
                with self.lock:self.active+=1;self.peak=max(self.peak,self.active)
                time.sleep(.06)
                text=prompt.split('【待译原文】')[-1].split('【需保留')[0].strip()
                with self.lock:self.active-=1
                return {'translation':'Screen count '+re.search(r'\d+',text)[0]}
        model=Model();result=translate_file(str(path),self.g,None,self.root/'output',client=model)
        self.assertGreaterEqual(model.peak,2);out=Presentation(result['path'])
        for i,slide in enumerate(out.slides):
            box=slide.shapes[0];self.assertEqual((box.left,box.top,box.width,box.height),frames[i]);self.assertEqual(box.text,f'Screen count {11+i}')
        self.assertEqual(store.fingerprint(path),before)
        value=comparison.read(result['comparison_id']);self.assertEqual(len(value['rows']),3);self.assertEqual(value['rows'][2]['source'],'屏幕数量 13')
    def test_lan_sessions_paths_csrf_and_manual_replacement(self):
        app=lan_service.create_app(True);first=app.test_client();second=app.test_client()
        csrf1=first.get('/shared/session').json['csrf'];csrf2=second.get('/shared/session').json['csrf']
        self.assertNotEqual(csrf1,csrf2)
        self.assertEqual(first.get('/api/model/settings').status_code,404)
        self.assertEqual(first.post('/shared/jobs').status_code,403)
        self.assertEqual(first.get('/',headers={'Host':'evil.invalid'}).status_code,400)
        self.assertEqual(first.post('/shared/jobs',headers={'Origin':'http://evil.invalid','X-Star-CSRF':csrf1}).status_code,403)
        data=Path(self.doc('source.docx','dot spacing')).read_bytes()
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            response=first.post('/shared/jobs',headers={'X-Star-CSRF':csrf1},data={'options':json.dumps({'kind':'check','use_llm':False,'paths':['C:/private.docx']}),'files':(io.BytesIO(data),'source.docx')})
            self.assertEqual(response.status_code,200,response.json);job=self.wait(response.json['id'])
        job_id=job['id'];self.assertEqual(second.get('/shared/jobs/'+job_id).status_code,400)
        public=first.get('/shared/jobs/'+job_id).json
        self.assertNotIn(str(self.root),json.dumps(public));self.assertFalse(public['result']['findings'][0]['applied'])
        saved=first.post(f'/shared/jobs/{job_id}/apply',headers={'X-Star-CSRF':csrf1},json={'file_index':0,'edits':[{'index':0,'text':'pixel pitch'}]})
        self.assertEqual(saved.status_code,200,saved.json);self.assertTrue(saved.json['result']['findings'][0]['applied'])
        download=first.get(f'/shared/jobs/{job_id}/download/0');self.assertEqual(download.status_code,200)
        self.assertEqual(Document(io.BytesIO(download.data)).paragraphs[0].text,'pixel pitch')
        download.close()
        self.assertEqual(second.get(f'/shared/jobs/{job_id}/download/0').status_code,400)
    def test_lan_rejects_lock_and_non_office_file(self):
        client=lan_service.create_app(True).test_client();csrf=client.get('/shared/session').json['csrf']
        for name in ['~$source.pptx','test.exe','fake.pptx']:
            response=client.post('/shared/jobs',headers={'X-Star-CSRF':csrf},data={'options':'{"kind":"check"}','files':(io.BytesIO(b'test'),name)})
            self.assertEqual(response.status_code,400)
    def test_same_filename_outputs_cannot_overwrite_each_other(self):
        paths=[self.doc(f'{i}/same.docx','屏幕数量 11') for i in range(2)]
        class Model:
            def chat_json(self,*args,**kw):return {'translation':'Screen count 11'}
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(2) as pool:outputs=list(pool.map(lambda p:translate_file(p,self.g,None,self.root/'output',client=Model()),paths))
        self.assertNotEqual(outputs[0]['path'],outputs[1]['path'])
        self.assertTrue(all(Path(o['path']).is_file() for o in outputs))

    def test_request_slots_are_bounded_across_all_tasks(self):
        from concurrent.futures import ThreadPoolExecutor
        active,peak=[0],[0];lock=threading.Lock()
        def request(_):
            with task_queue.request_slot():
                with lock:active[0]+=1;peak[0]=max(peak[0],active[0])
                time.sleep(.03)
                with lock:active[0]-=1
        with ThreadPoolExecutor(12) as pool:list(pool.map(request,range(16)))
        self.assertEqual(peak[0],4);self.assertEqual(active[0],0)

    def test_cleanup_requires_ownership_and_retains_uploaders_original(self):
        app=lan_service.create_app(True);client=app.test_client();foreign=app.test_client()
        csrf=client.get('/shared/session').json['csrf'];other=foreign.get('/shared/session').json['csrf']
        original=self.doc('original.docx','sample')
        with patch('backend.workflow.store.load_glossary',return_value=self.g):
            response=client.post('/shared/jobs',headers={'X-Star-CSRF':csrf},data={'options':'{"kind":"check","use_llm":false}','files':(io.BytesIO(Path(original).read_bytes()),'sample.docx')})
            job=self.wait(response.json['id'])
        job_id=job['id'];folder=next(Path(s['folders'][job_id]) for s in lan_service.SESSIONS.values() if job_id in s.get('folders',{}))
        self.assertEqual(foreign.post(f'/shared/jobs/{job_id}/cleanup',headers={'X-Star-CSRF':other}).status_code,400)
        self.assertTrue(folder.is_dir())
        diagnostics=client.get(f'/shared/jobs/{job_id}/diagnostics');self.assertEqual(diagnostics.status_code,200);self.assertNotIn('segments',diagnostics.json);diagnostics.close()
        self.assertEqual(client.post(f'/shared/jobs/{job_id}/cleanup',headers={'X-Star-CSRF':csrf}).status_code,200)
        self.assertFalse(folder.exists());self.assertTrue(Path(original).is_file());self.assertNotIn(job_id,workflow.REPORTS)


if __name__=='__main__':unittest.main()
