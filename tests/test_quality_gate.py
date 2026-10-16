"""Regressions for content gates, workbook cell identity and live export feedback."""
import os,sys,tempfile,threading,time,unittest
from pathlib import Path
from unittest.mock import patch
from io import BytesIO
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
if os.environ.get('QA_PYTHON_LIB'):sys.path.append(os.environ['QA_PYTHON_LIB'])
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from openpyxl import load_workbook
from docx import Document
from backend import workflow,report_exports
from backend.extract.office import open_office
from backend.replacer.pptx_package import save_office
from backend.translation import translate_file
from backend.glossary.loader import Glossary
from backend.matcher.residual import residual_findings
from backend.models import Segment,Finding
from backend.extract import extract
from backend.glossary.store import fingerprint

class QualityGateTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.addCleanup(self.tmp.cleanup)
        env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.root/'profile')});env.start();self.addCleanup(env.stop)
    def chart(self,series=1):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);data=CategoryChartData()
        data.categories=['显示屏','显示屏']
        for i in range(series):data.add_series('Screen count '+str(i),[10,20])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),Inches(1),Inches(6),Inches(3),data)
        path=self.root/'source.pptx';deck.save(path);return path
    def book(self,path):
        with ZipFile(path) as z:
            name=next(n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx'))
            return load_workbook(BytesIO(z.read(name)))
    def test_equal_source_strings_in_different_cells_keep_distinct_translations(self):
        source=self.chart();original=source.read_bytes();doc,paras,_=open_office(source)
        for p in next(iter(doc._star_workbooks.values())):
            if p.original=='显示屏':p.runs[0].text='Cinema screens' if p.cell[1]=='A2' else 'LED screens'
        target=self.root/'result.pptx';save_office(source,doc,target,[loc for loc,p,*_ in paras if getattr(p,'original',None)=='显示屏'])
        sheet=self.book(target).active
        self.assertEqual((sheet['A2'].value,sheet['A3'].value),('Cinema screens','LED screens'))
        self.assertEqual((sheet['B2'].value,sheet['B3'].value),(10,20))
        with ZipFile(source) as a,ZipFile(target) as b:self.assertEqual(a.read('ppt/slides/slide1.xml'),b.read('ppt/slides/slide1.xml'))
        self.assertEqual(source.read_bytes(),original)
    def test_two_chart_caches_for_one_cell_are_reconciled_without_aborting(self):
        source=self.chart(2);doc,paras,_=open_office(source)
        same=[p for p in next(iter(doc._star_workbooks.values())) if p.cell==('Sheet1','A2')]
        self.assertEqual(len(same),2)
        same[0].runs[0].text='Cinema screens';same[1].runs[0].text='Cinema displays'
        target=self.root/'result.pptx';save_office(source,doc,target,[loc for loc,p,*_ in paras if p in same])
        self.assertEqual(self.book(target).active['A2'].value,'Cinema screens')
        self.assertEqual({p.text for p in same},{'Cinema screens'})
        self.assertEqual(self.book(target).active['A3'].value,'显示屏')
    def test_chart_label_update_keeps_embedded_spreadsheet_formulas(self):
        source=self.chart();book=self.book(source);book.active['A2']='="显示屏"';data=BytesIO();book.save(data)
        staged=self.root/'formula.pptx'
        with ZipFile(source) as a,ZipFile(staged,'w') as b:
            for info in a.infolist():b.writestr(info,data.getvalue() if info.filename.startswith('ppt/embeddings/') and info.filename.endswith('.xlsx') else a.read(info.filename))
        doc,paras,_=open_office(staged);labels=next(iter(doc._star_workbooks.values()))
        for p in labels:
            if p.cell==('Sheet1','A2'):p.runs[0].text='Cinema screens'
        output=self.root/'result.pptx';save_office(staged,doc,output,[loc for loc,p,*_ in paras if getattr(p,'cell',None)==('Sheet1','A2')])
        self.assertEqual(self.book(output).active['A2'].value,'="显示屏"')
    def test_chart_point_label_fits_plot_space_without_losing_number_or_emphasis(self):
        from lxml import etree
        from backend.replacer.pptx_layout import capture_layout,fit_document
        source=self.chart();deck=Presentation(source);chart=deck.slides[0].shapes[0].chart
        ns='http://schemas.openxmlformats.org/drawingml/2006/chart';a='http://schemas.openxmlformats.org/drawingml/2006/main'
        series=chart._chartSpace.find('.//{'+ns+'}ser')
        series.append(etree.fromstring(('<c:dLbls xmlns:c="'+ns+'" xmlns:a="'+a+'" xmlns:c15="http://schemas.microsoft.com/office/drawing/2012/chart"><c:dLbl><c:idx val="1"/><c:tx><c:rich><a:bodyPr vertOverflow="ellipsis" lIns="38100" rIns="38100" tIns="19050" bIns="19050"/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="2000"/></a:pPr><a:r><a:rPr sz="2000" b="1"/><a:t>预计500+</a:t></a:r><a:endParaRPr sz="2000"/></a:p></c:rich></c:tx><c:extLst><c:ext uri="fixture"><c15:layout><c:manualLayout><c:w val="0.225"/><c:h val="0.097"/></c:manualLayout></c15:layout></c:ext></c:extLst></c:dLbl></c:dLbls>').encode()))
        deck.save(source);doc,paras,_=open_office(source)
        loc,p,*_=next(item for item in paras if item[1].text=='预计500+')
        snapshots=capture_layout(doc,[loc]);p.runs[0].text='Estimated 500+'
        fit_document(doc,[loc],snapshots)
        self.assertEqual(p.text,'Estimated 500+');self.assertTrue(p.runs[0].font.bold)
        size=p.runs[0].font.size.pt;self.assertLessEqual(size,11);self.assertGreaterEqual(size,9)
        self.assertEqual(int(p._p.find('{'+a+'}endParaRPr').get('sz')),round(size*100))
        output=self.root/'chart-fit.pptx';save_office(source,doc,output,[loc])
        self.assertEqual(self.book(output).active['B2'].value,10)
    def test_empty_residual_suggestion_can_be_manually_edited(self):
        findings,incomplete=residual_findings(Segment('s','source.pptx','显示控制器','slide1/shape0/p0'),Glossary([]))
        self.assertTrue(incomplete);self.assertTrue(findings[0].editable);self.assertEqual(findings[0].suggested,'')
    def test_whole_sentence_and_term_choices_merge_without_losing_bold_words(self):
        source=self.root/'English.docx';doc=Document();p=doc.add_paragraph()
        p.add_run('Many ');p.add_run('cabinets').bold=True;p.add_run(' has errors.');doc.save(source)
        segment=extract(str(source))[0]
        findings=[Finding(segment.id,source.name,None,1,'text',old,suggested=new,file_path=str(source),method='llm') for old,new in [('cabinets','display cabinets'),('Many cabinets has errors.','Many cabinets have errors.')]]
        workflow.REPORTS['quality-merge']={'segments':[segment],'hashes':{str(source):fingerprint(source)},'findings':findings}
        self.addCleanup(lambda:workflow.REPORTS.pop('quality-merge',None))
        result=workflow.apply_report(dict(report_id='quality-merge',path=str(source),edits=[{'index':i,'text':f.suggested} for i,f in enumerate(findings)]))
        p=Document(result['output']).paragraphs[0]
        self.assertEqual(p.text,'Many display cabinets have errors.')
        self.assertTrue(next(r for r in p.runs if 'cabinets' in r.text).bold)
        self.assertTrue(all(f.applied for f in findings));self.assertEqual(result['failed'],[])
    def test_failed_translation_gets_a_full_paragraph_repair_before_completion(self):
        source=self.root/'source.docx';doc=Document();doc.add_paragraph('亮度500尼特');doc.add_paragraph('刷新率60Hz');doc.save(source)
        class Model:
            def __init__(self):self.responses=iter(['Brightness 500 nits','still Chinese','still Chinese','60 Hz'])
            def chat_json(self,prompt,**kw):return {'translation':next(self.responses)}
        result=translate_file(str(source),Glossary([]),None,client=Model(),workers=1)
        self.assertTrue(result['translation_complete']);self.assertEqual(result['untranslated'],0)
        self.assertEqual([p.text for p in Document(result['path']).paragraphs],['Brightness 500 nits','60 Hz'])
    def test_incomplete_draft_is_reported_and_can_be_saved_without_bulk_retranslation(self):
        source=self.root/'source.docx';doc=Document();doc.add_paragraph('中文原稿');doc.save(source)
        draft=self.root/'draft.docx';doc.save(draft)
        result=dict(path=str(draft),source=str(source),translation_complete=False,translated=0,total=1,untranslated=1,status='needs_review',needs_review=1,review_items=[{'locator':extract(str(draft))[0].locator}],warnings=[])
        job=dict(id='quality-incomplete',kind='translate_review',cancel=threading.Event(),status='running')
        class Model:
            def chat_json(self,*args,**kwargs):raise AssertionError('Known untranslated Chinese must not be sent to English proofreading')
        with patch('backend.workflow.translate_file',return_value=result):
            workflow.run(job,dict(kind='translate_review',paths=[str(source)],use_llm=True,_client=Model()))
        self.assertFalse(job['result']['completed_paths']);self.assertEqual(result['review_status'],'partial')
        report=job['result']['report'];self.assertEqual(len(report['findings']),1)
        self.assertIn('翻译待核对',report['findings'][0]['reason'])
        saved=workflow.apply_report(dict(report_id=job['id'],path=str(draft),save_unchanged=True))
        self.assertEqual(Document(saved['output']).paragraphs[0].text,'中文原稿')
        self.assertTrue(saved['content_warnings'])
    def test_language_detection_parse_failure_uses_normal_task_error_handling(self):
        from zipfile import BadZipFile
        source=self.root/'invalid.pptx';source.write_bytes(b'not an office archive')
        job=dict(id='quality-invalid-source',kind='check',cancel=threading.Event(),status='running')
        with patch('backend.workflow.predominantly_chinese',side_effect=BadZipFile('invalid archive')):
            workflow.run(job,dict(kind='check',paths=[str(source)],use_llm=True))
        self.assertEqual(job['status'],'failed');self.assertTrue(job['error'])
    def test_equal_selected_suggestions_are_all_marked_applied(self):
        source=self.root/'same.docx';doc=Document();doc.add_paragraph('dot spacing');doc.save(source);segment=extract(str(source))[0]
        findings=[Finding(segment.id,source.name,None,1,'text','dot spacing',suggested='pixel pitch',file_path=str(source),method=method) for method in ('rule','llm')]
        workflow.REPORTS['quality-duplicate']={'segments':[segment],'hashes':{str(source):fingerprint(source)},'findings':findings}
        result=workflow.apply_report(dict(report_id='quality-duplicate',path=str(source),edits=[dict(index=i,text='pixel pitch') for i in range(2)]))
        self.assertEqual(result['failed'],[]);self.assertTrue(all(f.applied for f in findings));self.assertEqual(Document(result['output']).paragraphs[0].text,'pixel pitch')
    def test_chinese_original_in_proofreading_is_translated_first_and_history_matches(self):
        from backend import history
        source=self.root/'source.docx';doc=Document();doc.add_paragraph('这是一份完整的中文技术文档原稿');doc.save(source)
        draft=self.root/'draft.docx';doc=Document();doc.add_paragraph('This is a technical document.');doc.save(draft)
        class Model:
            def chat_json(self,*args,**kwargs):return []
        data=dict(kind='check',paths=[str(source)],use_llm=True,_client=Model())
        job=dict(id='quality-chinese-routing',cancel=threading.Event(),status='running')
        history.begin(job['id'],data,'queued')
        result=dict(path=str(draft),source=str(source),translation_complete=True,translated=1,total=1,untranslated=0,status='complete',needs_review=0,review_items=[],warnings=[])
        with patch('backend.workflow.translate_file',return_value=result):workflow.run(job,data)
        self.assertEqual(job['kind'],'translate_review');self.assertEqual(job['result']['report']['stats']['residual_chinese'],0)
        row=history.record(job['id']);self.assertEqual(row['kind'],'translate_review');self.assertEqual(row['resume_options']['kind'],'translate_review')
    def test_background_export_exposes_live_phase_without_blocking(self):
        entered=threading.Event();release=threading.Event();workflow.REPORTS['quality-export']={}
        self.addCleanup(lambda:workflow.REPORTS.pop('quality-export',None))
        def apply(data,progress=None,client=None):
            progress('render',1,4,'正在逐页检查版式');entered.set();release.wait(5);return {'output':'fixture.docx'}
        with patch('backend.workflow.apply_report',side_effect=apply):
            task=report_exports.start({'report_id':'quality-export'})
            try:
                self.assertTrue(entered.wait(2));live=report_exports.read(task['id'],'quality-export')
                from backend import task_queue
                self.assertGreater(task_queue.ACTIVE_EXPORTS,0)
                self.assertEqual((live['phase'],live['done'],live['total']),('render',1,4))
                with self.assertRaises(ValueError):report_exports.read(task['id'],'other-report')
            finally:release.set()
            for _ in range(100):
                if report_exports.read(task['id'])['status']=='complete':break
                time.sleep(.01)
        self.assertEqual(report_exports.read(task['id'])['result']['output'],'fixture.docx')
    def test_shared_background_save_keeps_session_ownership_and_updates_report(self):
        import io
        from backend import lan_service
        from backend.models import GlossaryEntry
        service=lan_service.create_app(True);owner=service.test_client();other=service.test_client()
        csrf=owner.get('/shared/session').json['csrf'];other.get('/shared/session')
        doc=Document();doc.add_paragraph('dot spacing');data=BytesIO();doc.save(data)
        g=Glossary([GlossaryEntry('点间距','pixel pitch',wrong_variants=['dot spacing'],priority=1)])
        with patch('backend.workflow.store.load_glossary',return_value=g):
            response=owner.post('/shared/jobs',headers={'X-Star-CSRF':csrf},data={'options':'{"kind":"check","use_llm":false}','files':(io.BytesIO(data.getvalue()),'source.docx')})
            job_id=response.json['id']
            for _ in range(200):
                if workflow.JOBS[job_id]['status'] not in ('queued','running'):break
                time.sleep(.01)
        response=owner.post(f'/shared/jobs/{job_id}/exports',headers={'X-Star-CSRF':csrf},json={'file_index':0,'edits':[{'index':0,'text':'pixel pitch'}]})
        self.assertEqual(response.status_code,200,response.json);export_id=response.json['id']
        url=f'/shared/jobs/{job_id}/exports/{export_id}'
        self.assertEqual(other.get(url).status_code,400)
        for _ in range(200):
            task=owner.get(url).json
            if task['status'] not in ('queued','running'):break
            time.sleep(.01)
        self.assertEqual(task['status'],'complete',task)
        self.assertTrue(owner.get(f'/shared/jobs/{job_id}').json['result']['findings'][0]['applied'])
    def test_incomplete_layout_check_is_disclosed_but_does_not_block_saving(self):
        from backend.models import GlossaryEntry
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)).text='dot spacing'
        source=self.root/'English.pptx';deck.save(source)
        job=dict(id='quality-format-warning',kind='check',status='running',cancel=threading.Event())
        g=Glossary([GlossaryEntry('点间距','pixel pitch',wrong_variants=['dot spacing'],priority=1)])
        with patch('backend.workflow.store.load_glossary',return_value=g):workflow.run(job,dict(kind='check',paths=[str(source)],use_llm=False))
        with patch('backend.workflow.review_pptx',return_value=([{'slide':1,'reason':'文档排版处理未完成。'}],[])):
            result=workflow.apply_report(dict(report_id=job['id'],path=str(source),edits=[dict(index=0,text='pixel pitch')]))
        self.assertTrue(Path(result['output']).is_file());self.assertFalse(result['layout_checked']);self.assertTrue(result['layout_warnings'])

    def test_calendar_dates_and_named_plan_are_equal_but_wrong_dates_still_fail(self):
        from backend.writing_policy import validate_translation
        examples=[('2025 年 1 月，建立中心。','A center was established in January 2025.'),
            ('2024 年 12 月 23 日，中国北京','December 23, 2024, Beijing, China'),
            ('2024 年 12 月 23 日，中国北京','23 December 2024, Beijing, China'),
            ('2021《“十四五”规划》支持 LED 电影屏。','The 2021 14th Five-Year Plan supports LED cinema screens.')]
        for source,target in examples:validate_translation(source,target)
        for source,target in [('2025 年 1 月','February 2025'),('2024 年 12 月 23 日','December 24, 2024'),
                ('2021《“十四五”规划》','The 2021 15th Five-Year Plan'),('支持 LED 电影屏','Supports 14 LED cinema screens')]:
            with self.assertRaises(ValueError):validate_translation(source,target)
    def test_complex_ppt_emphasis_uses_one_reasoning_request_before_style_loss(self):
        from types import SimpleNamespace
        from backend.translation import mapped_translation
        class Model:
            supports_thinking_mode=True
            def __init__(self):self.calls=[]
            def chat_json(self,prompt,**kw):
                self.calls.append(kw)
                return {'translation':'Important information','runs':[{'id':0,'text':'Important '},{'id':1,'text':'information'}]}
        model=Model();groups=[{'id':0,'text':'重要','key':'large-bold'},{'id':1,'text':'说明','key':'body'}]
        runs,approx=mapped_translation(model,'fixture','重要说明',groups,None,segment=SimpleNamespace(file_path='fixture.pptx',file='fixture.pptx'))
        self.assertFalse(approx);self.assertEqual(len(model.calls),1);self.assertTrue(model.calls[0]['thinking'])

    def test_compact_calendar_dates_match_named_dates_without_weakening_quantity_checks(self):
        from backend.writing_policy import validate_translation,repair_numeric_translation
        valid=[
            ('2025.6 影院通知：设立专项资金。','June 2025 cinema notice: Dedicated funding.'),
            ('2025.06 影院通知：设立专项资金。','Jun. 2025 cinema notice: Dedicated funding.'),
            ('截至2025.6，支持300尼特。','As of June 2025, supports 300 nits.'),
            ('于2025年6月17日发布，亮度300尼特。','Released on 17 June 2025, with a brightness of 300 nits.'),
            ('发布日期：2025-06-17','Release date: June 17th, 2025'),
            ('2025.6.17 发布通知','Notice released on 17 June 2025'),
            ('2025.6 影院通知','2025.6 Cinema Notice'),
            ('2025年5月通知','May 2025 notice'),
            ('2025年6月与2025年7月','June 2025 and July 2025'),
            ('2025.6nits亮度，版本2025.6','2025.6 nits brightness, version 2025.6')]
        for source,target in valid:
            with self.subTest(source=source,target=target):
                validate_translation(source,target)
                self.assertEqual(repair_numeric_translation(source,target),target)
        invalid=[
            ('2025.6 影院通知','July 2025 cinema notice'),
            ('2025.6 影院通知','June 2026 cinema notice'),
            ('2025年6月17日','June 6, 2025 and 17'),
            ('2025年6月17日','June 2025'),
            ('2025年6月与2025年7月','June 2025 and June 2025'),
            ('2025.6 影院通知，亮度300尼特','June 2025 cinema notice, brightness 301 nits'),
            ('数值2025.6','June 2025'),
            ('2025.6尼特亮度通知','June 2025 brightness notice'),
            ('版本2025.6','June 2025'),
            ('型号2025.6.17','June 17, 2025'),
            ('2025.6万元','June 2025'),
            ('2025.6','June 2025')]
        for source,target in invalid:
            with self.subTest(source=source,target=target):
                with self.assertRaises(ValueError):validate_translation(source,target)

    def test_dated_ppt_notice_completes_translation_and_starts_real_proofreading(self):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1))
        shape.text='2025.6 影院通知：设立专项资金。'
        source=self.root/'dated-notice.pptx';deck.save(source);original=source.read_bytes()
        class Model:
            def __init__(self):self.translations=0;self.reviews=0
            def chat_json(self,prompt,**kw):
                if kw.get('system','').startswith('你是 LED 显示行业的中英技术翻译'):
                    self.translations+=1
                    return {'translation':'June 2025 cinema notice: Dedicated funding.'}
                self.reviews+=1
                return []
        model=Model();job=dict(id='quality-compact-date-review',kind='translate_review',cancel=threading.Event(),status='running')
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])),patch('backend.workflow.review_pptx',return_value=([],[])):
            workflow.run(job,dict(kind='translate_review',paths=[str(source)],use_llm=True,_client=model,workers=1))
        self.assertTrue(job['result']['outputs'][0]['translation_complete'])
        self.assertEqual(model.translations,1);self.assertGreater(model.reviews,0)
        self.assertGreater(job['result']['report']['stats']['judged'],0)
        self.assertEqual(source.read_bytes(),original)

    def test_one_unfixable_numeric_paragraph_does_not_block_review_or_final_export(self):
        source=self.root/'partial-source.docx';doc=Document()
        doc.add_paragraph('分辨率为3840×2160。');doc.add_paragraph('亮度500尼特。');doc.save(source)
        original=source.read_bytes()
        class Model:
            def __init__(self):self.reviews=0;self.translations=0
            def chat_json(self,prompt,**kw):
                if kw.get('system','').startswith('你是 LED 显示行业的中英技术翻译'):
                    self.translations+=1
                    piece=prompt.split('【待译原文】\n',1)[1].split('\n【',1)[0]
                    return {'translation':'1920 x 1080.' if '3840' in piece else 'Brightness 500 nits.'}
                self.reviews+=1;return []
        client=Model();job=dict(id='quality-one-local-failure',kind='translate_review',status='running',cancel=threading.Event())
        with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
            workflow.run(job,dict(kind='translate_review',paths=[str(source)],combined_output='final',output_dir=str(self.root/'output'),_client=client,workers=1))
        output=job['result']['outputs'][0];report=job['result']['report']
        self.assertEqual(job['status'],'partial');self.assertTrue(job['can_resume'])
        self.assertEqual(report['stats']['judged'],1);self.assertEqual(report['stats']['unreviewed'],1)
        self.assertEqual(client.reviews,1);self.assertLessEqual(client.translations,5)
        saved=workflow.apply_report(dict(report_id=job['id'],path=output['path'],save_unchanged=True))
        self.assertEqual([p.text for p in Document(saved['output']).paragraphs],['分辨率为3840×2160。','Brightness 500 nits.'])
        self.assertTrue(saved['content_warnings']);self.assertEqual(source.read_bytes(),original)
        issue=next(i for i,f in enumerate(report['findings']) if f['reason'].startswith('翻译待核对'))
        fixed=workflow.apply_report(dict(report_id=job['id'],path=output['path'],edits=[dict(index=issue,text='Resolution: 3840 x 2160.')]))
        self.assertEqual(fixed['output'],saved['output']);self.assertFalse(fixed['content_warnings'])
        self.assertEqual(Document(fixed['output']).paragraphs[0].text,'Resolution: 3840 x 2160.')

    def test_total_model_outage_still_saves_a_marked_document(self):
        from backend.llm.client import LLMUnavailable
        source=self.root/'offline-source.docx';doc=Document();doc.add_paragraph('中文原稿');doc.save(source)
        original=source.read_bytes()
        class Model:
            def __init__(self):self.calls=0
            def chat_json(self,*args,**kwargs):self.calls+=1;raise LLMUnavailable('连接中断')
        model=Model();result=translate_file(str(source),Glossary([]),None,client=model,workers=1)
        self.assertFalse(result['translation_complete']);self.assertTrue(result['review_items'])
        self.assertEqual(Document(result['path']).paragraphs[0].text,'中文原稿')
        self.assertEqual(model.calls,1);self.assertEqual(source.read_bytes(),original)

    def test_unfinished_text_with_a_changed_locator_is_reported_on_its_actual_paragraph(self):
        source=self.root/'locator-source.docx';doc=Document();doc.add_paragraph('原稿内容');doc.save(source)
        output=self.root/'locator-output.docx';doc=Document();doc.add_paragraph('未完成中文');doc.add_paragraph('Completed English.');doc.save(output)
        result=dict(path=str(output),source=str(source),translation_complete=False,review_items=[dict(locator='old-locator',translation='未完成中文',reason='该处未完成')],warnings=[])
        class Model:
            def chat_json(self,*args,**kwargs):return []
        review=workflow.auto_review(result,Glossary([]),None,threading.Event(),None,client=Model(),workers=1)
        self.assertEqual(review['stats']['judged'],1);self.assertEqual(review['stats']['unreviewed'],1)
        finding=next(f for f in review['local']['findings'] if f.reason.startswith('翻译待核对'))
        self.assertEqual(finding.found_text,'未完成中文');self.assertEqual(finding.segment_id,extract(str(output))[0].id)
    def test_incomplete_translation_is_partial_and_directly_resumable(self):
        source=self.root/'partial.docx';doc=Document();doc.add_paragraph('完整中文原稿');doc.save(source)
        draft=self.root/'draft.docx';doc.save(draft)
        result=dict(path=str(draft),source=str(source),translation_complete=False,translated=0,total=1,
            untranslated=1,status='needs_review',needs_review=1,review_items=[dict(locator=extract(str(draft))[0].locator,reason='连接暂不可用')],warnings=[])
        job=dict(id='quality-direct-resume',kind='translate_review',cancel=threading.Event(),status='running')
        with patch('backend.workflow.translate_file',return_value=result):
            workflow.run(job,dict(kind='translate_review',paths=[str(source)]))
        self.assertEqual(job['status'],'partial');self.assertTrue(job['can_resume'])
        self.assertEqual(job['result']['stats']['model_status'],'partial')
    def test_combined_review_runs_format_comparison_after_text_review(self):
        source=self.root/'source.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='Technical information';deck.save(source)
        result=dict(path=str(source),source=str(source),format_review_deferred=True,review_items=[],warnings=[])
        order=[]
        class Model:
            def chat_json(self,*args,**kwargs):order.append('text');return []
        def visual(*args,**kwargs):
            order.append('format');kwargs['progress']('正在逐页检查版式',1,1);kwargs['audit'].update(pages=1,image_pages=1);return [],[]
        with patch('backend.workflow.review_pptx',side_effect=visual):
            review=workflow.auto_review(result,Glossary([]),None,threading.Event(),lambda *a:None,
                client=Model(),format_progress=lambda *a:order.append('progress'))
        self.assertLess(order.index('text'),order.index('format'));self.assertEqual(result['format_review']['image_pages'],1)
        self.assertTrue(result['layout_checked']);self.assertEqual(review['stats']['format_pages'],1)

if __name__=='__main__':unittest.main()
