"""Layout-aware proofreading regressions; no network or real model required."""
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'dist/win-unpacked/resources/app.asar.unpacked' if '--packaged' in sys.argv else ROOT))
if '--packaged' in sys.argv: sys.argv.remove('--packaged')
from pptx import Presentation
from pptx.util import Inches, Pt
from backend.extract import extract
from backend.document_context import DocumentContext
from backend.models import Finding, Segment
from backend.glossary.loader import Glossary
from backend.glossary import store
from backend.matcher import check_segments
from backend.references import ReferenceIndex
from backend import workflow

TITLE = 'Global Screen Total Exceeds'
NUMBER = '214,000 screens, a huge market.'
BAD = dict(found=TITLE, suggested=TITLE+' 214,000', kind='grammar', reason='exceeds needs an object')

class Model:
    def __init__(self): self.prompts = []
    def chat_json(self, prompt, **kw):
        self.prompts.append((prompt,kw['system']))
        target = prompt.split('【待审校段落】\n',1)[1].split('\n\n输出 JSON',1)[0].strip()
        if target == TITLE: return [BAD]
        if 'flexibile' in target: return [dict(found='flexibile',suggested='flexible',kind='typo',reason='spelling')]
        return []

class ContextTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.env = patch.dict(os.environ, NOVATERM_DATA_DIR=str(self.folder/'data'))
        self.env.start()
        self.glossary = Glossary([])
    def tearDown(self):
        self.env.stop(); self.temp.cleanup()
    def deck(self, name='input.pptx', title=TITLE, same_frame=False, distant=False):
        p = Presentation(); slide = p.slides.add_slide(p.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Pt(30))
        box.text = title
        if same_frame: box.text_frame.add_paragraph().text = NUMBER
        else: slide.shapes.add_textbox(Inches(6 if distant else 1), Inches(2)+Pt(30), Inches(4), Pt(30)).text = NUMBER
        slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Pt(30)).text = 'A flexibile display.'
        path = self.folder/name; p.save(path); return str(path)
    def test_touching_separate_boxes_are_context_not_two_incomplete_sentences(self):
        path = self.deck(); segments = extract(path); model = Model(); streamed = []
        findings,stats = check_segments(segments,self.glossary,client=model,on_finding=streamed.extend)
        self.assertEqual([f.found_text for f in findings],['flexibile'])
        self.assertEqual([f['found_text'] for f in streamed],['flexibile'])
        self.assertEqual(stats['discarded_context_hits'],1)
        self.assertEqual(stats['model_status'],'complete'); self.assertEqual(stats['completed_files'],[path])
        prompt,system = model.prompts[0]
        local = prompt.split('【当前文本框及紧邻且对齐的文字',1)[1].split('【同页整体内容',1)[0]
        self.assertIn(TITLE,local); self.assertIn(NUMBER,local)
        self.assertIn('当前段落 [slide1/shape0/p0]',local)
        self.assertIn('相邻段落 [slide1/shape1/p0]',local)
        self.assertIn('不要求每个文本框',system)
    def test_same_frame_continuation_blocks_numeric_copy(self):
        s = extract(self.deck(same_frame=True)); c = DocumentContext(s)
        self.assertTrue(c.duplicates_neighbor_number(s[0],TITLE,TITLE+' 214000'))
        self.assertEqual(c.report_context(s[0]),TITLE+'\n'+NUMBER)
    def test_other_card_with_same_number_is_not_structural_evidence(self):
        s = extract(self.deck(distant=True)); c = DocumentContext(s)
        self.assertFalse(c.duplicates_neighbor_number(s[0],TITLE,TITLE+' 214,000'))
        self.assertNotIn(NUMBER,c.report_context(s[0]))
    def test_existing_numeric_correction_is_allowed(self):
        s = extract(self.deck(title=TITLE+' 241,000')); c = DocumentContext(s)
        self.assertFalse(c.duplicates_neighbor_number(s[0],s[0].text,TITLE+' 214,000'))
    def test_partial_fragment_and_year_heading_cannot_import_neighbor_count(self):
        s = extract(self.deck(title='2025 '+TITLE)); c = DocumentContext(s)
        self.assertTrue(c.duplicates_neighbor_number(s[0],'Exceeds','Exceeds 214,000'))
        self.assertTrue(c.duplicates_neighbor_number(s[0],s[0].text,s[0].text+' 214,000'))
    def test_actual_user_card_geometry_has_connected_reading_order(self):
        actual = ROOT.parent/'测试文档_English_checked.pptx'
        if not actual.exists(): self.skipTest('actual user fixture absent')
        before = actual.read_bytes(); s = extract(str(actual))
        title = next(x for x in s if x.text.strip() == TITLE+' 214,000')
        title.text = TITLE
        c = DocumentContext(s)
        self.assertIn(NUMBER,c.report_context(title))
        self.assertTrue(c.duplicates_neighbor_number(title,TITLE,TITLE+' 214,000'))
        self.assertEqual(before,actual.read_bytes())
    def test_original_number_in_exact_paired_position_can_be_restored(self):
        path = self.deck(); original = self.deck('original.pptx', title='全球银幕总量突破214,000')
        refs = ReferenceIndex([original],self.glossary,{path:original})
        hits,_ = check_segments(extract(path),self.glossary,client=Model(),references=refs)
        title = next(f for f in hits if f.found_text == TITLE)
        self.assertTrue(title.numeric_addition_verified)
    def test_number_only_in_next_source_box_does_not_authorize_copy(self):
        path = self.deck(); original = self.deck('original.pptx',title='全球银幕总量突破')
        refs = ReferenceIndex([original],self.glossary,{path:original})
        hits,_ = check_segments(extract(path),self.glossary,client=Model(),references=refs)
        self.assertEqual([f.found_text for f in hits],['flexibile'])
    def test_export_revalidates_stale_bad_model_suggestion_and_applies_real_error(self):
        path = self.deck(); segments = extract(path)
        title,number,typo = segments
        findings = [Finding(title.id,title.file,1,1,'text',TITLE,suggested=BAD['suggested'],file_path=path),
                    Finding(typo.id,typo.file,1,1,'text','flexibile',suggested='flexible',file_path=path)]
        report = dict(segments=segments,findings=findings,hashes={path:store.fingerprint(path)})
        with patch.object(workflow.officecli_bridge,'inspect',return_value=None),patch.object(workflow.officecli_bridge,'compare',return_value=[]):
            result = workflow._apply_report(dict(path=path,selected={title.id:[[TITLE,BAD['suggested']]],typo.id:[['flexibile','flexible']]}),report)
        self.assertEqual(result['applied_count'],1); self.assertEqual(len(result['failed']),1)
        p = Presentation(result['output']); self.assertEqual(p.slides[0].shapes[0].text,TITLE)
        self.assertEqual(p.slides[0].shapes[1].text,NUMBER)
        self.assertEqual(p.slides[0].shapes[2].text,'A flexible display.')
        self.assertFalse(findings[0].applied)
    def test_export_does_not_change_explicit_user_rewrite(self):
        path = self.deck(); s = extract(path); f = Finding(s[0].id,s[0].file,1,1,'text',TITLE,suggested=BAD['suggested'],file_path=path)
        report = dict(segments=s,findings=[f],hashes={path:store.fingerprint(path)})
        with patch.object(workflow.officecli_bridge,'inspect',return_value=None),patch.object(workflow.officecli_bridge,'compare',return_value=[]):
            result = workflow._apply_report(dict(path=path,edits=[dict(index=0,text='Global Cinema Screen Count')]),report)
        self.assertEqual(Presentation(result['output']).slides[0].shapes[0].text,'Global Cinema Screen Count')
        self.assertIn(NUMBER,result['finding_updates'][0]['finding']['context'])
    def test_group_coordinates_do_not_connect_different_parent_groups(self):
        s = extract(self.deck()); s[1].layout['container'] = 'slide1/other-group'
        self.assertFalse(DocumentContext(s).duplicates_neighbor_number(s[0],TITLE,BAD['suggested']))
    def test_multifile_non_ppt_ids_do_not_change_neighbor_position(self):
        a = [Segment(str(i),'a','a'+str(i),'body/'+str(i),file_path='a') for i in range(30)]
        b = [Segment('0','b','b','body/0',file_path='b')]
        self.assertIn(a[0],DocumentContext(a+b).neighbors(a[0]))

if __name__ == '__main__': unittest.main()
