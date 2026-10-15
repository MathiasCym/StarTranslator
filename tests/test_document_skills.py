"""Offline integration checks for writing rules, inline styles and real OfficeCLI."""
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from backend import officecli_bridge, workflow
from backend.extract import extract
from backend.glossary.loader import Glossary
from backend.models import GlossaryEntry
from backend.matcher.judge import SYSTEM_PROMPT
from backend.translation import translate_file, SYSTEM
from backend.llm.client import LLMUnavailable, LLMResponseInvalid, LLMTruncated
from backend.writing_policy import validate_translation
from backend.replacer.styled_text import style_groups, slice_groups


ROOT = Path(__file__).resolve().parents[1]


def fixture(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('14243B')
    title = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(8), Inches(.7))
    title.text = 'LED 显示方案'
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    body = slide.shapes.add_textbox(Inches(.5), Inches(1.5), Inches(8), Inches(1.1))
    p = body.text_frame.paragraphs[0]
    p.space_after = Pt(6)
    for text, bold, color in [('采用', False, 'FFFFFF'), ('逐点校正', True, '48CCFF'),
                               ('，亮度为500nits。', False, 'FFFFFF')]:
        run = p.add_run(); run.text = text
        run.font.name, run.font.size, run.font.bold = 'Arial', Pt(22), bold
        run.font.color.rgb = RGBColor.from_string(color)
        if bold:
            run.hyperlink.address = 'https://www.novastar.tech/'
    table = slide.shapes.add_table(1, 1, Inches(.5), Inches(3), Inches(6), Inches(.6)).table
    table.cell(0, 0).text = '分辨率为3840×2160。'
    table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    # An untouched picture, notes and a separate slide must survive the round trip.
    slide.shapes.add_picture(str(ROOT / 'build/icon.png'), Inches(8.7), Inches(.4), width=Inches(.8))
    slide.notes_slide.notes_text_frame.text = '点间距为2.5毫米。'
    other = deck.slides.add_slide(deck.slide_layouts[6])
    other.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = 'Existing English content'
    deck.save(path)
    with ZipFile(path, 'a') as archive:
        archive.writestr('qa/unknown-extension.bin', b'Unreferenced vendor extension: do not discard.\x00\xff')
    return path


class Model:
    def __init__(self, bad=False):
        self.prompts, self.bad = [], bad

    def chat_json(self, prompt, **kwargs):
        self.prompts.append((prompt, kwargs))
        if self.bad:
            return {'translation': 'LED display solution with 99 pixels.'}
        source = prompt.split('【待译原文】\n', 1)[1].split('\n【', 1)[0]
        responses = {
            'LED 显示方案': {'translation': 'LED display solution'},
            '采用逐点校正，亮度为500nits。': {
                'translation': 'Use pixel-level calibration at a brightness of 500 nits.',
                'runs': [{'id': 0, 'text': 'Use '}, {'id': 1, 'text': 'pixel-level calibration'},
                         {'id': 2, 'text': ' at a brightness of 500 nits.'}]},
            '分辨率为3840×2160。': {'translation': 'The resolution is 3840 x 2160.'},
            '点间距为2.5毫米。': {'translation': 'The pixel pitch is 2.5 mm.'},
        }
        return responses[source]


class DocumentSkillTests(unittest.TestCase):
    def setUp(self):
        # These tests exercise text/package behavior. Image review has its own suite.
        for target in ('backend.translation.review_pptx','backend.workflow.review_pptx'):
            stub=patch(target,return_value=([],[]));stub.start();self.addCleanup(stub.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.environment = patch.dict(os.environ, {'NOVATERM_DATA_DIR': str(self.root / 'profile')})
        self.environment.start()
        self.glossary = Glossary([GlossaryEntry('点间距', 'pixel pitch', wrong_variants=['dot spacing'], priority=1)])

    def tearDown(self):
        self.environment.stop()
        self.temp.cleanup()

    def test_real_officecli_export_keeps_mixed_styles_links_geometry_and_unedited_parts(self):
        source = fixture(self.root / 'styled.pptx')
        original = source.read_bytes()
        model = Model()
        result = translate_file(str(source), self.glossary, None, client=model)
        output = Path(result['path']); deck = Presentation(output)
        p = deck.slides[0].shapes[1].text_frame.paragraphs[0]
        self.assertEqual(p.text, 'Use pixel-level calibration at a brightness of 500 nits.')
        self.assertEqual([r.font.bold for r in p.runs], [False, True, False])
        self.assertEqual([r.font.color.rgb for r in p.runs], [RGBColor(255,255,255), RGBColor.from_string('48CCFF'), RGBColor(255,255,255)])
        self.assertEqual([r.font.size for r in p.runs], [Pt(22)] * 3)
        self.assertEqual(p.runs[1].hyperlink.address, 'https://www.novastar.tech/')
        self.assertEqual(p.space_after, Pt(6))
        self.assertEqual(deck.slides[0].shapes[1].text_frame.auto_size, MSO_AUTO_SIZE.NONE)
        before_deck = Presentation(source)
        for before, after in zip(before_deck.slides[0].shapes, deck.slides[0].shapes):
            self.assertEqual((before.left, before.top, before.width, before.height),
                             (after.left, after.top, after.width, after.height))
        changed = []
        with ZipFile(source) as before, ZipFile(output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if before.read(name) != after.read(name):
                    changed.append(name)
        self.assertEqual(set(changed), {'ppt/slides/slide1.xml', 'ppt/notesSlides/notesSlide1.xml'})
        self.assertEqual(source.read_bytes(), original)
        self.assertTrue(result['layout_checked'])
        self.assertEqual(result['translated'], 4)
        self.assertEqual(officecli_bridge.command('validate', output)['errors'], [])
        self.assertTrue(all('Write clear, natural technical English' in kwargs['system'] for _, kwargs in model.prompts))
        self.assertIn('【需保留的原文样式片段】', model.prompts[1][0])

    def test_persistently_invalid_facts_are_retried_but_never_published_as_english(self):
        source = fixture(self.root / 'bad.pptx'); original = source.read_bytes(); model = Model(bad=True)
        with self.assertRaises(LLMUnavailable):translate_file(str(source), self.glossary, None, client=model)
        self.assertEqual(list(self.root.glob('*English*')),[])
        self.assertGreaterEqual(len(model.prompts), 8)
        self.assertIn('重新生成要求', model.prompts[1][0])
        self.assertEqual(source.read_bytes(), original)

    def test_unresolved_style_mapping_still_exports_best_available_source_styles(self):
        source = fixture(self.root / 'bad-style.pptx')
        model = Model(); original_call = model.chat_json
        def no_styles(prompt, **kwargs):
            answer = original_call(prompt, **kwargs); answer.pop('runs', None); return answer
        model.chat_json = no_styles
        original=source.read_bytes()
        result=translate_file(str(source),self.glossary,None,client=model)
        self.assertEqual(source.read_bytes(),original)
        output=Presentation(result['path'])
        self.assertIn('Use pixel-level calibration',output.slides[0].shapes[1].text)
        # An unmapped Chinese emphasis must not be assigned to an arbitrary
        # English word simply to make a style-coverage check pass.
        self.assertTrue(result['layout_warnings'])
        self.assertTrue(any('强调' in item['reason'] for item in result['layout_warnings']))

    def test_whitespace_and_punctuation_in_runs_are_aligned_without_retranslation(self):
        source = fixture(self.root / 'spaces.pptx'); model = Model(); call = model.chat_json
        def response(prompt, **kwargs):
            result = call(prompt, **kwargs)
            if 'runs' in result:
                result['runs'][0]['text'] = ' Use'
                result['runs'][-1]['text'] = ' at a brightness of 500 nits.\n'
            return result
        model.chat_json = response
        result = translate_file(str(source), self.glossary, None, client=model)
        p = Presentation(result['path']).slides[0].shapes[1].text_frame.paragraphs[0]
        self.assertEqual(p.text, 'Use pixel-level calibration at a brightness of 500 nits.')
        self.assertTrue(any(r.font.bold and r.text == 'pixel-level calibration' for r in p.runs))
        self.assertEqual(len(model.prompts), 4)
        self.assertEqual(result['layout_warnings'], [])

    def test_style_repair_keeps_the_verified_translation_without_calling_for_new_content(self):
        source = fixture(self.root / 'repair.pptx'); model = Model(); call = model.chat_json
        def response(prompt, **kwargs):
            result = call(prompt, **kwargs)
            if 'runs' in result and '【已确定英文译文】' not in prompt:
                result['runs'][1]['text'] = 'a different wording'
            return result
        model.chat_json = response
        result = translate_file(str(source), self.glossary, None, client=model)
        self.assertEqual(result['layout_warnings'], [])
        self.assertEqual(len(model.prompts), 5)
        p = Presentation(result['path']).slides[0].shapes[1].text_frame.paragraphs[0]
        self.assertEqual(p.text, 'Use pixel-level calibration at a brightness of 500 nits.')
        self.assertTrue(p.runs[1].font.bold)

    def test_generated_json_failure_retries_content_and_does_not_report_file_read_failure(self):
        doc = Document(); doc.add_paragraph('LED 显示方案'); source = self.root / 'retry.docx'; doc.save(source)
        for error in (LLMResponseInvalid('JSON 无效'), LLMTruncated('输出被截断')):
            model = Model(); call = model.chat_json
            answers = iter([error, {'translation': 'LED display solution'}])
            def response(prompt, **kwargs):
                answer = next(answers)
                if isinstance(answer, Exception):
                    raise answer
                return answer
            model.chat_json = response
            result = translate_file(str(source), self.glossary, None, client=model)
            self.assertEqual(Document(result['path']).paragraphs[0].text, 'LED display solution')

    def test_proofing_metadata_does_not_split_one_visual_style_into_many(self):
        doc = Document(); p = doc.add_paragraph()
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        for text, language in [('采用', 'zh-CN'), ('LED', 'en-US'), ('显示方案', 'zh-CN')]:
            run = p.add_run(text); node = OxmlElement('w:lang'); node.set(qn('w:val'), language); run._r.get_or_add_rPr().append(node)
        self.assertEqual(len(style_groups(p)), 1)

    def test_validation_findings_are_a_baseline_but_io_failures_remain_errors(self):
        from subprocess import CompletedProcess
        finding = {'type': 'Schema', 'description': 'WPS extension'}
        response = {'success': False, 'data': {'errors': [finding], 'count': 1}}
        with patch.dict(os.environ, {'OFFICECLI_SKIP_UPDATE': '0'}), patch('backend.officecli_bridge.subprocess.run', return_value=CompletedProcess([], 1, json.dumps(response), '')) as run:
            self.assertEqual(officecli_bridge.command('validate', 'example.pptx')['errors'], [finding])
            self.assertEqual(run.call_args.kwargs['env']['OFFICECLI_SKIP_UPDATE'], '1')
            self.assertEqual(run.call_args.kwargs['env']['OFFICECLI_NO_AUTO_RESIDENT'], '1')
        response['error'] = {'code': 'io_error', 'error': 'Denied'}
        with patch('backend.officecli_bridge.subprocess.run', return_value=CompletedProcess([], 1, json.dumps(response), '')):
            with self.assertRaisesRegex(ValueError, 'Denied'):
                officecli_bridge.command('validate', 'example.pptx')

    @unittest.skipUnless(os.name == 'nt', 'Windows sharing-mode regression')
    def test_open_pptx_is_inspected_via_snapshot_without_competing_with_editing_lock(self):
        import ctypes
        from ctypes import wintypes
        source = fixture(self.root / 'locked.pptx'); original = source.read_bytes()
        create = ctypes.windll.kernel32.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        close = ctypes.windll.kernel32.CloseHandle; close.argtypes = [wintypes.HANDLE]
        handle = create(str(source), 0xC0000000, 1, None, 3, 0x80, None)
        self.assertNotEqual(handle, wintypes.HANDLE(-1).value)
        try:
            with self.assertRaises(ValueError):
                officecli_bridge.command('validate', source)
            self.assertEqual(officecli_bridge.inspect(source)['errors'], [])
            self.assertEqual(source.read_bytes(), original)
        finally:
            close(handle)

    def test_new_structure_error_prevents_overwrite_and_cleans_temporary_output(self):
        source = self.root / 'review.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        p = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.paragraphs[0]
        for text, bold in [('Prefix ', True), ('dot ', False), ('spacing', False), (' suffix', True)]:
            r = p.add_run(); r.text = text; r.font.bold = bold
        deck.save(source); original = source.read_bytes()
        job = {'id': 'skills-qa', 'cancel': threading.Event(), 'status': 'running'}
        with patch('backend.workflow.store.load_glossary', return_value=self.glossary):
            workflow.run(job, {'kind':'check', 'paths':[str(source)], 'use_llm':False})
        finding = job['result']['findings'][0]
        data = {'report_id': job['result']['report_id'], 'path': str(source), 'overwrite': True,
                'selected': {finding['segment_id']: [[finding['found_text'], finding['suggested']]]}}
        before = {'errors': [], 'issues': []}
        after = {'errors': [{'description': 'New XML error'}], 'issues': []}
        with patch('backend.workflow.officecli_bridge.inspect', side_effect=[before, after]):
            with self.assertRaisesRegex(ValueError, '结构问题'):
                workflow.apply_report(data)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(list(self.root.glob('.star-review-*')), [])
        # A real validated second attempt retains unaffected styles and changes only one slide.
        result = workflow.apply_report(data)
        self.assertTrue(result['layout_checked'])
        p = Presentation(source).slides[0].shapes[0].text_frame.paragraphs[0]
        self.assertEqual(p.text, 'Prefix pixel pitch suffix')
        self.assertTrue(p.runs[0].font.bold); self.assertTrue(p.runs[-1].font.bold)
        self.assertFalse(p.runs[1].font.bold)
        from io import BytesIO
        with ZipFile(BytesIO(original)) as old, ZipFile(source) as new:
            self.assertEqual([n for n in old.namelist() if old.read(n) != new.read(n)], ['ppt/slides/slide1.xml'])

    def test_word_mixed_styles_remain_and_inline_drawing_is_not_silently_deleted(self):
        path = self.root / 'word.docx'; doc = Document(); p = doc.add_paragraph()
        for text, bold in [('采用', False), ('逐点校正', True), ('，亮度为500nits。', False)]:
            run = p.add_run(text); run.bold = bold
        q = doc.add_paragraph('中文图片说明'); q.runs[0].add_picture(str(ROOT / 'build/icon.png'), width=Inches(.5))
        doc.save(path)
        result = translate_file(str(path), self.glossary, None, client=Model())
        output = Document(result['path'])
        self.assertEqual([r.bold for r in output.paragraphs[0].runs], [False, True, False])
        self.assertEqual(len(output.inline_shapes), 1)
        self.assertEqual(output.paragraphs[1].text, '中文图片说明')
        self.assertGreater(result['skipped'], 0)

    def test_chunk_style_slicing_and_language_reordering(self):
        doc = Document(); p = doc.add_paragraph(); p.add_run('abc').bold = True; p.add_run('def').italic = True
        groups = style_groups(p)
        self.assertEqual(slice_groups(groups, 2, 5), [{'id':0, 'text':'c'}, {'id':1, 'text':'de'}])
        from backend.replacer.styled_text import response_runs, write_styled
        runs = [{'id':1, 'text':'second '}, {'id':0, 'text':'first'}]
        accepted = response_runs({'runs':runs}, 'second first', groups)
        write_styled(p, groups, accepted)
        self.assertTrue(p.runs[0].italic); self.assertTrue(p.runs[1].bold)

    def test_factual_guard_accepts_dimensions_codes_and_rejects_lost_numbers_or_codes(self):
        validate_translation('VX1000输出3840×2160的4K画面，支持HDR/SDR。',
                             'VX1000 outputs 3840 x 2160 in 4K and supports HDR / SDR.')
        validate_translation('亮度为1,000nits。', 'The brightness is 1000 nits.')
        for source, text in [('支持4K。', 'Supports 4 pixels.'), ('VX1000支持HDR。','VX1000 supports color.'),
                              ('亮度为500nits。','The brightness is 1000 nits.')]:
            with self.assertRaises(ValueError): validate_translation(source, text)

    def test_baseline_issues_do_not_become_new_failures_and_overflow_has_readable_location(self):
        error = {'description': 'Existing template error'}
        old = {'errors':[error], 'issues':[]}
        new = {'errors':[error], 'issues':[{'id':'O1','path':'/slide[3]/shape[@id=71]', 'message':'text overflow: too tall'}]}
        warning = officecli_bridge.compare(old, new)
        self.assertEqual(warning[0]['slide'], 3); self.assertEqual(warning[0]['text'], '文本框 71')
        self.assertEqual(officecli_bridge.compare(new, new), [])

    def test_embedded_rules_keep_proofreading_schema_and_license_attribution(self):
        self.assertIn('Humanizer', SYSTEM); self.assertIn('OfficeCLI', SYSTEM)
        self.assertIn('只输出 JSON 数组', SYSTEM_PROMPT)
        self.assertIn('Do not\nreport correct synonyms', SYSTEM_PROMPT)
        self.assertTrue((ROOT/'backend/skills/humanizer/LICENSE').is_file())
        self.assertTrue((ROOT/'backend/skills/officecli/NOTICE').is_file())
        self.assertIn('Apache', (ROOT/'backend/skills/officecli/LICENSE').read_text())


if __name__ == '__main__':
    unittest.main()
