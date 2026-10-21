"""Word-generated labels and splittable body rows survive English export."""
import sys,os,tempfile,unittest
from unittest.mock import patch
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from backend.extract.word import word_paragraphs
from backend.extract.word_numbering import english_numbering, used_levels
from backend.replacer.word_layout import fit_word_tables
from backend.replacer.pptx_package import save_office
from backend.translation_audit import saved_residuals


class WordNumberingTests(unittest.TestCase):
    def fixture(self,label='%1、'):
        doc=Document();p=doc.add_paragraph('Meeting background')
        root=doc.part.numbering_part.element
        abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'99')
        lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
        for name,value in [('start','3'),('numFmt','chineseCounting'),('lvlText',label)]:
            n=OxmlElement('w:'+name);n.set(qn('w:val'),value);lvl.append(n)
        abstract.append(lvl);root.append(abstract)
        num=OxmlElement('w:num');num.set(qn('w:numId'),'99');ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),'99');num.append(ref);root.append(num)
        numpr=OxmlElement('w:numPr');ref=OxmlElement('w:numId');ref.set(qn('w:val'),'99');numpr.append(ref);p._p.get_or_add_pPr().append(numpr)
        return doc,p

    def test_used_numbering_is_audited_converted_and_saved_without_changing_start(self):
        with tempfile.TemporaryDirectory() as tmp:
            source=Path(tmp)/'source.docx';output=Path(tmp)/'English.docx'
            doc,_=self.fixture();doc.save(source)
            self.assertTrue(any(i['source']=='numbering' for i in saved_residuals(source)))
            word_paragraphs(doc);english_numbering(doc);save_office(source,doc,output,[])
            self.assertEqual(saved_residuals(output),[])
            _,levels=used_levels(Document(output));self.assertEqual(len(levels),1)
            self.assertEqual(levels[0].find(qn('w:start')).get(qn('w:val')),'3')
            self.assertEqual(levels[0].find(qn('w:numFmt')).get(qn('w:val')),'decimal')
            self.assertEqual(levels[0].find(qn('w:lvlText')).get(qn('w:val')),'%1. ')

    def test_label_words_use_model_mapping_and_preserve_numbering_placeholders(self):
        doc,_=self.fixture('第%1章');paragraphs,_=word_paragraphs(doc)
        label=next(p for loc,p,*_ in paragraphs if loc.startswith('numbering/'))
        with self.assertRaises(ValueError):label.write_styled([], [{'text':'Chapter 1'}])
        label.write_styled([], [{'text':'Chapter %1'}]);english_numbering(doc)
        self.assertEqual(label.text,'Chapter %1')

    def test_translation_pipeline_includes_words_in_automatic_numbering(self):
        from backend.translation import translate_file
        from backend.glossary.store import Glossary
        class Model:
            def chat_json(self,*args,**kwargs):return {'translation':'Chapter %1'}
        with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,{'NOVATERM_DATA_DIR':tmp}):
            source=Path(tmp)/'source.docx';doc,_=self.fixture('第%1章');doc.save(source)
            result=translate_file(source,Glossary([]),None,client=Model(),workers=1)
            self.assertTrue(result['translation_complete']);self.assertEqual(saved_residuals(result['path']),[])

    def test_inherited_numbering_and_unused_styles_are_not_confused(self):
        doc,p=self.fixture();style=doc.styles.add_style('Numbered meeting',1)
        props=p._p.get_or_add_pPr();number=props.find(qn('w:numPr'));props.remove(number)
        style.element.get_or_add_pPr().append(number);p.style=style
        self.assertEqual(len(used_levels(doc)[1]),1)
        word_paragraphs(doc);english_numbering(doc)
        p.style=doc.styles['Normal'];self.assertEqual(used_levels(doc)[1],[])

    def test_only_large_flowing_rows_relax_minimum_not_fixed_or_form_rows(self):
        doc=Document();table=doc.add_table(rows=4,cols=2)
        for i,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr();height=OxmlElement('w:trHeight');height.set(qn('w:val'),'10000');height.set(qn('w:hRule'),'exact' if i==1 else 'atLeast');trpr.append(height)
            if i!=3:
                for _ in range(6):row.cells[0].add_paragraph('A long paragraph to preserve in the translated document. '*8)
            if i==2:trpr.append(OxmlElement('w:cantSplit'))
        self.assertEqual(len(fit_word_tables(doc)),1)
        self.assertIsNone(table.rows[0]._tr.find('w:trPr/w:trHeight',table.rows[0]._tr.nsmap))
        for row in table.rows[1:]:self.assertIsNotNone(row._tr.find('w:trPr/w:trHeight',row._tr.nsmap))


if __name__=='__main__':unittest.main()
