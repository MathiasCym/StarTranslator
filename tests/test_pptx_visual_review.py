"""Page cohorts, original geometry and the private image-review delivery gate."""
import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches,Pt
from backend import pptx_visual_review as visual
from backend.replacer.pptx_layout import capture_layout,fit_document,_frames,_paragraph_data,FONT_STEPS
from backend.replacer.pptx_package import save_office
from backend.llm.client import LLMClient
from backend.translation import translate_file,translation_system,Cancelled
from backend.glossary.loader import Glossary


class VisualTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'NOVATERM_DATA_DIR':str(self.folder/'profile')});self.env.start()
        self.addCleanup(self.env.stop);self.addCleanup(self.temp.cleanup)
    def deck(self,pages=2):
        deck=Presentation()
        for number in range(pages):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            for x,y,w,text in ((.5,.5,4,'简短标题'),(.5,2,2,'技术说明')):
                shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(.7))
                shape.text=text
                run=shape.text_frame.paragraphs[0].runs[0];run.font.name='Arial';run.font.size=Pt(24)
        source=self.folder/'source.pptx';deck.save(source)
        return deck,source
    def renderer(self,path,folder,cancel=None):
        folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
        from PIL import Image
        pages=[]
        for i in range(len(Presentation(path).slides)):
            png=folder/f'page-{i+1}.png';Image.new('RGB',(128,72),'white').save(png)
            pages.append({'image':png,'file':png.name,'shapes':[]})
        return {'pages':pages,'native':False}
    def test_a_crowded_frame_does_not_shrink_spacious_peers_or_other_pages(self):
        deck,source=self.deck();original=capture_layout(deck,[])
        short=deck.slides[0].shapes[0];long=deck.slides[0].shapes[1]
        short.text_frame.paragraphs[0].runs[0].text='Title'
        long.text_frame.paragraphs[0].runs[0].text='Full technical explanation for cinema display processing'
        changed=['slide1/shape1/p0'];fit_document(deck,changed,original)
        sizes=[s.text_frame.paragraphs[0].runs[0].font.size.pt for s in deck.slides[0].shapes]
        self.assertEqual(sizes[0],24);self.assertIn(sizes[1],FONT_STEPS);self.assertLess(sizes[1],24)
        self.assertEqual(deck.slides[1].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt,24)
        self.assertNotIn('slide1/shape0/p0',changed)
        output=self.folder/'output.pptx';save_office(source,deck,output,changed)
        with ZipFile(source) as a,ZipFile(output) as b:
            self.assertEqual(a.read('ppt/slides/slide2.xml'),b.read('ppt/slides/slide2.xml'))
        self.assertEqual(long.text,'Full technical explanation for cinema display processing')
    def test_group_dimensions_scale_but_font_points_do_not(self):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);group=slide.shapes.add_group_shape()
        shape=group.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));shape.text='Example'
        shape.text_frame.paragraphs[0].runs[0].font.size=Pt(20)
        xf=group._element.grpSpPr.xfrm;xf.ext.cx=xf.chExt.cx//2;xf.ext.cy=xf.chExt.cy//2
        frame=list(_frames(deck,['slide1/shape0/shape0/p0']))[0]
        self.assertEqual(frame[4],shape.width//2);self.assertEqual(frame[5],shape.height//2)
        self.assertEqual(_paragraph_data(shape.text_frame,shape,slide,deck)[0][1][0][1][1],20)
    def test_tracking_is_measured_in_hundredths_of_a_point(self):
        deck,_=self.deck(1);shape=deck.slides[0].shapes[0]
        shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set('spc','1000')
        self.assertEqual(_paragraph_data(shape.text_frame,shape,deck.slides[0],deck)[0][1][0][1][4],10)
    def test_translated_tracking_is_adapted_before_font_reduction_without_moving_textbox(self):
        deck,_=self.deck(1);shape=deck.slides[0].shapes[0];run=shape.text_frame.paragraphs[0].runs[0]
        run._r.get_or_add_rPr().set('spc','300');snapshot=capture_layout(deck,['slide1/shape0/p0'])
        geometry=(shape.left,shape.top,shape.width,shape.height);run.text='A clear English title'
        fit_document(deck,['slide1/shape0/p0'],snapshot)
        self.assertEqual((shape.left,shape.top,shape.width,shape.height),geometry)
        self.assertLessEqual(float(run._r.rPr.get('spc'))/100,.72)
        self.assertEqual(run.text,'A clear English title')
    def test_package_guard_refuses_alignment_and_margin_changes(self):
        from pptx.enum.text import PP_ALIGN
        for mode in ('margin','alignment'):
            deck,source=self.deck(1);shape=deck.slides[0].shapes[0]
            if mode=='margin':shape.text_frame.margin_left=Pt(30)
            else:shape.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
            with self.assertRaisesRegex(ValueError,'对象结构'):
                save_office(source,deck,self.folder/'bad.pptx',['slide1/shape0/p0'])
    def test_scaled_numbers_keep_special_font_size_and_bold_on_page_two(self):
        from backend.replacer.styled_text import style_groups,anchored_runs,write_styled,response_runs,StyleMappingError
        deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(.5),Inches(.5),Inches(7),Inches(1.2));para=shape.text_frame.paragraphs[0]
        r=para.add_run();r.text='21.4万';r.font.name='Arial';r.font.size=Pt(24);r.font.bold=True
        r=para.add_run();r.text='块，市场体量庞大';r.font.name='Arial';r.font.size=Pt(14);r.font.bold=False
        groups=style_groups(para);snapshot=capture_layout(deck,['slide2/shape0/p0'])
        translated='214,000 screens offer a large market.'
        styled,unresolved=anchored_runs(translated,groups)
        self.assertFalse(unresolved);self.assertEqual(styled[0]['text'],'214,000')
        write_styled(para,groups,styled);fit_document(deck,['slide2/shape0/p0'],snapshot)
        self.assertEqual(para.runs[0].font.size.pt,24);self.assertTrue(para.runs[0].font.bold)
        self.assertEqual(para.runs[0].font.name,'Arial');self.assertEqual(para.runs[1].font.size.pt,14)
        with self.assertRaises(StyleMappingError):
            response_runs({'runs':[{'id':1,'text':'214,000 screens offer a '},{'id':0,'text':'large market.'}]},translated,groups)
    def test_package_guard_refuses_geometry_changes(self):
        deck,source=self.deck();deck.slides[0].shapes[0].left+=Inches(1)
        with self.assertRaisesRegex(ValueError,'对象结构'):
            save_office(source,deck,self.folder/'bad.pptx',['slide1/shape0/p0'])
    def test_model_can_enlarge_readable_text_into_free_space_and_rerender(self):
        deck,source=self.deck(1);shape=deck.slides[0].shapes[1];snapshot=capture_layout(deck,[])
        from pptx.dml.color import RGBColor
        run=shape.text_frame.paragraphs[0].runs[0];run.text='English information';run.font.size=Pt(12);run.font.bold=True;run.font.color.rgb=RGBColor(20,180,240)
        changed=['slide1/shape1/p0'];out=self.folder/'grown.pptx';save_office(source,deck,out,changed)
        calls=[]
        class Vision:
            def chat_images_json(_,prompt,**kw):
                calls.append(prompt)
                return {'reviewed':True,'issues':[],'adjustments':[{'shape_id':shape.shape_id,'font_size':18,'width':216,'height':54,'reason':'整页图像显示右侧有空白，可以放大'}] if len(calls)==1 else []}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,fixes=visual.review_pptx(source,out,deck,changed,snapshot,Vision())
        self.assertFalse(issues);self.assertEqual(len(calls),2);self.assertTrue(any(f.get('model_readability') for f in fixes))
        result=Presentation(out).slides[0].shapes[1]
        self.assertEqual(result.width,Pt(216));self.assertEqual(result.height,Pt(54));self.assertEqual(result.left,shape.left)
        self.assertEqual(result.text,'English information');r=result.text_frame.paragraphs[0].runs[0]
        self.assertEqual(r.font.size,Pt(18));self.assertTrue(r.font.bold);self.assertEqual(r.font.color.rgb,RGBColor(20,180,240))
        self.assertIn('"font_sizes": [18.0]',calls[-1])
        self.assertIn('"source_font_sizes": [24.0]',calls[-1])
    def test_model_expansion_cannot_cover_an_image_or_other_text_or_leave_page(self):
        from backend.replacer.pptx_readability import apply_actions
        from PIL import Image
        deck,source=self.deck(1);slide=deck.slides[0];shape=slide.shapes[1];shape.text='Text';shape.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
        image=self.folder/'picture.png';Image.new('RGB',(40,40),'blue').save(image)
        slide.shapes.add_picture(str(image),Inches(3),Inches(2),Inches(1),Inches(1))
        before=(shape.width,shape.height,shape.text_frame.paragraphs[0].runs[0].font.size)
        for width in (300,10000):
            fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=18,width=width,height=54,reason='model')],[],[{'shapes':[]}])
            self.assertEqual(fixes,[]);self.assertEqual((shape.width,shape.height,shape.text_frame.paragraphs[0].runs[0].font.size),before)
    def test_font_growth_inside_unchanged_overlapping_frames_cannot_cover_text(self):
        from backend.replacer.pptx_readability import apply_actions
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        a=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(2));a.text='Readable text in this area.'
        b=slide.shapes.add_textbox(Inches(1),Inches(1.55),Inches(4),Inches(.5));b.text='Neighbor text'
        for shape in (a,b):shape.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
        feedback=[]
        fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape0',font_size=32,reason='留白')],[],[{'shapes':[]}],feedback)
        self.assertFalse(fixes);self.assertTrue(feedback)
        self.assertEqual(a.text_frame.paragraphs[0].runs[0].font.size,Pt(10))

    def test_model_can_recover_tiny_text_without_a_multiplier_ceiling(self):
        from backend.replacer.pptx_readability import apply_actions
        deck,_=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='Readable title'
        shape.width=Inches(5);shape.height=Inches(1)
        shape.text_frame.paragraphs[0].runs[0].font.size=Pt(6)
        fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=32,reason='恢复可读标题')],[],[{'shapes':[]}])
        self.assertTrue(fixes);self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size,Pt(32))

    def test_model_can_expand_or_wrap_without_reducing_or_increasing_font(self):
        from backend.replacer.pptx_readability import apply_actions
        deck,_=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='English information'
        fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=24,width=300,height=60,wrap=True,reason='使用旁边空白')],[],[{'shapes':[]}])
        self.assertTrue(fixes);self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size,Pt(24))
        self.assertEqual(shape.width,Pt(300))

    def test_model_can_expand_left_to_preserve_a_centered_label_and_guards_position(self):
        from backend.replacer.pptx_readability import apply_actions
        from pptx.enum.text import PP_ALIGN
        deck,source=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='English information'
        shape.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
        deck.save(source);original_left=shape.left
        changed=[]
        fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=24,x=18,y=144,width=240,height=60,reason='向两侧扩展标签')],changed,[{'shapes':[]}])
        self.assertTrue(fixes);self.assertLess(shape.left,original_left)
        output=self.folder/'moved.pptx';save_office(source,deck,output,changed)
        self.assertEqual(Presentation(output).slides[0].shapes[1].left,Pt(18))
        shape.left+=Pt(1)
        with self.assertRaisesRegex(ValueError,'位置无法验证'):
            save_office(source,deck,self.folder/'bad-move.pptx',changed)

    def test_model_move_cannot_cross_page_or_cover_other_text(self):
        from backend.replacer.pptx_readability import apply_actions
        deck,_=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='Information'
        before=(shape.left,shape.top,shape.width,shape.height)
        for x,y in ((-20,144),(18,36),(10000,144)):
            fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=24,x=x,y=y,width=240,height=60,reason='model')],[],[{'shapes':[]}])
            self.assertFalse(fixes);self.assertEqual((shape.left,shape.top,shape.width,shape.height),before)

    def test_visual_followup_does_not_mechanically_refit_unconstrained_tiny_text(self):
        deck,_=self.deck(1);a,b=list(deck.slides[0].shapes)
        b.text='Unchanged small note';b.text_frame.paragraphs[0].runs[0].font.size=Pt(9)
        snapshot=capture_layout(deck,[])
        fit_document(deck,['slide1/shape0/p0','slide1/shape1/p0'],snapshot,constraints={'slide1/shape0':True})
        self.assertEqual(b.text_frame.paragraphs[0].runs[0].font.size,Pt(9))
        fit_document(deck,['slide1/shape0/p0','slide1/shape1/p0'],snapshot,constraints={'slide1/shape1':True})
        self.assertEqual(b.text_frame.paragraphs[0].runs[0].font.size,Pt(9))

    def test_approved_expansion_does_not_let_another_geometry_change_through(self):
        from backend.replacer.pptx_readability import apply_actions
        deck,source=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='Text';shape.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
        changed=[]
        self.assertTrue(apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=18,width=216,height=54,reason='model')],changed,[{'shapes':[]}]))
        deck.slides[0].shapes[0].left+=Inches(1)
        with self.assertRaisesRegex(ValueError,'对象结构'):save_office(source,deck,self.folder/'unsafe.pptx',changed)
    def test_unfitting_model_growth_gets_measurement_feedback_instead_of_shrinking(self):
        deck,source=self.deck(1);shape=deck.slides[0].shapes[1];snapshot=capture_layout(deck,[])
        run=shape.text_frame.paragraphs[0].runs[0];run.text='An English explanation with more information';run.font.size=Pt(12)
        out=self.folder/'feedback.pptx';changed=['slide1/shape1/p0'];save_office(source,deck,out,changed);calls=[]
        class Vision:
            def chat_images_json(_,prompt,**kw):
                calls.append(prompt)
                return {'reviewed':True,'issues':[{'shape_id':shape.shape_id,'kind':'crowded','reason':'文字太小，需要利用空白放大'}] if len(calls)<3 else [],
                    'adjustments':[dict(shape_id=shape.shape_id,font_size=24,width=216 if len(calls)==1 else 360,height=54 if len(calls)==1 else 100,wrap=True,reason='利用右侧及下方空白')] if len(calls)<3 else []}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            warnings,fixes=visual.review_pptx(source,out,deck,changed,snapshot,Vision())
        self.assertFalse(warnings);self.assertEqual(len(calls),3);self.assertIn('measured_min_height',calls[1])
        self.assertTrue(any(f.get('model_readability') for f in fixes));self.assertEqual(Presentation(out).slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size,Pt(24))
    def test_model_rounding_does_not_discard_already_available_height(self):
        from backend.replacer.pptx_readability import apply_actions
        deck,_=self.deck(1);shape=deck.slides[0].shapes[1];shape.text='Readable text';shape.text_frame.paragraphs[0].runs[0].font.size=Pt(9)
        height=shape.height
        fixes=apply_actions(deck,[dict(kind='readability',slide=1,locator='slide1/shape1',font_size=18,width=216,height=40,reason='留白允许放大')],[],[{'shapes':[]}])
        self.assertTrue(fixes);self.assertEqual(shape.height,height);self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size,Pt(18))
    def test_every_page_is_reviewed_and_private_images_are_removed(self):
        deck,source=self.deck(3);out=self.folder/'out.pptx';save_office(source,deck,out,[])
        calls=[]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):
                calls.append((prompt,list(images)))
                self.assertEqual(len(images),2);self.assertTrue(all(Path(p).is_file() for p in images))
                return {'reviewed':True,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,fixes=visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision())
        self.assertEqual(issues,[]);self.assertEqual(len(calls),3)
        self.assertTrue(all(not Path(p).exists() for _,images in calls for p in images))
        self.assertNotIn('image',json.dumps(issues))
    def test_visual_repair_changes_only_typography_and_rereviews_affected_page(self):
        deck,source=self.deck();changed=['slide1/shape0/p0'];snapshot=capture_layout(deck,changed)
        out=self.folder/'out.pptx';save_office(source,deck,out,[]);text=deck.slides[0].shapes[0].text
        answers=iter([{'reviewed':True,'issues':[{'shape_id':deck.slides[0].shapes[0].shape_id,'kind':'crowded','reason':'文字拥挤'}]},
                      {'reviewed':True,'issues':[]},{'reviewed':True,'issues':[]}]);calls=[]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):calls.append(prompt);return next(answers)
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,fixes=visual.review_pptx(source,out,deck,changed,snapshot,Vision())
        self.assertFalse(issues);self.assertEqual(len(calls),3);self.assertIn('第 1 页',calls[-1])
        self.assertEqual(deck.slides[0].shapes[0].text,text)
        self.assertEqual(deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt,22)
        self.assertEqual(deck.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size.pt,22)
        self.assertTrue(fixes)
    def test_unavailable_or_invalid_image_review_cannot_claim_success(self):
        deck,source=self.deck();out=self.folder/'out.pptx';save_office(source,deck,out,[])
        class Vision:
            def chat_images_json(_,prompt,**kwargs):return {'reviewed':False,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,_=visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision())
        self.assertTrue(issues)
        self.assertTrue(visual.review_pptx(source,out,deck,[],{},object())[0])
    def test_invalid_page_response_is_retried_without_rolling_back_other_pages(self):
        deck,source=self.deck(2);out=self.folder/'retry.pptx';save_office(source,deck,out,[])
        calls=[];audit={}
        class Vision:
            def chat_images_json(_,prompt,**kwargs):
                calls.append(prompt)
                if len(calls)==1:return {'reviewed':True,'issues':[{'kind':'small_font','reason':'格式不符合约定'}]}
                return {'reviewed':True,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,_=visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision(),audit=audit)
        self.assertFalse(issues);self.assertEqual(audit['status'],'complete');self.assertEqual(audit['image_pages'],2)
        self.assertEqual(len(calls),3);self.assertIn('上次结果不符合格式',calls[1])

    def test_persistent_invalid_page_is_local_and_other_pages_still_finish(self):
        deck,source=self.deck(2);out=self.folder/'partial.pptx';save_office(source,deck,out,[])
        calls=[];audit={}
        class Vision:
            def chat_images_json(_,prompt,**kwargs):
                calls.append(prompt)
                return {'reviewed':True,'issues':[{'kind':'other','shape_id':{},'reason':'无效定位'}]} if prompt.startswith('第 1 页') else {'reviewed':True,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,_=visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision(),audit=audit)
        self.assertEqual(audit['status'],'partial');self.assertEqual(audit['image_pages'],1)
        self.assertEqual([i['slide'] for i in issues],[1]);self.assertEqual(len(calls),3)
        with ZipFile(out) as archive:self.assertIsNone(archive.testzip())

    def test_worse_visual_repairs_keep_the_best_complete_file(self):
        deck,source=self.deck(1);out=self.folder/'out.pptx';save_office(source,deck,out,[])
        original=out.read_bytes();sid=deck.slides[0].shapes[0].shape_id;count=[0]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):
                count[0]+=1
                return {'reviewed':True,'issues':[{'shape_id':sid,'kind':'overlap','reason':'重叠'}]*count[0]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,fixes=visual.review_pptx(source,out,deck,['slide1/shape0/p0'],capture_layout(deck,[]),Vision())
        self.assertEqual(out.read_bytes(),original);self.assertEqual(len(issues),1)
        self.assertEqual(count[0],3);self.assertEqual(fixes,[])
    def test_a_worse_page_does_not_discard_safe_font_recovery_on_another_page(self):
        deck,source=self.deck(2);shape=deck.slides[0].shapes[1];shape.text='English information'
        shape.text_frame.paragraphs[0].runs[0].font.size=Pt(12);deck.save(source)
        snapshot=capture_layout(deck,[]);out=self.folder/'per-page.pptx';save_office(source,deck,out,[])
        sid=deck.slides[1].shapes[0].shape_id;calls={1:0,2:0}
        class Vision:
            def chat_images_json(_,prompt,**kwargs):
                page=1 if prompt.startswith('第 1 页') else 2;calls[page]+=1
                if page==1:
                    return {'reviewed':True,'issues':[],'adjustments':[dict(shape_id=shape.shape_id,font_size=18,width=216,height=54,reason='恢复清晰文字')] if calls[page]==1 else []}
                return {'reviewed':True,'issues':[dict(shape_id=sid,kind='overlap',reason='未解决的局部问题')]*calls[page]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,fixes=visual.review_pptx(source,out,deck,['slide1/shape1/p0','slide2/shape0/p0'],snapshot,Vision())
        self.assertEqual(Presentation(out).slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size,Pt(18))
        self.assertTrue(any(i.get('model_readability') for i in fixes));self.assertEqual([i['slide'] for i in issues],[2])
        with ZipFile(source) as a,ZipFile(out) as b:self.assertEqual(a.read('ppt/slides/slide2.xml'),b.read('ppt/slides/slide2.xml'))

    def test_in_place_staged_review_keeps_a_readonly_package_before_repair(self):
        deck,source=self.deck(1);out=self.folder/'same.pptx';save_office(source,deck,out,[])
        sid=deck.slides[0].shapes[0].shape_id;count=[0]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):
                count[0]+=1
                return {'reviewed':True,'issues':[{'shape_id':sid,'kind':'crowded','reason':'文字拥挤'}] if count[0]==1 else []}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,_=visual.review_pptx(out,out,deck,['slide1/shape0/p0'],capture_layout(deck,[]),Vision(),reference=source)
        self.assertFalse(issues);self.assertGreater(count[0],1)
        with ZipFile(out) as archive:self.assertIsNone(archive.testzip())
    def test_native_table_cell_and_left_clipping_map_to_the_correct_frame(self):
        old=dict(id='10/1/1',shape_id=10,row=0,col=0,x=0,y=0,width=100,height=20,bound_x=0,bound_y=0,bound_width=50,bound_height=15)
        new={**old,'bound_x':-5}
        found=visual._native_issues({'shapes':[old]},{'shapes':[new]},1,[{'shape_id':10,'locator':'slide1/shape2'}])
        self.assertEqual(len(found),1);self.assertEqual(found[0]['locator'],'slide1/shape2/table/0/0/')
    def test_native_review_detects_new_ink_overlap_inside_unchanged_frames(self):
        a=dict(id=1,x=0,y=0,width=200,height=100,bound_x=0,bound_y=0,bound_width=100,bound_height=15)
        b=dict(id=2,x=0,y=0,width=200,height=100,bound_x=0,bound_y=40,bound_width=100,bound_height=15)
        objects=[{'shape_id':1,'locator':'slide1/shape0'},{'shape_id':2,'locator':'slide1/shape1'}]
        original={'shapes':[a,b]}
        found=visual._native_issues(original,{'shapes':[{**a,'bound_height':50},b]},1,objects)
        self.assertEqual({v['locator'] for v in found if v['kind']=='overlap'},{'slide1/shape0','slide1/shape1'})
        self.assertFalse(visual._native_issues(original,original,1,objects))

    def test_visual_object_inventory_does_not_create_text_in_covered_merged_cells(self):
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        table=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(6),Inches(1)).table
        table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='Merged label'
        cell=table.cell(0,1);body=cell._tc.find('{http://schemas.openxmlformats.org/drawingml/2006/main}txBody')
        cell._tc.remove(body);before=slide._element.xml
        self.assertTrue(visual._objects(deck)[0]);self.assertEqual(slide._element.xml,before)
    def test_cancellation_exits_and_cleans_private_images(self):
        deck,source=self.deck();out=self.folder/'out.pptx';save_office(source,deck,out,[]);cancel=threading.Event();paths=[]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):paths.extend(images);cancel.set();return {'reviewed':True,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer),self.assertRaises(Cancelled):
            visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision(),cancel)
        self.assertTrue(all(not p.exists() for p in paths))
    def test_page_model_checks_use_bounded_parallelism_without_parallel_office_writes(self):
        deck,source=self.deck(6);out=self.folder/'parallel.pptx';save_office(source,deck,out,[])
        lock=threading.Lock();gate=threading.Event();active=[0];peak=[0];calls=[]
        class Vision:
            def chat_images_json(_,prompt,images,**kwargs):
                with lock:
                    active[0]+=1;peak[0]=max(peak[0],active[0]);calls.append(prompt)
                    if active[0]==3:gate.set()
                    if len(calls)==1:
                        active[0]-=1
                        return {'reviewed':True,'issues':[]}
                if not gate.wait(3):raise ValueError('Page checks did not run in parallel')
                with lock:active[0]-=1
                return {'reviewed':True,'issues':[]}
        with patch.object(visual,'render_pages',side_effect=self.renderer):
            issues,_=visual.review_pptx(source,out,deck,[],capture_layout(deck,[]),Vision(),workers=3)
        self.assertEqual(issues,[]);self.assertEqual(len(calls),6);self.assertEqual(peak[0],3)
    def test_incomplete_image_review_still_exports_without_exposing_internal_files(self):
        deck,source=self.deck(1)
        class Model:
            def chat_json(_,prompt,**kwargs):return {'translation':'Technical explanation'}
        original=source.read_bytes()
        result=translate_file(str(source),Glossary([]),None,client=Model())
        self.assertEqual(source.read_bytes(),original)
        self.assertTrue(Path(result['path']).is_file())
        self.assertEqual(result['status'],'complete')
        self.assertFalse(result['layout_checked'])
        self.assertNotIn('.png',json.dumps(result));self.assertTrue(result['layout_warnings'])
    def test_translation_modes_are_explicit_and_isolated_per_call(self):
        self.assertIn('不得为适应版面',translation_system('faithful'))
        self.assertIn('用户选择文案压缩',translation_system('compressed'))
        with self.assertRaises(ValueError):translation_system('automatic')
    def test_vision_payload_supports_both_configured_api_formats(self):
        png=self.folder/'p.png'
        from PIL import Image
        Image.new('RGB',(10,10),'white').save(png)
        for fmt in ('openai','anthropic'):
            client=LLMClient.__new__(LLMClient);client.api_format=fmt;client.provider='compatible';client.model='fixture';client.base_url='https://fixture.invalid/v1'
            payloads=[]
            def post(endpoint,body):
                payloads.append(body)
                content='{"reviewed":true,"issues":[]}'
                return {'content':[{'type':'text','text':content}],'stop_reason':'end_turn'} if fmt=='anthropic' else {'choices':[{'message':{'content':content},'finish_reason':'stop'}]}
            client._post=post
            self.assertTrue(client.chat_images_json('Compare pages',[png,png])['reviewed'])
            contents=payloads[0]['messages'][0]['content'];self.assertEqual(len(contents),3)
            self.assertEqual(contents[1]['type'],'image' if fmt=='anthropic' else 'image_url')


if __name__=='__main__':unittest.main()
