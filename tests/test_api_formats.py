"""Model protocol contracts and real local HTTP integration, with synthetic text only."""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import model_profiles, workflow
from backend.glossary.loader import Glossary
from backend.llm.client import LLMClient, LLMTruncated, LLMResponseInvalid, LLMUnavailable
from backend.models import Segment
from backend.matcher.judge import judge_text
from docx import Document


def response(blocks, reason='end_turn'):
    return io.BytesIO(json.dumps({'content': blocks, 'stop_reason': reason}).encode())


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.env = patch.dict(os.environ, {'NOVATERM_DATA_DIR': self.temp.name}, clear=True)
        self.env.start(); self.addCleanup(self.env.stop)

    def client(self, **options):
        return LLMClient(provider='compatible', base_url='https://relay.invalid', model='fixture',
                         api_key='fixture-key', api_format='anthropic', retries=0, **options)

    def test_anthropic_request_response_and_separate_authentication(self):
        blocks = [{'type':'thinking','thinking':'Do not expose reasoning'},
                  {'type':'text','text':'{"translation":'}, {'type':'text','text':'"LED display"}'}]
        for auth in ('auto','bearer','x-api-key'):
            with self.subTest(auth=auth), patch('backend.llm.client._urlopen', return_value=response(blocks)) as send:
                result = self.client(auth_type=auth).chat_json('显示屏',system='Preserve all numbers',max_tokens=2048)
            self.assertEqual(result, {'translation':'LED display'})
            request = send.call_args.args[0]; payload = json.loads(request.data)
            self.assertEqual(request.full_url,'https://relay.invalid/v1/messages')
            self.assertEqual(payload, {'model':'fixture','messages':[{'role':'user','content':'显示屏'}],
                'system':'Preserve all numbers','max_tokens':2048,'stream':False})
            self.assertEqual(request.get_header('Anthropic-version'),'2023-06-01')
            if auth == 'bearer':
                self.assertEqual(request.get_header('Authorization'),'Bearer fixture-key'); self.assertIsNone(request.get_header('X-api-key'))
            else:
                self.assertEqual(request.get_header('X-api-key'),'fixture-key'); self.assertIsNone(request.get_header('Authorization'))

    def test_endpoint_prefixes_and_complete_endpoints_do_not_duplicate_paths(self):
        for base in ('https://relay.invalid','https://relay.invalid/v1','https://relay.invalid/v1/messages'):
            client = LLMClient(provider='compatible',base_url=base,model='fixture',api_key='key',api_format='anthropic')
            self.assertEqual(client._endpoint('/messages'),'https://relay.invalid/v1/messages')
            self.assertEqual(client._endpoint('/models'),'https://relay.invalid/v1/models')
        client = LLMClient(provider='compatible',base_url='https://relay.invalid/claude',model='fixture',api_format='anthropic')
        self.assertEqual(client._endpoint('/messages'),'https://relay.invalid/claude/v1/messages')
        client = LLMClient(provider='compatible',base_url='https://relay.invalid/custom/messages',model='fixture',api_format='anthropic')
        self.assertEqual(client._endpoint('/messages'),'https://relay.invalid/custom/messages')
        client = LLMClient(provider='compatible',base_url='https://relay.invalid/v1/chat/completions',model='fixture')
        self.assertEqual(client._endpoint('/chat/completions'),'https://relay.invalid/v1/chat/completions')

    def test_thinking_empty_tool_or_refusal_is_not_a_successful_translation(self):
        for blocks, reason in [([], 'end_turn'),([{'type':'thinking','thinking':'x'}],'end_turn'),
                              ([{'type':'text','text':'partial'}],'tool_use'),
                              ([{'type':'text','text':'refusal'}],'refusal')]:
            with self.subTest(reason=reason), patch('backend.llm.client._urlopen',return_value=response(blocks,reason)):
                with self.assertRaises(LLMResponseInvalid): self.client().chat_json('test')
        with patch('backend.llm.client._urlopen',return_value=response([{'type':'text','text':'partial'}],'max_tokens')):
            with self.assertRaises(LLMTruncated) as error: self.client().chat_json('test')
            self.assertEqual(error.exception.text,'partial')

    def test_truncated_complete_findings_are_recoverable(self):
        text='[{"found":"bad","suggested":"good","reason":"typo"},{"found":"unfinished'
        with patch('backend.llm.client._urlopen',return_value=response([{'type':'text','text':text}],'max_tokens')):
            items,truncated=self.client().chat_json_recoverable('test')
        self.assertTrue(truncated); self.assertEqual(items[0]['suggested'],'good')

    def test_protocol_edits_require_use_and_old_records_keep_openai(self):
        payload={'provider':'compatible','base_url':'https://relay.invalid/v1','model':'fixture','api_key':'old-key'}
        saved=model_profiles.save_profile(payload)['profile']; ident=saved['id']
        self.assertEqual(saved['api_format'],'openai'); self.assertEqual(saved['auth_type'],'auto')
        model_profiles.activate(ident)
        model_profiles.save_profile({**payload,'id':ident,'api_format':'anthropic','auth_type':'bearer'})
        self.assertEqual(LLMClient().api_format,'openai'); self.assertTrue(model_profiles.list_profiles()['profiles'][0]['pending'])
        model_profiles.activate(ident); client=LLMClient()
        self.assertEqual(client.api_format,'anthropic'); self.assertEqual(client.auth_type,'bearer'); self.assertEqual(client.api_key,'old-key')
        self.assertFalse(model_profiles.list_profiles()['profiles'][0]['pending'])
        # Omitted protocol fields on a later update retain the user's selection.
        model_profiles.save_profile({**payload,'id':ident,'name':'renamed'})
        self.assertEqual(model_profiles.list_profiles()['profiles'][0]['api_format'],'anthropic')
        legacy=Path(self.temp.name)/'model-profiles.json'; data=json.loads(legacy.read_text())
        for record in [*data['profiles'],data['applied']]:
            record.pop('api_format',None);record.pop('auth_type',None)
        legacy.write_text(json.dumps(data)); model_profiles.activate(ident)
        self.assertEqual(LLMClient().api_format,'openai')

    def test_invalid_protocol_or_intranet_format_never_sends_a_request(self):
        for options in ({'api_format':'bogus'}, {'auth_type':'bogus'}, {'provider':'intranet','api_format':'anthropic'}):
            values={'provider':'compatible','base_url':'http://llm.novastar.tech:31000/v1','model':'fixture','api_key':'key',**options}
            with self.subTest(options=options), patch('backend.llm.client._urlopen') as send, self.assertRaises(LLMUnavailable):
                LLMClient(**values).chat('test')
            send.assert_not_called()

    def test_local_http_protocol_switch_translation_and_review(self):
        calls=[]
        class Handler(BaseHTTPRequestHandler):
            def log_message(self,*args): pass
            def reply(self,body):
                self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(json.dumps(body).encode())
            def do_GET(self):
                calls.append((self.path,dict(self.headers),None));self.reply({'data':[{'id':'fixture'}]})
            def do_POST(self):
                body=json.loads(self.rfile.read(int(self.headers['Content-Length'])));calls.append((self.path,dict(self.headers),body))
                prompt=body['messages'][-1]['content']
                text='{"translation":"The dot spacing is 2.5 mm."}' if '【待译原文】' in prompt else '[{"found":"dot spacing","suggested":"pixel pitch","kind":"terminology","reason":"术语"}]'
                if self.path.endswith('/messages'): self.reply({'content':[{'type':'text','text':text}],'stop_reason':'end_turn'})
                else: self.reply({'choices':[{'message':{'content':text},'finish_reason':'stop'}]})
        server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        self.addCleanup(server.server_close);self.addCleanup(server.shutdown)
        base=f'http://127.0.0.1:{server.server_port}/v1'
        source=Path(self.temp.name)/'source.docx'; doc=Document();doc.add_paragraph('点间距为2.5毫米。');doc.save(source)
        original=source.read_bytes()
        for protocol in ('openai','anthropic'):
            saved=model_profiles.save_profile({'provider':'compatible','base_url':base,'model':'fixture','api_key':'fixture-only',
                'api_format':protocol})['profile'];model_profiles.activate(saved['id'])
            self.assertEqual(LLMClient().connection_status()['state'],'connected')
            job={'id':'format-'+protocol,'status':'running','cancel':threading.Event()}
            with patch('backend.workflow.store.load_glossary',return_value=Glossary([])):
                workflow.run(job,{'kind':'translate_review','paths':[str(source)]})
            self.assertEqual(job['status'],'complete',job.get('error'));self.assertEqual(job['result']['report']['stats']['model_status'],'complete')
            self.assertTrue(job['result']['report']['findings']);self.assertTrue(all(not f['applied'] for f in job['result']['report']['findings']))
            self.assertEqual(Document(job['result']['outputs'][0]['path']).paragraphs[0].text,'The dot spacing is 2.5 mm.')
        self.assertEqual(source.read_bytes(),original)
        self.assertEqual([route for route,_,_ in calls],['/v1/models','/v1/chat/completions','/v1/chat/completions','/v1/models','/v1/messages','/v1/messages'])
        for route,headers,body in calls:
            if route == '/v1/messages':
                headers={key.lower():value for key,value in headers.items()}
                self.assertEqual(headers.get('x-api-key'),'fixture-only');self.assertEqual(headers.get('anthropic-version'),'2023-06-01')
                self.assertIn('system',body);self.assertNotIn('authorization',headers)


if __name__ == '__main__': unittest.main()
