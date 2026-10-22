"""Keep PDF illustrations in the reflowed Word; no shared Office mutations."""
import io
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

MAX_IMAGE_PIXELS = 40_000_000
MAX_PAGE_IMAGE_BYTES = 64*1024*1024


def append_images(document, page, number):
    from PIL import Image
    count, total, failures = 0, 0, []
    def failed(reason,index=0):
        failures.append(reason)
        if not hasattr(document,'_star_pdf_image_failures'):document._star_pdf_image_failures=[]
        document._star_pdf_image_failures.append(dict(locator=f'pdf-image/{number}/{index}',slide=number,
            para=None,source_text='',translation='',reason=reason,image=True,missing_image=True,retained_source=True))
    try:
        images=page.images
        keys=images.keys() if hasattr(images,'keys') else range(len(images))
    except Exception:
        failed(f'PDF 第 {number} 页图片读取失败；原 PDF 中的图像仍需对照。')
        return 0, failures
    for index, key in enumerate(keys, 1):
        try:
            # Decode one image at a time. A damaged XObject must not prevent
            # the other figures on the same page from being copied.
            image=images[key]
            raw = image.data
            total += len(raw)
            if total > MAX_PAGE_IMAGE_BYTES: raise ValueError('page image limit')
            with Image.open(io.BytesIO(raw)) as picture:
                width, height = picture.size
                if width*height > MAX_IMAGE_PIXELS: raise ValueError('image pixel limit')
                # Word has no JP2 reader. Convert the decoded pixels, preserving
                # transparency; JPEG/PNG can retain their original encoded bytes.
                if picture.format not in {'JPEG','PNG'}:
                    output = io.BytesIO()
                    picture.convert('RGBA' if 'A' in picture.getbands() else 'RGB').save(output, format='PNG')
                    raw = output.getvalue()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(6)
            section = document.sections[-1]
            available = (section.page_width-section.left_margin-section.right_margin)/12700
            scale = min(available/max(width,1), 270/max(height,1), 1)
            shape = paragraph.add_run().add_picture(io.BytesIO(raw), width=Pt(width*scale), height=Pt(height*scale))
            shape._inline.docPr.set('descr', f'Original PDF page {number}, image {index}')
            if not hasattr(document,'_star_pdf_images'):document._star_pdf_images=[]
            document._star_pdf_images.append({'page':number,'index':index,'shape':shape,'data':raw})
            count += 1
        except Exception:
            # Keep the rest of the valid document, but never report a missing
            # illustration as successfully preserved.
            failed(f'PDF 第 {number} 页第 {index} 张图片未能保留；请对照原 PDF。',index)
    return count, failures
