"""Retain PDF annotation destinations in its reflowed Word translation."""
from urllib.parse import urlsplit
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


def mark_page(paragraph,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),f'_StarPDFPage_{number}')
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0,start);paragraph._p.append(end)


def destinations(reader, page):
    external, internal = [], []
    for ref in page.get('/Annots', []):
        annotation = ref.get_object()
        if annotation.get('/Subtype') != '/Link': continue
        action = annotation.get('/A') or {}
        if hasattr(action,'get_object'): action=action.get_object()
        uri = action.get('/URI')
        if uri is not None:
            value=str(uri)
            if value and value not in external: external.append(value)
            continue
        target = annotation.get('/Dest') or action.get('/D')
        if target is None: continue
        if isinstance(target,str): target=reader.named_destinations.get(target)
        if target is None: continue
        try:
            from pypdf.generic import Destination,ArrayObject,Fit
            if not isinstance(target,Destination):
                if hasattr(target,'get_object'):target=target.get_object()
                if not isinstance(target,(list,ArrayObject)) or len(target)<2:continue
                target=Destination('Source page',target[0],Fit(str(target[1]),target[2:]))
            number=reader.get_destination_page_number(target)
            if number is not None and number>=0 and number+1 not in internal:internal.append(number+1)
        except (ValueError,TypeError,KeyError,IndexError):continue
    return external, internal


def append_page_links(document,reader,page):
    external,internal=destinations(reader,page)
    if not external and not internal:return
    def paragraph(label):
        p=document.add_paragraph();p.paragraph_format.space_after=Pt(2)
        r=p.add_run(label);r.font.name='Arial';r.font.size=Pt(9)
        return p
    def link(p,text,uri=None,anchor=None):
        node=OxmlElement('w:hyperlink')
        if uri is not None:node.set(qn('r:id'),document.part.relate_to(uri,RT.HYPERLINK,is_external=True))
        if anchor is not None:node.set(qn('w:anchor'),anchor)
        run=OxmlElement('w:r');props=OxmlElement('w:rPr')
        style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');props.append(style)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'18');props.append(size);run.append(props)
        value=OxmlElement('w:t');value.text=text;run.append(value);node.append(run);p._p.append(node)
    for uri in external:
        p=paragraph('Source link: ')
        try: scheme=urlsplit(uri).scheme.lower()
        except ValueError:scheme=''
        if scheme in {'http','https','mailto'}:link(p,uri,uri=uri)
        else:p.add_run(uri)
    if internal:
        p=paragraph('Original PDF pages: ')
        for i,number in enumerate(internal):
            if i:p.add_run(', ')
            link(p,str(number),anchor=f'_StarPDFPage_{number}')
