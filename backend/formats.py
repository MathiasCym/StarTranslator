"""Structure-preserving text adapters; Office keeps its existing implementation.

Only prose is exposed to the model. Writers patch the original text spans or
package parts, rather than letting a model regenerate file syntax.
"""
import html
import json
import re
from pathlib import Path
from html.parser import HTMLParser
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Pt

OFFICE = {'.docx', '.pptx'}
EXTRA = {'.pdf', '.xlsx', '.txt', '.md', '.markdown', '.html', '.htm', '.json', '.epub', '.srt', '.ass'}
SUPPORTED = OFFICE | EXTRA
MAX_TEXT = 32 * 1024 * 1024
MAX_SEGMENTS = 40000
def parse_xml(raw):
    return etree.fromstring(raw,etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False))


def output_extension(path):
    return '.docx' if Path(path).suffix.lower() == '.pdf' else Path(path).suffix


def archive_check(archive):
    entries = archive.infolist()
    if (len(entries) > 20000 or len({i.filename for i in entries}) != len(entries)
            or sum(i.file_size for i in entries) > 512 * 1024 * 1024
            or any(i.file_size > 128 * 1024 * 1024 for i in entries)):
        raise ValueError('文档展开后的内容过大或包含重复文件，请拆分后处理。')


def read_text(path):
    if Path(path).stat().st_size > MAX_TEXT: raise ValueError('文本内容超过 32 MB，请拆分后处理。')
    raw = Path(path).read_bytes()
    if len(raw) > MAX_TEXT: raise ValueError('文本内容超过 32 MB，请拆分后处理。')
    encodings = ('utf-8-sig',) if raw.startswith(b'\xef\xbb\xbf') else ('utf-16',) if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8', 'gb18030')
    for encoding in encodings:
        try: value = raw.decode(encoding)
        except UnicodeError: continue
        if '\x00' not in value: return value, encoding
    raise ValueError('文件不是可读取的文本，请使用 UTF-8、UTF-16 或 GB18030 编码。')


class TextDocument:
    def __init__(self, path):
        self.path = str(path)
        self.document = Document()
        self.paragraphs, self.blocked, self.values = [], [], {}
        self.writer = None
        self.normalize = literal
        self.note = '保留文件结构；仅修改可读正文。'

    def add(self, locator, value, source='text', page=None):
        from .extract.common import is_meaningful
        if not is_meaningful(value): return None
        if len(self.paragraphs) >= MAX_SEGMENTS: raise ValueError('文档段落过多，请拆分后处理。')
        paragraph = self.document.add_paragraph()
        paragraph.add_run(value)
        self.paragraphs.append((locator, paragraph, source, page, len(self.paragraphs)+1))
        self.values[locator] = value
        return paragraph

    def texts(self):
        return {key: paragraph.text for key, paragraph, *_ in self.paragraphs}

    def save(self, destination):
        for locator,paragraph,*_ in self.paragraphs:
            if paragraph.text != self.values[locator]:
                normalized = self.normalize(paragraph.text)
                if normalized != paragraph.text: paragraph.text = normalized
        self.writer(Path(destination), self.texts())


def spans_document(path, text, encoding, spans, note='保留编码、段落及非正文内容。'):
    document = TextDocument(path)
    replacements = []
    for index, (start, end, value, render) in enumerate(spans):
        locator = f'text/{index}'
        if document.add(locator, value) is not None:
            replacements.append((start, end, locator, render))
    def writer(destination, values):
        result = text
        for start, end, locator, render in reversed(replacements):
            if values[locator] != document.values[locator]:
                result = result[:start] + render(values[locator]) + result[end:]
        destination.write_bytes(result.encode(encoding))
    document.writer, document.note = writer, note
    return document


def literal(value): return value


def subtitle_text(value):
    # Newlines added by a model must never create a new cue or ASS record.
    return re.sub(r'[\r\n]+',' ',value)


def plain_spans(text, protected=()):
    """Emit prose between immutable syntax ranges, keeping whitespace untouched."""
    cursor = 0
    for start, end in sorted(protected):
        if end <= cursor: continue
        if start > cursor: yield cursor, start, text[cursor:start], literal
        cursor = max(cursor, end)
    if cursor < len(text): yield cursor, len(text), text[cursor:], literal


def line_spans(text):
    offset = 0
    for line in text.splitlines(keepends=True):
        content = line.rstrip('\r\n')
        yield offset, offset+len(content), content, literal
        offset += len(line)


def markdown_spans(text):
    protected = []
    # YAML metadata, code fences, code spans, URLs, link targets and markup.
    fence = None; offset = 0
    for line in text.splitlines(keepends=True):
        mark = re.match(r'^\s{0,3}(`{3,}|~{3,})', line)
        if fence:
            protected.append((offset, offset+len(line)))
            if mark and mark[1][0] == fence[0] and len(mark[1]) >= len(fence): fence = None
        elif mark:
            fence = mark[1]; protected.append((offset, offset+len(line)))
        elif line.startswith(('    ', '\t')):
            protected.append((offset, offset+len(line)))
        offset += len(line)
    for pattern in (r'\A---\r?\n.*?\r?\n---(?:\r?\n|\Z)', r'(`+)[^\n]*?\1',
                    r'!?\[|\]\([^\r\n]*?\)|\](?=\[)', r'(?<=\])\[[^\]\n]*\]',
                    r'https?://[^\s<>]+', r'<(script|style|pre|code)\b[^>]*>.*?</\1\s*>', r'<[^>]*>',
                    r'(?m)^\s{0,3}\[[^\]\n]+\]:[^\n]*', r'(?m)^\s{0,3}(?:#{1,6}\s+|>\s*|[-+*]\s+|\d+[.)]\s+)',
                    r'[*_~|]+'):
        protected.extend(m.span() for m in re.finditer(pattern, text, re.S))
    # Split prose by lines, so a translation cannot swallow Markdown boundaries.
    protected.extend(m.span() for m in re.finditer(r'\r?\n', text))
    yield from plain_spans(text, protected)


class HTMLSpans(HTMLParser):
    def __init__(self, text):
        super().__init__(convert_charrefs=False)
        self.text, self.spans, self.stack, self.pending = text, [], [], None
        self.lines = [0] + [m.end() for m in re.finditer('\n', text)]

    def position(self):
        row, col = self.getpos(); return self.lines[row-1]+col

    def flush(self):
        if self.pending:
            start, end = self.pending
            self.spans.append((start, end, html.unescape(self.text[start:end]), lambda s: html.escape(s, quote=False)))
            self.pending = None

    def handle_starttag(self, tag, attrs):
        self.flush()
        if tag not in {'area','base','br','col','embed','hr','img','input','link','meta','param','source','track','wbr'}: self.stack.append(tag)

    def handle_startendtag(self, tag, attrs): self.flush()

    def handle_endtag(self, tag):
        self.flush()
        if tag in self.stack:
            self.stack = self.stack[:len(self.stack)-1-self.stack[::-1].index(tag)]

    def handle_data(self, data):
        if any(t in {'script','style','code','pre','svg','math'} for t in self.stack): return
        start = self.position()
        if self.pending and self.pending[1] == start: self.pending = (self.pending[0], start+len(data))
        else: self.flush(); self.pending = (start, start+len(data))

    def entity(self, raw):
        if not self.text.startswith(raw,self.position()): raw = raw[:-1]
        self.handle_data(raw)

    def handle_entityref(self, name): self.entity('&'+name+';')
    def handle_charref(self, name): self.entity('&#'+name+';')
    def handle_comment(self, data): self.flush()
    def handle_decl(self, data): self.flush()
    def handle_pi(self, data): self.flush()


def html_spans(text):
    parser = HTMLSpans(text); parser.feed(text); parser.close(); parser.flush()
    return parser.spans


def json_document(path, text, encoding):
    def pairs(items):
        if len({k for k, _ in items}) != len(items): raise ValueError('JSON 存在重复键名，请先检查文件。')
        return dict(items)
    json.loads(text, object_pairs_hook=pairs, parse_constant=lambda _: (_ for _ in ()).throw(ValueError('JSON 数值无效。')))
    spans = []
    for token in re.finditer(r'"(?:[^"\\]|\\.)*"', text):
        if text[token.end():].lstrip().startswith(':'): continue
        value = json.loads(token[0])
        if re.match(r'^(?:[a-z]+://|[A-Za-z]:[\\/]|[/\\])', value, re.I): continue
        spans.append((*token.span(), value, lambda s: json.dumps(s, ensure_ascii=False)))
    return spans_document(path, text, encoding, spans, '仅处理 JSON 字符串值；键名、数字、布尔值及结构保留，路径和网址跳过。')


def subtitle_spans(text, kind):
    spans, offset = [], 0
    if kind == '.srt':
        blocks = re.split(r'\r?\n\s*\r?\n', text.strip())
        for block in blocks:
            lines = block.splitlines()
            if len(lines) < 3 or not re.fullmatch(r'\d+', lines[0]) or not re.match(r'^\d{2,}:\d{2}:\d{2}[,.]\d{3}\s+-->\s+\d{2,}:\d{2}:\d{2}[,.]\d{3}', lines[1]):
                raise ValueError('SRT 序号或时间轴格式不正确，请先确认字幕能正常打开。')
        for line in text.splitlines(keepends=True):
            content = line.rstrip('\r\n')
            if content.strip() and not content.strip().isdigit() and '-->' not in content:
                for a,b,value,render in plain_spans(content, [m.span() for m in re.finditer(r'<[^>]*>|\{[^}]*\}', content)]):
                    spans.append((offset+a,offset+b,html.unescape(value),lambda s: html.escape(subtitle_text(s), quote=False)))
            offset += len(line)
    else:
        events, fields = False, None
        for line in text.splitlines(keepends=True):
            content = line.rstrip('\r\n')
            if re.match(r'^\[.*\]$', content): events = content.lower() == '[events]'
            if events and content.lower().startswith('format:'):
                fields = [v.strip().lower() for v in content.split(':',1)[1].split(',')]
                if not fields or fields[-1] != 'text': raise ValueError('ASS 的 Text 字段必须位于最后一列。')
            if events and content.lower().startswith('dialogue:'):
                if not fields: raise ValueError('ASS 缺少 Events Format 字段。')
                pieces = content.split(':',1)[1].split(',',len(fields)-1)
                if len(pieces) != len(fields): raise ValueError('ASS 对话字段不完整。')
                dialogue = pieces[-1]; start = offset+len(content)-len(dialogue)
                # Drawing events are geometry, not prose.
                if not re.search(r'\\p[1-9]', dialogue):
                    for a,b,value,render in plain_spans(dialogue,[m.span() for m in re.finditer(r'\{[^}]*\}|\\[Nnh]', dialogue)]):
                        spans.append((start+a,start+b,value,lambda s: subtitle_text(s).replace('\\','＼').replace('{','｛').replace('}','｝')))
            offset += len(line)
    return spans


def xlsx_document(path):
    document = TextDocument(path)
    namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    q = lambda name: '{'+namespace+'}'+name
    with ZipFile(path) as archive:
        archive_check(archive)
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            strings = [''.join(t.text or '' for t in node.findall('.//'+q('t')) if not any(p.tag == q('rPh') for p in t.iterancestors())) for node in parse_xml(archive.read('xl/sharedStrings.xml'))]
        workbook = parse_xml(archive.read('xl/workbook.xml'))
        relations = parse_xml(archive.read('xl/_rels/workbook.xml.rels'))
        rels = {r.get('Id'):r.get('Target') for r in relations if r.get('TargetMode') != 'External'}
        cells, roots = {}, {}
        for page, sheet in enumerate(workbook.findall('.//'+q('sheet')),1):
            target = rels.get(sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'),'')
            name = target.lstrip('/') if target.startswith('/') else 'xl/'+target
            import posixpath
            name = posixpath.normpath(name)
            if name not in archive.namelist(): raise ValueError('Excel 工作表关系不完整。')
            root = roots[name] = parse_xml(archive.read(name))
            for cell in root.findall('.//'+q('sheetData')+'/'+q('row')+'/'+q('c')):
                if cell.find(q('f')) is not None: continue
                value, kind = cell.find(q('v')), cell.get('t')
                if kind == 's' and value is not None:
                    try: text = strings[int(value.text)]
                    except (IndexError,ValueError,TypeError) as e: raise ValueError('Excel 共享文本索引无效。') from e
                elif kind == 'inlineStr': text = ''.join(t.text or '' for t in cell.findall('.//'+q('t')))
                elif kind == 'str' and value is not None: text = value.text or ''
                else: continue
                locator = f'xlsx/{page}/{cell.get("r")}'
                if document.add(locator,text,'table',page) is not None: cells[locator] = (name,cell)
        if any(n.startswith(('xl/drawings/','xl/charts/')) for n in archive.namelist()):
            document.blocked.append(('xlsx/drawings','[Excel 图表、图片和绘图对象]','image',None,'图表及图片保留；其内部文字暂不自动改写。'))
    def writer(destination, values):
        changed = set()
        for locator, (name,cell) in cells.items():
            if values[locator] == document.values[locator]: continue
            for child in list(cell):
                if child.tag in {q('v'),q('is')}: cell.remove(child)
            cell.set('t','inlineStr'); node = etree.Element(q('is')); text = etree.SubElement(node,q('t'))
            text.text = values[locator]; text.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
            cell.insert(0,node); changed.add(name)
        with ZipFile(path) as original, ZipFile(destination,'w') as output:
            output.comment = original.comment
            for info in original.infolist():
                output.writestr(info, etree.tostring(roots[info.filename],encoding='utf-8',xml_declaration=True) if info.filename in changed else original.read(info.filename))
    document.writer = writer
    document.note = '保留工作表、公式、数字、单元格样式及图表；仅改写文本单元格，较长译文请检查列宽与行高。'
    return document


def epub_document(path):
    document = TextDocument(path); contents = {}
    with ZipFile(path) as archive:
        archive_check(archive)
        container = parse_xml(archive.read('META-INF/container.xml'))
        opf = container.find('.//{*}rootfile')
        if opf is None: raise ValueError('EPUB 缺少内容清单。')
        opf_name = opf.get('full-path'); root = parse_xml(archive.read(opf_name))
        manifest = {item.get('id'):item for item in root.findall('.//{*}manifest/{*}item')}
        ordered = [ref.get('idref') for ref in root.findall('.//{*}spine/{*}itemref')]
        ordered += [key for key in manifest if key not in ordered]
        import posixpath
        from urllib.parse import unquote
        for page,key in enumerate(ordered,1):
            item = manifest[key]
            if item.get('media-type') not in {'application/xhtml+xml','text/html'}: continue
            name = posixpath.normpath(posixpath.join(posixpath.dirname(opf_name), unquote(item.get('href','').split('#')[0])))
            if name.startswith('../') or name.startswith('/'): raise ValueError('EPUB 内容路径无效。')
            raw = archive.read(name)
            try: text = raw.decode('utf-8-sig')
            except UnicodeError as e: raise ValueError('EPUB 正文需要 UTF-8 编码。') from e
            if len(raw) > MAX_TEXT: raise ValueError('EPUB 章节过大，请拆分处理。')
            encoding = 'utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8'
            replacements = []
            for index,(start,end,value,render) in enumerate(html_spans(text)):
                locator = f'epub/{page}/{index}'
                if document.add(locator,value,'text',page) is not None: replacements.append((start,end,locator,render))
            contents[name] = (text,encoding,replacements)
    def writer(destination,values):
        changed = {}
        for name,(text,encoding,replacements) in contents.items():
            result = text
            for start,end,locator,render in reversed(replacements):
                if values[locator] != document.values[locator]: result = result[:start]+render(values[locator])+result[end:]
            if result != text: changed[name] = result.encode(encoding)
        with ZipFile(path) as original, ZipFile(destination,'w') as output:
            output.comment = original.comment
            for info in original.infolist(): output.writestr(info, changed.get(info.filename,original.read(info.filename)))
    document.writer = writer; document.note = '保留 EPUB 章节、目录链接、样式、图片及打包结构；仅处理可读正文。'
    return document


def pdf_document(path):
    from pypdf import PdfReader
    from .pdf_reflow import extract_page, page_blocks, compact, join_link_wraps, fit_contents
    from .pdf_links import mark_page, append_page_links, destinations
    from .pdf_images import append_images
    from collections import Counter
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(''): raise ValueError('PDF 已加密，请先提供可直接打开的未加密文件。')
    if len(reader.pages) > 2000: raise ValueError('PDF 页数过多，请拆分处理。')
    document = TextDocument(path)
    document._star_pdf_headers=set()
    section=document.document.sections[0]
    section.page_width=Pt(595.3);section.page_height=Pt(841.9)
    section.top_margin=section.bottom_margin=Pt(48)
    section.left_margin=section.right_margin=Pt(48)
    missing, pages, contents = [], [], []
    for page, item in enumerate(reader.pages,1):
        stream = item.get_contents()
        if stream is not None and len(stream.get_data()) > MAX_TEXT: raise ValueError(f'PDF 第 {page} 页内容过大，请拆分处理。')
        pages.append(extract_page(item))
    headers=Counter(compact(line) for text in pages for line in text.splitlines()[:1] if 0<len(line.strip())<100)
    repeated={line for line,n in headers.items() if n>1}
    for page, (item,text) in enumerate(zip(reader.pages,pages),1):
        if not text.strip():
            missing.append(page)
            # Keep the original-page anchor without inserting an obsolete OCR
            # instruction into a picture whose text can now be reviewed below.
            paragraph=document.document.add_paragraph()
            paragraph.paragraph_format.page_break_before=page>1
            mark_page(paragraph,page)
            count,failures=append_images(document.document,item,page)
            if not count and not failures and item.get_contents() is not None and item.get_contents().get_data().strip():
                if not hasattr(document.document,'_star_pdf_image_failures'):document.document._star_pdf_image_failures=[]
                reason=f'PDF 第 {page} 页有内容但没有可读文字层或可保留位图，可能含矢量轮廓文字；请对照原稿。'
                document.document._star_pdf_image_failures.append(dict(locator=f'pdf-image/{page}/unreadable',slide=page,para=None,
                    source_text='',translation='',reason=reason,image=True,unreadable_page=True,retained_source=True))
            for reason in failures:document.blocked.append((f'pdf/image/{page}',reason,'image',page,reason))
            append_page_links(document.document,reader,item)
            continue
        first = True
        text=join_link_wraps(text,destinations(reader,item)[0])
        for block in page_blocks(text,repeated):
            locator = f'body/{len(document.paragraphs)}'
            paragraph = document.add(locator,block.text,'text',page)
            # Pure dates/page numbers need no language check, but must still be
            # present in the output after separating them from translatable text.
            if paragraph is None:paragraph=document.document.add_paragraph(block.text)
            if paragraph is not None:
                paragraph.paragraph_format.page_break_before = first and page > 1
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.runs[0].font.name = 'Arial'
                if block.kind=='header':document._star_pdf_headers.add(locator)
                paragraph.runs[0].font.size = Pt(9 if block.kind in {'header','footer','reference'} else 13 if block.kind=='heading' else 11)
                if block.kind=='toc':contents.append(paragraph)
                if block.kind=='heading':
                    paragraph.runs[0].bold=True
                    paragraph.paragraph_format.keep_with_next=True
                if first:
                    mark_page(paragraph,page)
                first = False
        _,failures=append_images(document.document,item,page)
        for index,reason in enumerate(failures):document.blocked.append((f'pdf/image/{page}/{index}',reason,'image',page,reason))
        # PDF annotations are not returned by extract_text(). Preserve their
        # targets separately instead of guessing which translated English words
        # match a source rectangle. The output is explicitly a reflowed Word.
        append_page_links(document.document,reader,item,visible_text=text)
    if not document.paragraphs and not getattr(document.document,'_star_pdf_images',[]): raise ValueError('PDF 未发现可读取的文字层，可能是扫描件。请先进行 OCR，或使用原始 Word / PPT 文档。')
    if missing: document.blocked.append(('pdf/scans','[无可读文字层的页面：'+','.join(map(str,missing))+']','image',None,'这些页面未计入正文翻译；保留的位图另行检查，复杂扫描版式与矢量轮廓文字需对照原 PDF。'))
    document.blocked.append(('pdf/layout','[PDF 原始版式]','image',None,'PDF 文字与图片按原页归组重新排入 Word；图片中文独立识别与核验，未能安全处理的位置单独提示，复杂表格与原始版式不自动重建。'))
    def writer(destination,values):
        fit_contents(document.document,contents)
        document.document.save(str(destination))
    document.writer = writer
    document.note = 'PDF 文字层按原页顺序输出为可编辑 Word，采用重新排版；原图像在所属页文字之后保留，图片中文独立识别、翻译并对照核验；原链接与页间跳转集中保留。不能安全覆盖的图片文字、复杂表格与无文字层页面会保留并提示。'
    return document


def open_text_document(path):
    kind = Path(path).suffix.lower()
    if kind == '.xlsx': document = xlsx_document(path)
    elif kind == '.epub': document = epub_document(path)
    elif kind == '.pdf': document = pdf_document(path)
    else:
        text,encoding = read_text(path)
        if kind == '.json': document = json_document(path,text,encoding)
        else:
            spans = markdown_spans(text) if kind in {'.md','.markdown'} else html_spans(text) if kind in {'.html','.htm'} else subtitle_spans(text,kind) if kind in {'.srt','.ass'} else line_spans(text)
            note = '保留字幕时间轴、序号及样式标记，仅处理字幕正文。' if kind in {'.srt','.ass'} else '保留标记、链接及代码，处理可读正文。' if kind != '.txt' else '保留文本编码、段落及换行。'
            document = spans_document(path,text,encoding,spans,note)
            if kind == '.srt': document.normalize = subtitle_text
            elif kind == '.ass': document.normalize = lambda s: subtitle_text(s).replace('\\','＼').replace('{','｛').replace('}','｝')
    return document,document.paragraphs,document.blocked
