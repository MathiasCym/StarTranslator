"""Visible master/layout text and chart labels with their embedded workbook."""
from io import BytesIO
from copy import deepcopy
from collections import Counter
from pathlib import PurePosixPath
import posixpath
from zipfile import ZipFile
from lxml import etree
from ..text_integrity import contains_chinese
from pptx.text.text import TextFrame

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS = {'a':A, 'c':C, 'p':P}


class Owner:
    def __init__(self, part):
        self.part = part


class XmlRun:
    def __init__(self, element, parent):
        self._r, self._parent = element, parent

    @property
    def text(self):
        return self._r.text or ''

    @text.setter
    def text(self, value):
        self._r.text = value


class XmlParagraph:
    def __init__(self, node):
        self.node, self._p = node, node.getparent()
        self.original = node.text or ''

    @property
    def runs(self):
        # write_styled replaces the original node, so always locate its successor.
        return [XmlRun(self._p.find(f'{{{C}}}v'), self)]

    @property
    def text(self):
        return self.runs[0].text


def _cell_reference(node):
    """Resolve a string-cache point to its workbook cell, not its display text."""
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    point = node.getparent()
    cache = point.getparent()
    if cache.tag != f'{{{C}}}strCache': return None
    formula = cache.getparent().findtext(f'{{{C}}}f') or ''
    if '!' not in formula: return None
    sheet, address = formula.rsplit('!', 1)
    sheet = sheet.strip("'").replace("''", "'")
    if '[' in sheet or ',' in address: return None
    try:
        left, top, right, bottom = range_boundaries(address.replace('$', ''))
        index = int(point.get('idx', '0'))
        width = right-left+1
        row, col = top+index//width, left+index%width
        if row > bottom or index < 0: return None
        return sheet, f'{get_column_letter(col)}{row}'
    except (TypeError, ValueError): return None


def extra_paragraphs(path, document):
    paragraphs, parts, workbooks, blocked = [], {}, {}, []
    package = {str(part.partname).lstrip('/'):part for part in document.part.package.iter_parts()}
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        for name in sorted(names):
            if not name.endswith('.xml') or not name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/charts/')):
                continue
            if '/_rels/' in name or name not in package:
                continue
            part = package[name]
            root = getattr(part, '_element', None)
            if root is None:
                root = etree.fromstring(archive.read(name))
            parts[name] = root
            source = 'chart' if name.startswith('ppt/charts/') else 'template'
            if source == 'template':
                nodes = root.xpath('.//p:sp[not(p:nvSpPr/p:nvPr/p:ph)]/p:txBody', namespaces=NS) if type(root) is etree._Element else root.xpath('.//p:sp[not(p:nvSpPr/p:nvPr/p:ph)]/p:txBody')
            else:
                nodes = root.xpath('.//a:txBody | .//c:rich', namespaces=NS) if type(root) is etree._Element else root.xpath('.//a:txBody | .//c:rich')
            for i, body in enumerate(nodes):
                for j, paragraph in enumerate(TextFrame(body, Owner(part)).paragraphs):
                    paragraphs.append((f'part:{name}/body{i}/p{j}', paragraph, source, None, j+1))
            if source != 'chart':
                continue
            rel_name = posixpath.join(posixpath.dirname(name), '_rels', PurePosixPath(name).name + '.rels')
            workbook = None
            if rel_name in names:
                relationships = etree.fromstring(archive.read(rel_name))
                for rel in relationships:
                    if rel.get('Type', '').endswith('/package') and rel.get('TargetMode') != 'External':
                        target = posixpath.normpath(posixpath.join(posixpath.dirname(name), rel.get('Target', '')))
                        if target.startswith('ppt/embeddings/') and target.endswith('.xlsx') and target in names:
                            workbook = target
                            break
            labels = []
            for node in root.iter(f'{{{C}}}v'):
                parent = node.getparent()
                # Never touch numeric caches or formula/axis references.
                ancestors = {ancestor.tag for ancestor in node.iterancestors()}
                if f'{{{C}}}strLit' not in ancestors and not (workbook and f'{{{C}}}strCache' in ancestors):
                    if f'{{{C}}}strCache' in ancestors:
                        if contains_chinese(node.text or ''):
                            blocked.append((f'part:{name}/linked-label{len(blocked)}',node.text,'chart',None,
                                '该图表关联外部数据，需更新关联数据中的中文标签。'))
                    continue
                proxy = XmlParagraph(node)
                proxy.cell = _cell_reference(node)
                proxy.part_name = name
                labels.append(proxy)
                paragraphs.append((f'part:{name}/label{len(labels)-1}', proxy, 'chart', None, len(labels)))
            if workbook:
                workbooks.setdefault(workbook, []).extend(labels)
    document._star_extra_parts = parts
    document._star_workbooks = workbooks
    document._star_extra_blocked = blocked
    return paragraphs


def changed_extra_parts(document, changed, original):
    result = {}
    names = set()
    for locator in changed:
        if locator.startswith('part:'):
            names.add(locator[5:].split('.xml/',1)[0] + '.xml')
    for name, paragraphs in getattr(document, '_star_workbooks', {}).items():
        by_cell = {}
        for paragraph in paragraphs:
            if paragraph.cell:
                by_cell.setdefault(paragraph.cell, []).append(paragraph)
        replacements = {}
        for cell, labels in by_cell.items():
            edited = [p.text for p in labels if p.text != p.original]
            if not edited: continue
            # Several chart caches can reference one cell. Keep them coherent;
            # equal source strings in *different* cells remain independent.
            canonical = Counter(edited).most_common(1)[0][0]
            replacements[cell] = canonical
            for label in labels:
                if label.text != canonical:
                    label.runs[0].text = canonical
                names.add(label.part_name)
        if not replacements:
            continue
        target = BytesIO()
        with ZipFile(BytesIO(original.read(name))) as workbook, ZipFile(target, 'w') as output:
            output.comment = workbook.comment
            sheets = etree.fromstring(workbook.read('xl/workbook.xml'))
            rels = etree.fromstring(workbook.read('xl/_rels/workbook.xml.rels'))
            targets = {r.get('Id'):r.get('Target') for r in rels}
            sheet_names = {}
            for sheet in sheets.iter(f'{{{S}}}sheet'):
                target_name = targets.get(sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'), '')
                if target_name:
                    sheet_names[posixpath.normpath(posixpath.join('xl', target_name)).lstrip('/')] = sheet.get('name')
            shared = etree.fromstring(workbook.read('xl/sharedStrings.xml')) if 'xl/sharedStrings.xml' in workbook.namelist() else []
            for info in workbook.infolist():
                data = workbook.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    edited = False
                    for cell in root.iter(f'{{{S}}}c'):
                        text = replacements.get((sheet_names[info.filename], cell.get('r')))
                        if text is None: continue
                        if cell.find(f'{{{S}}}f') is not None:
                            # Preserve spreadsheet calculations; only update the
                            # displayed cached string consumed by the chart.
                            cell.set('t','str')
                            cached=cell.find(f'{{{S}}}v')
                            if cached is None:cached=etree.SubElement(cell,f'{{{S}}}v')
                            cached.text=text;edited=True
                            continue
                        container = cell.find(f'{{{S}}}is')
                        value = cell.find(f'{{{S}}}v')
                        if cell.get('t') == 's' and value is not None:
                            container = deepcopy(shared[int(value.text)])
                            container.tag = f'{{{S}}}is'
                        elif container is not None: container = deepcopy(container)
                        else: container = etree.Element(f'{{{S}}}is')
                        texts = list(container.iter(f'{{{S}}}t'))
                        if not texts: texts = [etree.SubElement(container, f'{{{S}}}t')]
                        texts[0].text = text
                        texts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        for node in texts[1:]: node.text = ''
                        # Do not change a shared string used by unrelated cells.
                        for child in list(cell):
                            if child.tag in {f'{{{S}}}v', f'{{{S}}}is'}: cell.remove(child)
                        cell.set('t', 'inlineStr'); cell.append(container)
                        edited = True
                    if edited:
                        data = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
                output.writestr(info, data)
        result[name] = target.getvalue()
    # Synchronizing a shared workbook cell can also change an otherwise untouched
    # chart. Serialize only after every linked cache has its final text.
    for name in names:
        root = document._star_extra_parts[name]
        result[name] = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
    return result
