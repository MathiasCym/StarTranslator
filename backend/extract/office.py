"""One traversal shared by extraction and export, including tables and groups."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .common import is_meaningful
from ..models import Segment, RunInfo
from ..document_input import validate_document_input
from ..replacer.text_runs import text_runs, editable_paragraph


def open_office(path):
    validate_document_input(path)
    if Path(path).suffix.lower() not in {'.docx', '.pptx'}:
        from ..formats import open_text_document
        return open_text_document(path)
    paragraphs, blocked, seen, layouts = [], [], set(), {}
    if Path(path).suffix.lower() == '.docx':
        document = Document(path)
        from .word import word_paragraphs
        paragraphs, blocked = word_paragraphs(document)
        images = len(document.element.body.findall('.//' + qn('w:drawing')))
        if images:
            blocked.append(('drawings', f'[{images} 个图片/绘图对象]', 'image', None, '图片内文字未做 OCR，需人工检查。'))
    else:
        document = Presentation(path)
        from ..replacer.pptx_geometry import walk_shapes, bounds
        def shapes(items, prefix, slide):
            for key,shape,transform in walk_shapes(items,prefix):
                if shape.has_table:
                    for r,row in enumerate(shape.table.rows):
                        for c,cell in enumerate(row.cells):
                            if cell.is_spanned:
                                continue
                            for p,para in enumerate(cell.text_frame.paragraphs):
                                paragraphs.append((f'{key}/table/{r}/{c}/{p}',para,'table',slide,p+1))
                elif shape.has_text_frame:
                    for p,para in enumerate(shape.text_frame.paragraphs):
                        locator = f'{key}/p{p}'
                        paragraphs.append((locator,para,'text',slide,p+1))
                        x,y,w,h = bounds(shape,transform)
                        layouts[locator] = dict(container=key.rsplit('/',1)[0], x=x/12700, y=y/12700,
                                               width=w/12700, height=h/12700)
                elif shape.has_chart:
                    blocked.append((key,'[图表对象]','image',slide,'可编辑的图表文字已纳入检查，图形内其他标注需人工检查。'))
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    blocked.append((key,'[图片/图表]','image',slide,'图片或图表内文字需人工检查。'))
                else:
                    # Preserve unsupported embedded diagrams without claiming coverage.
                    if 'graphicFrame' in shape._element.tag:
                        blocked.append((key,'[嵌入图示]','image',slide,'嵌入图示内文字暂不支持自动处理。'))
        for s,slide in enumerate(document.slides,1):
            shapes(slide.shapes,f'slide{s}',s)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                for p,para in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                    paragraphs.append((f'slide{s}/notes/p{p}',para,'notes',s,p+1))
        from .pptx_extra import extra_paragraphs
        paragraphs.extend(extra_paragraphs(path, document))
        blocked.extend(document._star_extra_blocked)
    if Path(path).suffix.lower() == '.pptx': document._star_text_layouts = layouts
    return document, paragraphs, blocked


def extract_office(path):
    document, paragraphs, blocked = open_office(path)
    result = []
    for locator,para,source,slide,number in paragraphs:
        text = para.text
        if not is_meaningful(text):
            continue
        runs = [RunInfo(r.text,i) for i,r in enumerate(text_runs(para))]
        editable = editable_paragraph(para)
        reason = getattr(para, 'not_editable_reason', '')
        if reason: editable = False
        result.append(Segment(id=str(len(result)),file=Path(path).name,text=text,locator=locator,
            source=source,slide=slide,para=number,runs=runs,editable=editable,
            layout=getattr(document,'_star_text_layouts',{}).get(locator,{}),
            not_editable_reason='' if editable else reason or '含动态域、嵌入对象或特殊文字结构，需人工处理。'))
    for locator,text,source,slide,reason in blocked:
        result.append(Segment(id=str(len(result)),file=Path(path).name,text=text,locator=locator,
            source=source,slide=slide,editable=False,not_editable_reason=reason))
    return result
