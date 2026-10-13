"""PPTX 提取：形状文本框、占位符、备注。

不可提取的（只报告不替换，或在第二阶段用 VL 读图）：
  - 图片内文字 → 标记 not_editable
  - 图表数据
  - SmartArt / 组合形状内的深层文本（结构特殊，风险高）
"""
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..models import RunInfo, Segment
from .common import is_meaningful, needs_merge


def _shape_locator(slide_idx: int, shape_idx: int, para_idx: int) -> str:
    return f"slide{slide_idx + 1}/shape{shape_idx}/p{para_idx}"


def extract_pptx(path: str) -> list[Segment]:
    prs = Presentation(path)
    file_name = os.path.basename(path)
    segments: list[Segment] = []
    counter = 0

    def next_id() -> str:
        nonlocal counter
        counter += 1
        return f"seg-{counter:04d}"

    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            # 图片：文字在像素里，标注为不可替换
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                segments.append(Segment(
                    id=next_id(), file=file_name,
                    text="[图片]",
                    locator=f"slide{slide_idx + 1}/shape{shape_idx}",
                    source="image", slide=slide_idx + 1,
                    editable=False,
                    not_editable_reason="图片内文字需第二阶段用视觉模型提取",
                ))
                continue

            # 组合形状：结构特殊，首版不深入
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                segments.append(Segment(
                    id=next_id(), file=file_name,
                    text=f"[组合形状：{shape.name}]",
                    locator=f"slide{slide_idx + 1}/shape{shape_idx}",
                    source="group", slide=slide_idx + 1,
                    editable=False,
                    not_editable_reason="组合形状内文本结构特殊，首版只报告不替换",
                ))
                continue

            if not shape.has_text_frame:
                continue

            for para_idx, para in enumerate(shape.text_frame.paragraphs):
                text = para.text
                if not is_meaningful(text):
                    continue
                runs = [RunInfo(text=r.text, index=i) for i, r in enumerate(para.runs)]
                run_texts = [r.text for r in runs]
                segments.append(Segment(
                    id=next_id(), file=file_name,
                    text=text,
                    locator=_shape_locator(slide_idx, shape_idx, para_idx),
                    source="text", slide=slide_idx + 1,
                    runs=runs,
                    editable=True,
                    needs_merge=needs_merge(run_texts),
                ))

        # 演讲者备注
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            for para_idx, para in enumerate(notes_tf.paragraphs):
                text = para.text
                if not is_meaningful(text):
                    continue
                runs = [RunInfo(text=r.text, index=i) for i, r in enumerate(para.runs)]
                segments.append(Segment(
                    id=next_id(), file=file_name,
                    text=text,
                    locator=f"slide{slide_idx + 1}/notes/p{para_idx}",
                    source="notes", slide=slide_idx + 1,
                    runs=runs,
                    editable=True,
                    needs_merge=needs_merge([r.text for r in runs]),
                ))

    return segments
