"""DOCX 提取：正文段落、表格、页眉页脚。

用 python-docx 遍历。locator 记录 (body_type, container_index, para_index)，
替换层按同样的路径回找到对应段落对象。
"""
import os

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..models import RunInfo, Segment
from .common import is_meaningful, needs_merge

# python-docx 的 Table 不在 document.paragraphs 里，需要手动按文档顺序遍历。


def _iter_block_items(parent):
    """按文档顺序产出 Paragraph 和 Table。"""
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _para_to_segment(para, seg_id: str, file_name: str, locator: str,
                     source: str, para_idx: int) -> Segment | None:
    text = para.text
    if not is_meaningful(text):
        return None
    runs = [RunInfo(text=r.text, index=i) for i, r in enumerate(para.runs)]
    run_texts = [r.text for r in runs]
    return Segment(
        id=seg_id,
        file=file_name,
        text=text,
        locator=locator,
        source=source,
        para=para_idx,
        runs=runs,
        editable=True,
        needs_merge=needs_merge(run_texts),
    )


def extract_docx(path: str) -> list[Segment]:
    doc = Document(path)
    file_name = os.path.basename(path)
    segments: list[Segment] = []
    counter = 0

    def next_id() -> str:
        nonlocal counter
        counter += 1
        return f"seg-{counter:04d}"

    # --- 正文：段落 + 表格 ---
    for block_idx, block in enumerate(_iter_block_items(doc)):
        if isinstance(block, Paragraph):
            seg = _para_to_segment(
                block, next_id(), file_name,
                locator=f"body/{block_idx}/p", source="text", para_idx=block_idx,
            )
            if seg:
                segments.append(seg)
        elif isinstance(block, Table):
            for row_idx, row in enumerate(block.rows):
                for cell_idx, cell in enumerate(row.cells):
                    for p_idx, para in enumerate(cell.paragraphs):
                        seg = _para_to_segment(
                            para, next_id(), file_name,
                            locator=f"body/{block_idx}/tbl/{row_idx}/{cell_idx}/{p_idx}",
                            source="table", para_idx=block_idx,
                        )
                        if seg:
                            segments.append(seg)

    # --- 页眉页脚 ---
    for sec_idx, section in enumerate(doc.sections):
        for hf_name, hf in (
            ("header", section.header), ("footer", section.footer),
        ):
            if hf is None:
                continue
            for p_idx, para in enumerate(hf.paragraphs):
                seg = _para_to_segment(
                    para, next_id(), file_name,
                    locator=f"sec/{sec_idx}/{hf_name}/{p_idx}",
                    source=hf_name, para_idx=p_idx,
                )
                if seg:
                    segments.append(seg)

    return segments
