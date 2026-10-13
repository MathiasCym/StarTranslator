"""Excel presentation shared by glossary exports and editable shared tables."""
import json
import math
import re
from copy import copy
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree as ET


def presentation_note(text):
    """Reflow short fragments copied from slide text into readable paragraphs."""
    text=str(text or '')
    lines=text.splitlines()
    # Normal manually authored notes keep their deliberate line breaks.
    if len(lines)<4 or sum(0<len(line.strip())<28 for line in lines)<3:return text
    paragraphs=[];short=[]
    def flush():
        if short:paragraphs.append(' '.join(short));short.clear()
    for raw in lines:
        line=raw.strip()
        if not line:flush();continue
        if len(line)>=40 or re.match(r'^(?:中文定义来自|规则种子[：:]|待专业确认|待确认|20\d{2}[-/]\d{1,2}[-/]\d{1,2}|\d+[.、)]\s*)',line) or len(line)>=20 and re.search(r'[。！？]$',line):
            flush();paragraphs.append(line)
        else:short.append(line)
    flush()
    return '\n\n'.join(paragraphs)


@lru_cache(maxsize=1)
def preset():
    return json.loads(Path(__file__).with_name('excel-style.json').read_text(encoding='utf-8'))


def apply_style(cell, field, header=False):
    from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection
    spec=preset()['columns'].get(field,preset()['columns']['note'])
    style=spec['header' if header else 'body']
    for name,cls in [('font',Font),('fill',PatternFill),('border',Border),('alignment',Alignment),('protection',Protection)]:
        setattr(cell,name,cls.from_tree(ET.fromstring(style[name])))
    cell.number_format=style['number_format']


@lru_cache(maxsize=64)
def _font(name, size, bold):
    # Use the bundled Pillow with Windows fonts where available; never require Excel.
    from PIL import ImageFont
    fonts=Path('C:/Windows/Fonts')
    files={'微软雅黑':('msyh.ttc','msyhbd.ttc'),'Microsoft YaHei':('msyh.ttc','msyhbd.ttc'),
           'Calibri':('calibri.ttf','calibrib.ttf'),'Arial':('arial.ttf','arialbd.ttf')}
    pair=files.get(name,files['微软雅黑'])
    for filename in (pair[int(bool(bold))], 'msyh.ttc'):
        try:return ImageFont.truetype(str(fonts/filename),max(1,round(size*96/72)))
        except OSError:pass
    return None


def wrapped_lines(text, width, font):
    text=str(text or '').replace('\r\n','\n').replace('\r','\n')
    measure=_font(font.name,float(font.sz or 11),font.b)
    pixels=max(10,float(width or 13)*7-12)
    def length(value):
        if measure:return measure.getlength(value)
        return sum(14 if ord(c)>255 else 7 for c in value)*(float(font.sz or 11)/11)
    count=0
    # Keep whole English words when possible, permit CJK and long tokens to wrap.
    for paragraph in text.split('\n'):
        line='';count+=1
        for token in re.findall(r'[A-Za-z0-9_]+(?:[-\u2019\u0027][A-Za-z0-9_]+)*|[^A-Za-z0-9_]',paragraph):
            if line and length(line+token)>pixels:
                count+=1;line=''
            if length(token)>pixels:
                for char in token:
                    if line and length(line+char)>pixels:count+=1;line=''
                    line+=char
            else:line+=token
    return count


def fit_rows(sheet, count, width):
    """Wrap visually without inserting delimiters into glossary values."""
    from openpyxl.utils import get_column_letter
    for row in range(1,count+2):
        height=22 if row>1 else 24
        for col in range(1,width+1):
            cell=sheet.cell(row,col)
            align=copy(cell.alignment);align.wrap_text=True;align.shrink_to_fit=False;cell.alignment=align
            if sheet.column_dimensions[get_column_letter(col)].hidden:continue
            lines=wrapped_lines(cell.value,sheet.column_dimensions[get_column_letter(col)].width,cell.font)
            height=max(height,lines*float(cell.font.sz or 11)*1.2+6)
        sheet.row_dimensions[row].height=min(409,math.ceil(height))


def workbook(rows,columns,header,existing=None):
    """Create the standard format, or retain an existing table's custom styles."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    from .writer import _header_map, TableTooComplex
    preserve=existing is not None and Path(existing).is_file()
    book=openpyxl.load_workbook(existing,keep_vba=Path(existing).suffix.lower()=='.xlsm') if preserve else openpyxl.Workbook()
    sheet=book.active
    if len(book.worksheets)!=1 or sheet.merged_cells:
        book.close();raise TableTooComplex('术语表含多工作表或合并单元格，请整理为单张表后保存。')
    width=len(columns);original_max=sheet.max_row;old_cols=_header_map([c.value for c in sheet[1]]) if preserve else None
    row_styles={};body_styles={};old_headers={}
    style_key=lambda field,index:field if field else ('custom',index)
    if preserve and old_cols:
        for i,field in enumerate(old_cols,1):
            old_headers[style_key(field,i)]=copy(sheet.cell(1,i)._style)
            body_styles[style_key(field,i)]=copy(sheet.cell(2,i)._style)
        for cells in sheet.iter_rows(min_row=2):
            identity=tuple(str(cells[old_cols.index(f)].value or '').strip() for f in ('zh','en'))
            row_styles[identity]={style_key(field,i+1):(copy(cells[i]._style),copy(cells[i].comment),copy(cells[i].hyperlink)) for i,field in enumerate(old_cols)}
    if not preserve:
        book.loaded_theme=preset()['theme'].encode('utf-8');sheet.title=preset()['sheet'];sheet.freeze_panes='C2'
    for col,field in enumerate(columns,1):
        letter=get_column_letter(col)
        if not preserve or not old_cols or field not in old_cols:
            sheet.column_dimensions[letter].width=preset()['columns'].get(field,{'width':24})['width']
            sheet.column_dimensions[letter].hidden=field in ('priority','source')
        cell=sheet.cell(1,col,header[col-1]);cell.data_type='s'
        key=style_key(field,col)
        if key in old_headers:cell._style=copy(old_headers[key])
        else:apply_style(cell,field,True)
    for number,values in enumerate(rows,2):
        identity=tuple(str(values[columns.index(f)]).strip() for f in ('zh','en'))
        original=row_styles.get(identity,{})
        for col,(field,value) in enumerate(zip(columns,values),1):
            cell=sheet.cell(number,col);cell.value=presentation_note(value) if field=='note' else value;cell.data_type='s';cell.comment=None;cell.hyperlink=None
            key=style_key(field,col)
            if key in original:
                cell._style=copy(original[key][0]);cell.comment=copy(original[key][1]);cell.hyperlink=copy(original[key][2])
            elif key in body_styles:cell._style=copy(body_styles[key])
            else:apply_style(cell,field)
    if original_max>len(rows)+1:
        sheet.delete_rows(len(rows)+2,original_max-len(rows)-1)
        for number in list(sheet.row_dimensions):
            if number>len(rows)+1:del sheet.row_dimensions[number]
    for table in sheet.tables.values():
        table.ref=f'A1:{get_column_letter(width)}{max(2,len(rows)+1)}'
        table.tableColumns=[TableColumn(id=i,name=str(h)) for i,h in enumerate(header,1)]
        if table.autoFilter:table.autoFilter.ref=table.ref
    if not sheet.tables:sheet.auto_filter.ref=f'A1:{get_column_letter(width)}{len(rows)+1}'
    fit_rows(sheet,len(rows),width)
    return book
