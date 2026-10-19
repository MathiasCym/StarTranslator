"""Measure real Windows fonts and fit whole frames without moving slide objects."""
from collections import Counter
from contextlib import AbstractContextManager
import ctypes
from ctypes import wintypes
import os
import re
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt
from types import SimpleNamespace

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
EMU_PT = 12700


class Metrics(AbstractContextManager):
    """GDI works offline on Windows, including fallback fonts and bold faces."""
    def __init__(self):
        if os.name != 'nt':
            raise OSError('此环境暂不支持字体测量。')
        self.gdi = ctypes.WinDLL('gdi32', use_last_error=True)
        self.gdi.CreateCompatibleDC.argtypes = [wintypes.HDC]
        self.gdi.CreateCompatibleDC.restype = wintypes.HDC
        self.gdi.CreateFontW.argtypes = [ctypes.c_int]*5 + [wintypes.DWORD]*8 + [wintypes.LPCWSTR]
        self.gdi.CreateFontW.restype = wintypes.HANDLE
        self.gdi.SelectObject.argtypes = [wintypes.HDC, wintypes.HANDLE]
        self.gdi.SelectObject.restype = wintypes.HANDLE
        self.gdi.DeleteObject.argtypes = [wintypes.HANDLE]
        self.gdi.DeleteDC.argtypes = [wintypes.HDC]
        self.gdi.GetTextExtentPoint32W.argtypes = [wintypes.HDC, wintypes.LPCWSTR, ctypes.c_int, ctypes.POINTER(wintypes.SIZE)]
        self.gdi.GetTextFaceW.argtypes = [wintypes.HDC, ctypes.c_int, wintypes.LPWSTR]
        self.dc = self.gdi.CreateCompatibleDC(None)
        if not self.dc: raise OSError('无法初始化字体测量。')
        self.fonts, self.cache, self.initial, self.available = {}, {}, None, {}

    def measure(self, text, style, scale=1):
        name, size, bold, italic, tracking = style
        pixels = max(1, round(size * scale * 8))
        key = (name, pixels, bold, italic)
        cache_key = (text, key, tracking, round(scale, 4))
        if cache_key in self.cache: return self.cache[cache_key]
        if key not in self.fonts:
            self.fonts[key] = self.gdi.CreateFontW(-pixels, 0, 0, 0, 700 if bold else 400,
                int(italic), 0, 0, 1, 0, 0, 4, 0, name)
            if not self.fonts[key]: raise OSError('无法加载字体。')
        previous = self.gdi.SelectObject(self.dc, self.fonts[key])
        if self.initial is None: self.initial = previous
        extent = wintypes.SIZE()
        length = len(text.encode('utf-16-le')) // 2
        if not self.gdi.GetTextExtentPoint32W(self.dc, text or ' ', length or 1, ctypes.byref(extent)):
            raise OSError('无法测量文字。')
        value = (extent.cx / 8 + max(0, len(text)-1) * tracking * scale,
                 max(size * scale, extent.cy / 8) * 1.025)
        self.cache[cache_key] = value
        return value

    def has_font(self, name):
        if name in self.available:return self.available[name]
        self.measure('Display', (name,12,False,False,0))
        self.gdi.SelectObject(self.dc,self.fonts[(name,96,False,False)])
        face=ctypes.create_unicode_buffer(128)
        self.gdi.GetTextFaceW(self.dc,128,face)
        self.available[name]=face.value.casefold()==name.casefold()
        return self.available[name]

    def __exit__(self, *args):
        if self.initial: self.gdi.SelectObject(self.dc, self.initial)
        for font in self.fonts.values(): self.gdi.DeleteObject(font)
        self.gdi.DeleteDC(self.dc)


def _node(parent, tag):
    return None if parent is None else parent.find(f'{{{A}}}{tag}')


def _placeholder_chain(shape, slide):
    result = []
    if shape is None or not getattr(shape, 'is_placeholder', False): return result
    idx = shape.placeholder_format.idx
    kind = shape.placeholder_format.type
    for owner in (slide.slide_layout, slide.slide_layout.slide_master):
        candidates = list(owner.placeholders)
        match = next((p for p in candidates if p.placeholder_format.idx == idx), None)
        if match is None: match = next((p for p in candidates if p.placeholder_format.type == kind), None)
        if match is not None and match.has_text_frame: result.append(match.text_frame._txBody)
    return result


def _defaults(paragraph, frame, shape, slide, document):
    direct = paragraph._p.find(f'{{{A}}}pPr')
    level = int(direct.get('lvl','0')) if direct is not None else 0
    sources = [direct]
    for body in [frame._txBody, *_placeholder_chain(shape, slide)]:
        lst = _node(body, 'lstStyle')
        sources.append(_node(lst, f'lvl{level+1}pPr'))
        if body is not frame._txBody:
            paras = body.findall(f'{{{A}}}p')
            if paras: sources.append(_node(paras[min(level, len(paras)-1)], 'pPr'))
    if slide is not None:
        master = slide.slide_layout.slide_master
        txstyles = master._element.find(f'{{{P}}}txStyles')
        kind = 'otherStyle'
        if shape is not None and getattr(shape, 'is_placeholder', False):
            kind = 'titleStyle' if int(shape.placeholder_format.type) in {1, 3} else 'bodyStyle'
        textstyle = None if txstyles is None else txstyles.find(f'{{{P}}}{kind}')
        sources.append(_node(textstyle, f'lvl{level+1}pPr'))
    default = document._element.find(f'{{{P}}}defaultTextStyle')
    sources.append(_node(default, f'lvl{level+1}pPr'))
    return [s for s in sources if s is not None]


def _font_name(value, slide, document):
    if not value or not value.startswith('+'): return value or 'Arial'
    part = slide.slide_layout.slide_master.part if slide is not None else document.part
    try:
        theme = next(rel.target_part for rel in part.rels.values() if rel.reltype.endswith('/theme'))
        from lxml import etree
        root = etree.fromstring(theme.blob)
        kind = 'majorFont' if value.startswith('+mj') else 'minorFont'
        name = root.find(f'.//{{{A}}}{kind}/{{{A}}}latin').get('typeface')
        return name or 'Arial'
    except (StopIteration, AttributeError): return 'Arial'


def _style(run, defaults, slide, document, font_scale):
    sources = [_node(run._r, 'endParaRPr' if run._r.tag == f'{{{A}}}p' else 'rPr'), *[_node(p, 'defRPr') for p in defaults]]
    sources = [s for s in sources if s is not None]
    def attr(name, fallback):
        return next((s.get(name) for s in sources if s.get(name) is not None), fallback)
    name = next((n.get('typeface') for s in sources for n in [_node(s,'latin')] if n is not None), '+mn-lt')
    return (_font_name(name, slide, document), float(attr('sz',1800))/100 * font_scale,
            attr('b','0') in {'1','true','on'}, attr('i','0') in {'1','true','on'}, float(attr('spc',0))/100)


def _paragraph_data(frame, shape, slide, document):
    from .text_runs import text_runs, BreakRun
    norm = _node(frame._txBody.bodyPr, 'normAutofit')
    font_scale = float(norm.get('fontScale','100000')) / 100000 if norm is not None else 1
    result = []
    class BlankRun:
        def __init__(self,paragraph): self._r,self.paragraph=paragraph._p,paragraph
        text=' '
        @property
        def font(self):
            from pptx.text.text import Font
            return Font(self.paragraph._p.get_or_add_endParaRPr())
    for paragraph in frame.paragraphs:
        defaults = _defaults(paragraph, frame, shape, slide, document)
        def attr(name, fallback=0):
            return float(next((p.get(name) for p in defaults if p.get(name) is not None), fallback)) / EMU_PT
        pieces = []
        last_style = None
        for run in text_runs(paragraph):
            if isinstance(run,BreakRun) and _node(run._r,'rPr') is None and last_style is not None:
                style = last_style
            else:
                style = _style(run,defaults,slide,document,font_scale)
            pieces.append((run, style)); last_style = style
        if not pieces:
            run=BlankRun(paragraph)
            pieces.append((run,_style(run,defaults,slide,document,font_scale)))
        spacing = {}
        for name in ['lnSpc','spcBef','spcAft']:
            node = next((node for p in defaults for node in [_node(p,name)] if node is not None), None)
            if node is not None and len(node):
                child = node[0]
                spacing[name] = (child.tag.rsplit('}',1)[-1], float(child.get('val','0')))
        bullet = any(_node(p, 'buChar') is not None or _node(p,'buAutoNum') is not None for p in defaults)
        result.append((paragraph, pieces, attr('marL'), attr('marR'), attr('indent'), spacing, bullet))
    return result


def estimate(data, width, metrics, scale=1, wrap=True):
    height, widest, clipped = 0, 0, False
    for paragraph, pieces, left, right, indent, spacing, bullet in data:
        chars = [(char, style) for run, style in pieces for char in run.text]
        if not chars: continue
        available = max(1, width-left-right)
        first = max(0, indent) + (max(style[1] for _,style in chars)*scale*.65 if bullet and not left else 0)
        used, line_height, lines = first, 0, []
        text = ''.join(char for char,_ in chars)
        for token in re.finditer(r'[\v\n]|[^\S\v\n]+|[^\s]+', text):
            value = token.group()
            if value in {'\v','\n'}:
                lines.append(line_height or max(style[1] for _,style in chars)*scale*1.2)
                widest = max(widest, used); used, line_height = 0, 0
                continue
            fragments = []
            for char, style in chars[token.start():token.end()]:
                if fragments and fragments[-1][1] == style:
                    fragments[-1] = (fragments[-1][0]+char,style)
                else: fragments.append((char,style))
            widths = [metrics.measure(s.replace('\t','    '),style,scale) for s,style in fragments]
            token_width, token_height = sum(w for w,h in widths), max(h for w,h in widths)
            if value.isspace():
                if used: used += token_width
                continue
            if wrap and used and used+token_width > available:
                lines.append(line_height or token_height)
                widest = max(widest, min(used,available)); used, line_height = 0, 0
            if wrap and token_width > available:
                # PowerPoint can break a long word; measure the actual fragments.
                for char,style in chars[token.start():token.end()]:
                    cw,ch = metrics.measure(char,style,scale)
                    if used and used+cw > available:
                        lines.append(line_height); widest=max(widest,used); used,line_height=0,0
                    used += cw; line_height=max(line_height,ch)
            else:
                used += token_width; line_height=max(line_height,token_height)
        lines.append(line_height or max(style[1] for _,style in chars)*scale*1.2)
        widest=max(widest,used)
        clipped |= used > available+.2
        kind, value = spacing.get('lnSpc',('spcPct',100000))
        line_heights = [max(h, value/100*scale) if kind == 'spcPts' else h*value/100000 for h in lines]
        height += sum(line_heights)
        for name in ['spcBef','spcAft']:
            kind,value = spacing.get(name,('spcPts',0))
            height += value/100 if kind == 'spcPts' else max(line_heights)*value/100000
    return height, widest, clipped


def _frames(document, changed):
    from .pptx_geometry import walk_shapes, dimensions
    def shapes(items,prefix,slide):
        for key,shape,transform in walk_shapes(items,prefix):
            actual_width,actual_height=dimensions(shape,transform)
            sx,sy=actual_width/max(1,shape.width),actual_height/max(1,shape.height)
            if shape.has_table:
                table=shape.table
                for r,row in enumerate(table.rows):
                    for c,cell in enumerate(row.cells):
                        loc=f'{key}/table/{r}/{c}/'
                        if cell.is_spanned or not any(v.startswith(loc) for v in changed): continue
                        width=sum(table.columns[k].width for k in range(c,c+cell.span_width))
                        height=sum(table.rows[k].height for k in range(r,r+cell.span_height))
                        margins=(cell.margin_left,cell.margin_right,cell.margin_top,cell.margin_bottom)
                        yield loc,cell.text_frame,None,slide,round(width*sx),round(height*sy),margins
            elif shape.has_text_frame and any(v.startswith(key+'/p') for v in changed):
                frame=shape.text_frame
                yield key,frame,shape,slide,actual_width,actual_height,(frame.margin_left,frame.margin_right,frame.margin_top,frame.margin_bottom)
            elif shape.has_chart:
                name=str(shape.chart.part.partname).lstrip('/')
                root=getattr(document,'_star_extra_parts',{}).get(name)
                if root is None: continue
                from pptx.text.text import TextFrame
                from ..extract.pptx_extra import Owner
                for i,body in enumerate(root.xpath('.//a:txBody | .//c:rich')):
                    loc=f'part:{name}/body{i}/'
                    if not any(v.startswith(loc) for v in changed): continue
                    owner=next((p for p in body.iterancestors() if p.tag in
                        {f'{{{C}}}dLbl',f'{{{C}}}title'}),None)
                    if owner is None: continue
                    layout=owner.find(f'{{{C}}}layout/{{{C}}}manualLayout')
                    if layout is None:
                        # Office/WPS stores point-label bounds in the c15 extension.
                        layout=next(iter(owner.findall(f'.//{{{C}}}manualLayout')),None)
                    def dimension(tag):
                        value=layout.find(f'{{{C}}}{tag}') if layout is not None else None
                        try: return float(value.get('val')) if value is not None else None
                        except (TypeError,ValueError): return None
                    w,h=dimension('w'),dimension('h')
                    title=owner.tag==f'{{{C}}}title'
                    if not w or not 0<w<=1:
                        count=max((len(cat.findall(f'.//{{{C}}}pt')) for cat in root.findall(f'.//{{{C}}}cat')),default=5)
                        w=.9 if title else .78/max(1,count)
                    if not h or not 0<h<=1: h=.19 if title else .14
                    if not title:
                        # Point-label fractions are resolved inside the plot,
                        # which also reserves axes and title space. Treating
                        # them as the full chart hides the numeric suffix behind
                        # PowerPoint's ellipsis even when our estimate fits.
                        w*=.8;h*=.8
                    frame=TextFrame(body,Owner(shape.chart.part))
                    yield loc,frame,shape,slide,round(actual_width*w),round(actual_height*h),(frame.margin_left,frame.margin_right,frame.margin_top,frame.margin_bottom)
    for index,slide in enumerate(document.slides,1):
        yield from shapes(slide.shapes,f'slide{index}',slide)
    # Shared template objects are measured once; their geometry is retained.
    for name,root in getattr(document,'_star_extra_parts',{}).items():
        if not name.startswith(('ppt/slideMasters/','ppt/slideLayouts/')): continue
        for index,body in enumerate(root.xpath('.//p:sp[not(p:nvSpPr/p:nvPr/p:ph)]/p:txBody')):
            key=f'part:{name}/body{index}/'
            if not any(v.startswith(key) for v in changed): continue
            from pptx.text.text import TextFrame
            from ..extract.pptx_extra import Owner
            frame=TextFrame(body,Owner(None))
            ext=body.getparent().find(f'{{{P}}}spPr/{{{A}}}xfrm/{{{A}}}ext')
            if ext is not None:
                yield key,frame,None,None,int(ext.get('cx')),int(ext.get('cy')),(frame.margin_left,frame.margin_right,frame.margin_top,frame.margin_bottom)


def _all_locators(document):
    from .pptx_geometry import walk_shapes
    result=[]
    for index,slide in enumerate(document.slides,1):
        for key,shape,_ in walk_shapes(slide.shapes,f'slide{index}'):
            if shape.has_text_frame: result.append(key+'/p0')
            elif shape.has_table:
                result.extend(f'{key}/table/{r}/{c}/0' for r,row in enumerate(shape.table.rows)
                              for c,cell in enumerate(row.cells) if not cell.is_spanned)
    return result


def capture_layout(document, locators):
    """Freeze all slide peers so each source style can share a page font size."""
    from .pptx_geometry import walk_shapes, bounds
    snapshots={}
    rects={}
    for index,slide in enumerate(document.slides,1):
        for key,shape,transform in walk_shapes(slide.shapes,f'slide{index}'):
            rects[key]=bounds(shape,transform)
    for key,frame,shape,slide,width,height,margins in _frames(document,list(dict.fromkeys([*locators,*_all_locators(document)]))):
        data=_paragraph_data(frame,shape,slide,document)
        frozen=[(None,[(SimpleNamespace(text=run.text),style) for run,style in pieces],left,right,indent,dict(spacing),bullet)
                for _,pieces,left,right,indent,spacing,bullet in data]
        page=document.slides.index(slide)+1 if slide is not None else key.split('/body')[0]
        snapshots[key]={'data':frozen,'wrap':frame.word_wrap is not False,'auto_size':frame.auto_size,
                        'width':width,'height':height,'margins':margins,'page':page,'rect':rects.get(key),
                        'obstacles':[r for k,r in rects.items() if k!=key and k.startswith(f'slide{page}/')]}
    return snapshots


def _variant(data, compact=False, font=None):
    result=[]
    for paragraph,pieces,left,right,indent,spacing,bullet in data:
        spacing=dict(spacing)
        if compact:
            kind,value=spacing.get('lnSpc',('spcPct',100000))
            # Keep a full glyph-height baseline; remove excess leading only.
            spacing['lnSpc']=('spcPct',100000) if kind=='spcPts' else ('spcPct',min(value,100000))
            for name in ['spcBef','spcAft']:
                kind,value=spacing.get(name,('spcPts',0))
                spacing[name]=(kind,value*.6)
        if font:
            pieces=[(run,(font,*style[1:])) for run,style in pieces]
        result.append((paragraph,pieces,left,right,indent,spacing,bullet))
    return result


def _write_typography(data,scale,font_changed,spacing_changed,padding,baked_scale):
    from pptx.oxml.xmlchemy import OxmlElement
    for paragraph,pieces,_,_,_,spacing,_ in data:
        for run,style in pieces:
            size=round(style[1]*scale,2)
            if not hasattr(run,'font'):
                rpr=_node(run._r,'rPr')
                if rpr is None:
                    rpr=OxmlElement('a:rPr'); run._r.insert(0,rpr)
                rpr.set('sz',str(round(size*100)))
                rpr.set('spc',str(round(style[4]*100)))
                if font_changed:
                    latin=_node(rpr,'latin')
                    if latin is None: latin=OxmlElement('a:latin'); rpr.append(latin)
                    latin.set('typeface',style[0])
                continue
            explicit=run.font.size.pt if run.font.size is not None else None
            if scale<.999 or baked_scale or explicit is not None and abs(explicit-size)>.005:
                run.font.size=Pt(size)
            if font_changed: run.font.name=style[0]
            rpr=run.font._rPr
            rpr.set('spc',str(round(style[4]*100)))
        if paragraph._p.getparent().tag==f'{{{C}}}rich' and pieces:
            # Chart engines use paragraph/end-marker size for line height even
            # when every visible run has already been fitted smaller.
            size=str(round(max(st[1] for _,st in pieces)*scale*100))
            paragraph._p.get_or_add_pPr().get_or_add_defRPr().set('sz',size)
            paragraph._p.get_or_add_endParaRPr().set('sz',size)
        if spacing_changed or scale<.999:
            if 'lnSpc' in spacing:
                kind,value=spacing['lnSpc']
                paragraph.line_spacing=Pt(value/100*scale) if kind=='spcPts' else value/100000
            if spacing_changed:
                for name,attr in [('spcBef','space_before'),('spcAft','space_after')]:
                    kind,value=spacing.get(name,('spcPts',0))
                    if kind=='spcPts': setattr(paragraph,attr,Pt(value/100))
                    else:
                        ppr=paragraph._p.get_or_add_pPr(); old=_node(ppr,name)
                        if old is not None:ppr.remove(old)
                        node=OxmlElement('a:'+name); child=OxmlElement('a:spcPct');child.set('val',str(round(value)));node.append(child);ppr.append(node)
        if padding>.05:
            old=paragraph.space_after.pt if paragraph.space_after is not None else 0
            paragraph.space_after=Pt(old+padding)
        ppr=_node(paragraph._p,'pPr')
        if ppr is not None and paragraph._p.index(ppr):
            paragraph._p.remove(ppr); paragraph._p.insert(0,ppr)


# Newly reduced sizes use clean, editable point sizes. Untouched source sizes stay exact.
FONT_STEPS=(6,7,8,9,10,11,12,14,16,18,20,22,24,26,28,30,32,36,40,44,48,54,60,64,72,80,88,96)


def _cohort(page,style):
    return (page,style[0].casefold(),round(style[1],2),style[2],style[3])


def _step(size,original,visual=False):
    # Measured overflow may require a substantial fit. Additional visual-model
    # requests alone must not repeatedly collapse an already fitted heading.
    floor=min(original,max(12,original*.8)) if visual else min(original,max(6,min(9,original*.65)))
    return next((v for v in reversed(FONT_STEPS) if floor<=v<size-.005),size)


def _sized(data,page,targets):
    return [(p,[(r,(st[0],targets[_cohort(page,st)],*st[2:])) for r,st in pieces],l,rr,i,sp,b)
            for p,pieces,l,rr,i,sp,b in data]


def _reconcile_sizes(plans,targets):
    """Unify close fits, without a cramped label shrinking the entire page."""
    groups={}
    for _,_,_,data,scope,*_ in plans:
        for _,pieces,*_ in data:
            for run,style in pieces:
                if not run.text.strip():continue
                key=_cohort(scope,style)
                groups.setdefault((scope[0],*key[1:]),set()).add(key)
    for keys in groups.values():
        remaining=sorted(keys,key=lambda k:targets[k],reverse=True)
        while remaining:
            largest=targets[remaining[0]]
            cluster=[k for k in remaining if targets[k]>=largest*.85]
            common=min(targets[k] for k in cluster)
            for k in cluster:targets[k]=common
            remaining=[k for k in remaining if k not in cluster]


def _preserve_emphasis(plans,targets):
    """Keep source size contrast within mixed-size paragraphs, across page peers."""
    pairs=set()
    for _,_,_,data,page,*_ in plans:
        for _,pieces,*_ in data:
            styles={_cohort(page,st) for r,st in pieces if r.text.strip()}
            pairs.update((big,small) for big in styles for small in styles if big[2]>small[2]+.5)
    for _ in range(20):
        changed=False
        for big,small in sorted(pairs,key=lambda p:p[0][2]):
            ratio=big[2]/small[2]*.9
            if targets[big]+.01>=targets[small]*ratio:continue
            floor=min(small[2],max(6,min(9,small[2]*.65)))
            candidates=[v for v in FONT_STEPS if floor<=v<=targets[big]/ratio+.005 and v<targets[small]]
            if candidates:targets[small]=max(candidates);changed=True
            else:
                needed=targets[small]*ratio
                bigger=next((v for v in FONT_STEPS if needed<=v<=big[2]),big[2])
                if bigger>targets[big]:targets[big]=bigger;changed=True
        if not changed:break


def fit_document(document, changed, snapshots=None, constraints=None):
    """Plan fits, reconcile per-page source-style cohorts, then verify final sizes.

    constraints are private visual-review frame requests for one further font step.
    They can change typography, never content, object geometry or font families.
    """
    warnings,adjustments=[],[]
    snapshots=snapshots or capture_layout(document,changed)
    requested=list(changed)
    active_pages={v['page'] for k,v in snapshots.items() if any(c.startswith(k.rstrip('/')+'/') for c in requested)}
    peers=[k+('0' if k.endswith('/') else '/p0') for k,v in snapshots.items() if v['page'] in active_pages]
    targets={}
    plans=[]
    origins=getattr(document,'_star_cohort_origins',{})
    document._star_cohort_origins=origins
    from .pptx_style_manifest import run_key
    style_origins=getattr(document,'_star_run_styles',{})
    document._star_run_styles=style_origins
    rendered_sizes={}
    try:
        with Metrics() as metrics:
            for key,frame,shape,slide,w,h,margins in _frames(document,list(dict.fromkeys([*requested,*peers]))):
                original=snapshots.get(key)
                page=original['page'] if original else document.slides.index(slide)+1 if slide is not None else key.split('/body')[0]
                scope=(page,key)
                data=_paragraph_data(frame,shape,slide,document)
                touched=any(c.startswith(key.rstrip('/')+'/') for c in requested) and (not constraints or key in constraints)
                if not any(run.text.strip() for _,pieces,*_ in data for run,_ in pieces):continue
                restored=[]
                for p,pieces,l,r,i,sp,b in data:
                    values=[]
                    for run,style in pieces:
                        if hasattr(run,'_r'):style_origins.setdefault(run._r,run_key(run))
                        current_size=style[1]
                        style=origins.setdefault(run._r,style)
                        c=_cohort(scope,style)
                        rendered_sizes[c]=min(rendered_sizes.get(c,current_size),current_size)
                        values.append((run,style))
                    restored.append((p,values,l,r,i,sp,b))
                data=restored
                # Chinese tracking applied unchanged to many more Latin letters
                # creates artificial overflow. Keep modest spacing in translated
                # paragraphs before reducing font size; intentional English stays.
                if original:
                    adapted=[]
                    for index,(p,pieces,l,r,i,sp,b) in enumerate(data):
                        source_text=''.join(run.text for run,_ in original['data'][index][1]) if index<len(original['data']) else ''
                        translated=re.search(r'[\u3400-\u9fff]',source_text) and not re.search(r'[\u3400-\u9fff]',''.join(run.text for run,_ in pieces))
                        if translated:
                            cap=0 if constraints and key in constraints else None
                            pieces=[(run,(*st[:4],min(st[4],cap if cap is not None else st[1]*.03))) for run,st in pieces]
                        adapted.append((p,pieces,l,r,i,sp,b))
                    data=adapted
                location={'slide':page if isinstance(page,int) else None,'text':f'文本框 {shape.shape_id}' if shape is not None else '表格或共享文字','locator':key}
                props=frame._txBody.bodyPr
                if any(n.tag==f'{{{A}}}fld' for n in frame._txBody.iter()):
                    continue
                if props.get('vert','horz')!='horz' or int(props.get('numCol','1'))>1:
                    if any(c.startswith(key.rstrip('/')+'/') for c in requested):
                        warnings.append({**location,'reason':'该处文字方向或分栏需要确认。'})
                    continue
                for _,pieces,*_ in data:
                    for _,style in pieces: targets.setdefault(_cohort(scope,style),style[1])
                width=(w-margins[0]-margins[1])/EMU_PT*.975
                height=(h-margins[2]-margins[3])/EMU_PT*.97
                wrap=frame.word_wrap is not False
                natural=False
                if (width<=0 or height<=0) and original and original['auto_size']==MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT:
                    natural=True
                    if width<=0:
                        _,width,_=estimate(original['data'],10000,metrics,wrap=False)
                        width+=max((left+right for _,_,left,right,*_ in original['data']),default=0)+.5
                        wrap=False
                    if height<=0: height=estimate(original['data'],width,metrics,wrap=wrap)[0]+.5
                if width<=0 or height<=0:
                    if touched:warnings.append({**location,'reason':'原稿文字空间过小，请确认该处显示。'})
                    continue
                # Text frames may extend below their visible Chinese text into a picture.
                # Reserve that picture/next object if the source ink did not overlap it.
                if original and original.get('rect') and props.get('anchor','t')=='t':
                    x,y,rw,_=original['rect']; source_h=estimate(original['data'],width,metrics,wrap=wrap)[0]
                    ink_bottom=y/EMU_PT+margins[2]/EMU_PT+source_h
                    for ox,oy,ow,oh in original.get('obstacles',[]):
                        overlap=min(x+rw,ox+ow)-max(x,ox)
                        if ow>0 and oh>0 and overlap>min(rw,ow)*.25 and oy/EMU_PT>=ink_bottom+1:
                            height=min(height,(oy-y-margins[2])/EMU_PT-2)
                height=max(.1,height)
                budgets=None
                # Paragraphs share the actual frame height. Reserving each
                # Chinese line's old height wastes space after compression and
                # forces English into tiny fonts despite unused space below.
                units=[[item] for item in data] if budgets else [data]
                unit_plans=[]
                for index,unit in enumerate(units):
                    budget=budgets[index] if budgets else height
                    chosen=unit; compact=False
                    if touched:
                        if estimate(chosen,width,metrics,wrap=wrap)[0]>budget:
                            chosen=_variant(unit,compact=True);compact=True
                        cohorts={_cohort(scope,st) for _,pieces,*_ in chosen for _,st in pieces}
                        local={c:targets[c] for c in cohorts}
                        if constraints and key in constraints:
                            for c in cohorts: local[c]=min(local[c],_step(rendered_sizes.get(c,local[c]),c[2],visual=True))
                        def fits():
                            measured=_sized(chosen,scope,local)
                            mh,_,clip=estimate(measured,width,metrics,wrap=wrap)
                            return mh<=budget and not clip
                        while not fits():
                            reduced={c:_step(v,c[2],visual=bool(constraints)) for c,v in local.items()}
                            if reduced==local:break
                            local=reduced
                        if not fits() and not wrap and not natural:
                            wrap=True;frame.word_wrap=True
                            # With wrapping, retry from current page requirements.
                            local={c:targets[c] for c in cohorts}
                            while not fits():
                                reduced={c:_step(v,c[2],visual=bool(constraints)) for c,v in local.items()}
                                if reduced==local:break
                                local=reduced
                        for c,v in local.items(): targets[c]=min(targets[c],v)
                    unit_plans.append((chosen,budget,compact))
                plans.append((key,frame,location,data,scope,width,height,wrap,natural,budgets,unit_plans,touched))
            # Close peers share a clean size; unrelated crowded frames remain local.
            _reconcile_sizes(plans,targets)
            _preserve_emphasis(plans,targets)
            for key,frame,location,data,page,width,height,wrap,natural,budgets,units,touched in plans:
                any_change=False
                for index,(unit,budget,compact) in enumerate(units):
                    chosen=_sized(unit,page,targets)
                    measured,_,clip=estimate(chosen,width,metrics,wrap=wrap)
                    reduced=any(targets[_cohort(page,st)]<st[1]-.005 for _,pieces,*_ in unit for _,st in pieces)
                    ok=measured<=budget+.2 and not clip
                    padding=max(0,budget-measured-.1) if touched and budgets and index<len(units)-1 and ok else 0
                    if touched or reduced:
                        _write_typography(chosen,1,False,compact,padding,True)
                        any_change=True
                    if touched and not ok:
                        warnings.append({**location,'reason':f'第 {index+1} 段空间不足，已保留完整文字，请确认该处排版。'})
                    if reduced or compact:
                        sizes=sorted({st[1] for _,pieces,*_ in chosen for _,st in pieces})
                        adjustments.append({**location,'paragraph':index+1,'font_size':max(sizes),'font_sizes':sizes,
                                            'font':None,'spacing_compacted':compact,'original_bounds':natural})
                if any_change:
                    if not natural: frame.auto_size=MSO_AUTO_SIZE.NONE
                    # Include untouched English peers in the minimal-part saver.
                    locator=key+('0' if key.endswith('/') else '/p0')
                    if locator not in changed:changed.append(locator)
    except OSError as error:
        warnings.append({'slide':None,'text':'','reason':'部分文字排版检查未完成，请确认文档显示。'})
    return adjustments,warnings
