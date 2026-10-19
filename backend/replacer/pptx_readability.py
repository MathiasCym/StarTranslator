"""Apply image-model readability proposals only after local collision/fit checks."""
import math
from .pptx_geometry import walk_shapes,bounds,IDENTITY
from .pptx_layout import Metrics,_paragraph_data,_write_typography,estimate,EMU_PT

A='http://schemas.openxmlformats.org/drawingml/2006/main'

def _area(a,b):
    return max(0,min(a[0]+a[2],b[0]+b[2])-max(a[0],b[0]))*max(0,min(a[1]+a[3],b[1]+b[3])-max(a[1],b[1]))

def _number(value):
    return isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value)

def _ink(shape,transform,slide,document,metrics,data=None,width=None,height=None,wrap=None,origin=None):
    """Conservative text ink, rather than an empty overlapping frame."""
    frame=shape.text_frame
    x,y,w,h=[v/EMU_PT for v in bounds(shape,transform)]
    if origin is not None:x,y=origin
    w,h=width if width is not None else w,height if height is not None else h
    left,right,top,bottom=[v/EMU_PT for v in (frame.margin_left,frame.margin_right,frame.margin_top,frame.margin_bottom)]
    inner=max(.1,w-left-right)
    measured,widest,_=estimate(data or _paragraph_data(frame,shape,slide,document),inner*.975,metrics,
        wrap=frame.word_wrap is not False if wrap is None else wrap)
    y+=top
    anchor=frame._txBody.bodyPr.get('anchor','t')
    if anchor in {'ctr','b'}:y+=max(0,h-top-bottom-measured)*(.5 if anchor=='ctr' else 1)
    x+=left
    alignments={p.alignment for p in frame.paragraphs if p.text.strip()}
    from pptx.enum.text import PP_ALIGN
    if alignments=={PP_ALIGN.CENTER}:x+=max(0,inner-widest)/2
    elif alignments=={PP_ALIGN.RIGHT}:x+=max(0,inner-widest)
    return (x,y,widest,measured)


def apply_actions(document,issues,changed,native_pages,feedback=None):
    """The model chooses the size and expansion; this code never invents either."""
    fixes=[];modified={}
    with Metrics() as metrics:
        for item in issues:
            if item.get('kind')!='readability':continue
            number=item['slide'];slide=document.slides[number-1]
            shapes=list(walk_shapes(slide.shapes,f'slide{number}'))
            match=next((s for s in shapes if s[0]==item.get('locator')),None)
            if not match:continue
            key,shape,transform=match
            if not shape.has_text_frame or not shape.text.strip() or shape.rotation%360:continue
            # Group scaling is supported; rotated/flipped groups require an
            # inverse polygon fit rather than changing axis-aligned extents.
            if transform[1] or transform[2] or transform[0]<=0 or transform[3]<=0:continue
            x,y,w,h=[v/EMU_PT for v in bounds(shape,transform)]
            new_x,new_y=item.get('x',x),item.get('y',y)
            width,height=item.get('width',w),item.get('height',h)
            if not all(_number(v) for v in (new_x,new_y,width,height,item.get('font_size'))):continue
            # A model may round an unchanged side down. Preserve the space
            # already available instead of rejecting a useful expansion on the
            # other side; no extra area beyond old/model bounds is invented.
            width,height=max(w,width),max(h,height)
            if new_x<0 or new_y<0:continue
            if new_x+width>document.slide_width/EMU_PT-1 or new_y+height>document.slide_height/EMU_PT-1:continue
            props=shape.text_frame._txBody.bodyPr
            if props.get('vert','horz')!='horz' or int(props.get('numCol','1'))!=1:continue
            data=_paragraph_data(shape.text_frame,shape,slide,document)
            current=max((st[1] for _,pieces,*_ in data for r,st in pieces if r.text.strip()),default=0)
            target=item['font_size']
            if not current or not min(current,12)<=target<=96:continue
            factor=target/current
            candidate=[(p,[(r,(st[0],st[1]*factor,*st[2:])) for r,st in pieces],l,rr,i,sp,b)
                for p,pieces,l,rr,i,sp,b in data]
            frame=shape.text_frame
            wrap=item.get('wrap',frame.word_wrap is not False)
            if type(wrap) is not bool:continue
            if (abs(target-current)<=.25 and width<=w+.25 and height<=h+.25
                    and abs(new_x-x)<=.25 and abs(new_y-y)<=.25 and wrap==(frame.word_wrap is not False)):continue
            inner_w=(width-(frame.margin_left+frame.margin_right)/EMU_PT)*.975
            inner_h=(height-(frame.margin_top+frame.margin_bottom)/EMU_PT)*.97
            if inner_w<=0 or inner_h<=0:continue
            measured,widest,clipped=estimate(candidate,inner_w,metrics,wrap=wrap)
            if clipped or measured>inner_h:
                if feedback is not None:feedback.append({'slide':number,'shape_id':shape.shape_id,'proposed_font_size':target,
                    'proposed_width':width,'proposed_height':height,'reason':'该字号在提议的文字框中放不下，请重新选择尺寸；不要进一步缩小当前文字。',
                    'measured_min_height':round(measured/.97+(frame.margin_top+frame.margin_bottom)/EMU_PT+1,2),
                    'wrapping_enabled':wrap,'measured_min_width_without_wrapping':round(widest/.975+(frame.margin_left+frame.margin_right)/EMU_PT+1,2),
                    'hint':'不换行时必须提供足够宽度；需要多行时明确指定 wrap:true 并留出足够高度。'})
                continue
            before=(x,y,w,h);blocked=False
            old_ink=_ink(shape,transform,slide,document,metrics)
            new_ink=_ink(shape,transform,slide,document,metrics,candidate,width,height,wrap,(new_x,new_y))
            native={v['id']:v for v in native_pages[number-1].get('shapes',[]) if isinstance(v.get('id'),int)}
            for peer_key,peer,matrix in shapes:
                if peer_key==key:continue
                rect=tuple(v/EMU_PT for v in bounds(peer,matrix))
                if peer.has_text_frame:
                    if not peer.text.strip():continue
                    info=native.get(peer.shape_id) if matrix==IDENTITY else None
                    rect=modified.get((number,peer_key)) or ((info['bound_x'],info['bound_y'],info['bound_width'],info['bound_height'])
                        if info else _ink(peer,matrix,slide,document,metrics))
                elif not (peer.has_table or peer.has_chart or peer.shape_type==13):continue
                # A picture already containing this entire text box is an
                # existing background. Foreground objects remain obstacles.
                elif _area(before,rect)>=w*h-.5:continue
                if (_area((new_x,new_y,width,height),rect)>_area(before,rect)+.5
                        or _area(new_ink,rect)>_area(old_ink,rect)+.5):blocked=True;break
            if blocked:
                if feedback is not None:feedback.append({'slide':number,'shape_id':shape.shape_id,
                    'reason':'提议的文字区域会新增遮挡其他文字或前景对象，请观察图像后重新选择可用区域。'})
                continue
            xf=shape._element.find(f'{{http://schemas.openxmlformats.org/presentationml/2006/main}}spPr/{{{A}}}xfrm')
            ext=None if xf is None else xf.find(f'{{{A}}}ext')
            off=None if xf is None else xf.find(f'{{{A}}}off')
            if ext is None or off is None:continue
            before_extent=(int(ext.get('cx')),int(ext.get('cy')))
            before_offset=(int(off.get('x')),int(off.get('y')))
            after_extent=(round(width*EMU_PT/transform[0]),round(height*EMU_PT/transform[3]))
            after_offset=(round((new_x*EMU_PT-transform[4])/transform[0]),round((new_y*EMU_PT-transform[5])/transform[3]))
            # Only the model-approved text shape may change its validated
            # position/extents. Pictures, tables and all other guards stay.
            approved=getattr(document,'_star_visual_geometry',{})
            part=str(slide.part.partname).lstrip('/')
            entries=approved.setdefault(part,{})
            previous=entries.get(shape.shape_id,{})
            entries[shape.shape_id]={'before':previous.get('before',before_extent),'after':after_extent,
                'before_offset':previous.get('before_offset',before_offset),'after_offset':after_offset}
            document._star_visual_geometry=approved
            shape.width,shape.height=after_extent
            shape.left,shape.top=after_offset
            frame.word_wrap=wrap
            _write_typography(candidate,1,False,False,0,True)
            from pptx.enum.text import MSO_AUTO_SIZE
            frame.auto_size=MSO_AUTO_SIZE.NONE
            modified[(number,key)]=new_ink
            for index in range(len(frame.paragraphs)):
                loc=f'{key}/p{index}'
                if loc not in changed:changed.append(loc)
            fixes.append({'slide':number,'locator':key,'font_size':target,'model_readability':True,
                'x':new_x,'y':new_y,'width':width,'height':height,'reason':item['reason']})
    return fixes
